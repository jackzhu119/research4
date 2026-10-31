"""Synthetic tests; these fixtures are not research observations or results."""
import csv
import gzip
import json
from pathlib import Path
import sys
import tempfile
import unittest

import numpy as np
from sklearn.metrics import accuracy_score, balanced_accuracy_score, confusion_matrix, f1_score

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import export_cross_plant_transfer_20261031 as export


class TransferExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def write(self, path, value):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(value), encoding="utf-8")

    def csv_gz(self, path, rows):
        path.parent.mkdir(parents=True, exist_ok=True)
        with gzip.open(path, "wt", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)

    def metric(self, y, p):
        matrix = confusion_matrix(y, p, labels=export.CLASSES)
        return {"accuracy": accuracy_score(y, p), "macro_f1": f1_score(y, p, labels=export.CLASSES, average="macro", zero_division=0),
                "balanced_accuracy": balanced_accuracy_score(y, p), "correct": int(matrix.trace()), "total": len(y),
                "confusion_matrix": matrix.tolist(), "class_order": list(export.CLASSES)}

    def prediction_fixture(self):
        y = np.asarray([c for c in export.CLASSES for _ in range(export.COUNTS[c])])
        rows, pooled, folds, deltas = [], [], [], []
        look = {}
        for role in export.ROLES:
            for seed in export.SEEDS:
                p = y.copy()
                count = {"A": 20, "Aextra": 15, "B": 10, "C": 30, "source_policy": 10}[role] + seed - 42
                p[:count] = "ZY02"
                current = [{"role": role, "seed": seed, "record_id": f"synthetic{i:04d}", "outer_fold": i % 5,
                            "y_true": true, "y_predicted": pred} for i, (true, pred) in enumerate(zip(y, p))]
                rows.extend(current)
                m = self.metric(y, p)
                look[role, seed] = m
                pooled.append({"role": role, "seed": seed, "metrics": m})
                for fold in range(5):
                    mask = np.arange(855) % 5 == fold
                    folds.append({"role": role, "seed": seed, "outer_fold": fold, "metrics": self.metric(y[mask], p[mask])})
                if role != "A":
                    b = look["A", seed]
                    diff = m["correct"] - b["correct"]
                    deltas.append({"role": role, "seed": seed, "corrected": max(diff, 0), "broken": max(-diff, 0),
                                   **{k + "_delta": m[k] - b[k] for k in export.METRICS}})
        self.csv_gz(self.root / "scored_outer_predictions.csv.gz", rows)
        self.write(self.root / "pooled_metrics_by_seed.json", pooled)
        self.write(self.root / "paired_scratch_deltas.json", deltas)
        self.write(self.root / "fold_metrics.json", folds)
        return pooled

    def test_release_must_follow_complete_hash_bound_scoring(self):
        output = self.root / "report"
        release_path = self.root / "release.json"
        self.write(release_path, {"allow_aggregate_export": True, "created_at_utc": "synthetic"})
        with self.assertRaisesRegex(export.ExportError, "Scoring is incomplete"):
            export.export(self.root, output, release_path)
        self.assertFalse(output.exists())
        self.write(self.root / "scoring_manifest.json", {"synthetic": True})
        release = {"allow_aggregate_export": True, "created_at_utc": "synthetic", "scoring_manifest_sha256": export.sha(self.root / "scoring_manifest.json")}
        export.verify_release(self.root, release)
        self.write(self.root / "scoring_manifest.json", {"synthetic": "tampered"})
        with self.assertRaisesRegex(export.ExportError, "not bound"):
            export.verify_release(self.root, release)

    def test_all_case_metrics_confusions_and_pairs_recomputed(self):
        self.prediction_fixture()
        groups = export.prediction_groups(self.root)
        rows, confusions, classes = export.classification_rows(self.root, groups)
        self.assertEqual((len(rows), len(confusions), len(classes)), (15, 735, 105))
        paired = export.paired_rows(self.root, rows, groups)
        gain = next(r for r in paired if r["baseline"] == "A" and r["role"] == "B" and r["seed"] == 42)
        self.assertEqual((gain["corrected"], gain["broken"], gain["net_correct"]), (10, 0, 10))
        self.assertAlmostEqual(gain["accuracy_change_percentage_points"], 1000 / 855)
        self.assertAlmostEqual(gain["accuracy_relative_change_percent"], 1000 / 835)
        self.assertEqual(len(export.paired_predictions(groups)), 2565)
        self.assertEqual(len(export.fold_rows(self.root, groups)), 75)
        down = next(r for r in paired if r["baseline"] == "Aextra" and r["role"] == "C" and r["seed"] == 42)
        self.assertEqual(down["net_correct"], -15)
        self.assertLess(down["accuracy_delta"], 0)

    def test_bad_metrics_or_nan_never_receive_false_recalculation_pass(self):
        pooled = self.prediction_fixture()
        pooled[0]["metrics"]["macro_f1"] = float("nan")
        self.write(self.root / "pooled_metrics_by_seed.json", pooled)
        with self.assertRaisesRegex(export.ExportError, "recalculation disagrees"):
            export.classification_rows(self.root)

    def test_unbound_fold_numbers_are_checked_from_actual_predictions(self):
        self.prediction_fixture()
        groups = export.prediction_groups(self.root)
        rows = export.read(self.root / "fold_metrics.json")
        rows[0]["metrics"]["correct"] += 1
        self.write(self.root / "fold_metrics.json", rows)
        with self.assertRaisesRegex(export.ExportError, "fold metrics disagree"):
            export.fold_rows(self.root, groups)

    def test_missing_seed_or_conflicting_truth_rejected(self):
        self.prediction_fixture()
        with gzip.open(self.root / "scored_outer_predictions.csv.gz", "rt", encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle))
        rows = [r for r in rows if not (r["role"] == "C" and int(r["seed"]) == 44)]
        self.csv_gz(self.root / "scored_outer_predictions.csv.gz", rows)
        with self.assertRaisesRegex(export.ExportError, "All five roles"):
            export.prediction_groups(self.root)
        self.prediction_fixture()
        groups = export.prediction_groups(self.root)
        self.assertTrue(all(len(current) == 855 for current in groups.values()))
        one = groups["B", 42][0]
        one["y_true"] = "ZY07"
        self.csv_gz(self.root / "scored_outer_predictions.csv.gz", [r for current in groups.values() for r in current])
        with self.assertRaises(export.ExportError):
            export.prediction_groups(self.root)

    def test_seed_mean_sd_and_stored_summary_independently_checked(self):
        self.prediction_fixture()
        rows, _, _ = export.classification_rows(self.root)
        summary = export.summarize(rows)
        a = next(r for r in summary if r["role"] == "A")
        self.assertAlmostEqual(a["accuracy_sample_sd"], 1 / 855)
        self.assertEqual(a["correct_each_seed"], "42:835/855; 43:834/855; 44:833/855")
        stored = [{"role": r["role"], "observations_per_seed": 855, "seeds": list(export.SEEDS),
                   **{m + s: r[m + s] for m in export.METRICS for s in ("_mean", "_sample_sd")}} for r in summary]
        self.write(self.root / "seed_summary.json", stored)
        export.verify_seed_summary(self.root, summary)
        stored[0]["accuracy_mean"] += .01
        self.write(self.root / "seed_summary.json", stored)
        with self.assertRaisesRegex(export.ExportError, "summary disagrees"):
            export.verify_seed_summary(self.root, summary)

    def model_fixture(self):
        selected, source, outer = [], [], []
        for fold in range(5):
            selected.append({"outer_fold": fold, "shared_profile": {"id": "p48", "epochs": 48}, "policy_arm": "B",
                             "profile_chosen_only_from_A": True, "outer_used_for_selection": False, "source_means_at_shared_profile": {}})
            for phase in ("source", "outer"):
                for inner in ((0, 1) if phase == "source" else (None,)):
                    for profile in (("p48", "p96") if phase == "source" else ("p48",)):
                        for seed in export.SEEDS:
                            for arm in export.ARMS:
                                split = f"fold{fold}.inner{inner}"
                                directory = self.root / "models" / phase / split / profile / str(seed) / arm
                                directory.mkdir(parents=True)
                                (directory / "pipeline.joblib").write_bytes(b"Synthetic not a pickle; never deserialize")
                                (directory / "batch_trace.jsonl.gz").write_bytes(b"synthetic trace byte fixture")
                                self.write(directory / "batch_schedule.json", {"synthetic": True})
                                n = 342 if phase == "source" else 684
                                counts = dict(zip(export.CLASSES, (140, 2, 40, 40, 40, 40, 40)))
                                if n == 684:
                                    counts = {k: 2 * v for k, v in counts.items()}
                                epochs = 48 if profile == "p48" else 96
                                pools = {} if arm == "A" else {"target": 120} if arm == "Aextra" else {"ginseng": 120} if arm == "B" else {"huangqi": 40, "tianma": 40, "ginseng": 40}
                                meta = {"arm": arm, "seed": seed, "profile": {"id": profile, "epochs": epochs, "pretrain_batch_size": 32},
                                    "source_class_counts": counts, "initial_encoder_sha256": "same encoder",
                                    "initial_zy_head_sha256": "same paired head", "fresh_target_optimizer_initial_sha256": "same new optimizer",
                                    "target_order_sha256": "same target order", "BN_after_reset": {"mean": [0.], "var": [1.], "counter": 0},
                                    "target_class_weights": [n / (7 * counts[c]) for c in export.CLASSES],
                                    "target_steps": epochs * (6 if phase == "source" else 11), "pretrain_steps": sum(pools.values()),
                                    "pretrain_seconds": .02 if pools else 0., "target_training_seconds": .08, "elapsed_seconds": .1,
                                    "target_model_parameters": 100, "pretrain_pool_steps": pools, "pretrain_cache_reuse": False,
                                    "target_epoch_loss": [1. / (i + 1) for i in range(epochs)], "pretrain_loss": [1.] * sum(pools.values()),
                                    "batch_trace_sha256": export.sha(directory / "batch_trace.jsonl.gz"),
                                    "batch_schedule_sha256": export.sha(directory / "batch_schedule.json")}
                                self.write(directory / "training.json", meta)
                                record = {"outer_fold": fold, "inner_fold": inner, "arm": arm, "profile": profile, "seed": seed,
                                          "split_id": split, "model_directory": str(directory.relative_to(self.root)),
                                          "training_sha256": export.sha(directory / "training.json"), "pipeline_sha256": export.sha(directory / "pipeline.joblib")}
                                (source if phase == "source" else outer).append(record)
        self.write(self.root / "source_records.json", source)
        self.write(self.root / "outer_model_records.json", outer)
        return {"C_task_order": ["huangqi", "tianma", "ginseng"]}, {"outer_folds": selected}

    def test_all_actual_metadata_costs_and_paired_finetune_states(self):
        protocol, selection = self.model_fixture()
        costs, losses, pools, paired, selected = export.model_metadata_rows(self.root, protocol, selection)
        self.assertEqual((len(costs), len(paired), len(selected)), (300, 75, 5))
        self.assertEqual(sum(r["pretrain_steps"] for r in costs), 27000)
        self.assertTrue(all(r["presentations_are_not_independent_N"] for r in pools))
        totals = export.computational_summary(costs)
        self.assertEqual(sum(r["actual_pipeline_records"] for r in totals), 300)
        source = export.read(self.root / "source_records.json")
        row = next(r for r in source if r["arm"] == "B")
        path = self.root / row["model_directory"] / "training.json"
        m = export.read(path)
        m["target_order_sha256"] = "unpaired order"
        self.write(path, m)
        row["training_sha256"] = export.sha(path)
        self.write(self.root / "source_records.json", source)
        with self.assertRaisesRegex(export.ExportError, "Paired fine-tune state/order differs"):
            export.model_metadata_rows(self.root, protocol, selection)

    def test_path_boundary_and_historical_context_not_a_gain_baseline(self):
        with self.assertRaises(export.ExportError):
            export.safe_member(self.root, "../outside")
        self.assertAlmostEqual(export.HISTORICAL["accuracy"], 821 / 855)
        self.assertIn("not the matched transfer baseline", export.HISTORICAL["interpretation"])

    def complete_fixture(self):
        self.prediction_fixture()
        protocol, selection = self.model_fixture()
        protocol.update(seeds=list(export.SEEDS), arms=list(export.ARMS), roles=list(export.ROLES), classes=list(export.CLASSES),
                        source_pipeline_budget=240, outer_pipeline_budget=60)
        self.write(self.root / "protocol.json", protocol)
        self.write(self.root / "preparation.json", {"synthetic": True})
        manifest = {"protocol_sha256": export.sha(self.root / "protocol.json"),
                    "preparation_sha256": export.sha(self.root / "preparation.json"), "source_sha256": {}}
        self.write(self.root / "run_manifest.json", manifest)
        self.write(self.root / "source_release_snapshot.json", {"allow_source_training": True, **manifest})
        (self.root / "source_predictions.csv.gz").write_bytes(b"synthetic source snapshot; export does not score it")
        self.write(self.root / "candidate_summaries.json", [{"synthetic": True}])
        selection.update(source_pipeline_fits=240, outer_used_for_selection=False,
                         **{k: export.sha(self.root / name) for k, name in (("source_release_sha256", "source_release_snapshot.json"),
                             ("run_manifest_sha256", "run_manifest.json"), ("source_records_sha256", "source_records.json"),
                             ("source_predictions_sha256", "source_predictions.csv.gz"), ("candidate_summaries_sha256", "candidate_summaries.json"))})
        self.write(self.root / "selection_lock.json", selection)
        self.write(self.root / "outer_release_snapshot.json", {"allow_outer_evaluation": True,
            **{k: export.sha(self.root / name) for k, name in (("selection_lock_sha256", "selection_lock.json"),
                ("protocol_sha256", "protocol.json"), ("run_manifest_sha256", "run_manifest.json"))}})
        self.write(self.root / "outer_role_aliases.json", [{"synthetic": True}] * 15)
        (self.root / "outer_predictions.csv.gz").write_bytes(b"synthetic locked pre-truth byte snapshot")
        lock = {"outer_truth_parsed": False, "actual_outer_models": 60, "alias_cases": 15, "prediction_rows": 12825,
            **{k: export.sha(self.root / name) for k, name in (("prediction_sha256", "outer_predictions.csv.gz"),
                ("outer_model_records_sha256", "outer_model_records.json"), ("outer_role_aliases_sha256", "outer_role_aliases.json"),
                ("selection_lock_sha256", "selection_lock.json"), ("outer_release_sha256", "outer_release_snapshot.json"),
                ("run_manifest_sha256", "run_manifest.json"), ("protocol_sha256", "protocol.json"))}}
        self.write(self.root / "prediction_lock.json", lock)
        rows, _, _ = export.classification_rows(self.root)
        summary = export.summarize(rows)
        self.write(self.root / "seed_summary.json", [{"role": r["role"], "observations_per_seed": 855, "seeds": list(export.SEEDS),
            **{m + suffix: r[m + suffix] for m in export.METRICS for suffix in ("_mean", "_sample_sd")}} for r in summary])
        self.write(self.root / "scoring_manifest.json", {"created_at_utc": "synthetic", "new_training_fits": 0,
            "real_target_observations": 855, "three_seeds_are_not_new_specimens": True, "study_status": "SYNTHETIC_FIXTURE_ONLY",
            **{k: export.sha(self.root / name) for k, name in (("prediction_lock_sha256", "prediction_lock.json"),
                ("prediction_sha256", "outer_predictions.csv.gz"), ("metrics_sha256", "pooled_metrics_by_seed.json"),
                ("summary_sha256", "seed_summary.json"), ("paired_deltas_sha256", "paired_scratch_deltas.json"),
                ("scored_predictions_sha256", "scored_outer_predictions.csv.gz"))}})
        release = self.root / "root_export_release.json"
        self.write(release, {"allow_aggregate_export": True, "created_at_utc": "synthetic",
                            "scoring_manifest_sha256": export.sha(self.root / "scoring_manifest.json")})
        return release

    def test_complete_locked_synthetic_export_creates_workbook_without_model_execution(self):
        from openpyxl import load_workbook
        release = self.complete_fixture()
        output = self.root / "export"
        result = export.export(self.root, output, release, make_figures=False)
        self.assertEqual(result["tables"], 15)
        self.assertEqual(result["model_seed_cases"], 15)
        self.assertEqual(result["new_training_fits"], 0)
        book = load_workbook(output / "cross_plant_transfer_comparison.xlsx", read_only=True, data_only=True)
        self.assertEqual(len(book.sheetnames), 15)
        self.assertEqual(book["paired_predictions"].max_row, 2566)
        book.close()
        self.assertEqual(export.read(output / "export_manifest.json")["benchmark_score_calls"], 0)
        with self.assertRaises(FileExistsError):
            export.export(self.root, output, release, make_figures=False)
        (self.root / "source_predictions.csv.gz").write_bytes(b"tampered after source lock")
        with self.assertRaisesRegex(export.ExportError, "Frozen export dependency changed"):
            export.verify_scored(self.root)


if __name__ == "__main__":
    unittest.main()
