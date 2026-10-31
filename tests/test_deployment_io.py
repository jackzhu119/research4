import csv
import hashlib
from pathlib import Path
import tempfile
import unittest

import numpy as np
from openpyxl import Workbook, load_workbook

from herb_spectra.deployment_io import PredictionInputError, load_prediction_input, write_prediction_template


class DeploymentIOTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)
        self.axis = np.array([.1, .2, .3])
        self.x = np.array([[1., 2., 4.], [-1., 3., 2.]])

    def tearDown(self):
        self.temp.cleanup()

    def npz(self, name="spectra.npz", **changes):
        payload = {"X_alpha": self.x, "axis": self.axis, "record_ids": np.array(["ZY01-new", "ZY07-new"])}
        payload.update(changes)
        path = self.directory/name
        np.savez(path, **payload)
        return path

    def workbook(self, headers, rows, sheet="吸收系数", name="input.xlsx"):
        path = self.directory/name
        book = Workbook()
        book.active.title = sheet
        book.active.append(headers)
        for row in rows:
            book.active.append(row)
        book.save(path)
        book.close()
        return path

    def test_label_free_npz_keeps_negative_values_and_never_encodes_ids(self):
        p = self.npz()
        before = p.read_bytes()
        result = load_prediction_input(p, self.axis)
        np.testing.assert_array_equal(result["X_alpha"], self.x)
        self.assertEqual(result["X_alpha"].shape, (2, 3))
        self.assertEqual(result["record_ids"].tolist(), ["ZY01-new", "ZY07-new"])
        self.assertEqual(result["qc"]["sample_qc"][1]["negative_cells_selected"], 1)
        self.assertFalse(result["qc"]["negative_values_removed"])
        self.assertEqual(result["source_sha256"], hashlib.sha256(before).hexdigest())
        self.assertEqual(before, p.read_bytes())
        p2 = self.npz("renamed.npz", record_ids=np.array(["unrelated-a", "unrelated-b"]))
        np.testing.assert_array_equal(load_prediction_input(p2, self.axis)["X_alpha"], result["X_alpha"])

    def test_single_sample_is_one_row(self):
        p = self.npz(X_alpha=self.x[:1], record_ids=np.array(["NEW-001"]))
        self.assertEqual(load_prediction_input(p, self.axis)["X_alpha"].shape, (1, 3))

    def test_npz_rejects_ground_truth_unknown_metadata_and_disguised_arrays(self):
        for name, value in [("y", np.array(["ZY01", "ZY07"])), ("labels", [0, 1]),
                            ("class", [0, 1]), ("notes", np.array(["anything"]))]:
            with self.subTest(name=name):
                p = self.npz(**{name: value})
                with self.assertRaisesRegex(PredictionInputError, "exactly"):
                    load_prediction_input(p, self.axis)
        for changes in [{"X_alpha": self.x.astype(str)}, {"axis": np.array([False, True, True])},
                        {"record_ids": np.array([1, 2])}]:
            with self.subTest(changes=list(changes)):
                with self.assertRaises(ValueError):
                    load_prediction_input(self.npz(**changes), self.axis)

    def test_rejects_axis_order_shift_nonfinite_and_shape(self):
        for axis in [self.axis[::-1], np.array([.1, .1, .3]), self.axis+.001, np.array([.1, np.nan, .3])]:
            with self.subTest(axis=axis):
                with self.assertRaises(PredictionInputError):
                    load_prediction_input(self.npz(axis=axis), self.axis)
        for x in [self.x[:, :2], self.x[0], np.array([[np.inf, 1., 2.], [1., 2., 3.]])]:
            with self.subTest(shape=np.shape(x)):
                with self.assertRaises(PredictionInputError):
                    load_prediction_input(self.npz(X_alpha=x), self.axis)

    def test_explicit_full_axis_slice_is_contiguous_and_never_interpolates(self):
        axis = np.arange(2578, dtype=float)*.002 + .01
        wanted = axis[45:657]
        x = np.vstack([np.sin(axis), np.cos(axis)])
        path = self.npz(axis=axis, X_alpha=x)
        result = load_prediction_input(path, wanted)
        np.testing.assert_array_equal(result["X_alpha"], x[:, 45:657])
        self.assertEqual(result["qc"]["axis_handling"], "explicit_612_point_slice_from_2578")
        changed = axis.copy(); changed[100] += .00001
        with self.assertRaisesRegex(PredictionInputError, "exact contiguous"):
            load_prediction_input(self.npz(axis=changed, X_alpha=x), wanted)
        x[0, -1] = np.nan  # even unused cells are not silently repaired/dropped
        with self.assertRaises(PredictionInputError):
            load_prediction_input(self.npz(axis=axis, X_alpha=x), wanted)

    def test_duplicate_ids_including_thickness_aliases_are_rejected(self):
        for ids in [np.array(["new", "new"]), np.array(["NEW-001_1mm", "NEW-001_2mm"]), np.array([" ", "new"]), np.array(["label_1mm", "new"])]:
            with self.subTest(ids=ids):
                with self.assertRaises(PredictionInputError):
                    load_prediction_input(self.npz(record_ids=ids), self.axis)

    def test_flat_spectra_fail_before_snv(self):
        for row in [np.ones(3), np.array([1., 1.+1e-14, 1.])]:
            with self.subTest(row=row):
                x = self.x.copy(); x[0] = row
                with self.assertRaisesRegex(PredictionInputError, "flat"):
                    load_prediction_input(self.npz(X_alpha=x), self.axis)

    def test_teacher_workbook_excludes_reference_and_keeps_thickness_only_in_qc(self):
        p = self.workbook(["Thz", "参考值", "NEW-01_1.5mm", "NEW-02_2mm"],
                          [[a, "reference-is-not-a-feature", x, y] for a, x, y in zip(self.axis, *self.x)])
        result = load_prediction_input(p, self.axis)
        np.testing.assert_array_equal(result["X_alpha"], self.x)
        self.assertEqual(result["record_ids"].tolist(), ["NEW-01", "NEW-02"])
        self.assertEqual(result["qc"]["excluded_reference_headers"], ["参考值"])
        self.assertEqual(result["qc"]["sample_qc"][0]["reported_thickness_mm"], 1.5)

    def test_label_metadata_formula_and_unnamed_columns_are_rejected(self):
        cases = [(["frequency", "sample", "label"], [[a, x, 0] for a, x in zip(self.axis, self.x[0])]),
                 (["frequency", "sample", "thickness_mm"], [[a, x, 1] for a, x in zip(self.axis, self.x[0])]),
                 (["frequency", "sample"], [[.1, "=SUM(1,2)"], [.2, 2], [.3, 4]]),
                 (["frequency", None, "sample"], [[a, 5, x] for a, x in zip(self.axis, self.x[0])]),
                 (["frequency", "sample"], [[a, x, 8] for a, x in zip(self.axis, self.x[0])])]
        for i, (header, rows) in enumerate(cases):
            with self.subTest(case=i):
                with self.assertRaises(PredictionInputError):
                    load_prediction_input(self.workbook(header, rows), self.axis)

    def test_sheets_are_selected_explicitly_never_guessed(self):
        p = self.workbook(["frequency", "sample"], zip(self.axis, self.x[0]), sheet="Sheet")
        with self.assertRaisesRegex(PredictionInputError, "exactly one"):
            load_prediction_input(p, self.axis)
        book = load_workbook(p); book.active.title = "Spectra"; book.create_sheet("吸收系数"); book.save(p); book.close()
        with self.assertRaisesRegex(PredictionInputError, "exactly one"):
            load_prediction_input(p, self.axis)

    def test_csv_layout_supports_batch_and_rejects_incomplete_rows(self):
        p = self.directory/"input.csv"
        with p.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f); writer.writerow(["frequency", "NEW-01", "NEW-02"])
            writer.writerows(zip(self.axis, *self.x))
        np.testing.assert_array_equal(load_prediction_input(p, self.axis)["X_alpha"], self.x)
        with p.open("a") as f:
            f.write("0.4,1\n")
        with self.assertRaises(PredictionInputError):
            load_prediction_input(p, self.axis)

    def test_templates_are_blank_exclusive_and_can_be_filled_with_real_values(self):
        for suffix in [".csv", ".xlsx"]:
            with self.subTest(suffix=suffix):
                path = self.directory/("template"+suffix)
                info = write_prediction_template(path, self.axis, ("NEW-01", "NEW-02"))
                self.assertTrue(info["spectrum_cells_blank"])
                self.assertFalse(info["contains_synthetic_spectra"])
                with self.assertRaises(PredictionInputError):
                    load_prediction_input(path, self.axis)
                before = path.read_bytes()
                with self.assertRaises(FileExistsError):
                    write_prediction_template(path, self.axis)
                self.assertEqual(path.read_bytes(), before)
                if suffix == ".csv":
                    with path.open("w", encoding="utf-8-sig", newline="") as f:
                        writer = csv.writer(f); writer.writerow(["frequency", "NEW-01", "NEW-02"])
                        writer.writerows(zip(self.axis, *self.x))
                else:
                    book = load_workbook(path)
                    for row, values in enumerate(self.x.T, 2):
                        for column, value in enumerate(values, 2):
                            book["Spectra"].cell(row, column, float(value))
                    book.save(path); book.close()
                np.testing.assert_array_equal(load_prediction_input(path, self.axis)["X_alpha"], self.x)

    def test_actual_numbered_teacher_workbook_matches_frozen_model_matrix(self):
        root = Path(__file__).resolve().parents[1]
        with np.load(root/"data/multiplant_database/zaoren/aligned_modalities.npz", allow_pickle=False) as z:
            axis = z["axis_alpha"]
            expected = {rid: row for rid, row in zip(z["record_ids"], z["X_alpha"])}
        result = load_prediction_input(root/"data/processed/zaoren_20261031/样本2_编号版.xlsx", axis)
        self.assertEqual(result["X_alpha"].shape, (5, 612))
        for rid, row in zip(result["record_ids"], result["X_alpha"]):
            np.testing.assert_array_equal(row, expected[rid])
        self.assertEqual(result["qc"]["axis_handling"], "explicit_612_point_slice_from_2578")


if __name__ == "__main__":
    unittest.main()
