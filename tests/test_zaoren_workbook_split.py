import csv
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from herb_spectra.dataset_registry import file_sha256, select_workbooks


SPEC = importlib.util.spec_from_file_location("zaoren_workbook_split", Path(__file__).resolve().parents[1]
                                             / "scripts/split_zaoren_workbooks.py")
split = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(split)


def fixture(repo, counts=(5, 10), conflicting_physical=False):
    sources, rows = [], []
    for number, count in enumerate(counts, 1):
        group = f"ZY{number:02d}"
        path = repo / f"data/样本{number}_编号版.xlsx"; path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook(); workbook.remove(workbook.active)
        ids = [f"{group}-{i + 1:03d}" for i in range(count)]
        for title in ["时域", "频域", "折射率", "吸收系数"]:
            sheet = workbook.create_sheet(title)
            reference = ["参考值"] if title in {"时域", "频域"} else []
            headers = [identifier + ("_1.5mm" if title == "时域" else "") for identifier in ids]
            sheet.append(["时间" if title == "时域" else "Thz", *reference, *headers])
            for point in range(4):
                sheet.append([point * .25, *([3.14159265358979 + point] if reference else []),
                              *[number + index * .03125 + point * .000001 for index in range(count)]])
        workbook.save(path)
        sources.append({"path": str(path.relative_to(repo)), "sha256": file_sha256(path),
                        "status": "active", "dataset": "zaoren", "record_count": count, "origin_group_id": group})
        for identifier in ids:
            rows.append({"record_id": identifier, "physical_sample_id": "same" if conflicting_physical else identifier,
                         "origin_group_id": group, "source_workbook": str(path.relative_to(repo)),
                         "sheet": "吸收系数", "origin_group_verified": "true"})
    (repo / "registry.json").write_text(json.dumps({"version": "fixture", "effective_date": "2026-10-31",
                                                   "workbooks": sources}), encoding="utf-8")
    metadata = repo / "metadata.csv"
    with metadata.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader(); writer.writerows(rows)
    return select_workbooks(repo=repo, manifest=repo / "registry.json"), metadata


class ThreewayWorkbookSplitTests(unittest.TestCase):
    def test_deterministic_partition_complete_disjoint_and_stratified(self):
        with tempfile.TemporaryDirectory() as tmp:
            sources, metadata = fixture(Path(tmp))
            a = split.build_split_plan(sources, metadata); b = split.build_split_plan(sources, metadata)
            self.assertEqual(a, b)
            self.assertEqual(a["counts_by_origin"]["ZY01"], {"training": 3, "calibration": 1, "prediction": 1})
            self.assertEqual(a["role_counts"], {"training": 9, "calibration": 3, "prediction": 3})
            self.assertEqual(len({r["record_id"] for r in a["records"]}), 15)
            self.assertNotEqual(a["records"], split.build_split_plan(sources, metadata, 42)["records"])

    def test_repeated_physical_sample_rejected_even_with_different_record_ids(self):
        with tempfile.TemporaryDirectory() as tmp:
            sources, metadata = fixture(Path(tmp), conflicting_physical=True)
            with self.assertRaisesRegex(ValueError, "Duplicate record or physical"):
                split.build_split_plan(sources, metadata)

    def test_incomplete_or_overlapping_explicit_plan_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            sources, metadata = fixture(Path(tmp)); plan = split.build_split_plan(sources, metadata)
            plan["records"].pop()
            with self.assertRaisesRegex(ValueError, "completely partitioned"):
                split.validate_split_plan(plan, sources)
            plan = split.build_split_plan(sources, metadata)
            plan["records"][1]["physical_sample_id"] = plan["records"][0]["physical_sample_id"]
            with self.assertRaisesRegex(ValueError, "unique and disjoint"):
                split.validate_split_plan(plan, sources)

    def test_four_sheet_payloads_reference_values_suffixes_and_source_preserved(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); sources, metadata = fixture(root, counts=(5,))
            source = sources[0]; original_sha = file_sha256(source.path)
            plan = split.build_split_plan(sources, metadata)
            roles = {role: [r["record_id"] for r in plan["records"] if r["role"] == role] for role in split.ROLES}
            result = split.export_workbook(source, roles, root / "derived")
            original = load_workbook(source.path, data_only=True)
            try:
                for entry in result:
                    derived = load_workbook(root / "derived" / entry["path"], data_only=True)
                    try:
                        self.assertEqual(derived.sheetnames, original.sheetnames)
                        for name in derived.sheetnames:
                            before, after = original[name], derived[name]
                            kept = [i for i, cell in enumerate(before[1], 1)
                                    if not split.ID.fullmatch(str(cell.value))
                                    or split.ID.fullmatch(str(cell.value)).group(1) in roles[entry["role"]]]
                            expected = [tuple(row[i - 1] for i in kept) for row in before.iter_rows(values_only=True)]
                            self.assertEqual(list(after.iter_rows(values_only=True)), expected)
                            self.assertEqual(after.max_column, len(kept))
                            self.assertTrue(entry["verification"]["retained_cell_payloads_exact"])
                    finally:
                        derived.close()
            finally:
                original.close()
            self.assertEqual(file_sha256(source.path), original_sha)

    def test_sheet_id_mismatch_rejected_and_partial_exports_removed(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); sources, metadata = fixture(root, counts=(5,)); source = sources[0]
            plan = split.build_split_plan(sources, metadata)
            roles = {role: [r["record_id"] for r in plan["records"] if r["role"] == role] for role in split.ROLES}
            roles["training"][0] = "ZY01-999"
            with self.assertRaisesRegex(ValueError, "Four-sheet specimen IDs"):
                split.export_workbook(source, roles, root / "derived")
            self.assertEqual(list((root / "derived").rglob("*.partial")), [])
            self.assertEqual(list((root / "derived").rglob("*.xlsx")), [])

    def test_formula_or_column_references_fail_instead_of_corrupting(self):
        data = b'<worksheet><dimension ref="A1:B2"/><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>Thz</t></is></c><c r="B1" t="inlineStr"><is><t>ZY01-001</t></is></c></row><row r="2"><c r="A2"><v>0.1</v></c><c r="B2"><f>A2*2</f><v>0.2</v></c></row></sheetData></worksheet>'
        with self.assertRaisesRegex(ValueError, "Formula"):
            split._filter_sheet(data, [], {"ZY01-001"}, {"ZY01-001"})
        data = data.replace(b"<f>A2*2</f>", b"").replace(b"</worksheet>", b'<mergeCells><mergeCell ref="A1:B1"/></mergeCells></worksheet>')
        with self.assertRaisesRegex(ValueError, "column-addressed metadata"):
            split._filter_sheet(data, [], {"ZY01-001"}, {"ZY01-001"})

    def test_prediction_labels_are_separate_and_existing_exports_not_overwritten(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); sources, metadata = fixture(root, counts=(5,))
            plan = split.build_split_plan(sources, metadata); output = root / "derived"
            split.write_plan(output, plan)
            self.assertTrue((output / "evaluation/prediction_truth.csv").exists())
            self.assertFalse((output / split.ROLE_DIRS["prediction"] / "labels.csv").exists())
            roles = {role: [r["record_id"] for r in plan["records"] if r["role"] == role] for role in split.ROLES}
            split.export_workbook(sources[0], roles, output)
            with self.assertRaises(FileExistsError):
                split.export_workbook(sources[0], roles, output)


if __name__ == "__main__":
    unittest.main()
