"""Export only completed, independently scored auxiliary runs; no training."""
from pathlib import Path
import argparse,csv,gzip,hashlib,json
from collections import defaultdict
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook,load_workbook

LABELS={'full_svm':'Full SNV SVM','mcuve_svm':'MC-UVE + SVM','cnn':'CNN','resnet':'ResNet1D','se':'SE','fusion':'Fusion','inner_selected_neural':'Inner-selected NN'}

def digest(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def csv_write(p,rows):
    columns=list(dict.fromkeys(k for row in rows for k in row))
    with p.open('w',encoding='utf-8-sig',newline='') as h:
        writer=csv.DictWriter(h,fieldnames=columns);writer.writeheader();writer.writerows(rows)
def json_read(p):return json.loads(p.read_text())
def main(out,dest):
    if dest.exists():raise FileExistsError('Report exports may not overwrite a previous run')
    required=['pooled_metrics.json','seed_summary.json','outer_model_records.json','selection_lock.json','prediction_lock.json','scored_outer_predictions.csv.gz','inner_metrics.json']
    inputs={name:digest(out/name) for name in required}
    pooled=json_read(out/'pooled_metrics.json');summary=json_read(out/'seed_summary.json');outer=json_read(out/'outer_model_records.json');inner=json_read(out/'inner_metrics.json');selection=json_read(out/'selection_lock.json')
    with gzip.open(out/'scored_outer_predictions.csv.gz','rt',encoding='utf-8',newline='') as h:predictions=list(csv.DictReader(h))
    dest.mkdir(parents=True)
    all_metrics=[];per_class=[];confusion=[];cost=[];features=[];selections=[]
    for dataset,families in pooled.items():
        for family,seeds in families.items():
            for seed,m in seeds.items():
                all_metrics.append({'dataset_id':dataset,'family':family,'seed':int(seed),'accuracy':m['accuracy'],'macro_f1':m['macro_f1'],'balanced_accuracy':m['balanced_accuracy'],'correct':m['correct'],'N':m['total']})
                for c,v in m['per_class'].items():per_class.append({'dataset_id':dataset,'family':family,'seed':int(seed),'class':c,**v})
                classes=list(m['per_class'])
                for i,c in enumerate(classes):
                    for j,p in enumerate(classes):confusion.append({'dataset_id':dataset,'family':family,'seed':int(seed),'true_class':c,'predicted_class':p,'count':m['confusion_matrix'][i][j]})
        for family in LABELS:
            for phase,records in [('source_selection',inner),('outer_evaluation',outer)]:
                subset=[r for r in records if r['dataset_id']==dataset and r['family']==family]
                if not subset:continue
                cost.append({'dataset_id':dataset,'family':family,'phase':phase,'actual_model_fits':len(subset),'sum_trial_seconds':sum(r['elapsed_seconds'] for r in subset),
                    'sum_transform_fit_seconds':sum(r['metadata']['transform_fit_seconds'] for r in subset),'sum_estimator_fit_seconds':sum(r['metadata']['estimator_fit_seconds'] for r in subset),
                    'minimum_parameters':min((r['metadata']['parameters'] for r in subset if r['metadata']['parameters'] is not None),default=None),
                    'maximum_parameters':max((r['metadata']['parameters'] for r in subset if r['metadata']['parameters'] is not None),default=None),
                    'minimum_support_vectors':min((r['metadata']['support_vectors'] for r in subset if 'support_vectors' in r['metadata']),default=None),
                    'maximum_support_vectors':max((r['metadata']['support_vectors'] for r in subset if 'support_vectors' in r['metadata']),default=None),
                    'minimum_selected_features':min((r['metadata']['selected_features'] for r in subset if r['metadata']['selected_features'] is not None),default=None),
                    'maximum_selected_features':max((r['metadata']['selected_features'] for r in subset if r['metadata']['selected_features'] is not None),default=None)})
    for r in outer:
        if r['family']=='mcuve_svm':
            features.append({'dataset_id':r['dataset_id'],'outer_fold':r['outer_fold'],'seed':r['seed'],'candidate_id':r['candidate_id'],'n_selected':len(r['metadata']['selected_indices']),'alpha_indices_0based':json.dumps(r['metadata']['selected_indices'])})
    for c in selection['outer_folds']:
        for family,v in c['choices'].items():selections.append({'dataset_id':c['dataset_id'],'outer_fold':c['outer_fold'],'family':family,'candidate':json.dumps(v['selected'],sort_keys=True),'inner_macro_f1':v['mean_inner_metrics']['macro_f1']})
        v=c['inner_selected_neural'];selections.append({'dataset_id':c['dataset_id'],'outer_fold':c['outer_fold'],'family':'inner_selected_neural','candidate':json.dumps(v['candidate'],sort_keys=True),'architecture_family':v['family'],'inner_macro_f1':v['mean_inner_metrics']['macro_f1']})
    tables={'metrics_all_seeds':all_metrics,'seed_summary':summary,'per_class':per_class,'confusion_matrices':confusion,'costs':cost,'selected_features':features,'source_selected_candidates':selections}
    for name,rows in tables.items():csv_write(dest/(name+'.csv'),rows)
    rng=np.random.default_rng(20261031);examples=[]
    for dataset in pooled:
        ids=sorted({r['record_id'] for r in predictions if r['dataset_id']==dataset});chosen=set(rng.choice(ids,size=min(10,len(ids)),replace=False))
        examples.extend(r for r in predictions if r['dataset_id']==dataset and r['record_id'] in chosen and r['seed']=='42' and r['family'] in {'mcuve_svm','inner_selected_neural'})
    csv_write(dest/'preseeded_random_examples.csv',examples)
    validated=0
    for dataset in pooled:
        workbook=Workbook();workbook.remove(workbook.active)
        for name,rows in {**tables,'predictions':predictions}.items():
            subset=[r for r in rows if r['dataset_id']==dataset]
            if not subset:continue
            ws=workbook.create_sheet(name[:31]);columns=list(dict.fromkeys(k for r in subset for k in r));ws.append(columns)
            for row in subset:ws.append([json.dumps(row.get(k),ensure_ascii=False) if isinstance(row.get(k),(dict,list)) else row.get(k) for k in columns])
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        path=dest/(dataset+'_classification_evidence.xlsx');workbook.save(path);check=load_workbook(path,read_only=True,data_only=True)
        assert set(check.sheetnames)==set(workbook.sheetnames)
        for original in workbook:
            restored=check[original.title];assert restored.max_row==original.max_row and restored.max_column==original.max_column
            for a,b in zip(original.values,restored.values,strict=True):
                for x,y in zip(a,b,strict=True):
                    if isinstance(x,float):assert np.isclose(x,y,rtol=1e-14,atol=1e-14)
                    else:assert (None if x=='' else x)==y
                    validated+=1
        check.close()
        order=list(LABELS);x=np.arange(len(order));selected=[next(r for r in summary if r['dataset_id']==dataset and r['family']==f) for f in order]
        fig,ax=plt.subplots(figsize=(10.5,5));width=.38
        for offset,metric,label,color in [(-width/2,'accuracy','Accuracy','#2869A1'),(width/2,'macro_f1','Macro-F1','#D5842E')]:
            ax.bar(x+offset,[v[metric+'_mean'] for v in selected],width,yerr=[v[metric+'_sample_sd'] for v in selected],capsize=3,label=label,color=color)
        ax.set_xticks(x,[LABELS[f] for f in order],rotation=25,ha='right');ax.set_ylim(0,1.08);ax.set_ylabel('Score');ax.set_title(dataset.capitalize()+' four source groups: mean ± SD across fixed training seeds');ax.legend();ax.grid(axis='y',alpha=.2);fig.tight_layout()
        for ext in ['png','pdf']:fig.savefig(dest/(dataset+'_metrics.'+ext),dpi=180)
        plt.close(fig)
        fig,axes=plt.subplots(2,4,figsize=(12,6.7));axes=axes.ravel()
        for axis,family in zip(axes,order):
            m=pooled[dataset][family]['42'];matrix=np.asarray(m['confusion_matrix']);classes=list(m['per_class']);axis.imshow(matrix,cmap='Blues',vmin=0,vmax=matrix.max())
            for i in range(4):
                for j in range(4):axis.text(j,i,str(matrix[i,j]),ha='center',va='center',color='white' if matrix[i,j]>matrix.max()/2 else 'black',fontsize=8)
            axis.set_xticks(range(4),classes,rotation=40);axis.set_yticks(range(4),classes);axis.set_title(LABELS[family],fontsize=10);axis.set_xlabel('Predicted');axis.set_ylabel('True')
        axes[-1].axis('off');fig.suptitle(dataset.capitalize()+' pooled five-fold OOF confusion matrices, fixed seed42');fig.tight_layout()
        for ext in ['png','pdf']:fig.savefig(dest/(dataset+'_confusion_matrices.'+ext),dpi=180)
        plt.close(fig)
    assert inputs=={name:digest(out/name) for name in required}
    (dest/'export_validation.json').write_text(json.dumps({'status':'TECHNICAL_CHECK_PASS','source_files_sha256':inputs,'script_sha256':digest(__file__),'excel_cells_read_back':validated,'metrics_rows':len(all_metrics),'per_class_rows':len(per_class),'confusion_rows':len(confusion),'scored_prediction_rows':len(predictions),'random_example_seed':20261031,'labels':'anonymous separate four-source-group labels; named growth geography unverified','seed_error_bars':'sample SD over three fixed training seeds; not independent samples or population confidence intervals','independent_confirmation':False},indent=2)+'\n')
    print(json.dumps({'destination':str(dest),'excel_cells_read_back':validated,'metrics_rows':len(all_metrics)}))
if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path,required=True);parser.add_argument('--destination',type=Path,required=True);a=parser.parse_args();main(a.output,a.destination)
