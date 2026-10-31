"""Read-only multimodal spectral intake with explicit, non-inferred lineage.

Record IDs are representation-pair keys, not evidence of independent plants.
No filename, ID, thickness or label is placed in the numeric feature matrices.
"""
from dataclasses import dataclass
import hashlib
from pathlib import Path
import re
import warnings
import xml.etree.ElementTree as ET
import zipfile

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


MODALITIES = {"alpha": "吸收系数", "n": "折射率", "time": "时域", "frequency": "频域"}
ZAOREN_ID = re.compile(r"^(ZY\d{2}-\d{3,})(?:_(.*))?$")
SAMPLE_ID = re.compile(r"^(?:样品|样本)\s*(\d+)(?:_(.*))?$")


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def base_header(header, dataset_id):
    """Parse an explicit header identity; never use the column position."""
    value = str(header).strip()
    match = (ZAOREN_ID if dataset_id == "zaoren" else SAMPLE_ID).fullmatch(value)
    if match is None:
        raise ValueError(f"Unrecognized explicit sample header: {value!r}")
    return match.group(1)


def thickness_from_header(header):
    match = re.search(r"(?:^|_)(\d+(?:\.\d+)?)\s*mm(?:$|_)", str(header), re.I)
    return float(match.group(1)) if match else None


@dataclass(frozen=True)
class RawSheet:
    name: str
    axis_header: str
    axis: np.ndarray
    X: np.ndarray
    headers: tuple
    columns: tuple
    references: tuple
    reference_values: np.ndarray
    declared_dimension: str
    nonnumeric_cells: int


def read_raw_sheet(path, name):
    """Read all actual cells; preserve nonfinite cells for quality reporting."""
    warnings.filterwarnings("ignore", message="Workbook contains no default style")
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book[name]
        declared = sheet.calculate_dimension(force=False)
        sheet.reset_dimensions()
        iterator = iter(sheet.iter_rows(values_only=True))
        header = next(iterator, ())
        if not header:
            raise ValueError("Empty worksheet")
        columns = [i for i, h in enumerate(header) if i and h is not None and str(h).strip()
                   and "参考" not in str(h) and str(h).strip().lower() not in {"reference", "ref"}]
        refs = [i for i, h in enumerate(header) if i and h is not None and
                ("参考" in str(h) or str(h).strip().lower() in {"reference", "ref"})]
        if not columns:
            raise ValueError("No named spectral columns")
        values, nonnumeric = [], 0
        for row in iterator:
            if all(v is None for v in row):
                continue
            numeric = []
            for i in [0, *columns, *refs]:
                v = row[i] if i < len(row) else None
                try:
                    numeric.append(float(v) if v is not None else np.nan)
                except (TypeError, ValueError):
                    numeric.append(np.nan)
                    nonnumeric += 1
            values.append(numeric)
        if not values:
            raise ValueError("No numeric data rows")
        data = np.asarray(values, dtype=float)
        return RawSheet(name, str(header[0]), data[:, 0], data[:, 1:1+len(columns)].T,
                        tuple(str(header[i]) for i in columns),
                        tuple(get_column_letter(i+1) for i in columns),
                        tuple(str(header[i]) for i in refs), data[:, 1+len(columns):].T,
                        declared, nonnumeric)
    finally:
        book.close()


def sheet_members(path):
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    relns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    pkg = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    with zipfile.ZipFile(path) as archive:
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        mapping = {r.attrib["Id"]: r.attrib["Target"] for r in rels.findall(pkg+"Relationship")}
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        result = {}
        for sheet in workbook.findall(ns+"sheets/"+ns+"sheet"):
            target = mapping[sheet.attrib[relns+"id"]]
            result[sheet.attrib["name"]] = target.lstrip("/") if target.startswith("/") else "xl/"+target
        return result


def workbook_health(path):
    members, errors = sheet_members(path), {}
    with zipfile.ZipFile(path) as archive:
        for member in archive.infolist():
            try:
                archive.read(member)
            except Exception as exc:
                errors[member.filename] = {"error_type":type(exc).__name__, "error":str(exc)}
    return {"container_crc_pass":not errors,"bad_members":errors,"sheet_members":members}


def align_sheets(sheets, dataset_id, source_id):
    """Pair by recognized header identity and reject ambiguity or dropped IDs."""
    identities = {mode:[base_header(v, dataset_id) for v in sheet.headers] for mode,sheet in sheets.items()}
    for mode, identifiers in identities.items():
        if len(set(identifiers)) != len(identifiers):
            raise ValueError(f"Duplicate base IDs in {source_id}/{mode}; preparation ancestry unresolved")
    reference = identities.get("alpha", next(iter(identities.values())))
    if any(set(ids) != set(reference) for ids in identities.values()):
        raise ValueError("Modality sample-ID sets differ; no complete-case deletion allowed")
    ordered = {}
    for mode,sheet in sheets.items():
        lookup = {identifier:i for i,identifier in enumerate(identities[mode])}
        ordered[mode] = sheet.X[[lookup[i] for i in reference]]
    record_ids = reference if dataset_id == "zaoren" else [f"{source_id}-S{int(i):04d}" for i in reference]
    return list(record_ids), ordered, identities


def numeric_profile(sheet):
    axis = sheet.axis
    return {"record_count":len(sheet.X),"axis_points":len(axis),"axis_header":sheet.axis_header,
            "axis_unit":"header_only_THZ" if "thz" in sheet.axis_header.lower() else "not_explicitly_declared",
            "axis_min":float(np.nanmin(axis)),"axis_max":float(np.nanmax(axis)),
            "axis_finite":bool(np.isfinite(axis).all()),"axis_strictly_increasing":bool(np.all(np.diff(axis)>0)),
            "nonfinite_cells":int((~np.isfinite(sheet.X)).sum()),"negative_cells":int((sheet.X<0).sum()),
            "nonnumeric_cells":sheet.nonnumeric_cells,"declared_dimension":sheet.declared_dimension,
            "headers":list(sheet.headers),"excel_columns":list(sheet.columns),
            "reference_headers":list(sheet.references),"references_are_instrument_reference_not_study_average":"unknown",
            "physical_unit_verified":False,"axis_sha256":hashlib.sha256(np.asarray(axis,dtype="<f8").tobytes()).hexdigest()}


def load_modality(path):
    with np.load(path, allow_pickle=False) as archive:
        return {key:archive[key] for key in archive.files}


def assert_model_input_arrays(payload):
    X = payload["X"]
    if X.ndim != 2 or not np.issubdtype(X.dtype, np.number) or not np.isfinite(X).all():
        raise ValueError("Numeric finite spectra only; IDs and metadata are not model features")
    if len(X) != len(payload["record_ids"]) or X.shape[1] != len(payload["axis"]):
        raise ValueError("Axis or record identity does not match numeric matrix")
    return True


def numeric_row_hash(value):
    return hashlib.sha256(np.ascontiguousarray(value,dtype="<f8").tobytes()).hexdigest()


def validate_zaoren_row_bindings(payload, binding_csv):
    """Bind original-sheet correspondence to values, not merely shared IDs."""
    import csv
    with Path(binding_csv).open(encoding="utf-8-sig",newline="") as handle:
        rows=list(csv.DictReader(handle))
    ids=np.asarray(payload["record_ids"]).astype(str)
    if len(ids)!=len(set(ids.tolist())):
        raise ValueError("Repeated record IDs in multimodal payload")
    lookup={v:i for i,v in enumerate(ids)}
    for mode in MODALITIES:
        if "X_"+mode not in payload or "axis_"+mode not in payload:
            raise ValueError("Required modality is missing from the aligned payload")
        expected=[r for r in rows if r["modality"]==mode]
        if set(r["record_id"] for r in expected)!=set(lookup):
            raise ValueError("Binding IDs differ from multimodal input")
        values=np.asarray(payload["X_"+mode]);axis=np.asarray(payload["axis_"+mode])
        if len(values)!=len(ids):raise ValueError("Modality row count differs")
        for row in expected:
            if numeric_row_hash(values[lookup[row["record_id"]]])!=row["model_matrix_row_sha256"]:
                raise ValueError("Modality values do not match the original-sheet ID binding")
            if numeric_row_hash(axis)!=row["model_axis_sha256"]:
                raise ValueError("Modality axis changed")
    return {"status":"EXPLORATORY_VALUE_BINDINGS_VERIFIED","record_count":len(ids),"modality_rows":len(rows),
            "physical_independence_proved":False}
