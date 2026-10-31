"""Post-score report-only checks/addendum; no model training or benchmark score."""
from pathlib import Path
import csv
import hashlib
import json
import math
import sys
from datetime import datetime, timezone

from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "scripts"))
import export_cross_plant_transfer_20261031 as report

BASE = ROOT / "reports/cross-plant-deep-transfer-20261031-v1"
OUT = BASE / "comparison"
RUN = BASE / "neural"
sha = report.sha

release = report.read(BASE / "root_export_release.json")
assert sha(BASE / "root_export_release.json") == "f2668573c042a02d49145e0b8a25055f0544d50105b663718b62e52f51d2eed1"
assert sha(RUN / "scoring_manifest.json") == release["scoring_manifest_sha256"]
assert sha(ROOT / "scripts/export_cross_plant_transfer_20261031.py") == "a096cb63027f3edd420427e364ba20e8ea2c2a97c7ec2cc1cc799bf035e34609"
initial_readme_sha = sha(OUT / "README.md")
initial_index_sha = sha(OUT / "files_sha256.json")
initial_manifest_sha = sha(OUT / "export_manifest.json")
initial_readme = (OUT / "README.md").read_text(encoding="utf-8")
(BASE / "export-execution/initial_generated_README.md").write_text(initial_readme, encoding="utf-8")

# This supplementary whole-recipe contrast changes the ginseng dose120→40.
# It must not be interpreted as an isolated HQ/TM addition ablation.
groups = report.prediction_groups(RUN)
metrics, _, _ = report.classification_rows(RUN, groups)
metric_lookup = {(r["role"], r["seed"]): r for r in metrics}
contrasts = []
for seed in report.SEEDS:
    old = {r["record_id"]: r for r in groups["B", seed]}
    new = groups["C", seed]
    corrected = sum(old[r["record_id"]]["y_predicted"] != r["y_true"] and r["y_predicted"] == r["y_true"] for r in new)
    broken = sum(old[r["record_id"]]["y_predicted"] == r["y_true"] and r["y_predicted"] != r["y_true"] for r in new)
    a, b = metric_lookup["B", seed], metric_lookup["C", seed]
    contrasts.append({"baseline": "B", "role": "C", "seed": seed, "corrected": corrected, "broken": broken,
        "net_correct": corrected-broken, **{m + "_delta": b[m]-a[m] for m in report.METRICS},
        "accuracy_change_percentage_points": 100*(b["accuracy"]-a["accuracy"]),
        "accuracy_relative_change_percent": 100*(b["accuracy"]-a["accuracy"])/a["accuracy"], "denominator": 855,
        "interpretation": "Whole-recipe C-vs-B contrast; ginseng120→40steps plus HQ/TM40each; not isolated pool addition"})
    assert abs((corrected-broken)/855 - (b["accuracy"]-a["accuracy"])) < 1e-12
contrast_summary = report.summarize(contrasts, ("accuracy_delta", "macro_f1_delta", "balanced_accuracy_delta", "corrected", "broken", "net_correct", "accuracy_change_percentage_points", "accuracy_relative_change_percent"), ("baseline", "role"))
extra = {"external_C_vs_B_per_seed": contrasts, "external_C_vs_B_summary": contrast_summary}
for name, rows in extra.items():
    report.write_csv(OUT / (name + ".csv"), rows)
report.spreadsheet(OUT / "supplement_C_vs_B.xlsx", extra)

coverage_path = BASE / "pretraining-coverage/summary.json"
coverage = report.read(coverage_path)
coverage_rows = coverage["summary"]
eligible = {("B", "ginseng"): "42", ("C", "huangqi"): "403", ("C", "tianma"): "405",
            ("C", "ginseng"): "42", ("Aextra", "target"): "342或684（当前target source）"}
appendix = ["", "## 实际数据覆盖与额外对照", "",
    "以上850条是可用外部源池大小，不等于每条在每次固定120步中都进入梯度更新。StandardScaler使用当前允许源池全部观测；真实优化器覆盖来自原始batch trace的独立核算，见[覆盖证据](../pretraining-coverage/summary.json)。重复呈现不增加N。", "",
    "| 组／源任务 | 可用源池观测 | 每管线实际梯度观测范围 | 管线数 | 总呈现次数（非独立N） |",
    "|---|---|---:|---:|---:|"]
for r in coverage_rows:
    appendix.append(f"| {r['arm']}／{r['pool']} | {eligible[r['arm'],r['pool']]} | {r['unique_optimizer_observations_per_trajectory_min']}—{r['unique_optimizer_observations_per_trajectory_max']} | {r['training_trajectories']} | {r['presentations_total_not_independent_samples']} |")
summary = report.read(RUN / "seed_summary.json")
policy = next(r for r in summary if r["role"] == "source_policy")
a = next(r for r in summary if r["role"] == "A")
folds = report.read(RUN / "selection_lock.json")["outer_folds"]
policy_arms = [r["policy_arm"] for r in sorted(folds, key=lambda r:r["outer_fold"])]
assert policy_arms == ["Aextra", "Aextra", "A", "A", "A"]
delta = contrast_summary[0]
appendix += ["",
    f"保守source-policy在五外折依次选择{'、'.join(policy_arms)}，没有采用B或C。其Accuracy均值{100*policy['accuracy_mean']:.2f}%相对A的{100*a['accuracy_mean']:.2f}%增加{100*(policy['accuracy_mean']-a['accuracy_mean']):.2f}个百分点，来自当前酸枣仁source额外训练及回退规则，不能记为外部数据收益。",
    "",
    f"补充C−B的Accuracy平均变化{delta['accuracy_change_percentage_points_mean']:+.4f}个百分点，Macro-F1变化{delta['macro_f1_delta_mean']:+.6f}，三种子净正确数分别{[r['net_correct'] for r in contrasts]}。CSV及supplement_C_vs_B.xlsx保存全部配对纠正／损害。C相对B同时将人参预训练由120步改为40步并加入黄芪／天麻任务，因此这不是隔离黄芪／天麻单独作用的消融。",
    "", "这是预训练／微调对照，没有进行新的少样本校准或真实跨批次校准。"]
(OUT / "README.md").write_text(initial_readme + "\n".join(appendix) + "\n", encoding="utf-8")

def equal_cell(text, cell):
    if cell is None:
        return text == ""
    if isinstance(cell, bool):
        return text == str(cell)
    if isinstance(cell, (int, float)):
        try:
            return math.isfinite(float(text)) and math.isclose(float(text), cell, rel_tol=1e-12, abs_tol=1e-12)
        except ValueError:
            return False
    return text == str(cell)

books = {"main": load_workbook(OUT / "cross_plant_transfer_comparison.xlsx", read_only=True, data_only=True),
         "extra": load_workbook(OUT / "supplement_C_vs_B.xlsx", read_only=True, data_only=True)}
checks, compared_cells = [], 0
for path in sorted(OUT.glob("*.csv")):
    key = "extra" if path.stem.startswith("external_C_vs_B_") else "main"
    sheet = books[key][path.stem[:31]]
    count = 0
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for n, (texts, cells) in enumerate(zip(csv.reader(handle), sheet.iter_rows(values_only=True)), 1):
            assert len(texts) == len(cells), path.name
            for text, value in zip(texts, cells):
                assert equal_cell(text, value), (path.name, n, text, value)
                compared_cells += 1
            count += 1
    assert count == sheet.max_row, path.name
    checks.append({"csv":path.name,"sheet":sheet.title,"rows_including_header":count,"all_cells_match":True})
for book in books.values():
    book.close()
image_checks = []
for path in sorted(OUT.glob("*.png")):
    with Image.open(path) as image:
        size = image.size
        image.verify()
    pdf = path.with_suffix(".pdf")
    assert pdf.read_bytes().startswith(b"%PDF") and pdf.stat().st_size > 1000
    image_checks.append({"png":path.name,"size":size,"png_decodes":True,"pdf":pdf.name,"pdf_header":True})
assert len(image_checks) == 5
assert sha(RUN / "scoring_manifest.json") == release["scoring_manifest_sha256"]
assert sha(ROOT / "scripts/export_cross_plant_transfer_20261031.py") == "a096cb63027f3edd420427e364ba20e8ea2c2a97c7ec2cc1cc799bf035e34609"
validation = {"completed_at_utc":datetime.now(timezone.utc).isoformat(),"status":"POST_SCORE_EXPORT_VALIDATION_PASS",
    "root_release_sha256":sha(BASE / "root_export_release.json"),"source_scoring_sha256_unchanged":True,
    "source_code_frozen_sha256_unchanged":True,"new_training_fits":0,"benchmark_score_calls":0,
    "main_sheets":15,"supplemental_sheets":2,"csv_xlsx_checks":checks,"cells_compared":compared_cells,
    "five_png_pdf_pairs":image_checks,"root_coverage_summary_sha256":sha(coverage_path),
    "report_only_addendum":{"reason":"Actual eligible-pool vs gradient coverage and C-vs-B whole-recipe interpretation requested by root",
        "initial_generated_README_sha256":initial_readme_sha,"initial_generated_index_sha256":initial_index_sha,
        "original_export_manifest_sha256_unchanged":initial_manifest_sha,"final_README_sha256":sha(OUT / "README.md"),
        "metric_tables_or_scientific_results_changed":False}}
report.write_json(OUT / "export_validation.json", validation)
# Reseal this new report after the explicitly documented, report-only addendum.
(OUT / "files_sha256.json").write_text(json.dumps({p.name:{"sha256":sha(p),"bytes":p.stat().st_size}
    for p in sorted(OUT.iterdir()) if p.is_file() and p.name != "files_sha256.json"}, indent=2)+"\n")
print(json.dumps({"csv_tables":len(checks),"cells_compared":compared_cells,"images":len(image_checks),
    "model_source_or_score_changes":False,"C_vs_B_accuracy_mean_percentage_points":delta['accuracy_change_percentage_points_mean']}))
