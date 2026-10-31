"""Final full-length, all-cell and file-index check; report outputs only."""
from pathlib import Path
from itertools import zip_longest
import csv
import json
import math
import hashlib
from datetime import datetime, timezone

from openpyxl import load_workbook
from PIL import Image

BASE = Path(__file__).resolve().parents[1]
OUT = BASE / "comparison"

def sha(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()

def equals(text, value):
    if value is None:
        return text == ""
    if isinstance(value, bool):
        return text == str(value)
    if isinstance(value, (int, float)):
        try:
            return math.isfinite(float(text)) and math.isclose(float(text), value, rel_tol=1e-12, abs_tol=1e-12)
        except ValueError:
            return False
    return text == str(value)

release = json.loads((BASE / "root_export_release.json").read_text())
assert sha(BASE / "neural/scoring_manifest.json") == release["scoring_manifest_sha256"]
books = {"main": load_workbook(OUT / "cross_plant_transfer_comparison.xlsx", read_only=True, data_only=True),
         "extra": load_workbook(OUT / "supplement_C_vs_B.xlsx", read_only=True, data_only=True)}
checks, cells = [], 0
for path in sorted(OUT.glob("*.csv")):
    book = books["extra" if path.stem.startswith("external_C_vs_B_") else "main"]
    sheet = book[path.stem[:31]]
    rows = 0
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for n, (texts, values) in enumerate(zip_longest(csv.reader(handle), sheet.iter_rows(values_only=True)), 1):
            assert texts is not None and values is not None, (path.name, "row count", n)
            assert len(texts) == len(values), (path.name, "column count", n)
            for text, value in zip(texts, values):
                assert equals(text, value), (path.name, n, text, value)
                cells += 1
            rows += 1
    assert rows == sheet.max_row
    checks.append({"csv": path.name, "rows_including_header": rows, "every_cell_matches": True})
for book in books.values():
    book.close()
assert len(checks) == 17
plots = []
for path in sorted(OUT.glob("*.png")):
    with Image.open(path) as image:
        size = image.size
        image.verify()
    pdf = path.with_suffix(".pdf")
    assert pdf.read_bytes().startswith(b"%PDF") and pdf.stat().st_size > 1000
    plots.append({"png": path.name, "size": size, "png_decodes": True, "pdf_header_verified": True})
assert len(plots) == 5
assert sha(BASE / "neural/scoring_manifest.json") == release["scoring_manifest_sha256"]
result = {"completed_at_utc": datetime.now(timezone.utc).isoformat(), "status": "FINAL_ALL_LENGTH_CELL_AND_FIGURE_PASS",
    "all_csv_xlsx_cells_match": True, "csv_tables": len(checks), "xlsx_sheets": {"main": 15, "supplement": 2},
    "cells_compared": cells, "checks": checks, "plots": plots, "root_release_sha256": sha(BASE / "root_export_release.json"),
    "source_scoring_manifest_sha256_unchanged": release["scoring_manifest_sha256"],
    "actual_source_gradient_coverage": "../pretraining-coverage/summary.json", "new_research_fits": 0,
    "benchmark_score_calls": 0, "source_code_changed": False,
    "visual_inspection": ["All5roles/3metrics and SDcaption readable", "Paired B/C contrasts retain negative means and zero lines",
        "All15 seven-class confusion panels labelled N855, including rareZY02support5", "Actual48/96source losses and emptyouter48 accurately shown"]}
with (OUT / "final_validation.json").open("x", encoding="utf-8") as handle:
    json.dump(result, handle, ensure_ascii=False, indent=2)
    handle.write("\n")
# This is the final reseal after the explicitly recorded report-only addendum.
(OUT / "files_sha256.json").write_text(json.dumps({p.name: {"sha256": sha(p), "bytes": p.stat().st_size}
    for p in sorted(OUT.iterdir()) if p.is_file() and p.name != "files_sha256.json"}, indent=2) + "\n")
print(json.dumps({"tables": len(checks), "cells_compared": cells, "png_pdf_pairs": len(plots), "final_index_sha256": sha(OUT / "files_sha256.json")}))
