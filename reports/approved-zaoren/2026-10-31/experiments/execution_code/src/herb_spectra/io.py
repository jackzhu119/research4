"""Read scan columns without trusting the workbook's dimension declaration."""
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


@dataclass(frozen=True)
class Spectra:
    axis: np.ndarray
    values: np.ndarray  # scans x frequency/time points; scans are not physical samples
    scan_names: tuple[str, ...]
    column_numbers: tuple[int, ...]  # one-based Excel columns
    source: str
    sheet: str


def read_spectra(path, sheet="吸收系数", bounds=None):
    """Exclude coordinate/reference columns; reject missing or invalid numeric values.

    No provenance, class label, concentration, or physical-sample ID is inferred
    from file names or scan numbering. bounds must be given in the axis's units.
    """
    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        ws.reset_dimensions()
        rows = iter(ws.iter_rows(values_only=True))
        header = next(rows, None)
        if not header:
            raise ValueError(f"Empty worksheet: {path.name}/{sheet}")
        columns = [i for i, h in enumerate(header)
                   if i > 0 and h is not None and str(h).strip()
                   and "参考" not in str(h) and str(h).strip().lower() not in {"reference", "ref"}]
        if not columns:
            raise ValueError("No named spectrum columns")
        records = []
        for row in rows:
            if all(v is None for v in row):
                continue
            try:
                records.append([float(row[i]) for i in [0, *columns]])
            except (IndexError, TypeError, ValueError) as exc:
                raise ValueError("Incomplete or non-numeric spectrum row") from exc
        if not records:
            raise ValueError("No numeric spectrum rows")
        data = np.asarray(records, dtype=float)
        axis = data[:, 0]
        if not np.isfinite(axis).all() or not (np.diff(axis) > 0).all():
            raise ValueError("Coordinate axis must be finite and strictly increasing")
        mask = np.ones(len(axis), dtype=bool)
        if bounds is not None:
            low, high = bounds
            if not np.isfinite([low, high]).all() or low >= high:
                raise ValueError("bounds must be finite and increasing")
            mask = (axis >= low) & (axis <= high)
        values = data[mask, 1:].T
        if values.shape[1] < 2 or not np.isfinite(values).all():
            raise ValueError("Selected spectra require at least two finite points")
        return Spectra(axis[mask], values, tuple(str(header[i]) for i in columns),
                       tuple(i + 1 for i in columns), str(path), sheet)
    finally:
        wb.close()
