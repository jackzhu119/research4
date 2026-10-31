"""Audit external THz feasibility without fitting or merging plant labels."""
import argparse
from collections import Counter
import csv
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from herb_spectra.dataset_registry import select_workbooks
from herb_spectra.io import read_spectra
from herb_spectra.paper_c_external import load_ginseng_alpha_pellets, load_huangqi_alpha_records


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def task_audit(task):
    task.validate()
    x = task.spectra
    main = (task.axis >= .1) & (task.axis <= 1.5)
    return {
        "task_id": task.task_id, "spectrum_records": len(x),
        "verified_independent_physical_samples": None,
        "frequency_points": len(task.axis), "axis_range": [float(task.axis[0]), float(task.axis[-1])],
        "axis_unit": task.axis_unit, "physical_quantity": task.quantity, "value_unit": task.value_unit,
        "negative_cells_full": int((x < 0).sum()), "records_with_negative_cells": int((x < 0).any(axis=1).sum()),
        "negative_cells_in_zaoren_candidate_range": int((x[:, main] < 0).sum()),
        "exact_unique_spectrum_rows": int(len(np.unique(x, axis=0))),
        "scientific_labels": dict(Counter(task.labels)) if task.labels is not None else None,
        "independent_evaluation_blocked": bool(task.scientific_blocks()),
        "block_reasons": task.scientific_blocks(), "provenance": task.provenance,
        "duplicates_do_not_establish_independence": True,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--output", type=Path, default=Path("reports/paper-c/2026-10-31/external-data"))
    args = parser.parse_args()
    repo = args.repo.resolve()
    output = args.output if args.output.is_absolute() else repo / args.output
    output.mkdir(parents=True, exist_ok=True)
    if (output / "audit.json").exists():
        raise FileExistsError("Use a new output directory; do not overwrite a completed audit")
    inputs = [repo / "data/active_workbooks.json", repo / "data/metadata_template.csv"]
    manifest = json.loads(inputs[0].read_text())
    inputs += [repo / row["path"] for row in manifest["workbooks"]]
    ginseng_root = repo / "data/external/ginseng_hvbp5mdtxx_v2"
    inspection_path = ginseng_root / "annotated/inspection.json"
    inspection = json.loads(inspection_path.read_text())
    inputs += [inspection_path] + [ginseng_root / "original" / row["file"] for row in inspection["member_manifest"]]
    before = {str(path.relative_to(repo)): sha(path) for path in inputs}
    for row in inspection["member_manifest"]:
        assert sha(ginseng_root / "original" / row["file"]) == row["sha256"]
    hq, gs = load_huangqi_alpha_records(repo), load_ginseng_alpha_pellets(repo)
    selected = select_workbooks(repo=repo)
    zaoren = read_spectra(next(s.path for s in selected if s.dataset == "zaoren"), bounds=(.1, 1.5))
    assert hq.spectra.shape == (403, 612)
    assert gs.spectra.shape == (42, 65)
    assert np.array_equal(hq.axis, zaoren.axis)
    with (repo / "data/metadata_template.csv").open(encoding="utf-8-sig", newline="") as handle:
        metadata = [r for r in csv.DictReader(handle) if "/huangqi/" in r["source_workbook"]]
    metadata_fields = {field: sum(bool(r.get(field, "")) for r in metadata) for field in (
        "physical_sample_id", "origin", "origin_verified", "batch_id", "instrument_id", "species_verified")}
    sheet_inventory = []
    for source in selected:
        if source.dataset != "huangqi":
            continue
        workbook = load_workbook(source.path, read_only=True, data_only=True)
        sheets = {}
        for name in workbook.sheetnames:
            sheet = workbook[name]
            sheet.reset_dimensions()
            rows = iter(sheet.iter_rows(values_only=True))
            header = next(rows, None)
            first = next(rows, None)
            sheets[name] = {"nonempty_header": header is not None, "first_numeric_row_present": first is not None,
                            "coordinate_header": str(header[0]) if header else None}
        workbook.close()
        full_alpha = read_spectra(source.path)
        sheet_inventory.append({"source": source.relative_path, "sheets": sheets,
                                "full_alpha_shape": list(full_alpha.values.shape),
                                "full_alpha_axis_range": [float(full_alpha.axis[0]), float(full_alpha.axis[-1])],
                                "full_alpha_negative_cells": int((full_alpha.values < 0).sum())})
    interface_dir = output / "task-interfaces"
    interface_dir.mkdir()
    for task in (hq, gs):
        np.savez_compressed(interface_dir / f"{task.task_id}.npz", X_alpha=task.spectra, axis=task.axis,
                            record_ids=np.asarray(task.record_ids), physical_ids=np.asarray(task.physical_ids),
                            labels=np.asarray(task.labels or (), dtype=str))
    source_checks = [
        {"source_id": "ginseng_hvbp5mdtxx_v2", "doi": "10.17632/hvbp5mdtxx.2",
         "url": "https://data.mendeley.com/datasets/hvbp5mdtxx/2", "license": "CC BY 4.0",
         "current_web_metadata_checked": True, "local_numeric_files": True,
         "usable_now": "Technical reading/species-label mapping; independent external evaluation blocked",
         "independent_sample_count": None, "declared_tablets": 42,
         "block": "Parent plant, physical repeat and batch mapping unavailable; 141 negative alpha cells; no original waveforms"},
        {"source_id": "rice_wj2h5hn7kc_v3", "doi": "10.17632/wj2h5hn7kc.3",
         "url": "https://data.mendeley.com/datasets/compare/wj2h5hn7kc", "license": "CC BY 4.0",
         "current_web_metadata_checked": True, "local_numeric_files": False,
         "usable_now": "Source lead only: rice seed plumpness, not origin",
         "independent_sample_count": None, "block": "Cloud proxy denied; internal files, labels, units and seed identities unparsed"},
        {"source_id": "white_birch_seed_17258169", "doi": "10.5281/zenodo.17258169",
         "url": "https://zenodo.org/records/17258169", "license": None,
         "current_web_metadata_checked": True, "local_numeric_files": False,
         "usable_now": "Raw signal-processing lead only",
         "independent_sample_count": None, "declared_archive_bytes_approx": 1100000000,
         "published_md5": "82717c2ae282eaca71b006fe6356b9ef",
         "block": "Cloud proxy denied; 1.1GB archive not fetched; sample IDs/class labels/license not established"},
        {"source_id": "dryad_wheat_benzoic_acid", "doi": "10.5061/dryad.945c410",
         "url": "https://datadryad.org/dataset/doi:10.5061/dryad.945c410",
         "license": "CC0 according to prior Dryad policy registration; item-specific license not reverified in current parsed page",
         "current_web_metadata_checked": True, "local_numeric_files": False,
         "usable_now": "Quantitative additive-concentration task lead, not origin classification",
         "independent_sample_count": None,
         "listed_files": ["data.mat", "X.mat", "Y.mat", "fz.mat", "BP.m", "DemoBPNN.m", "DemoGRNN.m", "rms.m"],
         "direct_file_attempts": [
             {"url": "https://datadryad.org/downloads/file_stream/108442", "object": "data.mat", "web_tool_result": "403 Forbidden"},
             {"url": "https://datadryad.org/downloads/file_stream/108445", "object": "fz.mat", "web_tool_result": "403 Forbidden"}],
         "block": "Cloud proxy denied; web-tool file clicks403; numeric dimensions, specimen independence and units unparsed"},
    ]
    audit = {
        "protocol_id": "paper-c-external-feasibility-v1", "recorded_at_utc": datetime.now(timezone.utc).isoformat(),
        "main_task_unchanged": "855 anonymous-origin acid jujube specimens; external work is auxiliary",
        "model_fits": 0, "new_downloaded_numeric_datasets": 0,
        "huangqi": task_audit(hq), "ginseng": task_audit(gs),
        "huangqi_metadata_records": len(metadata), "huangqi_nonempty_scientific_metadata_fields": metadata_fields,
        "huangqi_sheet_inventory": sheet_inventory,
        "numeric_axis_compatibility": {"huangqi_equals_zaoren_612_axis": True,
                                       "ginseng_equals_zaoren_axis": False,
                                       "zaoren_points_within_ginseng_measured_support": int(((zaoren.axis >= gs.axis[0]) & (zaoren.axis <= gs.axis[-1])).sum()),
                                       "ginseng_points_in_zaoren_candidate_interval": int(((gs.axis >= .1) & (gs.axis <= 1.5)).sum()),
                                       "interpolation_performed": False,
                                       "upsampling_does_not_add_resolution": True},
        "external_source_checks": source_checks,
        "transfer_design_status": {
            "A": "Source-only jujube training remains available and is the mainline",
            "B": "Blocked: no fully verified compatible external biological task; no performance fabricated",
            "C": "Blocked: no verified multi-plant THz cohort; source overlap and units unresolved",
            "future_control": "A/B/C must share jujube train counts, outer samples, seeds and architecture; preserve any negative transfer"},
        "input_sha256": before,
        "code_sha256": {str(path.relative_to(repo)): sha(path) for path in (
            repo / "src/herb_spectra/paper_c_external.py", Path(__file__).resolve())},
    }
    for path in inputs:
        assert sha(path) == before[str(path.relative_to(repo))], "Input was mutated"
    (output / "audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n")
    lines = ["# 方案C：外部植物THz数据可行性实查", "", "2026-10-31。酸枣仁855份及当前最优候选保持主线；本次外部工作只读现有数据、核对来源并建立独立接口。新增模型拟合0、新下载数值数据集0。", "",
             "| 数据 | 实际可读取内容 | 当前科学验证状态 |", "| --- | --- | --- |",
             "| 自有黄芪 | 4份工作簿，403条吸收系数记录；主区403×612，与酸枣仁频轴精确相同 | 物理样品ID、产地核验、批次、仪器信息均空；α长度单位未声明，独立外部分类与迁移阻塞 |",
             "| 人参／西洋参 | 42片压片标签、α与n各65点；20亚洲人参／22西洋参；原18CSV哈希一致 | 物种二分类标签可解释，但植株／重复／批次对应未知；不能称42独立植物，独立验证阻塞 |",
             "| 水稻饱满度V3 | 发布者比较页核到THz任务及CC BY 4.0 | 云端代理拒绝，内部矩阵、单位、独立种子与标签未解析 |",
             "| 白桦种子 | 发布页列时域数据及1.1GB seed.zip | 未下载，许可、类别和独立种子信息未建立 |",
             "| 小麦粉苯甲酸 | Dryad文件清单明确data/X/Y/fz.mat及BP/GRNN代码 | 云端拒绝，网页工具点击data/fz文件也403；是浓度回归，不是产地标签 |", "",
             "## 黄芪的具体问题", "",
             "403是谱记录数；元数据的physical_sample_id、origin、origin_verified、batch_id、instrument_id、species_verified全部为空。四个文件名不自动成为四个经核验生长产地。甘肃工作簿“折射率”表为空，吸收系数表完整；不填补该缺失模态。主区612点无负α，完整导出高频区存在负值，未修改原数值。频轴相同只能说明数值网格兼容，不能证明单位、仪器及反演条件一致。", "",
             "## 人参的具体问题", "",
             "原始α共有141／2730负值，36／42谱含负值，在≤1.5THz区间仍有107个负格；全部保留。无完全相同谱不构成生物学独立证据。6个文件组也不是6批新样品。n与α是同42片的两个模态，不能累加样品数。原频轴0.2—1.6THz、65点，与酸枣仁网格不同；共同区内酸枣仁568点、人参60点。未插值、未外围补零、未把物种标签改为酸枣仁产地。", "",
             "## 真实访问与许可", "",
             "[人参V2](https://data.mendeley.com/datasets/hvbp5mdtxx/2)与[水稻V2/V3比较页](https://data.mendeley.com/datasets/compare/wj2h5hn7kc)本轮可由网页工具读取，均列CC BY 4.0；[白桦](https://zenodo.org/records/17258169)已读页面许可字段没有可核定值；[Dryad小麦](https://datadryad.org/dataset/doi:10.5061/dryad.945c410)具体数据许可本轮未重新核定，不把论文开放许可套用于数据。", "",
             "四个仓储落地页的云端合法GET均被继承代理CONNECT403拒绝，详见[network_attempts.json](network_attempts.json)。这是环境访问阻塞，不是数据不存在。未绕过代理或反复重试。网页可读元数据不等于已下载二进制。响应失败没有生成数据SHA；原用户上传文件的SHA另行保留。", "",
             "## 独立接口与迁移决策", "",
             "`src/herb_spectra/paper_c_external.py`按植物分别读取；接口区分record_ids、physical_ids和labels，并在身份、标签或单位不足时拒绝独立评价。黄芪不生成未经核验的产地标签，人参仅保留物种标签。`task-interfaces/`是现有记录的派生数值副本，不是新增样品。插值工具仅允许实测支撑范围内插值，不增加分辨率；当前没有执行插值。", "",
             "A（只用酸枣仁）可继续；B（其他植物预训练）与C（多植物预训练）按现有证据阻塞，未生成虚构准确率。补齐物理样品／重复／批次及单位后，预训练与从头训练必须用同一酸枣仁训练量、同一评价样品、相同架构与种子比较。跨植物负迁移结果同样保存。", "",
             "机器记录：[audit.json](audit.json)。重现命令：`.runtime/bin/python scripts/audit_paper_c_external.py --output <新的报告目录>`；脚本拒绝覆盖已有audit.json。", ""]
    (output / "README.md").write_text("\n".join(lines))
    files = {str(path.relative_to(output)): sha(path) for path in output.rglob("*") if path.is_file() and path.name != "files_sha256.json"}
    (output / "files_sha256.json").write_text(json.dumps(files, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"status": "pass", "huangqi_records": 403, "ginseng_pellets": 42,
                      "new_model_fits": 0, "independent_external_evaluation": "blocked_with_reasons",
                      "inputs_unchanged": len(before)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
