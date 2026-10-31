#!/usr/bin/env python3
"""Read-only, post-score exports for separately frozen public THz benchmarks.

This module neither trains nor calls benchmark scoring. A root release bound to
completed scoring manifests is mandatory. Fruit exports contain aggregate
metrics, aggregate costs and parameter/feature statistics only: never labels,
predictions, raw waveforms, fitted models or support vectors.
"""
from __future__ import annotations

import argparse
import collections
import csv
from datetime import datetime, timezone
import gzip
import hashlib
import itertools
import json
from pathlib import Path

import numpy as np

SEEDS = (42, 43, 44)
STAGES = ("R0", "R1", "R2", "R3", "R4_policy")
CONFIG_KEYS = {"id", "representation", "preprocessing", "kind", "family", "mc", "C",
               "gamma", "class_weight", "components", "hidden", "l2", "alpha", "epochs", "epsilon"}
HISTORICAL_CONTEXT = {
    "dataset": "acid_jujube_historical_only", "role": "R2_MCUVE_SVM",
    "seed": 42, "accuracy": 821/855, "correct": 821, "total": 855,
    "macro_f1": None, "balanced_accuracy": None,
    "protocol": "historical 855-observation anonymous-seven-class retrospective internal nested CV",
    "interpretation": "96.0234% immutable context; no cross-dataset gain subtraction; historical document supplies other exact metrics",
    "source": "docs/RESEARCH4_EXPERIMENT_AUDIT.md",
}


class ExportError(ValueError):
    pass


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def sha(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024*1024), b""):
            h.update(block)
    return h.hexdigest()


def safe_member(root, value):
    path = (root/str(value)).resolve()
    if not path.is_relative_to(root.resolve()):
        raise ExportError("Source record path escapes its benchmark directory")
    return path


def write_json(path, value):
    with Path(path).open("x", encoding="utf-8") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2, allow_nan=False)
        handle.write("\n")


def write_csv(path, rows):
    if not rows:
        return
    keys = list(dict.fromkeys(k for row in rows for k in row))
    with Path(path).open("x", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=keys)
        writer.writeheader()
        for row in rows:
            writer.writerow({k: json.dumps(v, ensure_ascii=False, sort_keys=True) if isinstance(v, (dict,list,tuple)) else v
                             for k,v in row.items()})


def verify_release(release, sources):
    if release.get("allow_aggregate_export") is not True or not release.get("created_at_utc"):
        raise ExportError("Explicit root post-score export release required")
    for kind, root in sources.items():
        path = root/"scoring_manifest.json"
        if not path.exists():
            raise ExportError(f"{kind}: score is incomplete; exporter must not score or expose partial results")
        if release.get(kind+"_scoring_manifest_sha256") != sha(path):
            raise ExportError(f"{kind}: root export release is not bound to this completed scoring manifest")


def verify_scored(root):
    manifest = read(root/"scoring_manifest.json")
    if not manifest.get("completed_at_utc") or manifest.get("source_model_fits_in_scoring") != 0:
        raise ExportError("Completed independent scoring manifest required")
    for field, name in (("prediction_lock_sha256", "prediction_lock.json"),
                        ("metrics_sha256", "pooled_metrics_by_seed.json"),
                        ("scored_predictions_sha256", "scored_outer_predictions.csv.gz")):
        if manifest.get(field) != sha(root/name):
            raise ExportError(f"Completed scoring artifact changed: {name}")
    prediction_lock = read(root/"prediction_lock.json")
    checks = {"prediction_sha256":"outer_predictions.csv.gz", "outer_model_records_sha256":"outer_model_records.json",
              "outer_role_aliases_sha256":"outer_role_aliases.json", "selection_lock_sha256":"selection_lock.json",
              "root_outer_release_sha256":"outer_release_snapshot.json", "run_manifest_sha256":"run_manifest.json",
              "protocol_sha256":"protocol.json"}
    if prediction_lock.get("outer_truth_parsed") is not False:
        raise ExportError("Predictions were not locked before truth parsing")
    for field, name in checks.items():
        if prediction_lock.get(field) != sha(root/name):
            raise ExportError(f"Prediction-bound artifact changed: {name}")
    if manifest.get("prediction_sha256") != prediction_lock.get("prediction_sha256"):
        raise ExportError("Scoring/prediction locks disagree")
    selection = read(root/"selection_lock.json")
    for field, name in (("inner_records_sha256","inner_records.json"),
                        ("inner_predictions_sha256","inner_predictions.csv.gz"),
                        ("candidate_summaries_sha256","candidate_summaries.json")):
        if selection.get(field) != sha(root/name):
            raise ExportError(f"Source-selection artifact changed: {name}")
    return manifest, read(root/"protocol.json"), selection


def require_finite(row, keys):
    if not all(np.isfinite(row[k]) for k in keys):
        raise ExportError("Nonfinite aggregate metric; preserve upstream failure rather than invent a value")


def safe_config(config):
    result={k:v for k,v in config.items() if k in CONFIG_KEYS}
    if result.get("mc") is not None:
        result["mc"]={k:v for k,v in result["mc"].items() if k in {"runs","components","ratio"}}
    return result


def classification_rows(root, protocol):
    metrics = read(root/"pooled_metrics_by_seed.json")
    expected_roles = set(protocol["outer_roles"])
    if set(metrics) != expected_roles:
        raise ExportError("Ginseng role inventory differs from frozen protocol")
    result, confusions = [], []
    for role in protocol["outer_roles"]:
        values = metrics[role]
        if set(map(int, values)) != set(SEEDS):
            raise ExportError("All 42/43/44 ginseng seeds required; no winner-seed export")
        for seed in SEEDS:
            m = values[str(seed)]
            if m["total"] != 42 or int(m["correct"]) != m["correct"] or abs(m["accuracy"]-m["correct"]/42)>1e-12:
                raise ExportError("Ginseng accuracy numerator/denominator mismatch")
            require_finite(m, ("accuracy","macro_f1","balanced_accuracy"))
            row = {"dataset":"ginseng_hvbp5mdtxx_v2", "task":"species_0American_1Asian", "role":role,
                   "seed":seed, **{k:m[k] for k in ("accuracy","macro_f1","balanced_accuracy","correct","total")},
                   "protocol":protocol["protocol_id"], "scientific_status":protocol["scientific_status"]}
            result.append(row)
            matrix = np.asarray(m["confusion_matrix"],dtype=int)
            if m["class_order"]!=["0","1"] or matrix.shape!=(2,2) or (matrix<0).any() or int(matrix.sum())!=42 or int(matrix.trace())!=m["correct"]:
                raise ExportError("Ginseng confusion matrix numerator/denominator mismatch")
            for i,true in enumerate(m["class_order"]):
                for j,pred in enumerate(m["class_order"]):
                    confusions.append({"role":role,"seed":seed,"true_class":str(true),"predicted_class":str(pred),"count":int(matrix[i,j])})
    # Independently recompute ginseng aggregates only. Fruit individual targets
    # and predictions are never semantically read by this exporter.
    from sklearn.metrics import accuracy_score, balanced_accuracy_score, f1_score
    grouped = collections.defaultdict(list)
    with gzip.open(root/"scored_outer_predictions.csv.gz","rt",encoding="utf-8-sig") as handle:
        for row in csv.DictReader(handle):
            grouped[(row["role"],int(row["seed"]))].append(row)
    if set(grouped)!={(role,seed) for role in expected_roles for seed in SEEDS}:
        raise ExportError("Ginseng scored predictions have unexpected role/seed keys")
    for row in result:
        current=grouped[(row["role"],row["seed"])]
        if len(current)!=42 or len({r["record_id"] for r in current})!=42:
            raise ExportError("Ginseng requires 42 unique predictions per role/seed")
        y=[r["y_true"] for r in current]; p=[r["y_predicted"] for r in current]
        recalculated={"accuracy":accuracy_score(y,p),"balanced_accuracy":balanced_accuracy_score(y,p),
                      "macro_f1":f1_score(y,p,labels=["0","1"],average="macro",zero_division=0)}
        if any(abs(row[k]-v)>1e-12 for k,v in recalculated.items()):
            raise ExportError("Independent ginseng metric recalculation disagrees")
    return result,confusions


def per_class_rows(confusion_rows):
    matrices=collections.defaultdict(lambda:np.zeros((2,2),dtype=int))
    for row in confusion_rows:
        matrices[(row["role"],row["seed"])][int(row["true_class"]),int(row["predicted_class"])]=row["count"]
    result=[]
    for (role,seed),matrix in sorted(matrices.items()):
        for i,label in enumerate(("0American","1Asian")):
            tp=int(matrix[i,i]);support=int(matrix[i].sum());predicted=int(matrix[:,i].sum())
            precision=tp/predicted if predicted else 0.;recall=tp/support if support else 0.
            result.append({"role":role,"seed":seed,"class":label,"precision":precision,"recall":recall,
                           "f1":2*precision*recall/(precision+recall) if precision+recall else 0.,
                           "support":support,"true_positive":tp,"predicted_count":predicted})
    return result


def regression_rows(root, protocol):
    metrics=read(root/"pooled_metrics_by_seed.json")
    expected={(task,role,seed) for task in protocol["tasks"] for role in protocol["outer_roles"] for seed in SEEDS}
    actual=[(m["task"],m["role"],int(m["seed"])) for m in metrics]
    if len(actual)!=len(expected) or set(actual)!=expected:
        raise ExportError("All fruit tasks/roles/seeds required; no winner-seed export")
    result=[]
    for m in metrics:
        require_finite(m,("rmse","mae","r2"))
        if m["observations"]!=40 or m["fruit_groups"]!=10:
            raise ExportError("Fruit observation/group denominator changed")
        result.append({"dataset":"princeton_fruit_local_only", "task":m["task"],"role":m["role"],"seed":int(m["seed"]),
                       **{k:m[k] for k in ("rmse","mae","r2","observations","fruit_groups")},
                       "target_units":"Excel reported numeric reading units; physical unit and assay unverified",
                       "protocol":protocol["protocol_id"], "scientific_status":protocol["scientific_status"]})
    return result


def summarize(rows, metrics):
    groups=collections.defaultdict(list)
    for row in rows:groups[(row["dataset"],row["task"],row["role"])].append(row)
    result=[]
    for (dataset,task,role),current in sorted(groups.items()):
        if sorted(r["seed"] for r in current)!=list(SEEDS):
            raise ExportError("Summary requires every predeclared seed exactly once")
        row={"dataset":dataset,"task":task,"role":role,"seeds":"42,43,44","seed_count":3,
             "protocol":current[0]["protocol"],"scientific_status":current[0]["scientific_status"]}
        for key in ("total","observations","fruit_groups","target_units"):
            if key in current[0]:row[key]=current[0][key]
        for name in metrics:
            vals=np.array([r[name] for r in current],dtype=float)
            row[name+"_mean"]=float(np.mean(vals));row[name+"_sample_sd"]=float(np.std(vals,ddof=1))
            row[name+"_mean_sd"]=f"{np.mean(vals):.6f} ± {np.std(vals,ddof=1):.6f}"
        if "correct" in current[0]:row["correct_each_seed"]=";".join(f"{r['seed']}:{r['correct']}/{r['total']}" for r in sorted(current,key=lambda r:r['seed']))
        result.append(row)
    return result


def within_protocol_rounds(rows):
    baselines={(r["task"],r["seed"]):r for r in rows if r["role"]=="R0"}
    result=[]
    for row in rows:
        if row["role"] not in STAGES:continue
        baseline=baselines[(row["task"],row["seed"])]
        if row["dataset"]!=baseline["dataset"] or row["protocol"]!=baseline["protocol"] or row["total"]!=baseline["total"]:
            raise ExportError("Cross-dataset/protocol gain subtraction is prohibited")
        result.append({**row,"accuracy_gain_percentage_points":100*(row["accuracy"]-baseline["accuracy"]),
                       "accuracy_relative_improvement_percent":100*(row["accuracy"]/baseline["accuracy"]-1) if baseline["accuracy"] else None,
                       "macro_f1_gain":row["macro_f1"]-baseline["macro_f1"],
                       "balanced_accuracy_gain_percentage_points":100*(row["balanced_accuracy"]-baseline["balanced_accuracy"]),
                       "net_more_correct":row["correct"]-baseline["correct"],"reference":"same-protocol R0 same seed only"})
    return result


def selected_models_and_costs(root, kind, protocol, selection):
    records=read(root/"outer_model_records.json"); aliases=read(root/"outer_role_aliases.json")
    inner=read(root/"inner_records.json")
    actual_by_path={r["estimator_path"]:r for r in records}
    if len(actual_by_path)!=len(records):raise ExportError("Duplicate actual estimator path")
    shared={(a.get("task","species"),a.get("outer",a.get("outer_group")),a["role"],int(a["seed"])):a for a in aliases}
    configurations=[];indices=collections.defaultdict(list);metadata_cache={}
    for fold in selection["outer_folds"]:
        task=fold.get("task","species");outer=fold.get("outer",fold.get("outer_group"))
        choices={role:choice["config"] for role,choice in fold["choices"].items()}
        if kind=="ginseng":choices["R0"]=protocol["R0"]
        else:choices["dummy_mean"]={"id":"dummy_mean","family":"dummy_mean","mc":None,"preprocessing":"raw"}
        for role in protocol["outer_roles"]:
            config=choices[role]
            for seed in SEEDS:
                alias=shared.get((task,outer,role,seed))
                if alias:rec=actual_by_path[alias["estimator_path"]]
                else:
                    candidates=[r for r in records if int(r["seed"])==seed and r.get("task","species")==task
                                and r.get("outer",r.get("outer_group"))==outer and r["config"]["id"]==config["id"]]
                    if len(candidates)!=1:raise ExportError("Could not resolve frozen role to actual estimator")
                    rec=candidates[0]
                path=safe_member(root,rec["estimator_path"]);meta_path=path.with_suffix(".json")
                if str(path) not in metadata_cache:
                    if rec["metadata_sha256"]!=sha(meta_path):raise ExportError("Selected model metadata changed")
                    if rec["estimator_sha256"]!=sha(path):raise ExportError("Selected model bytes changed")
                    metadata_cache[str(path)]=read(meta_path)
                meta=metadata_cache[str(path)]
                feature_indices=meta.get("retained_indices",meta.get("selected_indices"))
                if not isinstance(feature_indices,list):raise ExportError("Selected feature index metadata absent")
                dimension=len(feature_indices)
                row={"dataset":kind,"task":task,"outer_fold":outer,"role":role,"seed":seed,
                     "selected_input_features":dimension,"parameters":safe_config(config),
                     "recorded_trial_seconds":float(rec["elapsed_seconds"]),"model_bytes":path.stat().st_size,
                     "shared_existing_estimator":bool(alias),"scientific_status":protocol["scientific_status"]}
                # No fruit source/query IDs, per-record labels, scores, traces or
                # support vectors are copied from model metadata.
                if kind=="ginseng":row["selected_input_indices"]=feature_indices
                configurations.append(row)
                representation=config.get("representation","time_index_blocks_128")
                indices[(task,role,seed,representation)].append(set(feature_indices))
    stability=[]
    for (task,role,seed,representation),sets in sorted(indices.items()):
        pairs=[len(a&b)/len(a|b) if a|b else 1. for a,b in itertools.combinations(sets,2)]
        stability.append({"dataset":kind,"task":task,"role":role,"seed":seed,"representation":representation,
                          "folds_same_representation":len(sets),"fold_pairs":len(pairs),
                          "mean_pairwise_selected_column_jaccard":float(np.mean(pairs)) if pairs else None,
                          "min_selected_input_features":min(map(len,sets)),"max_selected_input_features":max(map(len,sets)),
                          "interpretation":"descriptive column stability within same representation; full-spectrum=1 is trivial; not biological peak or independent-batch evidence"})
    costs=[]
    for phase,current in (("source_search",inner),("outer_unique_models",records)):
        groups=collections.defaultdict(list)
        for record in current:
            stage=record["config"]["id"].split("_",1)[0]
            groups[(record.get("task","species"),int(record["seed"]),stage)].append(record)
        for (task,seed,stage),current in sorted(groups.items()):
            costs.append({"dataset":kind,"phase":phase,"task":task,"seed":seed,"stage":stage,
                          "actual_unique_estimator_records":len(current),
                          "sum_recorded_trial_seconds":float(sum(r["elapsed_seconds"] for r in current)),
                          "cost_note":"includes trial transforms/inference/serialization; source transform caches reused; not pure fit time or independent wall-clock"})
    return configurations,stability,costs


def total_cost_rows(costs):
    grouped=collections.defaultdict(list)
    for row in costs:grouped[(row["dataset"],row["phase"])].append(row)
    return [{"dataset":dataset,"phase":phase,
             "actual_unique_estimator_records":sum(r["actual_unique_estimator_records"] for r in current),
             "sum_recorded_trial_seconds":sum(r["sum_recorded_trial_seconds"] for r in current),
             "all_seeds":"42,43,44","cost_note":"All observed records, not winner-only cost; trial seconds are not wall-clock or pure fit time"}
            for (dataset,phase),current in sorted(grouped.items())]


def xlsx(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    book=Workbook();book.remove(book.active)
    for name,rows in sheets.items():
        if not rows:continue
        ws=book.create_sheet(name[:31]);keys=list(dict.fromkeys(k for r in rows for k in r));ws.append(keys)
        for row in rows:
            ws.append([json.dumps(row.get(k),ensure_ascii=False,sort_keys=True) if isinstance(row.get(k),(dict,list,tuple)) else row.get(k) for k in keys])
        ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(color="FFFFFF",bold=True);cell.fill=PatternFill("solid",fgColor="244A69")
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(65,max(14,len(str(col[0].value))+2))
    book.save(path)


def plot_exports(out, ginseng_summary, confusion_rows, fruit_summary, stability):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    def save(fig,name):
        fig.savefig(out/(name+".png"),dpi=180,bbox_inches="tight");fig.savefig(out/(name+".pdf"),bbox_inches="tight");plt.close(fig)
    if ginseng_summary:
        stages={r["role"]:r for r in ginseng_summary if r["role"] in STAGES}
        fig,axes=plt.subplots(1,3,figsize=(14,4),sharey=True)
        for ax,metric in zip(axes,("accuracy","macro_f1","balanced_accuracy")):
            ax.errorbar(range(5),[stages[k][metric+"_mean"] for k in STAGES],
                        yerr=[stages[k][metric+"_sample_sd"] for k in STAGES],fmt="o-",capsize=4)
            ax.set_xticks(range(5),["R0","R1 MC","R2 SVM","R3 prep","R4 policy"],rotation=25)
            ax.set_title(metric);ax.set_ylim(0,1.02);ax.grid(alpha=.2)
        fig.suptitle("Ginseng: fixed stages, all seeds 42/43/44; exploratory file-group CV")
        save(fig,"ginseng_four_rounds_all_seeds")
        cm={(r["role"],r["seed"]):np.zeros((2,2),dtype=int) for r in confusion_rows if r["role"] in STAGES}
        for r in confusion_rows:
            if r["role"] in STAGES:cm[(r["role"],r["seed"])][int(r["true_class"]),int(r["predicted_class"])]=r["count"]
        fig,axes=plt.subplots(3,5,figsize=(15,9))
        for i,seed in enumerate(SEEDS):
            for j,role in enumerate(STAGES):
                ax=axes[i,j];a=cm[(role,seed)];ax.imshow(a,vmin=0,vmax=42,cmap="Blues")
                for y,x in itertools.product(range(2),repeat=2):ax.text(x,y,str(a[y,x]),ha="center",va="center")
                ax.set_xticks([0,1],["American","Asian"],rotation=20);ax.set_yticks([0,1],["American","Asian"])
                ax.set_title(f"{role}, seed {seed}");ax.set_xlabel("Predicted");ax.set_ylabel("True")
        fig.suptitle("Ginseng OOF confusions: 42 observations in each panel, no seed selection")
        fig.tight_layout();save(fig,"ginseng_stage_confusions_all_seeds")
        feature_rows=[r for r in stability if r["dataset"]=="ginseng" and r["role"] in STAGES]
        fig,ax=plt.subplots(figsize=(9,4))
        for seed in SEEDS:
            selected=[r for r in feature_rows if r["seed"]==seed]
            for r in selected:
                x=STAGES.index(r["role"])+(seed-43)*.07
                ax.plot([x,x],[r["min_selected_input_features"],r["max_selected_input_features"]],linewidth=4,alpha=.7)
        ax.set_xticks(range(5),STAGES);ax.set_ylabel("Selected input columns, fold min to max")
        ax.set_title("Input feature counts; PCA latent dimensions are separate parameters")
        save(fig,"ginseng_selected_input_features")
    if fruit_summary:
        tasks=sorted({r["task"] for r in fruit_summary})
        for metric in ("rmse","mae","r2"):
            fig,axes=plt.subplots(2,2,figsize=(13,8));roles=("dummy_mean","PLSR","full_SVR","MLP","MC_SVR","source_policy")
            for ax,task in zip(axes.ravel(),tasks):
                vals={r["role"]:r for r in fruit_summary if r["task"]==task}
                ax.bar(range(6),[vals[k][metric+"_mean"] for k in roles],yerr=[vals[k][metric+"_sample_sd"] for k in roles],capsize=3)
                ax.set_xticks(range(6),roles,rotation=30,ha="right");ax.set_title(task);ax.set_ylabel(metric+" (reported reading units)" if metric!="r2" else "R2 (negative values retained)");ax.grid(axis="y",alpha=.2)
            fig.suptitle("Fruit: aggregate grouped OOF regression, all seeds; units/assay unverified")
            fig.tight_layout();save(fig,"fruit_aggregate_"+metric)


def export(sources, out, release_path):
    sources={k:Path(v).resolve() for k,v in sources.items()};out=Path(out).resolve()
    if not sources:raise ExportError("At least one scored benchmark source required")
    if out.exists() and any(out.iterdir()):raise FileExistsError("Refusing to overwrite an existing result export")
    if any(out.is_relative_to(root) for root in sources.values()):raise ExportError("Export must not mutate a frozen benchmark directory")
    release=read(release_path);verify_release(release,sources)
    classes=[];confusions=[];regression=[];models=[];stability=[];costs=[];provenance={}
    for kind,root in sources.items():
        scoring,protocol,selection=verify_scored(root)
        if tuple(protocol["seeds"])!=SEEDS:raise ExportError("Unexpected frozen seed protocol")
        if kind=="ginseng":classes,confusions=classification_rows(root,protocol)
        else:regression=regression_rows(root,protocol)
        m,s,c=selected_models_and_costs(root,kind,protocol,selection);models+=m;stability+=s;costs+=c
        provenance[kind]={"directory":str(root),"protocol":protocol["protocol_id"],"scientific_status":scoring["scientific_status"],
                          "scoring_manifest_sha256":sha(root/"scoring_manifest.json"),"metrics_sha256":sha(root/"pooled_metrics_by_seed.json"),
                          "prediction_lock_sha256":sha(root/"prediction_lock.json"),"source_failures":selection["source_failures"],
                          "per_record_numerical_data_exported":False}
    class_summary=summarize(classes,("accuracy","macro_f1","balanced_accuracy")) if classes else []
    fruit_summary=summarize(regression,("rmse","mae","r2")) if regression else []
    rounds=within_protocol_rounds(classes) if classes else []
    round_summary=summarize(rounds,("accuracy","macro_f1","balanced_accuracy","accuracy_gain_percentage_points","macro_f1_gain","balanced_accuracy_gain_percentage_points")) if rounds else []
    out.mkdir(parents=True,exist_ok=True)
    tables={"ginseng_per_seed":classes,"ginseng_summary":class_summary,"ginseng_four_rounds_per_seed":rounds,
            "ginseng_four_rounds_summary":round_summary,"ginseng_confusions":confusions,
            "ginseng_per_class":per_class_rows(confusions),
            "fruit_aggregate_per_seed":regression,"fruit_aggregate_summary":fruit_summary,
            "selected_model_parameters":models,"feature_stability":stability,"computational_costs":costs,
            "computational_cost_totals":total_cost_rows(costs),
            "historical_context_only":[HISTORICAL_CONTEXT]}
    for name,rows in tables.items():write_csv(out/(name+".csv"),rows)
    xlsx(out/"public_thz_comparison.xlsx",tables)
    plot_exports(out,class_summary,confusions,fruit_summary,stability)
    text="""# 公开太赫兹数据：冻结协议内的回顾性结果\n\n本目录由独立只读导出器在评分完成且获得 root 绑定哈希的导出许可后生成；不进行训练、再次评分或模型选择。\n\n## 解释边界\n\n- 酸枣仁历史 R2：821/855＝96.0234%，保持为原研究强基线。它仅出现在历史背景表，不能与人参/水果成绩相减作为准确率提升。\n- 人参为42片、物种二分类，0西洋参/1人参；六个导出文件组留一外层、五个剩余文件组内层。原料根与批次对应仍未知，文件组不是独立生物批次。全部42/43/44种子均列出；均值±样本标准差描述种子波动，不是置信区间。\n- R0→R1→R2→R3→R4是源内选择策略的对照，不能因外层某阶段较好而倒选配置；最终表中的最好分数也只属于回顾性描述，不构成新独立确认。准确率提升仅在相同人参协议与相同种子内相对R0计算，个百分点与相对百分比不同。\n- 水果分别为芒果Brix/DM、梨Brix、柿Brix四项回归；每任务40时点观察来自10个持续跟踪果实，外5/内3按果实分组。报告RMSE、MAE和R²，保留负R²，不能称为分类准确率。标签为原Excel读数，物理单位/测定方法未确认。其128个时域工程块不等于吸收系数或THz频率轴。\n- 水果原件、逐样本标签/预测、波形、模型和支持向量因许可不明保留本地；本目录只含聚合指标、参数和统计图。未复制逐样本数据。\n\n## 文件和统计口径\n\nCSV与XLSX提供人参各方法逐种子/均值±标准差、四轮提升、混淆矩阵、选参、输入特征数、描述性Jaccard稳定性和计算开销，以及水果聚合回归对照。PNG/PDF为独立图件。\n\n特征稳定性只比较同一表示中的原始输入列集合；全谱的Jaccard=1是平凡结果。PCA潜变量维数另见参数，不能把输入列数写成最终潜在维数。不同预处理下的列重合不是化学峰归属或跨批次稳定性证明。\n\n计算时间汇总实际trial记录，含部分变换、预测、序列化及共享缓存影响，不是纯fit耗时或端到端墙钟。source_search计全部实际内层模型，outer_unique_models只计真实拟合对象；角色别名复用不能重复宣称实际拟合次数。查询分数不参与成本或参数选择。\n\n`export_manifest.json`绑定源评分/预测锁与本次导出许可；`files_sha256.json`提供产物哈希。原始失败数保留于来源记录，未隐藏失败或挑选有利种子。\n"""
    (out/"README.md").write_text(text,encoding="utf-8")
    write_json(out/"export_manifest.json",{"created_at_utc":datetime.now(timezone.utc).isoformat(),"exporter_sha256":sha(Path(__file__)),
                  "release_sha256":sha(release_path),"source_provenance":provenance,"training_operations":0,
                  "benchmark_scoring_calls":0,"winner_seed_selected":False,"cross_dataset_gain_computed":False,
                  "fruit_per_record_data_exported":False,"ginseng_metrics_independently_recalculated":bool(classes),
                  "seed_dispersion":"sample SD ddof=1; not a confidence interval"})
    write_json(out/"files_sha256.json",{p.name:{"sha256":sha(p),"bytes":p.stat().st_size} for p in sorted(out.iterdir()) if p.is_file()})
    return {"output":str(out),"ginseng_role_seed_rows":len(classes),"fruit_aggregate_role_seed_rows":len(regression),"training_operations":0}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--ginseng",type=Path)
    parser.add_argument("--fruit",type=Path)
    parser.add_argument("--output",type=Path,required=True)
    parser.add_argument("--release",type=Path,required=True)
    args=parser.parse_args()
    sources={k:v for k,v in (("ginseng",args.ginseng),("fruit",args.fruit)) if v is not None}
    print(json.dumps(export(sources,args.output,args.release),ensure_ascii=False))


if __name__=="__main__":main()
