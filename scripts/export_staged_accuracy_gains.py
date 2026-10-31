"""Export accuracy gains from frozen OOF counts; no model fitting."""
import csv
import hashlib
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "reports/staged-nested-optimization/2026-10-31/tables/四轮OOF台账.csv"
DEST = ROOT / "reports/staged-accuracy-gains/2026-10-31"
MAIN = {
    "R0": "原方案（比较基准）",
    "R1": "第一轮：MC-UVE参数优化",
    "R2": "第二轮：继续优化SVM",
    "R3": "第三轮：继续选择预处理",
    "R4": "第四轮：继续选择组合家族",
}
DIAGNOSTICS = {
    "R3_diagnostic_raw": "预处理对照：Raw",
    "R3_diagnostic_snv": "预处理对照：SNV",
    "R3_diagnostic_sg0_snv": "预处理对照：SG平滑＋SNV",
    "R3_diagnostic_sg1_snv": "预处理对照：SG一阶＋SNV",
    "R3_diagnostic_sg2_snv": "预处理对照：SG二阶＋SNV",
    "R3_diagnostic_msc": "预处理对照：MSC",
    "R3_diagnostic_linear_detrend_snv": "预处理对照：去趋势＋SNV",
    "fixed_snv_C10_MC20": "固定SNV／C10／MC20锚定",
    "R4_diagnostic_ordinary_mcuve": "组合对照：普通MC-UVE",
    "R4_diagnostic_budget_mcuve": "组合对照：10倍MC预算",
    "R4_diagnostic_stable_mcuve": "组合对照：稳定MC-UVE",
    "R4_diagnostic_stable_cars": "组合对照：稳定CARS",
    "R4_diagnostic_stable_joint": "组合对照：MC-UVE／CARS联合稳定频率",
    "R4_diagnostic_cnn_softmax": "组合对照：1D-CNN原分类头",
    "R4_diagnostic_cnn_embedding_svm": "组合对照：同一CNN表示＋SVM",
}
HEADERS = ["方法代码", "方法", "正确数", "样品总数", "准确率", "较初始增加（百分点）", "相对初始准确率提升率", "净增加正确数"]


def main():
    source_sha = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    with SOURCE.open(encoding="utf-8-sig", newline="") as handle:
        records = list(csv.DictReader(handle))
    baseline = next(row for row in records if row["方法"] == "R0")
    total, original = int(baseline["独立样品数"]), int(baseline["正确数"])
    assert (total, original) == (855, 773)
    rows = []
    for row in records:
        code, count = row["方法"], int(row["正确数"])
        assert int(row["独立样品数"]) == total
        assert abs(float(row["准确率"]) - count / total) < 1e-14
        rows.append([code, (MAIN | DIAGNOSTICS)[code], count, total, count / total,
                     100 * (count - original) / total, (count - original) / original, count - original])
    assert len(rows) == len({r[0] for r in rows}) == 20
    main_rows = [next(row for row in rows if row[0] == code) for code in MAIN]
    assert [r[2] for r in main_rows] == [773, 791, 821, 816, 812]
    DEST.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    for name, values in [("四轮主结果", main_rows), ("全部20方法对照", rows)]:
        sheet = book.create_sheet(name)
        sheet.append(HEADERS)
        for values_row in values:
            sheet.append(values_row)
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244F70")
            cell.alignment = Alignment(wrap_text=True)
        for letter, width in {"A":42,"B":47,"C":12,"D":12,"E":16,"F":26,"G":30,"H":20}.items():
            sheet.column_dimensions[letter].width = width
        sheet.row_dimensions[1].height = 32
        for cells in sheet.iter_rows(min_row=2):
            cells[4].number_format = "0.00%"
            cells[5].number_format = "+0.00;-0.00;0.00"
            cells[6].number_format = "+0.00%;-0.00%;0.00%"
            cells[7].number_format = "+0;-0;0"
        csv_path = DEST / f"{name}.csv"
        with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle, lineterminator="\n")
            writer.writerow(HEADERS)
            writer.writerows(values)
    notes = [
        ("初始基准", "R0：773／855＝90.4093567251462%；显示为90.41%。"),
        ("增加百分点", "100×（当前正确数－773）／855；单位是百分点，不是相对百分比。"),
        ("相对提升率", "（当前准确率－初始准确率）／初始准确率＝（当前正确数－773）／773；Excel以百分数显示。"),
        ("数值精度", "全部由整数正确数计算，未用四舍五入后的90.41%作分母。负数表示下降。"),
        ("可比范围", "同一855份样品、同一5外折／2内折的OOF；不是513／171／171实验。"),
        ("诊断范围", "附表保留全部20条方法角色；同一预测可因控制目的出现多种角色，未增加独立样品。"),
        ("因果解释", "各轮均相对原R0比较；不是各步骤的单因素贡献，也不是相对前一轮的增幅。"),
        ("验证限制", "各产地只有一批、历史数据已用于探索；结果是回顾性批次内验证，不是跨批次准确率。"),
        ("数据来源", str(SOURCE.relative_to(ROOT))),
        ("来源SHA256", source_sha),
    ]
    sheet = book.create_sheet("口径与公式")
    for row in notes:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 115
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook = DEST / "模型准确率_相对初始提升对照.xlsx"
    book.save(workbook)
    check = load_workbook(workbook, data_only=True, read_only=True)
    checked = 0
    for name, expected in [("四轮主结果", main_rows), ("全部20方法对照", rows)]:
        actual = list(check[name].values)
        assert list(actual[0]) == HEADERS
        for a, e in zip(actual[1:], expected, strict=True):
            for av, ev in zip(a, e, strict=True):
                assert abs(av - ev) < 1e-12 if isinstance(ev, float) else av == ev
                checked += 1
    check.close()
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == source_sha
    table = ["| 方法 | 准确率 | 较初始增加（百分点） | 相对提升率 |", "| --- | ---: | ---: | ---: |"]
    table += [f"| {r[1]} | {100*r[4]:.2f}% | {r[5]:+.2f} | {100*r[6]:+.2f}% |" for r in main_rows]
    readme = "# 相对最初模型的准确率提升对照\n\n" + "\n".join(table)
    readme += "\n\n[下载Excel](模型准确率_相对初始提升对照.xlsx)：含四轮主结果、全部20方法对照及口径与公式。\n\n"
    readme += "\n\n".join(f"**{name}**：{value}" for name, value in notes)
    readme += "\n\n生成命令：`.runtime/bin/python scripts/export_staged_accuracy_gains.py`。本次仅计算已有结果的提升幅度，未重新训练或选参。\n"
    (DEST / "README.md").write_text(readme, encoding="utf-8")
    validation = {"status":"pass", "source":str(SOURCE.relative_to(ROOT)), "source_sha256":source_sha,
                  "baseline_correct":original,"samples":total,"method_roles":len(rows),
                  "excel_cells_read_back":checked,"new_model_fits":0,
                  "files_sha256":{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in DEST.iterdir() if p.is_file() and p.name != "validation.json"}}
    (DEST / "validation.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    print("\n".join(table))
    print(json.dumps({"status":"pass", "excel_cells_read_back":checked, "new_model_fits":0}, ensure_ascii=False))


if __name__ == "__main__":
    main()
