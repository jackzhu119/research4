"""Export completed image-guided experiments; no fitting or model selection."""
import argparse
from collections import Counter, defaultdict
import csv
import gzip
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


CLASSES = tuple(f"ZY{i:02d}" for i in range(1, 8))
HARD = ("accuracy", "balanced_accuracy", "macro_f1")
TRACKS = ("features", "probability", "fusion")
NAMES = {"features": "特征筛选与判别", "probability": "源OOF概率与判别调整", "fusion": "SVM与MLP概率融合"}


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def hard_metrics(truth, predicted):
    cm = np.zeros((7, 7), dtype=int)
    for actual, pred in zip(truth, predicted):
        cm[CLASSES.index(actual), CLASSES.index(pred)] += 1
    tp = cm.diagonal(); support = cm.sum(1); found = cm.sum(0)
    recall = np.divide(tp, support, out=np.zeros(7), where=support > 0)
    precision = np.divide(tp, found, out=np.zeros(7), where=found > 0)
    f1 = np.divide(2*tp, support+found, out=np.zeros(7), where=support+found > 0)
    return {"accuracy": float(tp.sum()/cm.sum()), "balanced_accuracy": float(recall[support > 0].mean()),
        "macro_f1": float(f1.mean()), "confusion_matrix": cm.tolist(), "class_recall": recall.tolist(),
        "class_counts": support.tolist(), "class_precision": precision.tolist(), "class_f1": f1.tolist(),
        "class_predicted_counts": found.tolist(), "class_true_positives": tp.tolist(),
        "class_false_positives": (found-tp).tolist(), "class_false_negatives": (support-tp).tolist(),
        "correct": int(tp.sum()), "total": int(cm.sum())}


def proper_metrics(probabilities, labels):
    p = np.asarray(probabilities, dtype=float)
    y = np.asarray([CLASSES.index(v) for v in labels], dtype=int)
    if p.shape != (len(y), 7) or not np.isfinite(p).all() or np.any(p < 0) or np.any(p > 1):
        raise AssertionError("Invalid saved seven-class probability vector")
    if not np.allclose(p.sum(1), 1, rtol=0, atol=1e-12):
        raise AssertionError("Saved probabilities do not sum to one")
    confidence = p.max(1); correct = p.argmax(1) == y
    bins = np.minimum((confidence*10).astype(int), 9)
    ece = sum(float(np.mean(bins == b))*abs(float(correct[bins == b].mean())-float(confidence[bins == b].mean()))
              for b in range(10) if np.any(bins == b))
    return {"nll": float(-np.log(np.maximum(p[np.arange(len(y)), y], 1e-15)).mean()),
        "brier": float(np.square(p-np.eye(7)[y]).sum(1).mean()), "ece": float(ece)}


def write_csv(path, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader(); writer.writerows(rows)


def add_sheet(book, name, rows):
    if not rows:
        rows = [{"说明": "该项没有已保存且完成的记录；未生成模拟结果"}]
    sheet = book.create_sheet(name); keys = list(rows[0]); sheet.append(keys)
    for row in rows:
        sheet.append([row.get(k) for k in keys])
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions; sheet.sheet_view.zoomScale = 80
    for cell in sheet[1]:
        cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="244F70")
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    for j, key in enumerate(keys, 1):
        sheet.column_dimensions[get_column_letter(j)].width = 20 if len(key) < 12 else 26
        if key in {"说明", "候选参数", "方法身份", "解释边界"}:
            sheet.column_dimensions[get_column_letter(j)].width = 62
        for row in sheet.iter_rows(min_row=2, min_col=j, max_col=j):
            cell = row[0]; cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F1F5F9")
            if any(token in key for token in ("准确率", "精确率", "召回率", "置信度", "选择比例", "正确率")):
                cell.number_format = "0.00%"
            elif key in {"宏F1", "NLL", "Brier", "ECE10", "平均Jaccard"} or key.endswith("F1"):
                cell.number_format = "0.0000"
    if "是否正确" in keys:
        j = keys.index("是否正确")
        for row in sheet.iter_rows(min_row=2):
            if row[j].value == 0:
                for cell in row:
                    cell.fill = PatternFill("solid", fgColor="FCE4E4")
    for i in range(1, len(rows)+2):
        sheet.row_dimensions[i].height = 42 if i == 1 else 48 if name in {"说明", "正式结果台账", "所有最终方法", "候选开发汇总"} else 26
    sheet.page_setup.orientation = "landscape"


def export(base, output, split_plan, input_path):
    base, output, split_plan, input_path = map(Path, (base, output, split_plan, input_path))
    gate_path = base/"all_development_locks.json"
    if not gate_path.exists():
        raise RuntimeError("三轨开发锁尚未统一释放，拒绝生成不完整或模拟结果")
    gate = read_json(gate_path); by_name = {r["name"]: r for r in gate.get("rounds", [])}
    if set(by_name) != set(TRACKS):
        raise RuntimeError("三轨统一开发锁不完整")
    runs = {}; checks = 0; maximum = 0.; provenance = {str(gate_path): digest(gate_path)}

    def close(actual, expected, label):
        nonlocal checks, maximum
        delta = float(np.max(np.abs(np.asarray(actual, dtype=float)-np.asarray(expected, dtype=float))))
        checks += 1; maximum = max(maximum, delta)
        if not np.isfinite(delta) or delta > 1e-12:
            raise AssertionError(f"Numerical discrepancy: {label}: {delta}")

    gate_checks = 0
    for name, record in by_name.items():
        run = Path(record["selection_lock"]).parent; runs[name] = run
        for p, expected in [(run/"selection_lock.json", record["selection_lock_sha256"]),
                (run/"run_manifest.json", record["run_manifest_sha256"]),
                (Path(record["protocol"]), record["protocol_sha256"])]:
            if not p.exists() or digest(p) != expected:
                raise RuntimeError("开发选择锁、执行清单或协议SHA已变化")
            gate_checks += 1
        if record.get("final_directory_existed_at_release", False):
            raise RuntimeError("统一放行前已有最终输出，需核查时序")
        if not (run/"status.json").exists() or read_json(run/"status.json").get("status") != "completed":
            raise RuntimeError(f"三轨最终评价尚未完成：{name}")
        if read_json(run/"status.json").get("target_support_labels_used", 0) != 0:
            raise RuntimeError("本轮应为零目标支持标签实验，不能混入其他标签预算")
        for filename in ("final_metrics.json", "summary.json", "all_development_metrics.json", "selection_lock.json"):
            if not (run/filename).exists():
                raise RuntimeError(f"缺少已完成证据：{name}/{filename}")
        if not list((run/"final").rglob("predictions.csv.gz")):
            raise RuntimeError(f"缺少{ name }最终预测记录")
    master = Path("data/research/image_guided_plan_20261031.json")
    if digest(master) != gate["master_plan_sha256"]:
        raise RuntimeError("总方案SHA改变")
    gate_checks += 1
    manifests = {name: read_json(run/"run_manifest.json") for name, run in runs.items()}
    if any(digest(input_path) != m["input_sha256"] for m in manifests.values()):
        raise RuntimeError("原始NPZ与运行清单不符")
    gate_checks += 1
    if any(digest(split_plan) != m["plan_sha256"] for m in manifests.values()):
        raise RuntimeError("固定用途分区与运行清单不符")
    gate_checks += 1
    plan = read_json(split_plan); truth = {r["record_id"]: r["origin_group_id"] for r in plan["records"]}
    with np.load(input_path, allow_pickle=False) as archive:
        full_axis = np.asarray(archive["axis"], dtype=float)
    expected_axis = full_axis[(full_axis >= .1) & (full_axis <= 1.5)]
    if len(expected_axis) != 612:
        raise AssertionError("Input no longer has the fixed612 frequency coordinates")
    final_ids = {r["record_id"] for r in plan["records"] if r["role"] == "prediction"}
    source_ids = {r["record_id"] for r in plan["records"] if r["role"] == "training"}
    if len(final_ids) != 171 or len(source_ids) != 513:
        raise AssertionError("Fixed source/prediction roles changed")
    random_ids = set(np.random.default_rng(20261102).choice(np.asarray(sorted(final_ids)), 20, replace=False).tolist())
    locks = {n: read_json(r/"selection_lock.json") for n, r in runs.items()}
    grids = {}
    for name, run in runs.items():
        if (run/"candidate_grid.json").exists():
            grid = read_json(run/"candidate_grid.json")
        else:
            grid = [{k: v for k, v in r.items() if k in {"id", "provider", "kind", "beta", "penalty"}}
                    for r in read_json(run/"summary.json")]
        grids[name] = {c["id"]: c for c in grid}
    if [len(grids[n]) for n in TRACKS] != [40, 21, 20]:
        raise AssertionError("Prespecified grid counts changed")

    cases = {}; raw_counts = Counter(); final_records = {}; dev_groups = defaultdict(list)
    for name, run in runs.items():
        for scope in ("development", "final"):
            files = sorted((run/scope).rglob("predictions.csv.gz"))
            for path in files:
                provenance[str(path)] = digest(path)
                with gzip.open(path, "rt", encoding="utf-8-sig", newline="") as handle:
                    for raw in csv.DictReader(handle):
                        checks += 4; identifier = raw["record_id"]
                        if identifier not in (final_ids if scope == "final" else source_ids):
                            raise AssertionError("Prediction contains specimen outside its fixed role")
                        if raw["y_true"] != truth[identifier] or raw["y_predicted"] not in CLASSES:
                            raise AssertionError("Recorded origin label differs from original split")
                        if int(raw["is_correct"]) != int(raw["y_true"] == raw["y_predicted"]):
                            raise AssertionError("Incorrect correctness marker")
                        if int(raw.get("target_support_labels_used", 0)) != 0:
                            raise AssertionError("Unexpected target-support label use")
                        episode = -1 if scope == "final" else int(raw["episode"])
                        seed = int(raw["seed"]); candidate = raw["candidate_id"]
                        method = raw.get("method", candidate)
                        key = (name, scope, episode, seed, method, candidate)
                        if key not in cases:
                            cases[key] = {}
                        if identifier in cases[key]:
                            raise AssertionError("Duplicate specimen in one prediction case")
                        cases[key][identifier] = raw; raw_counts[(name, scope)] += 1
            recorded = read_json(run/("final_metrics.json" if scope == "final" else "all_development_metrics.json"))
            for row in recorded:
                episode = -1 if scope == "final" else int(row["episode"])
                raw_method = row["method"] if name == "fusion" else row["candidate_id"]
                key = (name, scope, episode, int(row["seed"]), raw_method, row["candidate_id"])
                if key not in cases:
                    raise AssertionError("Metric JSON has no raw prediction case")
                raw_case = cases[key]; ids = sorted(raw_case)
                if len(ids) != (171 if scope == "final" else 103):
                    raise AssertionError("A saved prediction case is incomplete")
                actual = [raw_case[i]["y_true"] for i in ids]; predicted = [raw_case[i]["y_predicted"] for i in ids]
                computed = hard_metrics(actual, predicted)
                for metric in (*HARD, "confusion_matrix", "class_recall", "class_counts", "correct", "total"):
                    if metric in row:
                        close(computed[metric], row[metric], name+":"+metric)
                if name != "features":
                    columns = [f"p{i}" for i in range(7)] if name == "probability" else ["p_"+c for c in CLASSES]
                    p = np.asarray([[float(raw_case[i][c]) for c in columns] for i in ids])
                    for i, item in enumerate(ids):
                        if name == "probability":
                            if raw_case[item]["probability_argmax"] != CLASSES[p[i].argmax()]:
                                raise AssertionError("Probability argmax and saved probabilities disagree")
                            close(float(raw_case[item]["confidence"]), p[i].max(), "saved confidence")
                        elif predicted[i] != CLASSES[p[i].argmax()]:
                            raise AssertionError("Fusion decision is not its probability argmax")
                    proper = proper_metrics(p, actual)
                    for metric in ("nll", "brier", "ece"):
                        close(proper[metric], row["ece10" if name == "probability" and metric == "ece" else metric], metric)
                    computed.update(proper)
                    computed["decision_probability_disagreements"] = int(np.sum(np.asarray(predicted) != np.asarray(CLASSES)[p.argmax(1)]))
                role = row.get("method", "candidate_"+row["candidate_id"])
                if scope == "final":
                    final_records[(name, role, int(row["seed"]))] = {**row, **computed, "_raw": raw_case}
                else:
                    dev_groups[(name, row["candidate_id"], row["stage"])].append({**row, **computed})

    # Add predeclared probability roles as aliases, never as new raw model cases.
    for provider, choice in locks["probability"]["choices"].items():
        roles = {"raw": choice["raw"], "classification_selected": choice["selected"],
            "probability_selected": choice["probability_selected"], "diagnostic": choice["diagnostic"],
            "temperature_control": next(c for c in grids["probability"].values() if c["provider"] == provider and c["kind"] == "temperature")}
        for label, candidate in roles.items():
            for seed in (42, 43, 44):
                old = final_records[("probability", "candidate_"+candidate["id"], seed)]
                final_records[("probability", provider+"_"+label, seed)] = {**old, "method": provider+"_"+label, "_alias": True}
        for seed in (42, 43, 44):
            raw = final_records[("probability", provider+"_raw", seed)]["_raw"]
            temp = final_records[("probability", provider+"_temperature_control", seed)]["_raw"]
            if any(raw[i]["y_predicted"] != temp[i]["y_predicted"] or raw[i]["probability_argmax"] != temp[i]["probability_argmax"] for i in final_ids):
                raise AssertionError("Positive scalar temperature changed classification argmax")

    final_groups = defaultdict(list)
    for (name, method, seed), row in final_records.items():
        final_groups[(name, method)].append(row)
    summary_rows, seed_rows, class_rows, random_rows, sample_rows = [], [], [], [], []

    def identity(name, method):
        if name == "features":
            return "训练内部正式选择" if method == "selected" else "原谱参照" if method == "reference" else "开发预选家族诊断，不按最终分数重选"
        if name == "probability":
            return "开发锁定分类规则" if method.endswith("classification_selected") else "开发锁定概率规则" if method.endswith("probability_selected") else "预定义角色别名，非新增拟合" if not method.startswith("candidate_") else "全部固定候选；身份另由开发锁标记"
        return "开发正式选择" if method == "selected" else "开发预选诊断" if method == "diagnostic" else "单模型控制" if not method.startswith("candidate_") else "全部固定融合候选"

    for (name, method), values in sorted(final_groups.items()):
        candidate = values[0]["candidate_id"]; means = {m: float(np.mean([v[m] for v in values])) for m in HARD}
        entry = {"实验轨道": NAMES[name], "方法角色": method, "候选编号": candidate, "方法身份": identity(name, method),
            "训练种子数": len(values), "每种子预测样品": 171, "目标支持标签数": 0,
            "候选参数": json.dumps(grids[name][candidate], ensure_ascii=False, sort_keys=True),
            "准确率": means["accuracy"], "平衡准确率": means["balanced_accuracy"], "宏F1": means["macro_f1"],
            "NLL": None if name == "features" else float(np.mean([v["nll"] for v in values])),
            "Brier": None if name == "features" else float(np.mean([v["brier"] for v in values])),
            "ECE10": None if name == "features" else float(np.mean([v["ece"] for v in values])),
            "概率argmax与判别不一致数": None if name == "features" else int(sum(v["decision_probability_disagreements"] for v in values))}
        summary_rows.append(entry)
        for row in values:
            seed = int(row["seed"])
            seed_rows.append({"实验轨道": NAMES[name], "方法角色": method, "候选编号": candidate, "训练种子": seed,
                "正确数": row["correct"], "总数": row["total"], "准确率": row["accuracy"], "平衡准确率": row["balanced_accuracy"], "宏F1": row["macro_f1"],
                "NLL": row.get("nll"), "Brier": row.get("brier"), "ECE10": row.get("ece")})
        for j, c in enumerate(CLASSES):
            class_rows.append({"实验轨道": NAMES[name], "方法角色": method, "候选编号": candidate, "匿名产地组": c,
                "独立固定预测样品数": int(values[0]["class_counts"][j]),
                "三个种子平均精确率": float(np.mean([v["class_precision"][j] for v in values])),
                "三个种子平均召回率": float(np.mean([v["class_recall"][j] for v in values])),
                "三个种子平均F1": float(np.mean([v["class_f1"][j] for v in values])),
                "三个种子平均预测为该类数量": float(np.mean([v["class_predicted_counts"][j] for v in values])),
                "三个种子平均TP": float(np.mean([v["class_true_positives"][j] for v in values])),
                "三个种子平均FP": float(np.mean([v["class_false_positives"][j] for v in values])),
                "三个种子平均FN": float(np.mean([v["class_false_negatives"][j] for v in values]))})
        close(np.mean([r["三个种子平均F1"] for r in class_rows[-7:]]), means["macro_f1"], "classF1 macro aggregation")
        for row in values:
            close(sum(row["class_true_positives"]), row["correct"], "classTP correct count")
            close(sum(row["class_false_positives"]), row["total"]-row["correct"], "classFP wrong count")
            close(sum(row["class_false_negatives"]), row["total"]-row["correct"], "classFN wrong count")
        for identifier in sorted(final_ids):
            predicted_counts = Counter(v["_raw"][identifier]["y_predicted"] for v in values)
            sample_rows.append({"实验轨道": NAMES[name], "方法角色": method, "候选编号": candidate,
                "样品编号": identifier, "真实匿名产地": truth[identifier], "训练种子数": len(values),
                "正确种子数": predicted_counts[truth[identifier]], "跨种子正确率": predicted_counts[truth[identifier]]/len(values),
                "各种预测计数": json.dumps(dict(predicted_counts), ensure_ascii=False, sort_keys=True)})
            if identifier in random_ids:
                for row in values:
                    raw = row["_raw"][identifier]
                    random_rows.append({"实验轨道": NAMES[name], "方法角色": method, "候选编号": candidate,
                        "训练种子": int(row["seed"]), "样品编号": identifier, "真实匿名产地": truth[identifier],
                        "判别预测产地": raw["y_predicted"], "是否正确": int(raw["is_correct"]),
                        "概率argmax产地": raw.get("probability_argmax", raw["y_predicted"] if name == "fusion" else None),
                        "最大概率置信度": float(raw["confidence"]) if "confidence" in raw else max(float(raw["p_"+c]) for c in CLASSES) if name == "fusion" else None,
                        "说明": "固定20编号，保留错误；角色别名可能引用同一原始预测，非新增样品"})
    # Cross-check every producer summary against independent raw predictions.
    for name, run in runs.items():
        for row in read_json(run/"summary.json"):
            role = row["method"] if name != "probability" else "candidate_"+row["id"]
            values = final_groups[(name, role)]
            for metric in (*HARD, "nll", "brier", "ece10" if name == "probability" else "ece"):
                if metric+"_mean" in row:
                    key = "ece" if metric == "ece10" else metric
                    close(float(np.mean([v[key] for v in values])), row[metric+"_mean"], "summary:"+metric)

    development_rows = []
    for (name, candidate, stage), values in sorted(dev_groups.items()):
        development_rows.append({"实验轨道": NAMES[name], "候选编号": candidate,
            "开发阶段": "前4任务选择" if stage == "selection" else "后2任务检验", "实际评价数": len(values),
            "候选参数": json.dumps(grids[name][candidate], ensure_ascii=False, sort_keys=True),
            "准确率": float(np.mean([v["accuracy"] for v in values])), "平衡准确率": float(np.mean([v["balanced_accuracy"] for v in values])),
            "宏F1": float(np.mean([v["macro_f1"] for v in values])),
            "NLL": None if name == "features" else float(np.mean([v["nll"] for v in values])),
            "Brier": None if name == "features" else float(np.mean([v["brier"] for v in values])),
            "ECE10": None if name == "features" else float(np.mean([v["ece"] for v in values]))})

    oof_rows, temperature_rows = [], []
    for scope in ("development", "final"):
        paths = read_json(runs["probability"]/(scope+"_contexts.json"))
        for relative in paths:
            directory = runs["probability"]/relative; meta = read_json(directory/"context.json")
            policy = meta["oof_policy"]; seen = set(); source = set(meta["source_ids"])
            if digest(directory/"scores.npz") != meta["scores_npz_sha256"]:
                raise AssertionError("Source OOF score cache bytes changed")
            with np.load(directory/"scores.npz", allow_pickle=False) as archive:
                data = {k: archive[k] for k in archive.files}
            if (not source <= source_ids or len(source) != (513 if scope == "final" else 307)
                    or set(data["source_ids"].tolist()) != source
                    or source & set(meta["prediction_ids"])):
                raise AssertionError("Source OOF source/query roles differ from the fixed independent specimens")
            for identifier, label in zip(data["source_ids"], data["source_y"]):
                if CLASSES[int(label)] != truth[str(identifier)]:
                    raise AssertionError("Source OOF training labels differ from the original anonymous groups")
            excluded = set(policy["excluded_oof_ids"])
            if set(data["oof_ids"].tolist()) != source-excluded:
                raise AssertionError("Source OOF records do not match singleton exclusion")
            for fold in meta["folds"]:
                train = set(fold["source_train_ids"]); valid = set(fold["source_validation_ids"])
                if train & valid or train | valid != source or seen & valid or not excluded <= train:
                    raise AssertionError("OOF physical records overlap or singleton not anchored")
                seen |= valid; checks += 4
                oof_rows.append({"阶段": scope, "开发任务": meta["episode"], "种子": meta["seed"], "折号": fold["fold"],
                    "源样品数": len(source), "折内训练样品数": len(train), "折外验证样品数": len(valid),
                    "总OOF样品数": len(data["oof_ids"]), "排除校准拟合样品数": len(excluded),
                    "排除样品编号": json.dumps(sorted(excluded), ensure_ascii=False),
                    "排除类别": json.dumps([CLASSES[int(c)] if isinstance(c, int) else str(c) for c in policy["excluded_oof_classes"]], ensure_ascii=False),
                    "目标支持标签数": 0, "说明": "开发ZY02只有1份，始终在折训练，无其OOF行；最终源3份可正常分折"})
            if seen != source-excluded:
                raise AssertionError("OOF validation record coverage incomplete")
            for provider in ("svm", "mlp", "mlp_sup100"):
                scores = data["oof_scores_"+provider]; labels = [CLASSES[i] for i in data["oof_y"]]
                temperature = meta["temperatures"][provider]
                values = {}
                for mode, t in [("raw_oof", 1.), ("temperature_oof", temperature)]:
                    z = scores/t; z = z-z.max(1, keepdims=True); p = np.exp(z); p /= p.sum(1, keepdims=True)
                    proper = proper_metrics(p, labels)
                    for metric in ("nll", "brier", "ece"):
                        close(proper[metric], meta["temperature_fits"][provider][mode]["ece10" if metric == "ece" else metric], "sourceOOF:"+metric)
                    values[mode] = proper
                if not meta["temperature_fits"][provider]["argmax_unchanged"]:
                    raise AssertionError("Source temperature violates argmax invariant")
                temperature_rows.append({"阶段": scope, "开发任务": meta["episode"], "种子": meta["seed"], "源模型": provider,
                    "源OOF样品数": len(labels), "温度T": temperature, "原NLL": values["raw_oof"]["nll"],
                    "温度后NLL": values["temperature_oof"]["nll"], "原Brier": values["raw_oof"]["brier"],
                    "温度后Brier": values["temperature_oof"]["brier"], "原ECE10": values["raw_oof"]["ece"],
                    "温度后ECE10": values["temperature_oof"]["ece"], "说明": "温度只拟合源OOF；保持argmax，不等于目标小样本校准"})

    frequency_rows, stability_rows = [], []
    for name, candidate in grids["features"].items():
        sets = []; frequencies = {}
        contexts = set()
        for path in sorted((runs["features"]/"development").glob("episode_*/seed_*/"+name+".json")):
            meta = read_json(path); fit = meta["fit"]
            selected = fit["selected_indices"]; freq = fit["selected_frequencies"]
            if len(selected) != len(set(selected)) or len(selected) != len(freq) or any(i < 0 or i >= 612 for i in selected):
                raise AssertionError("Invalid selected frequency indices")
            close(freq, expected_axis[selected], "selected frequency coordinates")
            contexts.add((int(path.parent.parent.name.split("_")[-1]), int(path.parent.name.split("_")[-1])))
            sets.append(set(selected)); frequencies.update(zip(selected, freq)); checks += 1
        if len(sets) != 18 or contexts != {(e, s) for e in range(6) for s in (42, 43, 44)}:
            raise AssertionError("Feature stability requires all18 source-only contexts")
        all_jaccard = [len(a & b)/len(a | b) for i, a in enumerate(sets) for b in sets[i+1:]]
        count = Counter(i for selected in sets for i in selected)
        for index in sorted(count):
            stability_rows.append({"候选编号": name, "特征家族": candidate["family"], "预处理": candidate["preprocessing"],
                "原612频点索引0起": index, "原频率坐标": frequencies[index], "开发上下文数": 18,
                "被选择上下文数": count[index], "选择比例": count[index]/18, "平均Jaccard": float(np.mean(all_jaccard)),
                "说明": "重叠开发任务／算法种子的选择稳定性；全谱始终为1；不作化学峰归属"})
    for (name, method, seed), row in final_records.items():
        if name != "features":
            continue
        fit = read_json(runs[name]/"final"/f"seed_{seed}"/(row["candidate_id"]+".json"))["fit"]
        for index, freq in zip(fit["selected_indices"], fit["selected_frequencies"]):
            frequency_rows.append({"方法角色": method, "候选编号": row["candidate_id"], "训练种子": seed,
                "特征家族": grids[name][row["candidate_id"]]["family"], "原612频点索引0起": index,
                "原频率坐标": freq, "输入保留频点数": fit["selected_features"], "分类器输入维数": fit["classifier_features"],
                "说明": "所选模型变量；PCA-LDA的原频点列表不是主成分归属或化学机理证明"})

    complement_rows = []
    saved_complement = {(int(r["seed"]), r["mlp_provider"]): r for r in read_json(runs["fusion"]/"final_complementarity.json")}
    for seed in (42, 43, 44):
        a = final_records[("fusion", "svm_raw_softmax", seed)]["_raw"]
        for provider in ("mlp", "mlp_sup100"):
            b = final_records[("fusion", provider+"_raw_softmax", seed)]["_raw"]
            first = np.asarray([a[i]["y_predicted"] for i in sorted(final_ids)])
            second = np.asarray([b[i]["y_predicted"] for i in sorted(final_ids)])
            labels = np.asarray([truth[i] for i in sorted(final_ids)])
            ga, gb = first == labels, second == labels
            computed = {"n": 171, "both_correct": int(np.sum(ga & gb)), "svm_only_correct": int(np.sum(ga & ~gb)),
                "mlp_only_correct": int(np.sum(~ga & gb)), "both_wrong": int(np.sum(~ga & ~gb)),
                "different_predictions": int(np.sum(first != second)),
                "both_wrong_same_prediction": int(np.sum(~ga & ~gb & (first == second))),
                "oracle_accuracy_upper_bound": float(np.mean(ga | gb))}
            for metric, value in computed.items():
                close(value, saved_complement[(seed, provider)][metric], "complementarity:"+metric)
            complement_rows.append({"训练种子": seed, "MLP来源": provider, "独立预测样品数": 171,
                "两者都正确": computed["both_correct"], "仅SVM正确": computed["svm_only_correct"],
                "仅MLP正确": computed["mlp_only_correct"], "两者都错误": computed["both_wrong"],
                "预测不同样品数": computed["different_predictions"], "两者同样错": computed["both_wrong_same_prediction"],
                "真值Oracle准确率上界": computed["oracle_accuracy_upper_bound"],
                "说明": "Oracle需事先知道真值，不能用于实际预测／模型选择"})

    ledger = []
    for name in TRACKS:
        roles = [r for r in summary_rows if r["实验轨道"] == NAMES[name] and
                 (r["方法角色"] == "selected" or name == "probability" and r["方法角色"].endswith("classification_selected"))]
        for row in roles:
            ledger.append({"实验轨道": NAMES[name], "主候选设置数": len(grids[name]), "正式方法角色": row["方法角色"],
                "候选编号": row["候选编号"], "准确率": row["准确率"], "平衡准确率": row["平衡准确率"], "宏F1": row["宏F1"],
                "NLL": row["NLL"], "Brier": row["Brier"], "ECE10": row["ECE10"], "目标支持标签数": 0,
                "说明": "依开发锁决定；概率轨按三个源模型分别选择；不以最终最高分替换正式方法"})
    info = [{"事项": k, "说明": v} for k, v in [
        ("数据", "真实上传855份匿名光谱；固定513训练/171保留目标校准池/171预测；本轮目标校准池标签使用0"),
        ("当前范围", "0.1—1.5原频率坐标内612个吸收系数特征，尚无物理最优性证明"),
        ("三轨格点", "40特征+21概率判别+20融合=81主设置；角色别名和控制不代表新增模型或独立样品"),
        ("指标", "以宏F1开发选择，并约束准确率/平衡准确率；保留每个未通过和负结果"),
        ("逐类指标", "精确率/召回率/F1及预测数、TP/FP/FN分别先按每种子171份计算再取三种子均值；独立support数量不乘3"),
        ("预测次数", "每个最终源种子只预测固定171份一次；无目标支持重抽样，不添加100draw计数"),
        ("标签", "ZY01—ZY07为用户允许匿名组；一致性核对不等于外部核实真实地名/物种"),
        ("概率参照", "SVM原decision score不是概率；raw-softmax是未校准归一化参照"),
        ("温度", "温度仅源OOF拟合，正温度保持argmax；改善NLL不等于提高硬分类准确率"),
        ("判别偏置", "类别先验/坐标偏置用于决策，非posterior校准；概率argmax与硬决策分列"),
        ("proper指标", "NLL为均值负log p_true(1e-15下限)；Brier按七类平方误差求和再按样品平均，不除以7"),
        ("ECE", "10个固定等宽最大置信度桶，confidence1含于末桶；基于概率argmax正确性，不用偏置硬决策"),
        ("OOF限制", "开发307源中ZY02仅1，始终在折训练且排除其OOF校准行；其他306有真实heldout分数"),
        ("最终源OOF", "最终513源含ZY02三份，可普通三折，所有513具有heldout分数"),
        ("筛选", "CARS/MC-UVE/噪声UVE/单链SPA是明确说明的Python分类适配，不声称原MATLAB逐项复现"),
        ("特征频率", "索引与频率来自源模型元数据；稳定性不等于化学峰归属或独立批次稳定性"),
        ("LDA", "本轮LDA是分类器，未声称其提供低维transform；PCA-LDA先源PCA再分类"),
        ("融合", "按同一样品和同一类顺序融合概率；权重不由最终真值选择"),
        ("Oracle", "错误互补性Oracle使用真实标签，只是诊断上界，不能部署或当准确率主结果"),
        ("时序", "三轨开发锁齐并统一放行后才最终评价；本导出核查锁/协议/manifest及master/input/split SHA"),
        ("局限", "每产地一批、全部数据已有历史研究曝光；同批回顾性，三个种子不是独立批次"),
        ("错误", "随机20编号固定seed20261102；所有实际方法及角色错误保留，红色标记"),
    ]]
    sheets = {"说明": info, "正式结果台账": ledger, "所有最终方法": summary_rows,
        "候选开发汇总": development_rows, "全部种子": seed_rows, "逐产地召回": class_rows,
        "随机样品": random_rows, "逐样品统计": sample_rows, "源OOF分折": oof_rows,
        "温度拟合诊断": temperature_rows, "特征频点列表": frequency_rows, "特征选择稳定性": stability_rows,
        "融合错误互补": complement_rows,
        "概率可靠性指标": [r for r in summary_rows if r["NLL"] is not None]}
    output.mkdir(parents=True, exist_ok=True)
    book = Workbook(); book.remove(book.active)
    for name, rows in sheets.items():
        add_sheet(book, name, rows); write_csv(output/(name+".csv"), rows)
    workbook = output/"图片建议优化_完整结果与样品预测对照.xlsx"
    book.save(workbook)
    readback = load_workbook(workbook, read_only=True, data_only=True)
    for name, rows in sheets.items():
        if readback[name].max_row != len(rows)+1:
            raise AssertionError("Workbook readback row count differs")
        keys = list(rows[0])
        for saved, expected in zip(readback[name].iter_rows(min_row=2, values_only=True), rows):
            for key, actual in zip(keys, saved):
                value = expected.get(key)
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    close(actual, value, "workbook readback:"+name+":"+key)
                elif actual != value and not (actual is None and value == ""):
                    raise AssertionError("Workbook text/empty value changed on readback")
                else:
                    checks += 1
    readback.close()
    validation = {"status": "pass", "new_model_fits": 0, "hyperparameter_choices_from_final": 0,
        "global_gate_sha_checks": gate_checks, "checks_evaluated": checks,
        "maximum_numeric_difference": maximum,
        "main_configurations": {n: len(grids[n]) for n in TRACKS},
        "actual_run_directories": {n: str(runs[n]) for n in TRACKS},
        "raw_prediction_rows": {n+":"+s: count for (n, s), count in raw_counts.items()},
        "unique_raw_prediction_cases": len(cases), "final_role_rows_include_aliases": True,
        "sheet_data_rows": {name: len(rows) for name, rows in sheets.items()},
        "random_incorrect_rows_preserved": sum(r["是否正确"] == 0 for r in random_rows),
        "input_provenance": provenance,
        "output_files_sha256": {p.name: digest(p) for p in sorted(output.iterdir()) if p.is_file() and p.name != "table_export_validation.json"}}
    (output/"table_export_validation.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    print(json.dumps({"status": "pass", "workbook": str(workbook), "checks": checks,
        "maximum_numeric_difference": maximum, "sheets": len(sheets)}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-root", type=Path, default=Path("outputs/image-guided-20261031"))
    parser.add_argument("--output", type=Path, default=Path("outputs/image-guided-20261031/tables"))
    parser.add_argument("--split-plan", type=Path, default=Path("data/processed/zaoren_threeway_20261031_v1/split_plan.json"))
    parser.add_argument("--input", type=Path, default=Path("outputs/approved-zaoren-20261031/data-preparation/dataset.npz"))
    args = parser.parse_args()
    export(args.run_root, args.output, args.split_plan, args.input)
