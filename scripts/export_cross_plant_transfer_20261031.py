#!/usr/bin/env python3
"""Read-only, explicitly released exports for the paired jujube transfer run.

This script does not import the training module, load fitted models, fit a
transformer or call benchmark scoring. It verifies completed locks, independently
recalculates metrics from already-scored predictions and exports every fixed seed.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
from datetime import datetime, timezone
import gzip
import hashlib
import json
from pathlib import Path

import numpy as np
from sklearn.metrics import accuracy_score, balanced_accuracy_score, confusion_matrix, f1_score


SEEDS = (42, 43, 44)
ARMS = ("A", "Aextra", "B", "C")
ROLES = (*ARMS, "source_policy")
CLASSES = tuple(f"ZY{i:02d}" for i in range(1, 8))
COUNTS = dict(zip(CLASSES, (350, 5, 100, 100, 100, 100, 100)))
METRICS = ("accuracy", "macro_f1", "balanced_accuracy")
LABELS = {"A": "A scratch", "Aextra": "Aextra target-only extra steps",
          "B": "B ginseng pretraining", "C": "C multi-plant pretraining",
          "source_policy": "Source-selected guarded policy"}
HISTORICAL = {"role": "historical_MCUVE_SVM_R2", "seed": 42, "correct": 821,
              "total": 855, "accuracy": 821 / 855, "macro_f1": None,
              "balanced_accuracy": None, "source": "docs/RESEARCH4_EXPERIMENT_AUDIT.md",
              "interpretation": "Immutable exposed-development context only; not the matched transfer baseline"}


class ExportError(ValueError):
    pass


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def sha(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def safe_member(root, value):
    root = Path(root).resolve()
    path = (root / str(value)).resolve()
    if not path.is_relative_to(root):
        raise ExportError("Source evidence path escapes its run directory")
    return path


def write_json(path, value):
    with Path(path).open("x", encoding="utf-8") as handle:
        json.dump(value, handle, indent=2, ensure_ascii=False, allow_nan=False)
        handle.write("\n")


def write_csv(path, rows):
    if not rows:
        return
    keys = list(dict.fromkeys(key for row in rows for key in row))
    with Path(path).open("x", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=keys)
        writer.writeheader()
        for row in rows:
            writer.writerow({k: json.dumps(v, ensure_ascii=False, sort_keys=True)
                             if isinstance(v, (dict, list, tuple)) else v for k, v in row.items()})


def verify_release(root, release):
    if release.get("allow_aggregate_export") is not True or not release.get("created_at_utc"):
        raise ExportError("Dated root post-score export release required")
    path = Path(root) / "scoring_manifest.json"
    if not path.exists():
        raise ExportError("Scoring is incomplete; export must not train or score")
    if release.get("scoring_manifest_sha256") != sha(path):
        raise ExportError("Root export release is not bound to this completed scoring manifest")


def _hashes(root, record, checks):
    for key, name in checks.items():
        if record.get(key) != sha(Path(root) / name):
            raise ExportError("Frozen export dependency changed: " + name)


def verify_scored(root):
    """Validate all score/prediction/source locks before semantically reading outcomes."""
    root = Path(root)
    scored = read(root / "scoring_manifest.json")
    if (not scored.get("created_at_utc") or scored.get("new_training_fits") != 0
            or scored.get("real_target_observations") != 855
            or scored.get("three_seeds_are_not_new_specimens") is not True):
        raise ExportError("Complete 855-observation, zero-fit scoring manifest required")
    _hashes(root, scored, {"prediction_lock_sha256": "prediction_lock.json",
        "prediction_sha256": "outer_predictions.csv.gz", "metrics_sha256": "pooled_metrics_by_seed.json",
        "summary_sha256": "seed_summary.json", "paired_deltas_sha256": "paired_scratch_deltas.json",
        "scored_predictions_sha256": "scored_outer_predictions.csv.gz"})
    lock = read(root / "prediction_lock.json")
    if (lock.get("outer_truth_parsed") is not False or lock.get("actual_outer_models") != 60
            or lock.get("alias_cases") != 15 or lock.get("prediction_rows") != 12825):
        raise ExportError("Complete pre-truth prediction lock required")
    _hashes(root, lock, {"prediction_sha256": "outer_predictions.csv.gz",
        "outer_model_records_sha256": "outer_model_records.json", "outer_role_aliases_sha256": "outer_role_aliases.json",
        "selection_lock_sha256": "selection_lock.json", "outer_release_sha256": "outer_release_snapshot.json",
        "run_manifest_sha256": "run_manifest.json", "protocol_sha256": "protocol.json"})
    selection = read(root / "selection_lock.json")
    if selection.get("source_pipeline_fits") != 240 or selection.get("outer_used_for_selection") is not False:
        raise ExportError("All 240 source pipelines must precede outer use")
    _hashes(root, selection, {"source_release_sha256": "source_release_snapshot.json",
        "run_manifest_sha256": "run_manifest.json", "source_records_sha256": "source_records.json",
        "source_predictions_sha256": "source_predictions.csv.gz", "candidate_summaries_sha256": "candidate_summaries.json"})
    protocol = read(root / "protocol.json")
    if (tuple(protocol.get("seeds", ())) != SEEDS or tuple(protocol.get("arms", ())) != ARMS
            or tuple(protocol.get("roles", ())) != ROLES or tuple(protocol.get("classes", ())) != CLASSES
            or protocol.get("source_pipeline_budget") != 240 or protocol.get("outer_pipeline_budget") != 60):
        raise ExportError("Frozen all-arm, all-seed protocol differs")
    manifest = read(root / "run_manifest.json")
    _hashes(root, manifest, {"protocol_sha256": "protocol.json", "preparation_sha256": "preparation.json"})
    source_release, outer_release = read(root / "source_release_snapshot.json"), read(root / "outer_release_snapshot.json")
    if source_release.get("allow_source_training") is not True or outer_release.get("allow_outer_evaluation") is not True:
        raise ExportError("Explicit upstream stage releases missing")
    for key in ("protocol_sha256", "preparation_sha256", "source_sha256"):
        if source_release.get(key) != manifest.get(key):
            raise ExportError("Source release and frozen run disagree")
    _hashes(root, outer_release, {"selection_lock_sha256": "selection_lock.json",
        "protocol_sha256": "protocol.json", "run_manifest_sha256": "run_manifest.json"})
    return scored, protocol, selection


def prediction_groups(root):
    groups = defaultdict(list)
    with gzip.open(Path(root) / "scored_outer_predictions.csv.gz", "rt", encoding="utf-8-sig") as handle:
        for row in csv.DictReader(handle):
            try:
                key = row["role"], int(row["seed"])
            except (KeyError, ValueError) as exc:
                raise ExportError("Missing role/seed in scored predictions") from exc
            groups[key].append(row)
    if set(groups) != {(role, seed) for role in ROLES for seed in SEEDS}:
        raise ExportError("All five roles and all three fixed seeds required")
    shared_truth = None
    for key, rows in groups.items():
        if len(rows) != 855 or len({r["record_id"] for r in rows}) != 855:
            raise ExportError("Each role/seed requires exactly855 unique observations")
        truth = {r["record_id"]: r["y_true"] for r in rows}
        if Counter(truth.values()) != COUNTS or any(r["y_predicted"] not in CLASSES for r in rows):
            raise ExportError("Canonical seven-class counts/vocabulary changed")
        if shared_truth is not None and truth != shared_truth:
            raise ExportError("Paired role/seed IDs or true labels differ")
        shared_truth = truth
    return groups


def classification_rows(root, groups=None):
    groups = prediction_groups(root) if groups is None else groups
    stored = read(Path(root) / "pooled_metrics_by_seed.json")
    lookup = {(r["role"], int(r["seed"])): r["metrics"] for r in stored}
    if len(stored) != 15 or set(lookup) != set(groups):
        raise ExportError("Pooled model/seed inventory changed")
    rows, confusions, per_class = [], [], []
    for role in ROLES:
        for seed in SEEDS:
            current = groups[role, seed]
            y, p = [r["y_true"] for r in current], [r["y_predicted"] for r in current]
            matrix = confusion_matrix(y, p, labels=CLASSES)
            calc = {"accuracy": accuracy_score(y, p), "macro_f1": f1_score(y, p, labels=CLASSES, average="macro", zero_division=0),
                    "balanced_accuracy": balanced_accuracy_score(y, p), "correct": int(matrix.trace()), "total": 855}
            m = lookup[role, seed]
            if (m.get("class_order") != list(CLASSES) or not np.array_equal(m.get("confusion_matrix"), matrix)
                    or any(not np.isfinite(float(m[k])) or abs(float(m[k]) - v) > 1e-12 for k, v in calc.items())):
                raise ExportError("Independent prediction recalculation disagrees with stored metrics")
            rows.append({"role": role, "seed": seed, **calc})
            for i, true in enumerate(CLASSES):
                tp, actual, found = int(matrix[i, i]), int(matrix[i].sum()), int(matrix[:, i].sum())
                pr = {"role": role, "seed": seed, "class": true, "support": actual,
                      "precision": tp / found if found else 0., "recall": tp / actual,
                      "f1": 2 * tp / (actual + found) if actual + found else 0.}
                if "per_class" in m and any(not np.isfinite(m["per_class"][true][k]) or abs(m["per_class"][true][k] - pr[k]) > 1e-12 for k in ("precision", "recall", "f1", "support")):
                    raise ExportError("Stored per-class metric disagrees")
                per_class.append(pr)
                for j, pred in enumerate(CLASSES):
                    confusions.append({"role": role, "seed": seed, "true_class": true, "predicted_class": pred,
                                       "count": int(matrix[i, j])})
    return rows, confusions, per_class


def paired_rows(root, metrics, groups):
    lookup = {(r["role"], r["seed"]): r for r in metrics}
    rows = []
    for baseline in ("A", "Aextra"):
        for role in ROLES:
            if role == baseline:
                continue
            for seed in SEEDS:
                old = {r["record_id"]: r for r in groups[baseline, seed]}
                new = groups[role, seed]
                corrected = sum(old[r["record_id"]]["y_predicted"] != r["y_true"] and r["y_predicted"] == r["y_true"] for r in new)
                broken = sum(old[r["record_id"]]["y_predicted"] == r["y_true"] and r["y_predicted"] != r["y_true"] for r in new)
                a, b = lookup[baseline, seed], lookup[role, seed]
                deltas = {m + "_delta": b[m] - a[m] for m in METRICS}
                if abs(deltas["accuracy_delta"] - (corrected - broken) / 855) > 1e-12:
                    raise ExportError("Corrected/broken count does not reconcile to accuracy change")
                rows.append({"baseline": baseline, "role": role, "seed": seed,
                    "corrected": corrected, "broken": broken, "net_correct": corrected - broken,
                    **deltas, "accuracy_change_percentage_points": 100 * deltas["accuracy_delta"],
                    "accuracy_relative_change_percent": 100 * deltas["accuracy_delta"] / a["accuracy"] if a["accuracy"] else None,
                    "denominator": 855, "historical_96_used_for_gain": False})
    upstream = {(r["role"], int(r["seed"])): r for r in read(Path(root) / "paired_scratch_deltas.json")}
    expected = [r for r in rows if r["baseline"] == "A"]
    if set(upstream) != {(r["role"], r["seed"]) for r in expected}:
        raise ExportError("Upstream paired-scratch inventory mismatch")
    for r in expected:
        old = upstream[r["role"], r["seed"]]
        if any(not np.isfinite(float(old[k])) or abs(float(old[k]) - r[k]) > 1e-12 for k in ("corrected", "broken", *(m + "_delta" for m in METRICS))):
            raise ExportError("Stored paired-scratch contrast disagrees with actual predictions")
    return rows


def fold_rows(root, groups):
    """Recalculate the unbound fold table rather than trust its numeric contents."""
    stored = read(Path(root) / "fold_metrics.json")
    lookup = {(r["role"], int(r["seed"]), int(r["outer_fold"])): r["metrics"] for r in stored}
    expected = {(role, seed, fold) for role in ROLES for seed in SEEDS for fold in range(5)}
    if len(stored) != 75 or set(lookup) != expected:
        raise ExportError("All75 fold metric cases required")
    result, shared_folds = [], None
    for role in ROLES:
        for seed in SEEDS:
            rows = groups[role, seed]
            mapping = {r["record_id"]: int(r["outer_fold"]) for r in rows}
            if shared_folds is not None and mapping != shared_folds:
                raise ExportError("Paired record-to-outer-fold mapping differs")
            shared_folds = mapping
            for fold in range(5):
                part = [r for r in rows if int(r["outer_fold"]) == fold]
                if len(part) != 171:
                    raise ExportError("Each outer fold must have171 predictions")
                y, p = [r["y_true"] for r in part], [r["y_predicted"] for r in part]
                values = {"accuracy": accuracy_score(y, p), "macro_f1": f1_score(y, p, labels=CLASSES, average="macro", zero_division=0),
                          "balanced_accuracy": balanced_accuracy_score(y, p), "correct": sum(a == b for a, b in zip(y, p)), "total": 171}
                old = lookup[role, seed, fold]
                if any(not np.isfinite(old[k]) or abs(old[k] - value) > 1e-12 for k, value in values.items()):
                    raise ExportError("Stored fold metrics disagree with actual predictions")
                result.append({"role": role, "seed": seed, "outer_fold": fold, **values})
    return result


def verify_seed_summary(root, summary):
    old = read(Path(root) / "seed_summary.json")
    lookup = {r["role"]: r for r in old}
    if len(old) != 5 or set(lookup) != set(ROLES):
        raise ExportError("All5 seed summaries required")
    for row in summary:
        stored = lookup[row["role"]]
        if stored.get("observations_per_seed") != 855 or tuple(stored.get("seeds", ())) != SEEDS:
            raise ExportError("Seed summary denominator/seed inventory differs")
        for metric in METRICS:
            for suffix in ("_mean", "_sample_sd"):
                key = metric + suffix
                if not np.isfinite(stored[key]) or abs(stored[key] - row[key]) > 1e-12:
                    raise ExportError("Seed summary disagrees with paired prediction metrics")


def summarize(rows, metrics=METRICS, grouping=("role",)):
    groups = defaultdict(list)
    for row in rows:
        groups[tuple(row[k] for k in grouping)].append(row)
    result = []
    for key, current in groups.items():
        if len(current) != 3 or {r["seed"] for r in current} != set(SEEDS):
            raise ExportError("All fixed seeds must enter every summary")
        row = dict(zip(grouping, key))
        row.update(seeds="42/43/44", observations_each_seed=855)
        for metric in metrics:
            values = [r[metric] for r in current]
            if any(v is None for v in values):
                row[metric + "_mean"] = row[metric + "_sample_sd"] = None
            elif not np.isfinite(values).all():
                raise ExportError("Nonfinite metrics cannot be silently exported")
            else:
                row[metric + "_mean"], row[metric + "_sample_sd"] = float(np.mean(values)), float(np.std(values, ddof=1))
        if "correct" in current[0]:
            row["correct_each_seed"] = "; ".join(f"{r['seed']}:{r['correct']}/855" for r in sorted(current, key=lambda r: r["seed"]))
        result.append(row)
    return result


def paired_predictions(groups):
    result = []
    for seed in SEEDS:
        maps = {role: {r["record_id"]: r for r in groups[role, seed]} for role in ROLES}
        for rid, current in maps["A"].items():
            row = {"seed": seed, "record_id": rid, "true_label": current["y_true"]}
            for role in ROLES:
                r = maps[role][rid]
                row[role + "_prediction"] = r["y_predicted"]
                row[role + "_correct"] = int(r["y_predicted"] == r["y_true"])
            result.append(row)
    return result


def model_metadata_rows(root, protocol, selection):
    costs, losses, pools, pair_groups = [], [], [], defaultdict(list)
    for phase, name, expected in (("source", "source_records.json", 240), ("outer", "outer_model_records.json", 60)):
        records = read(Path(root) / name)
        if len(records) != expected:
            raise ExportError("Actual model record inventory incomplete")
        for record in records:
            directory = safe_member(root, record["model_directory"])
            training = directory / "training.json"
            if sha(training) != record["training_sha256"]:
                raise ExportError("Training metadata changed after score")
            m = read(training)
            if sha(directory / "pipeline.joblib") != record["pipeline_sha256"]:
                raise ExportError("Stored pipeline bytes changed after score")
            for field, filename in (("batch_trace_sha256", "batch_trace.jsonl.gz"), ("batch_schedule_sha256", "batch_schedule.json")):
                if sha(directory / filename) != m[field]:
                    raise ExportError("Actual optimizer trace/schedule bytes changed after score")
            profile = m["profile"]
            if m["arm"] != record["arm"] or m["seed"] != record["seed"] or profile["id"] != record["profile"]:
                raise ExportError("Model record and training metadata disagree")
            source_count = sum(m["source_class_counts"].values())
            if source_count != (342 if phase == "source" else 684):
                raise ExportError("Training source observation count changed")
            key = phase, record["split_id"], record["profile"], record["seed"]
            pair_groups[key].append(m)
            common = {"phase": phase, "outer_fold": record["outer_fold"], "inner_fold": record["inner_fold"],
                      "arm": record["arm"], "profile": profile["id"], "seed": record["seed"]}
            costs.append({**common, "source_observations": source_count, "target_epochs": profile["epochs"],
                "target_steps": m["target_steps"], "pretrain_steps": m["pretrain_steps"],
                "pretrain_seconds": m["pretrain_seconds"], "target_training_seconds": m["target_training_seconds"],
                "elapsed_seconds": m["elapsed_seconds"], "target_model_parameters": m["target_model_parameters"],
                "pretrain_pool_steps": m["pretrain_pool_steps"], "cache_reuse": m["pretrain_cache_reuse"]})
            for name, count in m["pretrain_pool_steps"].items():
                pools.append({**common, "pool": name, "steps": count, "batch_size": profile["pretrain_batch_size"],
                              "presentations": count * profile["pretrain_batch_size"], "presentations_are_not_independent_N": True})
            for epoch, loss in enumerate(m["target_epoch_loss"], 1):
                if not np.isfinite(loss):
                    raise ExportError("Nonfinite stored training loss")
                losses.append({**common, "stage": "target", "task_head": "ZY", "step_or_epoch": epoch, "loss": loss})
            tasks = ["target"] if m["arm"] == "Aextra" else ["ginseng"] if m["arm"] == "B" else protocol["C_task_order"]
            for step, loss in enumerate(m["pretrain_loss"]):
                losses.append({**common, "stage": "pretrain", "task_head": tasks[step % len(tasks)], "step_or_epoch": step + 1, "loss": loss})
    pair_checks = []
    for key, values in pair_groups.items():
        if {m["arm"] for m in values} != set(ARMS) or len(values) != 4:
            raise ExportError("Matched four-arm model group incomplete")
        for field in ("initial_encoder_sha256", "initial_zy_head_sha256", "fresh_target_optimizer_initial_sha256", "target_order_sha256", "BN_after_reset", "target_class_weights"):
            serialized = [json.dumps(m[field], sort_keys=True) for m in values]
            if len(set(serialized)) != 1:
                raise ExportError("Paired fine-tune state/order differs: " + field)
        pair_checks.append(dict(zip(("phase", "split_id", "profile", "seed"), key), matched_arms="A/Aextra/B/C",
                                same_encoder_initial=True, same_fresh_head=True, same_optimizer_initial=True,
                                same_target_order=True, same_reset_BN=True, same_class_weights=True))
    selected = [{"outer_fold": f["outer_fold"], "shared_profile": f["shared_profile"]["id"],
        "target_epochs_all_arms": f["shared_profile"]["epochs"], "policy_arm": f["policy_arm"],
        "profile_chosen_only_from_A": f["profile_chosen_only_from_A"], "outer_used_for_selection": f["outer_used_for_selection"],
        "source_means_at_shared_profile": f["source_means_at_shared_profile"]} for f in selection["outer_folds"]]
    return costs, losses, pools, pair_checks, selected


def computational_summary(costs):
    result = []
    groups = defaultdict(list)
    for row in costs:
        groups[row["phase"], row["arm"]].append(row)
    for (phase, arm), rows in groups.items():
        result.append({"phase": phase, "arm": arm, "actual_pipeline_records": len(rows),
            **{key + "_sum": sum(r[key] for r in rows) for key in ("pretrain_steps", "target_steps", "elapsed_seconds", "pretrain_seconds", "target_training_seconds")},
            "elapsed_seconds_mean": float(np.mean([r["elapsed_seconds"] for r in rows])),
            "cost_note": "Actual complete pipeline records; metadata elapsed includes hashing/normalization/I/O, not wall-clock or exact FLOPs"})
    return result


def spreadsheet(path, tables):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    book = Workbook()
    book.remove(book.active)
    for name, rows in tables.items():
        if not rows:
            continue
        sheet = book.create_sheet(name[:31])
        columns = list(dict.fromkeys(k for row in rows for k in row))
        sheet.append(columns)
        for row in rows:
            sheet.append([json.dumps(row.get(k), ensure_ascii=False, sort_keys=True)
                          if isinstance(row.get(k), (dict, list, tuple)) else row.get(k) for k in columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1D4E73")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(44, max(14, len(str(column[0].value)) + 2))
    book.save(path)


def figures(out, summary, paired_summary, confusions, losses, costs):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    order, lookup = list(ROLES), {r["role"]: r for r in summary}
    short = ("A", "Aextra", "B ginseng", "C multi", "Policy")
    fig, axes = plt.subplots(1, 3, figsize=(15, 4.8), constrained_layout=True)
    for ax, metric in zip(axes, METRICS):
        ax.bar(range(5), [lookup[r][metric + "_mean"] for r in order],
               yerr=[lookup[r][metric + "_sample_sd"] for r in order], capsize=4, color=("#4477AA", "#66CCEE", "#228833", "#CCBB44", "#AA3377"))
        ax.set_xticks(range(5), short, rotation=25)
        ax.set_ylim(0, 1.02)
        ax.set_title(metric)
        ax.grid(axis="y", alpha=.25)
    fig.suptitle("Jujube: matched ResNet transfer; all seeds42/43/44; N=855 per case\nError bars: seed sample SD, not confidence intervals")
    _save(fig, out, "matched_model_metrics", plt)
    fig, axes = plt.subplots(1, 3, figsize=(15, 4.8), constrained_layout=True)
    for ax, metric in zip(axes, METRICS):
        for j, baseline in enumerate(("A", "Aextra")):
            rows = [r for r in paired_summary if r["baseline"] == baseline and r["role"] in ("B", "C")]
            rows.sort(key=lambda r: r["role"])
            scale = 1 if metric == "macro_f1" else 100
            ax.bar(np.arange(2) + (j - .5) * .36, [scale * r[metric + "_delta_mean"] for r in rows], .34,
                   yerr=[scale * r[metric + "_delta_sample_sd"] for r in rows], capsize=4, label="vs " + baseline)
        ax.set_xticks(range(2), ("B ginseng", "C multi"))
        ax.axhline(0, color="black", linewidth=.8)
        ax.set_title(metric + (" delta" if metric == "macro_f1" else " delta (percentage points)"))
        ax.grid(axis="y", alpha=.25)
        ax.legend()
    fig.suptitle("Matched external-data contrasts; all fixed seeds; negative effects retained")
    _save(fig, out, "external_data_paired_gains", plt)
    fig, axes = plt.subplots(5, 3, figsize=(16, 23), constrained_layout=True)
    for i, role in enumerate(ROLES):
        for j, seed in enumerate(SEEDS):
            ax = axes[i, j]
            rows = [r for r in confusions if r["role"] == role and r["seed"] == seed]
            matrix = np.asarray([r["count"] for r in rows]).reshape(7, 7)
            ax.imshow(matrix, cmap="Blues", vmin=0, vmax=350)
            for a in range(7):
                for b in range(7):
                    ax.text(b, a, str(matrix[a, b]), ha="center", va="center", fontsize=8,
                            color="white" if matrix[a, b] > 175 else "black")
            ax.set_xticks(range(7), CLASSES, rotation=45, fontsize=8)
            ax.set_yticks(range(7), CLASSES, fontsize=8)
            ax.set_title(f"{role}, seed{seed}, N=855")
            ax.set_xlabel("Predicted")
            ax.set_ylabel("True")
    fig.suptitle("All pooled jujube confusion matrices: each record once per role/seed")
    _save(fig, out, "all_seed_confusions", plt)
    fig, axes = plt.subplots(2, 2, figsize=(13, 8), constrained_layout=True)
    for i, phase in enumerate(("source", "outer")):
        for j, profile in enumerate(("p48", "p96")):
            ax = axes[i, j]
            found = False
            for arm in ARMS:
                current = [r for r in losses if r["phase"] == phase and r["profile"] == profile and r["arm"] == arm and r["stage"] == "target"]
                points = defaultdict(list)
                for r in current:
                    points[r["step_or_epoch"]].append(r["loss"])
                if points:
                    found = True
                    ax.plot(sorted(points), [np.mean(points[e]) for e in sorted(points)], label=arm)
            ax.set_title(f"{phase} / {profile}")
            if found:
                ax.legend()
            else:
                ax.text(.5, .5, "No outer models selected at this profile", ha="center", va="center", transform=ax.transAxes)
            ax.set_xlabel("Epoch")
            ax.set_ylabel("Mean weighted minibatch training CE")
            ax.grid(alpha=.25)
    fig.suptitle("Actual target-source training losses; no outer loss used for checkpoint selection")
    _save(fig, out, "target_training_loss", plt)
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.5), constrained_layout=True)
    for ax, phase in zip(axes, ("source", "outer")):
        rows = {r["arm"]: r for r in costs if r["phase"] == phase}
        ax.bar(range(4), [rows[a]["elapsed_seconds_mean"] for a in ARMS])
        ax.set_xticks(range(4), ARMS)
        ax.set_title(phase)
        ax.set_ylabel("Mean recorded seconds / complete pipeline")
        ax.grid(axis="y", alpha=.25)
    fig.suptitle("Measured pipeline cost, including source normalization/hash/I/O; not exact FLOPs")
    _save(fig, out, "actual_pipeline_cost", plt)


def _save(fig, out, name, plt):
    fig.savefig(Path(out) / (name + ".png"), dpi=160, bbox_inches="tight")
    fig.savefig(Path(out) / (name + ".pdf"), bbox_inches="tight")
    plt.close(fig)


def report_text(summary, paired_summary, selected):
    text = ["# 外部植物预训练对酸枣仁识别的实际影响", "",
        "所有主要对照使用相同855条历史开发观测、五外折／二内折、128点α表示和小型ResNet1D；固定seed42／43／44全部报告。每个外折的共同48／96轮数只由从头训练A的source OOF选择，再用于所有组。外部标签使用独立任务头，不是ZY标签。", "",
        "Aextra只对当前酸枣仁source增加120个预训练步，用于检查收益是否只是额外计算。B使用人参42条；C使用黄芪403＋天麻405＋人参42共850条。预训练步骤和batch匹配，分类头不同，不能称精确FLOP匹配。", "",
        "单位／仪器／共同原料未核验，本轮是用户授权的工程形状迁移探索。三种子SD不是置信区间或新增样品；同批次、已曝光数据不能变成独立确认。水果、叶片、小麦及未下载的2026参未加入频率迁移。", "",
        "| 组 | Accuracy均值±SD | Macro-F1均值±SD | Balanced Accuracy | 每seed正确数 |",
        "|---|---:|---:|---:|---|"]
    for r in summary:
        text.append(f"| {LABELS[r['role']]} | {100*r['accuracy_mean']:.2f}% ± {100*r['accuracy_sample_sd']:.2f} | {r['macro_f1_mean']:.4f} ± {r['macro_f1_sample_sd']:.4f} | {100*r['balanced_accuracy_mean']:.2f}% | {r['correct_each_seed']} |")
    text += ["", "| 外部组与基线 | Accuracy变化（百分点） | 相对Accuracy变化率 | Macro-F1变化 | 改对／改错／净正确数（三种子均值） |",
             "|---|---:|---:|---:|---|"]
    for r in paired_summary:
        if r["role"] not in ("B", "C"):
            continue
        rel = r["accuracy_relative_change_percent_mean"]
        text.append(f"| {r['role']} − {r['baseline']} | {r['accuracy_change_percentage_points_mean']:+.2f} | {rel:+.2f}% | {r['macro_f1_delta_mean']:+.4f} | {r['corrected_mean']:.2f}／{r['broken_mean']:.2f}／{r['net_correct_mean']:+.2f} |" if rel is not None else f"| {r['role']} − {r['baseline']} | {r['accuracy_change_percentage_points_mean']:+.2f} | 分母为0，未定义 | {r['macro_f1_delta_mean']:+.4f} | {r['corrected_mean']:.2f}／{r['broken_mean']:.2f}／{r['net_correct_mean']:+.2f} |")
    text += ["", "## 对当前问题的直接回答", ""]
    paired_lookup = {(r["baseline"], r["role"]): r for r in paired_summary}
    for role in ("B", "C"):
        scratch, extra = paired_lookup["A", role], paired_lookup["Aextra", role]
        gain, adjusted = scratch["accuracy_change_percentage_points_mean"], extra["accuracy_change_percentage_points_mean"]
        answer = "观察到准确率均值提高" if gain > 1e-12 else "未观察到准确率均值提高"
        attribution = "也高于额外计算控制的观察均值" if adjusted > 1e-12 else "没有超过额外计算控制，不能把变化单独归因外部数据"
        text.append(f"- {role}：相对A为{gain:+.2f}个百分点，{answer}；相对Aextra为{adjusted:+.2f}个百分点，{attribution}。Macro-F1相对A变化{scratch['macro_f1_delta_mean']:+.4f}。")
    highest = max(summary, key=lambda r: r["accuracy_mean"])
    historical_answer = "观察均值高于" if highest["accuracy_mean"] > HISTORICAL["accuracy"] else "没有超过"
    text.append(f"- 本轮描述性最高组为{highest['role']}，Accuracy均值{100*highest['accuracy_mean']:.2f}%，{historical_answer}历史MC-UVE＋SVM的96.02%。这是开发成绩背景对照，不是迁移收益归因或新的确认结论。")
    text += ["", "百分点＝100×(新Accuracy−配对基线Accuracy)；相对变化率＝100×(新−基线)/基线，先逐种子配对计算，再汇总。负值全部保留。改对是原错误变正确，改错是原正确变错误；两者净差必须等于正确数变化。", "",
        "## 共同预算与保守策略", "",
        "| 外折 | 共同profile | 所有组目标epoch | source-policy |", "|---|---|---:|---|"]
    text += [f"| {r['outer_fold']} | {r['shared_profile']} | {r['target_epochs_all_arms']} | {r['policy_arm']} |" for r in selected]
    text += ["", "source-policy在训练内部按预先固定门槛选择或回退A，不按最终外层最好分数倒选。各单组行只作预声明组别比较，外层最高均值不自动变为确认赢家。", "",
        "## 与历史96.02%的关系", "",
        "历史MC-UVE＋SVM保持821/855＝96.0234%。本轮改变了表示、网络和训练预算，迁移收益的直接基线必须是配对A或Aextra，不能用新分数减96.02%归因外部数据。历史行单独列在historical_context.csv，不进入增益计算。", "",
        "## 文件与验证", "",
        "XLSX与CSV保存全部种子、逐类P/R/F1、混淆、逐样品配对预测、纠正／损害数、每fold／seed指标、共同选轮数、完整训练loss和实际成本。PNG/PDF为独立图件，负效应及零线均保留。", "",
        "导出只读取已经评分且SHA绑定的产物；不加载模型、不训练、不重新调用benchmark score。所有15个role/seed的指标及配对变化从实际预测独立复算，模型配对头／BN／优化器初始状态／目标批次顺序另查metadata。export_manifest.json绑定根放行及上游锁，files_sha256.json覆盖本目录其余产物。"]
    return "\n".join(text) + "\n"


def export(root, output, release_path, make_figures=True):
    root, output, release_path = Path(root), Path(output), Path(release_path)
    if output.exists():
        raise FileExistsError("Preserve earlier exports; choose a new output directory")
    release = read(release_path)
    verify_release(root, release)
    scored, protocol, selection = verify_scored(root)
    groups = prediction_groups(root)
    per_seed, confusions, per_class = classification_rows(root, groups)
    pairs = paired_rows(root, per_seed, groups)
    summary = summarize(per_seed)
    verify_seed_summary(root, summary)
    folds = fold_rows(root, groups)
    pair_summary = summarize(pairs, ("accuracy_delta", "macro_f1_delta", "balanced_accuracy_delta", "corrected", "broken", "net_correct", "accuracy_change_percentage_points", "accuracy_relative_change_percent"), ("baseline", "role"))
    costs, losses, pools, pairing, selected = model_metadata_rows(root, protocol, selection)
    cost_summary = computational_summary(costs)
    tables = {"model_per_seed": per_seed, "model_summary": summary, "paired_effects_per_seed": pairs,
        "paired_effects_summary": pair_summary, "per_class": per_class, "confusions": confusions,
        "paired_predictions": paired_predictions(groups), "fold_metrics": folds,
        "selected_shared_profiles": selected, "pipeline_costs": costs, "cost_summary": cost_summary,
        "training_loss_curves": losses, "pretrain_pool_steps": pools, "paired_training_states": pairing,
        "historical_context": [HISTORICAL]}
    output.mkdir(parents=True)
    for name, rows in tables.items():
        write_csv(output / (name + ".csv"), rows)
    spreadsheet(output / "cross_plant_transfer_comparison.xlsx", tables)
    if make_figures:
        figures(output, summary, pair_summary, confusions, losses, cost_summary)
    (output / "README.md").write_text(report_text(summary, pair_summary, selected), encoding="utf-8")
    if sha(root / "scoring_manifest.json") != release["scoring_manifest_sha256"]:
        raise ExportError("Scoring changed during read-only export")
    write_json(output / "export_manifest.json", {"created_at_utc": datetime.now(timezone.utc).isoformat(),
        "root_release_sha256": sha(release_path), "scoring_manifest_sha256": sha(root / "scoring_manifest.json"),
        "prediction_lock_sha256": sha(root / "prediction_lock.json"), "selection_lock_sha256": sha(root / "selection_lock.json"),
        "exporter_source_sha256": sha(__file__), "upstream_fold_metrics_sha256": sha(root / "fold_metrics.json"), "tables": {k: len(v) for k, v in tables.items()},
        "root_independent_evaluation_review_reference": release.get("independent_evaluation_review_sha256"),
        "new_training_fits": 0, "benchmark_score_calls": 0, "independent_prediction_recalculation": True,
        "historical_accuracy_is_not_transfer_gain_baseline": True, "study_status": scored["study_status"]})
    write_json(output / "files_sha256.json", {p.name: {"sha256": sha(p), "bytes": p.stat().st_size} for p in sorted(output.iterdir()) if p.is_file()})
    return {"output": str(output), "tables": len(tables), "model_seed_cases": len(per_seed), "prediction_rows_per_case": 855,
            "source_scoring_sha256": sha(root / "scoring_manifest.json"), "new_training_fits": 0}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--release", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(export(args.run, args.output, args.release), ensure_ascii=False))


if __name__ == "__main__":
    main()
