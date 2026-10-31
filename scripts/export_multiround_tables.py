"""Export all four completed research rounds; never fit, tune or select models."""
import argparse
from collections import Counter, defaultdict
import csv
import gzip
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROUNDS = ("neural-base", "support-svm", "anchored-ft", "meta-refinement")
ROUND_LABELS = {"neural-base": "第1轮：监督MLP基础模型", "support-svm": "第2轮：支持样品加权SVM",
    "anchored-ft": "第3轮：源锚定神经微调", "meta-refinement": "第4轮：加强元训练与微调"}
METRICS = ("accuracy", "balanced_accuracy", "macro_f1")
CLASSES = tuple(f"ZY{i:02d}" for i in range(1, 8))
METHOD_LABELS = {"source_zero": "源模型／零支持", "source_original": "原始源SVM",
    "matched_zero": "匹配零支持重拟合", "development_selected": "正式开发选择",
    "family_support_svm": "预选SVM家族诊断", "selected": "正式开发选择",
    "selected_matched_zero": "正式选择的匹配零支持", "selected_source_zero": "正式选择的自身零支持",
    "diagnostic": "开发预选诊断", "diagnostic_source_zero": "诊断的自身零支持",
    "diagnostic_matched_zero": "诊断的匹配零支持", "supervised_diagnostic": "监督继续训练诊断",
    "supervised_diagnostic_source_zero": "监督诊断的自身零支持",
    "supervised_diagnostic_matched_zero": "监督诊断的匹配零支持"}


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def metrics_from_cm(cm):
    diagonal = cm.diagonal()
    counts, predicted = cm.sum(1), cm.sum(0)
    recall = np.divide(diagonal, counts, out=np.zeros(len(CLASSES)), where=counts > 0)
    f1 = np.divide(2*diagonal, counts+predicted, out=np.zeros(len(CLASSES)), where=counts+predicted > 0)
    return {"accuracy": float(diagonal.sum()/cm.sum()), "balanced_accuracy": float(recall[counts > 0].mean()),
        "macro_f1": float(f1.mean()), "confusion_matrix": cm.tolist(), "class_recall": recall.tolist(),
        "class_counts": counts.tolist(), "correct": int(diagonal.sum()), "total": int(cm.sum())}


def _source_and_matched(round_name, method):
    if round_name == "neural-base":
        return method, method
    if round_name == "support-svm":
        return "source_original", "matched_zero"
    if round_name == "anchored-ft":
        if method == "source_zero":
            return method, method
        root = method.split("_matched_zero")[0]
        return "source_zero", root+"_matched_zero"
    root = method.removesuffix("_source_zero").removesuffix("_matched_zero")
    return root+"_source_zero", root+"_matched_zero"


def _selection(lock):
    return lock.get("choice", lock)


def _method_status(round_name, method, lock):
    choice = _selection(lock)
    if round_name == "neural-base":
        return "训练内部锁定源模型；无目标校准"
    if "zero" in method or method == "source_original":
        return "零支持对照；保留原始结果"
    if method in {"development_selected", "selected"}:
        selected = choice.get("selected")
        if selected is None or selected.get("family") == "identity":
            return "未通过开发门槛，正式方案回退零支持"
        return "通过训练内部开发选择及检验"
    return "开发预选诊断；不按最终预测分数替换正式选择"


def _sheet(book, name, rows, *, percents=(), decimals=()):
    sheet = book.create_sheet(name)
    keys = list(rows[0])
    sheet.append(keys)
    for row in rows:
        sheet.append([row.get(k) for k in keys])
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.zoomScale = 80
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="244F70")
        cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    for column, key in enumerate(keys, 1):
        sheet.column_dimensions[get_column_letter(column)].width = 18 if len(key) < 12 else 24
        if key in {"含义／边界", "方法身份", "固定参数", "说明", "候选参数"}:
            sheet.column_dimensions[get_column_letter(column)].width = 55
        for cells in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
            cell = cells[0]
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F1F5F9")
            if key in percents:
                cell.number_format = "0.00%"
            elif key in decimals:
                cell.number_format = "0.0000"
    sheet.row_dimensions[1].height = 42
    height = 58 if name == "说明" else 68 if name in {"四轮台账", "所有最终方法", "全部候选开发汇总"} else 28
    for row in range(2, len(rows)+2):
        sheet.row_dimensions[row].height = height
    if "是否正确" in keys:
        index = keys.index("是否正确")+1
        for row in sheet.iter_rows(min_row=2):
            if row[index-1].value == 0:
                for cell in row:
                    cell.fill = PatternFill("solid", fgColor="FCE4E4")
    sheet.page_setup.orientation = "landscape"
    return sheet


def _write_csv(path, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader(); writer.writerows(rows)


def export(base, output, split_plan):
    base, output, split_plan = Path(base), Path(output), Path(split_plan)
    gate_path = base/"all_development_locks.json"
    if not gate_path.is_file():
        raise RuntimeError("缺少四轮开发锁统一释放记录，不导出可能混入未锁定配置的结果")
    gate = read_json(gate_path)
    if gate.get("status") != "released_for_final_evaluation" or {r["name"] for r in gate.get("rounds", [])} != set(ROUNDS):
        raise RuntimeError("四轮统一释放记录不完整")
    gate_checks = 0
    for record in gate["rounds"]:
        run = base/record["name"]
        if record["final_directory_existed_at_release"]:
            raise RuntimeError("统一释放前已有该轮最终结果目录，需核查时序")
        for path, expected in ((run/"selection_lock.json", record["selection_lock_sha256"]),
                (run/"run_manifest.json", record["run_manifest_sha256"]),
                (Path(record["protocol"]), record["protocol_sha256"])):
            if not path.is_file() or digest(path) != expected:
                raise RuntimeError("统一释放后的开发锁／协议／执行清单hash改变")
            gate_checks += 1
    # Check completeness before creating any partial spreadsheet or inventing rows.
    for round_name in ROUNDS:
        run = base/round_name
        status_path = run/"status.json"
        if not status_path.exists():
            raise RuntimeError(f"四轮最终评价尚未全部完成：缺少{round_name}/status.json")
        status = read_json(status_path)
        if not (status.get("status") == "completed" or status.get("stage") == "final_complete"):
            raise RuntimeError(f"四轮最终评价尚未全部完成：{round_name}仍为{status.get('status', status.get('stage'))}")
        for filename in ("summary.json", "all_development_metrics.json", "candidate_grid.json", "selection_lock.json"):
            if not (run/filename).is_file():
                raise RuntimeError(f"最终评价证据不完整：{round_name}/{filename}")
        if not list((run/"final").rglob("predictions.csv.gz")):
            raise RuntimeError(f"缺少{round_name}最终逐样品预测")
    plan = read_json(split_plan)
    truth = {r["record_id"]: r["origin_group_id"] for r in plan["records"] if r["role"] == "prediction"}
    query_ids = set(truth)
    if len(query_ids) != 171 or set(truth.values()) != set(CLASSES):
        raise ValueError("需要原固定171份七类预测样品，不接受换分区")
    counts = Counter(truth.values())
    random_ids = set(np.random.default_rng(20261102).choice(np.asarray(sorted(truth)), 20, replace=False).tolist())
    checked = gate_checks; maximum_difference = 0.; provenance = {str(gate_path): digest(gate_path)}

    def close(actual, expected, label):
        nonlocal checked, maximum_difference
        difference = float(np.max(np.abs(np.asarray(actual, dtype=float)-np.asarray(expected, dtype=float))))
        maximum_difference = max(maximum_difference, difference); checked += 1
        if difference > 1e-12:
            raise AssertionError(f"数值不一致：{label} ({difference})")

    cases = {}; sample_stats = defaultdict(lambda: {"count": 0, "correct": 0, "predictions": Counter(), "first": None})
    random_rows, ledger, development_rows = [], [], []
    supplied_summaries = {}
    locks = {}; candidate_counts = {}
    prediction_counts = Counter()
    for round_name in ROUNDS:
        run = base/round_name
        lock = read_json(run/"selection_lock.json"); locks[round_name] = lock
        supplied_summaries[round_name] = read_json(run/"summary.json")
        grid = read_json(run/"candidate_grid.json"); candidate_counts[round_name] = len(grid)
        config = {c["id"]: c for c in grid}
        raw_development = read_json(run/"all_development_metrics.json")
        groups = defaultdict(list)
        for row in raw_development:
            groups[(row.get("strategy", "supervised" if round_name == "neural-base" else round_name), row["candidate_id"], row["stage"])].append(row)
        for (strategy, candidate, stage), values in sorted(groups.items()):
            record = {"轮次": ROUND_LABELS[round_name], "策略": strategy, "候选编号": candidate,
                "开发阶段": "前段选择" if stage == "selection" else "后段检验", "评价次数": len(values),
                "候选参数": json.dumps(config[candidate], ensure_ascii=False, sort_keys=True),
                "准确率": float(np.mean([r["accuracy"] for r in values])),
                "平衡准确率": float(np.mean([r["balanced_accuracy"] for r in values])),
                "宏F1": float(np.mean([r["macro_f1"] for r in values]))}
            for metric, label in (("accuracy", "准确率"), ("balanced_accuracy", "平衡准确率"), ("macro_f1", "宏F1")):
                if round_name == "neural-base":
                    source, matched = None, None
                else:
                    source = float(np.mean([r.get("source_"+metric, r[metric]-r.get(metric+"_gain_vs_source", 0.)) for r in values]))
                    matched = float(np.mean([r["matched_"+metric] for r in values]))
                record["自身零支持"+label] = source
                record["匹配零支持"+label] = matched
            choice = _selection(lock)
            formal = choice.get("selected")
            diagnostic = choice.get("diagnostic")
            record["开发身份"] = ("正式选中" if formal and formal.get("id") == candidate else
                "诊断候选" if diagnostic and diagnostic.get("id") == candidate else "完整保留的其他候选")
            development_rows.append(record)
        files = sorted((run/"final").rglob("predictions.csv.gz"))
        for path in files:
            provenance[str(path)] = digest(path)
            with gzip.open(path, "rt", encoding="utf-8", newline="") as handle:
                for raw in csv.DictReader(handle):
                    checked += 3
                    identifier = raw["record_id"]
                    if identifier not in query_ids or raw["y_true"] != truth[identifier]:
                        raise AssertionError("预测记录和固定分区的原始标签不一致")
                    correct = int(raw["is_correct"])
                    if correct != int(raw["y_true"] == raw["y_predicted"]):
                        raise AssertionError("预测正确标记与真实标签比对不一致")
                    if raw["y_predicted"] not in CLASSES:
                        raise AssertionError("预测类别超出七类词表")
                    seed = str(raw.get("seed", "确定模型"))
                    repeat = int(raw.get("repeat", 0)); method = raw.get("method", "source_zero")
                    strategy = raw.get("strategy", "supervised" if round_name == "neural-base" else round_name)
                    key = (round_name, seed, repeat, method)
                    if key not in cases:
                        cases[key] = {"cm": np.zeros((7, 7), dtype=int), "seen": set(),
                            "candidate_id": raw["candidate_id"], "strategy": strategy}
                    case = cases[key]
                    if identifier in case["seen"] or case["candidate_id"] != raw["candidate_id"]:
                        raise AssertionError("同一评价内样品重复或候选编号不一致")
                    case["seen"].add(identifier)
                    case["cm"][CLASSES.index(raw["y_true"]), CLASSES.index(raw["y_predicted"])] += 1
                    stats = sample_stats[(round_name, method, seed, identifier)]
                    stats["count"] += 1; stats["correct"] += correct; stats["predictions"][raw["y_predicted"]] += 1
                    if repeat == 0:
                        stats["first"] = raw["y_predicted"]
                    if identifier in random_ids and repeat == 0:
                        random_rows.append({"轮次": ROUND_LABELS[round_name], "方法": METHOD_LABELS.get(method, method),
                            "方法代码": method, "策略": strategy, "候选编号": raw["candidate_id"], "训练种子": seed,
                            "校准抽样序号": repeat, "样品编号": identifier, "真实匿名产地": raw["y_true"],
                            "预测匿名产地": raw["y_predicted"], "是否正确": correct})
                    prediction_counts[round_name] += 1
        metric_path = run/("final/metrics.json" if round_name == "support-svm" else "final_metrics.json")
        provenance[str(metric_path)] = digest(metric_path)
        recorded = read_json(metric_path)
        for row in recorded:
            key = (round_name, str(row.get("seed", "确定模型")), int(row.get("repeat", 0)), row.get("method", "source_zero"))
            if key not in cases:
                raise AssertionError("指标JSON找不到对应完整预测")
            case = cases[key]
            if case["seen"] != query_ids:
                raise AssertionError("某一次评价未覆盖全部固定171份预测")
            case["metrics"] = metrics_from_cm(case["cm"])
            for metric in (*METRICS, "confusion_matrix", "class_recall", "class_counts", "correct", "total"):
                if metric in row:
                    close(row[metric], case["metrics"][metric], "逐样品重算"+metric)
            case["recorded"] = row
        round_cases = {key: value for key, value in cases.items() if key[0] == round_name}
        if len(recorded) != len(round_cases):
            raise AssertionError("指标JSON与预测评价次数不同")
        observed_combinations = {(r.get("strategy", round_name), r["candidate_id"]) for r in raw_development}
        formal = _selection(lock).get("selected")
        ledger.append({"轮次": ROUND_LABELS[round_name], "主候选配置数": len(grid),
            "开发候选与策略组合数": len(observed_combinations), "开发评价次数": len(raw_development),
            "最终评价次数": len(round_cases), "最终逐样品预测行数": prediction_counts[round_name],
            "唯一预测样品数": len(query_ids), "训练种子数": 0 if round_name == "support-svm" else len({key[1] for key in round_cases}),
            "唯一源模型数": 1 if round_name == "support-svm" else 3,
            "每种子每方法校准抽样数": 1 if round_name == "neural-base" else 100,
            "目标支持样品数": 0 if round_name == "neural-base" else 7,
            "正式候选": formal["id"] if formal else "回退自身零支持",
            "选择锁SHA256": digest(run/"selection_lock.json"),
            "说明": "源模型改进，无目标校准" if round_name == "neural-base" else "含零支持／诊断对照；未通过者回退，七类每类1份"})
        for name in ("summary.json", "selection_lock.json", "candidate_grid.json", "all_development_metrics.json"):
            provenance[str(run/name)] = digest(run/name)
    # All primary metrics and gains below come from the raw class predictions.
    for key, case in cases.items():
        round_name, seed, repeat, method = key
        source_method, matched_method = _source_and_matched(round_name, method)
        source = cases[(round_name, seed, repeat, source_method)]["metrics"]
        matched = cases[(round_name, seed, repeat, matched_method)]["metrics"]
        case["source_metrics"], case["matched_metrics"] = source, matched
        recorded = case["recorded"]
        for metric in METRICS:
            for alias in (metric+"_gain_vs_source", metric+"_gain_over_source"):
                if alias in recorded:
                    close(recorded[alias], case["metrics"][metric]-source[metric], "自身零支持差值")
            for alias in (metric+"_gain_over_matched",):
                if alias in recorded:
                    close(recorded[alias], case["metrics"][metric]-matched[metric], "匹配零支持差值")
            if "source_"+metric in recorded:
                close(recorded["source_"+metric], source[metric], "自身零支持指标")
            if "matched_"+metric in recorded:
                close(recorded["matched_"+metric], matched[metric], "匹配零支持指标")
    summary_rows, seed_rows, class_rows = [], [], []
    groups = defaultdict(list)
    for key, case in cases.items():
        groups[(key[0], key[3])].append((key, case))
    for (round_name, method), values in sorted(groups.items(), key=lambda item: (ROUNDS.index(item[0][0]), item[0][1])):
        record = {"轮次": ROUND_LABELS[round_name], "方法": METHOD_LABELS.get(method, method), "方法代码": method,
            "策略": values[0][1]["strategy"], "候选编号": values[0][1]["candidate_id"],
            "方法身份": _method_status(round_name, method, locks[round_name]), "评价次数": len(values),
            "训练种子数": 0 if round_name == "support-svm" else len({key[1] for key, _ in values}),
            "唯一源模型数": 1 if round_name == "support-svm" else 3, "唯一预测样品数": 171}
        for metric, label in (("accuracy", "平均准确率"), ("balanced_accuracy", "平均平衡准确率"), ("macro_f1", "平均宏F1")):
            numbers = [c["metrics"][metric] for _, c in values]
            record[label] = float(np.mean(numbers))
        record.update({"准确率描述性标准差": float(np.std([c["metrics"]["accuracy"] for _, c in values])),
            "自身零支持准确率": float(np.mean([c["source_metrics"]["accuracy"] for _, c in values])),
            "匹配零支持准确率": float(np.mean([c["matched_metrics"]["accuracy"] for _, c in values])),
            "对自身零支持增益（百分点）": None if round_name == "neural-base" else 100*float(np.mean([c["metrics"]["accuracy"]-c["source_metrics"]["accuracy"] for _, c in values])),
            "对匹配零支持增益（百分点）": None if round_name == "neural-base" else 100*float(np.mean([c["metrics"]["accuracy"]-c["matched_metrics"]["accuracy"] for _, c in values]))})
        deltas = [c["metrics"]["accuracy"]-c["source_metrics"]["accuracy"] for _, c in values]
        record.update({"提高次数": sum(d > 1e-12 for d in deltas), "不变次数": sum(abs(d) <= 1e-12 for d in deltas),
            "下降次数": sum(d < -1e-12 for d in deltas)})
        supplied = supplied_summaries[round_name]
        entry = supplied if round_name == "neural-base" else next(r for r in supplied if r["method"] == method)
        for metric in METRICS:
            close(entry[metric+"_mean"], np.mean([c["metrics"][metric] for _, c in values]), "保存汇总与原始预测均值")
        summary_rows.append(record)
        for seed in sorted({key[1] for key, _ in values}):
            chosen = [c for key, c in values if key[1] == seed]
            seed_rows.append({"轮次": ROUND_LABELS[round_name], "方法": METHOD_LABELS.get(method, method), "方法代码": method,
                "候选编号": chosen[0]["candidate_id"], "训练种子": seed, "评价次数": len(chosen),
                "平均准确率": float(np.mean([c["metrics"]["accuracy"] for c in chosen])),
                "自身零支持准确率": float(np.mean([c["source_metrics"]["accuracy"] for c in chosen])),
                "匹配零支持准确率": float(np.mean([c["matched_metrics"]["accuracy"] for c in chosen])),
                "平均平衡准确率": float(np.mean([c["metrics"]["balanced_accuracy"] for c in chosen])),
                "平均宏F1": float(np.mean([c["metrics"]["macro_f1"] for c in chosen]))})
        for index, label in enumerate(CLASSES):
            class_rows.append({"轮次": ROUND_LABELS[round_name], "方法": METHOD_LABELS.get(method, method), "方法代码": method,
                "匿名产地组": label, "该类唯一预测样品数": counts[label], "重复评价次数": len(values),
                "平均召回率": float(np.mean([c["metrics"]["class_recall"][index] for _, c in values])),
                "自身零支持平均召回率": float(np.mean([c["source_metrics"]["class_recall"][index] for _, c in values])),
                "匹配零支持平均召回率": float(np.mean([c["matched_metrics"]["class_recall"][index] for _, c in values]))})
    sample_rows = []
    for (round_name, method, seed, identifier), value in sorted(sample_stats.items(), key=lambda item: (ROUNDS.index(item[0][0]), *item[0][1:])):
        expected = 1 if round_name == "neural-base" else 100
        if value["count"] != expected or value["first"] is None:
            raise AssertionError("逐样品统计缺少某次校准抽样")
        source_method, matched_method = _source_and_matched(round_name, method)
        source = sample_stats[(round_name, source_method, seed, identifier)]
        matched = sample_stats[(round_name, matched_method, seed, identifier)]
        mode = sorted(value["predictions"], key=lambda c: (-value["predictions"][c], c))[0]
        sample_rows.append({"轮次": ROUND_LABELS[round_name], "方法": METHOD_LABELS.get(method, method), "方法代码": method,
            "训练种子": seed, "样品编号": identifier, "真实匿名产地": truth[identifier],
            "评价次数": value["count"], "正确次数": value["correct"], "正确比例": value["correct"]/value["count"],
            "自身零支持正确比例": source["correct"]/source["count"], "匹配零支持正确比例": matched["correct"]/matched["count"],
            "首次抽样预测": value["first"], "最常预测组": mode,
            "各预测组次数": json.dumps(dict(sorted(value["predictions"].items())), ensure_ascii=False)})
    if sum(candidate_counts.values()) != 85:
        raise AssertionError("四轮主配置总数不是已固定的85；请核对是否混入其他实验")
    expected_random = sum((3 if r == "neural-base" else len({key[1] for key in cases if key[0] == r})*len({key[3] for key in cases if key[0] == r})) * 20 for r in ROUNDS)
    if len(random_rows) != expected_random:
        raise AssertionError("随机样品展示行数不完整")
    notes = [
        {"项目": "数据", "含义／边界": "用户上传的真实酸枣仁光谱；ZY01—ZY07为匿名产地组，未生成模拟光谱。"},
        {"项目": "固定角色", "含义／边界": "855份：513训练／171校准池／171预测；四轮使用相同171份预测，不新增独立样品。"},
        {"项目": "每产地批次", "含义／边界": "每产地目前仅一批，预测组已有历史研究曝光；全部属于同批次回顾性探索。"},
        {"项目": "候选记账", "含义／边界": "四轮85个主配置；元训练还对监督继续训练策略评价相同配置，控制不应冒充独立新算法。"},
        {"项目": "正式选择", "含义／边界": "全部根据训练内部开发与检验锁定；诊断和失败候选完整保留，未按最终最高分替换正式方案。"},
        {"项目": "第一轮", "含义／边界": "源MLP改进，不使用目标校准标签；3个训练种子各只预测一次，不能说成300次小样本校准。"},
        {"项目": "b006结构", "含义／边界": "锁定源候选b006隐藏层为64→32、无PCA、200轮、未加类别权重；不是单隐藏层。"},
        {"项目": "第二轮", "含义／边界": "单个确定SVM，100次支持抽样；冻结预处理/gamma/源classweight后refit副本，是参数适配，不是纯分数修正。"},
        {"项目": "第三／四轮", "含义／边界": "3种子×100次支持抽样；300情景共享固定171样品，不是300个独立批次。"},
        {"项目": "第4轮正式回退", "含义／边界": "目标校准未通过开发门槛，回退预先指定的100更新FO源模型；500更新FO与监督继续训练是诊断对照，不能按最终高分替换正式结果。"},
        {"项目": "支持量", "含义／边界": "七类每类1份、共7份；ZY02校准池仅1份，当前不支持七类均匀3／5-shot。"},
        {"项目": "两种零支持", "含义／边界": "自身源模型与匹配零支持（相同额外回放但无目标支持损失）分开；基础表示收益不记为校准标签收益。"},
        {"项目": "统计量", "含义／边界": "均值和标准差描述当前种子／重叠抽样，不提供独立新批次统计置信区间。"},
        {"项目": "零支持重复行", "含义／边界": "重复列出的零支持预测用于与每次校准配对；SVM只有1个确定源模型，神经方法每配置3个种子，不是重复训练100／300个源模型。"},
        {"项目": "随机展示", "含义／边界": "种子20261102固定20个样品；第0次支持抽样，错误行全部保留并标红。"},
        {"项目": "极少类", "含义／边界": "ZY02最终预测仅1份，召回率只有0或100%；应同时检查宏F1、平衡准确率和逐类召回。"},
        {"项目": "原论文关系", "含义／边界": "已有MAML／ProtoNet思想的分类工程适配，不是原NIR/MIR回归论文精确复现，也不据此声称发明算法。"},
        {"项目": "外部限制", "含义／边界": "未验证真实地名、跨批次或新仪器；本轮不使用Gaussian，也不宣称已连接Google Colab。"},
    ]
    output.mkdir(parents=True, exist_ok=True)
    book = Workbook(); book.remove(book.active)
    tables = [("说明", notes), ("四轮台账", ledger), ("所有最终方法", summary_rows), ("全部候选开发汇总", development_rows),
        ("全部种子", seed_rows), ("逐产地召回", class_rows), ("随机样品", random_rows), ("逐样品统计", sample_rows)]
    percent_keys = {"准确率", "平衡准确率", "自身零支持准确率", "匹配零支持准确率", "自身零支持平衡准确率", "匹配零支持平衡准确率",
        "平均准确率", "平均平衡准确率", "准确率描述性标准差",
        "平均召回率", "自身零支持平均召回率", "匹配零支持平均召回率", "正确比例", "自身零支持正确比例", "匹配零支持正确比例"}
    decimals = {"宏F1", "平均宏F1", "自身零支持宏F1", "匹配零支持宏F1", "对自身零支持增益（百分点）", "对匹配零支持增益（百分点）"}
    diagnostic_path = base/"meta-refinement"/"source_control_diagnostic.json"
    source_control_rows = []
    source_control_prediction_rows = 0
    if diagnostic_path.exists():
        diagnostic = read_json(diagnostic_path)
        provenance[str(diagnostic_path)] = digest(diagnostic_path)
        diagnostic_csv = diagnostic_path.with_suffix(".csv")
        if diagnostic.get("status") != "pass" or diagnostic.get("new_fits") != 0 or digest(diagnostic_csv) != diagnostic["csv_sha256"]:
            raise AssertionError("元训练源对照诊断尚未通过或CSV hash不一致")
        provenance[str(diagnostic_csv)] = digest(diagnostic_csv)
        source_predictions = defaultdict(dict); diagnostic_prediction_rows = []
        with diagnostic_csv.open(encoding="utf-8-sig", newline="") as handle:
            for row in csv.DictReader(handle):
                identifier = row["record_id"]; checked += 2
                if identifier not in truth or truth[identifier] != row["y_true"] or int(row["is_correct"]) != int(row["y_true"] == row["y_predicted"]):
                    raise AssertionError("源对照预测标签或正确标记不一致")
                key = (int(row["seed"]), row["source_model"])
                if identifier in source_predictions[key]:
                    raise AssertionError("源对照存在重复样品")
                source_predictions[key][identifier] = row["y_predicted"]
                diagnostic_prediction_rows.append({"训练种子": int(row["seed"]), "源模型": row["source_model"],
                    "样品编号": identifier, "真实匿名产地": row["y_true"], "预测匿名产地": row["y_predicted"],
                    "是否正确": int(row["is_correct"])})
                source_control_prediction_rows += 1
        source_computed = {}
        for record in diagnostic["source_model_metrics"]:
            key = (record["seed"], record["source_model"]); values = source_predictions[key]
            if set(values) != query_ids:
                raise AssertionError("源对照未覆盖全部171固定样品")
            cm = np.zeros((7, 7), dtype=int)
            for identifier, predicted in values.items():
                cm[CLASSES.index(truth[identifier]), CLASSES.index(predicted)] += 1
            computed = metrics_from_cm(cm); source_computed[key] = computed
            for metric in (*METRICS, "confusion_matrix", "class_recall", "correct", "total"):
                close(record[metric], computed[metric], "源对照逐样品重算")
            source_control_rows.append({"汇总层级": "每个种子", "源模型": record["source_model"], "训练种子": record["seed"],
                "固定预测样品数": 171, "准确率": computed["accuracy"], "平衡准确率": computed["balanced_accuracy"],
                "宏F1": computed["macro_f1"], "目标校准标签数": 0,
                "说明": "已存源模型只读推理；原任务额外训练，不算目标标签校准收益"})
        for record in diagnostic["source_model_means"]:
            values = [value for (seed, model), value in source_computed.items() if model == record["source_model"]]
            for metric in METRICS:
                close(record[metric+"_mean"], np.mean([r[metric] for r in values]), "源对照三种子均值")
            source_control_rows.append({"汇总层级": "三个种子均值", "源模型": record["source_model"], "训练种子": "42／43／44",
                "固定预测样品数": 171, "准确率": record["accuracy_mean"], "平衡准确率": record["balanced_accuracy_mean"],
                "宏F1": record["macro_f1_mean"], "目标校准标签数": 0,
                "说明": "源学习对照；相同预测标签不代表权重／分数相同，不作为新增正式校准方法"})
        disagreement_rows = []
        for record in diagnostic["pairwise_disagreement"]:
            first = source_predictions[(record["seed"], record["first"])]; second = source_predictions[(record["seed"], record["second"])]
            different = sum(first[i] != second[i] for i in query_ids)
            corrected = sum(first[i] != truth[i] and second[i] == truth[i] for i in query_ids)
            harmed = sum(first[i] == truth[i] and second[i] != truth[i] for i in query_ids)
            close(record["different_predictions"], different, "源对照预测分歧")
            close(record["second_corrected_first"], corrected, "源对照纠正错误")
            close(record["second_new_errors"], harmed, "源对照引入错误")
            disagreement_rows.append({"训练种子": record["seed"], "第一个源模型": record["first"], "第二个源模型": record["second"],
                "固定预测样品数": 171, "预测标签不同的样品数": different, "第二模型纠正数": corrected, "第二模型新增错误数": harmed,
                "最大logit差（原诊断记录）": record["max_absolute_logit_difference"],
                "说明": "逐样品标签差异已从源对照CSV重算；源学习变化不是校准标签收益"})
        notes.append({"项目": "元训练源对照", "含义／边界": "额外只读比较5个源模型：100／500更新FO与同任务监督继续训练在全部3种子上的预测标签分别完全相同。源提升不能归因FO独有机制，也不是目标校准收益。"})
        tables.append(("元训练源对照诊断", source_control_rows))
        tables.append(("源模型预测分歧", disagreement_rows))
        tables.append(("源对照逐样品", diagnostic_prediction_rows))
    csv_files = {}
    for name, rows in tables:
        _sheet(book, name, rows, percents=percent_keys, decimals=decimals)
        csv_path = output/(name+".csv")
        _write_csv(csv_path, rows); csv_files[csv_path.name] = digest(csv_path)
    workbook_path = output/"四轮优化_完整结果与样品预测对照.xlsx"
    book.save(workbook_path)
    loaded = load_workbook(workbook_path, read_only=True, data_only=True)
    readback = {}
    for name, rows in tables:
        sheet = loaded[name]
        if sheet.max_row != len(rows)+1:
            raise AssertionError("Excel导出丢失记录")
        keys = list(rows[0]); actual_rows = sheet.iter_rows(min_row=2, values_only=True)
        for expected, actual in zip(rows, actual_rows):
            for column, value in enumerate(actual):
                wanted = expected.get(keys[column])
                checked += 1
                if isinstance(wanted, (int, float)):
                    close(value, wanted, "Excel写入读回")
                elif value != wanted:
                    raise AssertionError("Excel文本写入读回不一致")
        readback[name] = len(rows)
    loaded.close()
    validation = {"status": "pass", "new_model_fits": 0, "parameter_or_seed_selection_performed": False,
        "checks_evaluated": checked, "maximum_numeric_difference": maximum_difference,
        "final_prediction_rows_checked": sum(prediction_counts.values()), "prediction_rows_by_round": dict(prediction_counts),
        "final_cases_checked": len(cases), "methods_including_controls": len(summary_rows),
        "main_configuration_count": sum(candidate_counts.values()), "main_configuration_counts_by_round": candidate_counts,
        "development_aggregate_rows": len(development_rows), "workbook_data_rows": readback,
        "fixed_query_unique_samples": 171, "random_sample_unique_ids": sorted(random_ids),
        "random_example_rows": len(random_rows), "random_incorrect_rows_preserved": sum(r["是否正确"] == 0 for r in random_rows),
        "workbook_sha256": digest(workbook_path), "csv_sha256": csv_files,
        "split_plan_sha256": digest(split_plan), "script_sha256": digest(Path(__file__)),
        "all_development_locks_sha256": digest(gate_path), "round_release_hash_checks": gate_checks,
        "meta_source_control_diagnostic_included": bool(source_control_rows),
        "source_control_prediction_rows_checked": source_control_prediction_rows,
        "input_artifact_sha256": provenance,
        "limits": ["Anonymous labels on real uploaded spectra", "Retrospective same-batch; repeated draws share fixed171query samples", "Source representation improvements are not support-label calibration gains", "All diagnostic/negative/zero-support results retained"]}
    (output/"table_export_validation.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    print(json.dumps({k: v for k, v in validation.items() if k not in {"input_artifact_sha256", "random_sample_unique_ids", "csv_sha256"}}, ensure_ascii=False))
    return validation


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base", type=Path, default=Path("outputs/multiround-20261031"))
    parser.add_argument("--output", type=Path, default=Path("outputs/multiround-20261031/tables"))
    parser.add_argument("--split-plan", type=Path, default=Path("data/processed/zaoren_threeway_20261031_v1/split_plan.json"))
    args = parser.parse_args()
    export(args.base, args.output, args.split_plan)


if __name__ == "__main__":
    main()
