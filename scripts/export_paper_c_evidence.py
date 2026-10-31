"""Create paper tables/figures from recorded predictions, without fitting models.

The fixed 171-record prediction experiment is deliberately excluded from the
855-record OOF model comparison.  New benchmark inputs are optional and must
already contain completed outer predictions.
"""
import argparse
from collections import Counter, defaultdict
import csv
import gzip
import hashlib
import json
from pathlib import Path
import subprocess

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

plt.rcParams.update({"pdf.fonttype": 42, "ps.fonttype": 42, "svg.fonttype": "none"})


ROOT = Path(__file__).resolve().parents[1]
HISTORY = ROOT / "reports/staged-nested-optimization/2026-10-31"
CLASSES = tuple(f"ZY{i:02d}" for i in range(1, 8))
STAGES = ("R0", "R1", "R2", "R3", "R4")
PRIMARY_METHODS = ("R2_reproduction", "full_svm", "shrinkage_lda", "mlp",
                   "strong_mlp_200", "strong_mlp_200_plus100_source_updates")
LABELS = {
    "R0": "Original MC-UVE + SVM",
    "R1": "MC-UVE search",
    "R2": "MC-UVE + SVM search",
    "R3": "Preprocessing selection",
    "R4": "Combination selection",
    "fixed_snv_C10_MC20": "Fixed SNV / C10 / MC20 anchor",
    "R3_diagnostic_raw": "Raw",
    "R3_diagnostic_snv": "SNV",
    "R3_diagnostic_sg0_snv": "SG smoothing + SNV",
    "R3_diagnostic_sg1_snv": "SG first derivative + SNV",
    "R3_diagnostic_sg2_snv": "SG second derivative + SNV",
    "R3_diagnostic_msc": "MSC",
    "R3_diagnostic_linear_detrend_snv": "Detrend + SNV",
    "R4_diagnostic_ordinary_mcuve": "Ordinary MC-UVE",
    "R4_diagnostic_budget_mcuve": "10x MC budget",
    "R4_diagnostic_stable_mcuve": "Stable MC-UVE",
    "R4_diagnostic_stable_cars": "Stable CARS",
    "R4_diagnostic_stable_joint": "MC-UVE / CARS stable joint",
    "R4_diagnostic_cnn_softmax": "1D-CNN original head",
    "R4_diagnostic_cnn_embedding_svm": "Same CNN embedding + SVM",
    "full_svm": "Full-spectrum SVM",
    "shrinkage_lda": "Shrinkage LDA",
    "mlp": "Supervised MLP",
    "R2_reproduction": "R2 MC-UVE + SVM replay",
    "R2_full_features": "R2 without MC selection",
    "R2_random_K": "R2 label-blind random K",
    "R2_no_C_gamma_search": "R2 fixed C10 / gamma scale",
    "R2_no_class_weight": "R2 without class weighting",
    "R2_fixed_MC20_LV2_K62": "R2 fixed MC20 / LV2 / K62",
    "strong_mlp_200": "Strong 64-32 MLP, 200 epochs",
    "strong_mlp_200_plus100_source_updates": "Same MLP + 100 source updates",
}
SCOPE = "855 OOF; original 5 outer / 2 inner folds; retrospective within-batch"


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_csv(path):
    opener = gzip.open if Path(path).suffix == ".gz" else open
    with opener(path, "rt", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def scores(rows):
    """Independent integer-confusion arithmetic; never reinterprets OVR scores."""
    matrix = np.zeros((7, 7), dtype=np.int64)
    ids = set()
    for row in rows:
        assert row["record_id"] not in ids
        ids.add(row["record_id"])
        a, p = CLASSES.index(row["y_true"]), CLASSES.index(row["y_predicted"])
        assert int(row["is_correct"]) == (a == p)
        matrix[a, p] += 1
    support, predicted, tp = matrix.sum(1), matrix.sum(0), matrix.diagonal()
    recall = np.divide(tp, support, out=np.zeros(7), where=support > 0)
    precision = np.divide(tp, predicted, out=np.zeros(7), where=predicted > 0)
    f1 = np.divide(2 * tp, support + predicted, out=np.zeros(7), where=support + predicted > 0)
    assert len(rows) == 855 and support.tolist() == [350, 5, 100, 100, 100, 100, 100]
    return {"accuracy": float(tp.sum() / len(rows)), "macro_f1": float(f1.mean()),
            "balanced_accuracy": float(recall.mean()), "correct": int(tp.sum()),
            "total": len(rows), "confusion_matrix": matrix.tolist(),
            "per_class": {c: {"precision": float(precision[i]), "recall": float(recall[i]),
                               "f1": float(f1[i]), "support": int(support[i])}
                          for i, c in enumerate(CLASSES)}}


def group_predictions(rows):
    grouped = defaultdict(list)
    for row in rows:
        grouped[row["method"]].append(row)
    baseline_ids = set(row["record_id"] for row in next(iter(grouped.values())))
    baseline_roles = {row["record_id"]: (row["y_true"], row["outer_fold"])
                      for row in next(iter(grouped.values()))}
    for method, values in grouped.items():
        assert set(row["record_id"] for row in values) == baseline_ids
        assert Counter(row["outer_fold"] for row in values) == {str(i): 171 for i in range(5)}, method
        assert {row["record_id"]: (row["y_true"], row["outer_fold"]) for row in values} == baseline_roles
    return grouped


def write_csv(path, rows):
    assert rows
    with Path(path).open("w", encoding="utf-8-sig", newline="") as handle:
        keys = list(dict.fromkeys(key for row in rows for key in row))
        writer = csv.DictWriter(handle, fieldnames=keys, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def metric_row(method, value, baseline, role):
    delta = value["accuracy"] - baseline["accuracy"]
    return {"method": method, "label": LABELS.get(method, method), "role": role,
            "correct": value["correct"], "total": value["total"],
            "accuracy": value["accuracy"], "macro_f1": value["macro_f1"],
            "balanced_accuracy": value["balanced_accuracy"],
            "accuracy_gain_pp_vs_R0": 100 * delta,
            "relative_accuracy_gain_vs_R0": delta / baseline["accuracy"],
            "dataset_version": "2026-10-31-zaoren-numbered-v1",
            "protocol_id": "zaoren-staged-nested-R4-stability-CNN-20261031-v1" if method.startswith("R4") else "zaoren-staged-nested-four-round-20261031",
            "protocol": SCOPE}


def paired_r2_change(rows, reference, value, r2_metrics):
    prior = {row["record_id"]: row for row in reference}
    corrected = harmed = changed_wrong = 0
    for row in rows:
        old = prior[row["record_id"]]
        assert row["y_true"] == old["y_true"]
        old_right = old["y_predicted"] == old["y_true"]
        new_right = row["y_predicted"] == row["y_true"]
        corrected += int(not old_right and new_right)
        harmed += int(old_right and not new_right)
        changed_wrong += int(not old_right and not new_right and old["y_predicted"] != row["y_predicted"])
    assert r2_metrics["correct"] + corrected - harmed == value["correct"]
    return {"accuracy_change_pp_vs_R2": 100*(value["accuracy"]-r2_metrics["accuracy"]),
            "macro_f1_change_vs_R2": value["macro_f1"]-r2_metrics["macro_f1"],
            "balanced_accuracy_change_pp_vs_R2": 100*(value["balanced_accuracy"]-r2_metrics["balanced_accuracy"]),
            "paired_corrected_vs_R2": corrected, "paired_harmed_vs_R2": harmed,
            "both_wrong_changed_label_vs_R2": changed_wrong, "net_correct_change_vs_R2": corrected-harmed}


def save_figure(fig, dest, stem):
    for extension in ("png", "pdf"):
        fig.savefig(dest / f"{stem}.{extension}", dpi=240, bbox_inches="tight")
    plt.close(fig)


def stage_figure(dest, metrics):
    colors = ["#607D8B", "#4C78A8", "#147D64", "#E49C3A", "#B56A57"]
    fig, axes = plt.subplots(1, 2, figsize=(11.2, 4.9))
    x = np.arange(5)
    for ax, key, maximum, title in [(axes[0], "accuracy", 100, "Accuracy (%)"),
                                    (axes[1], "macro_f1", 1, "Macro-F1")]:
        values = [metrics[m][key] * (100 if key == "accuracy" else 1) for m in STAGES]
        bars = ax.bar(x, values, color=colors, width=.62)
        ax.set_xticks(x, ["R0\nOriginal", "R1\nMC-UVE", "R2\nSVM tuning", "R3\nPreprocess", "R4\nCombination"])
        ax.set_ylim(0, maximum * 1.09)
        ax.set_yticks(np.linspace(0, maximum, 6))
        ax.set_ylabel(title)
        ax.grid(axis="y", alpha=.2)
        ax.set_axisbelow(True)
        for bar, value in zip(bars, values):
            ax.text(bar.get_x() + bar.get_width()/2, value + maximum*.018,
                    f"{value:.2f}" if key == "accuracy" else f"{value:.4f}",
                    ha="center", fontsize=10)
    fig.suptitle("Acid jujube kernel: five prespecified optimization stages", fontsize=14)
    fig.text(.5, .02, "855 pooled OOF predictions; retrospective within-batch evidence, not independent confirmation.\n"
             "R2 is the best observed prespecified route; later stages were not replaced using outer-test results.",
             ha="center", fontsize=9)
    fig.tight_layout(rect=[0, .11, 1, .93])
    save_figure(fig, dest, "historical_stage_accuracy_macro_f1")


def confusion_figure(dest, r2):
    matrix = np.asarray(r2["confusion_matrix"])
    support = matrix.sum(1)
    normalized = matrix / support[:, None]
    fig, axes = plt.subplots(1, 2, figsize=(13.2, 6.1))
    for ax, data, title in zip(axes, [matrix, normalized], ["OOF counts", "Row-normalized recall"]):
        image = ax.imshow(data, cmap="Blues", vmin=0, vmax=int(matrix.max()) if data is matrix else 1)
        ax.set_xticks(range(7), CLASSES)
        ax.set_yticks(range(7), [f"{c} (n={n})" for c, n in zip(CLASSES, support)])
        ax.set_xlabel("Predicted anonymous origin group")
        ax.set_ylabel("True anonymous origin group")
        ax.set_title(title)
        ax.get_yticklabels()[1].set_color("#A61C2B")
        ax.get_yticklabels()[1].set_fontweight("bold")
        for i in range(7):
            for j in range(7):
                value = data[i, j]
                label = str(int(value)) if data is matrix else f"{100*value:.1f}%"
                ax.text(j, i, label, ha="center", va="center", fontsize=9,
                        color="white" if value > (matrix.max() * .55 if data is matrix else .55) else "#17202A")
        fig.colorbar(image, ax=ax, shrink=.74, pad=.03)
    fig.suptitle("R2 MC-UVE + SVM: 821 / 855 correct (96.02%)", fontsize=14)
    fig.text(.5, .018, "ZY02: only 5 specimens, 3 correct; overall accuracy is not the accuracy of each class.\n"
             "Retrospective within-batch OOF; origins remain anonymous and batch-generalization is untested.",
             ha="center", fontsize=9)
    fig.tight_layout(rect=[0, .1, 1, .94])
    save_figure(fig, dest, "r2_confusion_counts_and_recall")


def feature_figure(dest, outer_metrics):
    r2 = sorted((row for row in outer_metrics if row["method"] == "R2"), key=lambda row: row["outer_fold"])
    assert len(r2) == 5
    sets = [set(row["selected_indices"]) for row in r2]
    assert [len(values) for values in sets] == [92, 62, 123, 123, 92]
    jac = np.mean([len(sets[i] & sets[j]) / len(sets[i] | sets[j]) for i in range(5) for j in range(i+1, 5)])
    frequency_count = Counter()
    for row in r2:
        frequency_count.update(row["selected_frequencies"])
    fig = plt.figure(figsize=(12.4, 7.5))
    gs = fig.add_gridspec(2, 2, width_ratios=[1, 2.2], height_ratios=[1, .72])
    count_ax, raster_ax, freq_ax = fig.add_subplot(gs[0, 0]), fig.add_subplot(gs[0, 1]), fig.add_subplot(gs[1, :])
    counts = [len(values) for values in sets]
    count_ax.bar(range(1, 6), counts, color="#147D64", width=.58)
    count_ax.set_xticks(range(1, 6)); count_ax.set_xlabel("Outer fold")
    count_ax.set_ylim(0, 150); count_ax.set_ylabel("Selected original frequency points")
    for x, n in enumerate(counts, 1):
        count_ax.text(x, n+3, str(n), ha="center")
    for row in r2:
        raster_ax.scatter(row["selected_frequencies"], [row["outer_fold"]+1] * row["selected_features"],
                          marker="|", s=92, color="#147D64", linewidths=1)
    raster_ax.set_yticks(range(1, 6)); raster_ax.set_ylim(.5, 5.5)
    raster_ax.set_xlim(.1, 1.5); raster_ax.set_ylabel("Outer fold")
    raster_ax.set_xlabel("Original frequency coordinate (THz)")
    raster_ax.set_title("Each fold selects only from its source data")
    coordinates = np.asarray(sorted(frequency_count))
    freq_ax.vlines(coordinates, 0, [frequency_count[float(x)] for x in coordinates], color="#4C78A8", linewidth=.8)
    freq_ax.set_xlim(.1, 1.5); freq_ax.set_ylim(0, 5.6); freq_ax.set_yticks(range(6))
    freq_ax.set_xlabel("Original frequency coordinate (THz)")
    freq_ax.set_ylabel("Number of folds selecting point")
    freq_ax.set_title(f"Mean K = {np.mean(counts):.1f} / 612; descriptive pairwise Jaccard = {jac:.4f}")
    fig.suptitle("R2 frequency selection: dimension and source-resampling description", fontsize=14)
    fig.text(.5, .015, "Overlapping fold-training samples do not form independent batches. Selection frequency is not a chemical-peak assignment.",
             ha="center", fontsize=9)
    fig.tight_layout(rect=[0, .055, 1, .94])
    save_figure(fig, dest, "r2_selected_frequency_count_and_stability")
    return r2, float(jac)


def flow_figure(dest):
    mermaid = '''flowchart TD
    D["酸枣仁主任务：855编号独立样品声明；七匿名产地组；每组仅一批"] --> P["固定数据与旧5外折/2内折；训练内选参"]
    P --> C["核心约70%：MC-UVE与SVM优化、消融、LDA/MLP公平比较"]
    C --> L["每外折冻结选择；外层一次评分；保留负结果"]
    L --> E["855份OOF指标、错误分析、频点稳定性与计算成本"]
    E --> X["辅助约20%：黄芪/人参等独立分类任务"]
    X --> G{"独立样品、标签、许可与物理轴兼容证据充分？"}
    G -- 否 --> B["当前阻塞：补标注/可访问数据，不生成虚构迁移分数"]
    G -- 是 --> T["辅助约10%：跨植物A/B/C迁移；相同酸枣仁query和从头训练对照"]
    E --> K["单列513/171/171：1-shot保守校准与零标签控制；非跨批次"]
    T --> N["真正独立新批次：未来外部确认与跨批次校准"]
    K --> N
'''
    (dest / "paper_c_experimental_flow.mmd").write_text(mermaid, encoding="utf-8")
    dot = '''digraph PaperC {
      graph [rankdir=TB,bgcolor="white",fontname="DejaVu Sans",nodesep=.36,ranksep=.48,pad=.15];
      node [shape=box,style="rounded,filled",fillcolor="#EFF5F8",color="#668497",fontname="DejaVu Sans",fontsize=11,margin=".16,.11"];
      edge [color="#627784",arrowsize=.7,fontname="DejaVu Sans",fontsize=10];
      data [label="Acid jujube kernel main task\n855 declared-independent specimens, 7 anonymous groups\nOne source batch per origin; no pristine confirmation set"];
      source [label="Fixed data + original 5 outer / 2 inner folds\nSource-only preprocessing, selection and tuning"];
      core [label="CORE: about 70% of planned effort\nMC-UVE + SVM optimization / ablations\nFair SVM, shrinkage LDA and MLP comparison",fillcolor="#DDEFE8",color="#147D64",penwidth=2];
      lock [label="Freeze each outer fold's inner choices\nThen score outer predictions; preserve negative results"];
      evidence [label="855 OOF accuracy, Macro-F1, balanced accuracy\nClass errors, frequency stability and measured cost"];
      ext [label="AUXILIARY: about 20%\nSeparate Huangqi / ginseng / plant THz tasks\nNo mixing of plant labels into jujube origin labels"];
      gate [shape=diamond,label="Sample identity, labels, license\nand physical-axis compatibility verified?",fillcolor="#FFF1D6"];
      blocked [label="Currently blocked where evidence is missing\nNo invented samples or transfer results",fillcolor="#FBE8E7",color="#A64842"];
      trans [label="AUXILIARY: about 10% transfer + calibration\nA: jujube only; B: other-plant pretraining\nC: multi-plant pretraining; same jujube evaluation"];
      calibration [label="Separate 513 / 171 / 171 protocol\n1-shot conservative calibration + zero-label controls\nRetrospective within-batch only"];
      future [label="Future independent new batches\nExternal confirmation + real cross-batch calibration",style="rounded,dashed",fillcolor="white"];
      data -> source -> core -> lock -> evidence;
      evidence -> ext -> gate;
      gate -> blocked [label="No"];
      gate -> trans [label="Yes; not yet established",style=dashed];
      evidence -> calibration;
      trans -> future [style=dashed]; calibration -> future [style=dashed];
    }'''
    (dest / "paper_c_experimental_flow.dot").write_text(dot, encoding="utf-8")
    for extension in ("svg", "png", "pdf"):
        subprocess.run(["dot", f"-T{extension}", str(dest / "paper_c_experimental_flow.dot"),
                        "-o", str(dest / f"paper_c_experimental_flow.{extension}")], check=True)


def benchmark_figures(dest, table):
    primary = [next(row for row in table if row["method"] == name) for name in PRIMARY_METHODS if any(row["method"] == name for row in table)]
    ablations = [row for row in table if row["method"].startswith("R2_")]
    for rows, stem, title in [(primary, "unified_classifier_comparison", "Classifiers under the same 855-record OOF split protocol"),
                              (ablations, "r2_conditional_ablation", "R2 conditional ablation controls (not independent causal effects)")]:
        fig, axes = plt.subplots(1, 2, figsize=(14.8, max(6, .61*len(rows)+2)))
        y = np.arange(len(rows))
        for ax, key, maximum, xlabel in [(axes[0], "accuracy", 100, "Accuracy (%)"),
                                         (axes[1], "macro_f1", 1, "Macro-F1")]:
            values = [row[key] * (100 if key == "accuracy" else 1) for row in rows]
            ax.barh(y, values, color=["#147D64" if row["method"] == "R2_reproduction" else "#4C78A8" for row in rows], height=.6)
            ax.set_yticks(y, [row["label"] for row in rows]); ax.invert_yaxis()
            ax.set_xlim(0, maximum * 1.12); ax.set_xlabel(xlabel)
            ax.set_xticks(np.linspace(0, maximum, 6)); ax.grid(axis="x", alpha=.2); ax.set_axisbelow(True)
            for pos, value in enumerate(values):
                ax.text(value + maximum*.012, pos, f"{value:.2f}" if key == "accuracy" else f"{value:.4f}", va="center", fontsize=9)
        fig.suptitle(title, fontsize=13)
        note = "Retrospective within-batch OOF, not independent confirmation. Search budgets / representations differ."
        if stem == "unified_classifier_comparison":
            note += "\nModel seed42 is primary; all three prespecified strong-MLP seeds are reported separately."
        fig.text(.5, .015, note, ha="center", fontsize=9)
        fig.tight_layout(rect=[0, .09, 1, .94])
        save_figure(fig, dest, stem)


def workbook(dest, tables, note_rows):
    book = Workbook()
    book.remove(book.active)
    expected = {}
    for name, rows in tables.items():
        sheet = book.create_sheet(name)
        keys = list(dict.fromkeys(key for row in rows for key in row))
        sheet.append(keys)
        for row in rows:
            sheet.append([row.get(key) for key in keys])
        expected[name] = [[row.get(key) for key in keys] for row in rows]
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 36
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244F70")
            cell.alignment = Alignment(wrap_text=True)
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 22 if col[0].value not in {"label", "protocol", "role"} else 54
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                key = keys[cell.column-1]
                if key in {"accuracy", "balanced_accuracy", "relative_accuracy_gain_vs_R0", "accuracy_mean", "accuracy_sd",
                           "balanced_accuracy_mean", "balanced_accuracy_sd", "support_evidence_rejection_rate"}:
                    cell.number_format = "0.00%"
                elif key in {"macro_f1", "precision", "recall", "f1", "macro_f1_mean", "macro_f1_sd"}:
                    cell.number_format = "0.0000"
                elif key == "accuracy_gain_pp_vs_R0" or key.endswith("_pp_vs_R2"):
                    cell.number_format = "+0.00;-0.00;0.00"
                elif key == "macro_f1_change_vs_R2":
                    cell.number_format = "+0.0000;-0.0000;0.0000"
                elif key.endswith("_seconds"):
                    cell.number_format = "0.000"
                if "metric" in keys and key in {"seed42", "mean_all3_seeds", "sample_std_all3_seeds"}:
                    cell.number_format = "0.0000" if row[keys.index("metric")].value == "macro_f1" else "0.00%"
    sheet = book.create_sheet("Notes")
    for row in note_rows:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 130
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    path = dest / "paper_c_model_tables.xlsx"
    book.save(path)
    check = load_workbook(path, data_only=True, read_only=True)
    cells = 0
    for name, values in expected.items():
        found = list(check[name].values)[1:]
        assert len(found) == len(values)
        for a, e in zip(found, values, strict=True):
            for actual, reference in zip(a, e, strict=True):
                assert abs(actual-reference) < 1e-12 if isinstance(reference, float) else actual == reference
                cells += 1
    check.close()
    return cells


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=ROOT / "reports/paper-c/2026-10-31/paper-assets")
    parser.add_argument("--benchmark-dir", type=Path)
    parser.add_argument("--strong-mlp-dir", type=Path)
    parser.add_argument("--calibration-dir", type=Path)
    args = parser.parse_args()
    if args.strong_mlp_dir and not args.benchmark_dir:
        parser.error("--strong-mlp-dir requires a completed --benchmark-dir for the joint classifier table")
    dest = args.output.resolve()
    dest.mkdir(parents=True, exist_ok=True)
    source_paths = [HISTORY / "core/outer_predictions.csv.gz", HISTORY / "combinations/outer_predictions.csv.gz",
                    HISTORY / "tables/四轮OOF台账.csv", HISTORY / "core/outer_metrics.json",
                    HISTORY / "core/protocol_snapshot.json", HISTORY / "combinations/protocol_snapshot.json",
                    Path(__file__).resolve()]
    inputs = {str(p.relative_to(ROOT)): sha(p) for p in source_paths}
    historical = group_predictions(read_csv(source_paths[0]) + read_csv(source_paths[1]))
    metrics = {method: scores(rows) for method, rows in historical.items()}
    ledger = {row["方法"]: row for row in read_csv(source_paths[2])}
    assert set(metrics) == set(ledger) and len(metrics) == 20
    metric_checks = 0
    for method, value in metrics.items():
        for key, original in [("accuracy", "准确率"), ("macro_f1", "宏F1"), ("balanced_accuracy", "平衡准确率")]:
            assert abs(value[key] - float(ledger[method][original])) < 1e-12
            metric_checks += 1
        assert value["correct"] == int(ledger[method]["正确数"])
    assert [metrics[m]["correct"] for m in STAGES] == [773, 791, 821, 816, 812]
    baseline = metrics["R0"]
    historical_table = [metric_row(method, metrics[method], baseline, "prespecified_stage" if method in STAGES else "prespecified_diagnostic") for method in ledger]
    core_id = json.loads(source_paths[4].read_text())["protocol_id"]
    combination_id = json.loads(source_paths[5].read_text())["protocol_id"]
    for row in historical_table:
        assert row["protocol_id"] == (combination_id if row["method"].startswith("R4") else core_id)
    stage_table = [next(row for row in historical_table if row["method"] == method) for method in STAGES]
    preprocess = [row for row in historical_table if row["method"].startswith("R3_diagnostic_")]
    combinations = [row for row in historical_table if row["method"].startswith("R4_diagnostic_")]
    class_table = [{"method": method, "class": c, **value["per_class"][c]} for method, value in metrics.items() for c in CLASSES]
    confusion_table = [{"true_class": c, **dict(zip(CLASSES, metrics["R2"]["confusion_matrix"][i]))} for i, c in enumerate(CLASSES)]
    stage_figure(dest, metrics)
    confusion_figure(dest, metrics["R2"])
    outer_metrics = json.loads(source_paths[3].read_text())
    r2_folds, jaccard = feature_figure(dest, outer_metrics)
    feature_table = [{"outer_fold_1based": row["outer_fold"] + 1, "selected_points": row["selected_features"],
                      "mc_runs": row["config"]["mc"]["runs"], "pls_components": row["config"]["mc"]["components"],
                      "retained_ratio": row["config"]["mc"]["ratio"], "C": row["config"]["svm"]["C"],
                      "gamma_rule": row["config"]["svm"]["gamma"], "effective_gamma": row["effective_gamma"],
                      "class_weight": row["config"]["svm"]["class_weight"], "preprocessing": row["config"]["preprocessing"]} for row in r2_folds]
    flow_figure(dest)
    tables = {"Historical_stages": stage_table, "Preprocessing_diagnostics": preprocess,
              "Combination_diagnostics": combinations, "Historical_all20": historical_table,
              "Class_metrics": class_table, "R2_confusion_counts": confusion_table,
              "R2_fold_parameters": feature_table}
    new_comparison = []
    if args.benchmark_dir:
        bench = args.benchmark_dir.resolve()
        paths = [bench / "outer_predictions.csv.gz", bench / "pooled_metrics.json", bench / "outer_metrics.json", bench / "protocol_snapshot.json"]
        assert all(p.is_file() for p in paths), "Only completed outer benchmark results may be plotted"
        inputs.update({str(p.relative_to(ROOT)): sha(p) for p in paths})
        recorded = json.loads(paths[1].read_text())
        benchmark_protocol = json.loads(paths[3].read_text())
        predicted = group_predictions(read_csv(paths[0]))
        assert set(recorded) == set(predicted)
        r0_rows = {r["record_id"]: r for r in historical["R0"]}
        for method, values in predicted.items():
            for row in values:
                assert row["outer_fold"] == r0_rows[row["record_id"]]["outer_fold"]
                assert row["y_true"] == r0_rows[row["record_id"]]["y_true"]
            computed = scores(values)
            for key in ("accuracy", "macro_f1", "balanced_accuracy", "correct", "total"):
                assert abs(computed[key] - recorded[method][key]) < 1e-12
                metric_checks += 1
            assert computed["confusion_matrix"] == recorded[method]["confusion_matrix"]
            role = "classifier_search" if method in {"full_svm", "shrinkage_lda", "mlp"} else ("reference_reproduction" if method == "R2_reproduction" else "conditional_ablation")
            row = metric_row(method, computed, baseline, role)
            row["protocol_id"] = benchmark_protocol["protocol_id"]
            row.update(paired_r2_change(values, historical["R2"], computed, metrics["R2"]))
            row.update(recorded[method]["cost"])
            row["cost_scope"] = ("Fresh source-inner search plus fresh outer fit" if role == "classifier_search" else
                                 "Historical per-fold source choices reused; zero new inner-search time is not zero original tuning cost")
            new_comparison.append(row)
            for c in CLASSES:
                class_table.append({"method": method, "class": c, **computed["per_class"][c]})
        reproduced = {r["record_id"]: r["y_predicted"] for r in predicted["R2_reproduction"]}
        assert reproduced == {r["record_id"]: r["y_predicted"] for r in historical["R2"]}
        tables["Conditional_ablations"] = [row for row in new_comparison if row["method"].startswith("R2_")]
    strong_rows = []
    strong_pair_rows = []
    if args.strong_mlp_dir:
        strong = args.strong_mlp_dir.resolve()
        paths = [strong / "outer_predictions.csv.gz", strong / "pooled_metrics_by_seed.json",
                 strong / "primary_seed42_metrics.json", strong / "seed_sensitivity.json", strong / "protocol_snapshot.json"]
        assert all(p.is_file() for p in paths), "Only completed strong-MLP outer results may be plotted"
        inputs.update({str(p.relative_to(ROOT)): sha(p) for p in paths})
        all_strong = read_csv(paths[0])
        recorded = json.loads(paths[1].read_text())
        primary = json.loads(paths[2].read_text())
        recorded_sensitivity = json.loads(paths[3].read_text())
        strong_protocol = json.loads(paths[4].read_text())
        r0_roles = {r["record_id"]: (r["y_true"], r["outer_fold"]) for r in historical["R0"]}
        assert {int(r["seed"]) for r in all_strong} == {42, 43, 44}
        for seed in (42, 43, 44):
            grouped = group_predictions([r for r in all_strong if int(r["seed"]) == seed])
            assert set(grouped) == set(recorded) == set(primary)
            for method, values in grouped.items():
                assert {r["record_id"]: (r["y_true"], r["outer_fold"]) for r in values} == r0_roles
                computed = scores(values)
                for key in ("accuracy", "macro_f1", "balanced_accuracy", "correct", "total"):
                    assert abs(computed[key] - recorded[method][str(seed)][key]) < 1e-12
                    if seed == 42:
                        assert abs(computed[key] - primary[method][key]) < 1e-12
                    metric_checks += 1
                assert computed["confusion_matrix"] == recorded[method][str(seed)]["confusion_matrix"]
                row = metric_row(method, computed, baseline, "predeclared_seed42_primary" if seed == 42 else "prespecified_seed_sensitivity")
                row["protocol_id"] = strong_protocol["protocol_id"]
                row["seed"] = seed
                row.update(paired_r2_change(values, historical["R2"], computed, metrics["R2"]))
                row.update(recorded[method][str(seed)].get("cost", {}))
                row["cost_scope"] = "Fixed strong-MLP protocol; paired200-only/continued models share base training; per-seed source diagnostics recorded separately"
                strong_rows.append(row)
                if seed == 42:
                    new_comparison.append(dict(row))
                for c in CLASSES:
                    class_table.append({"method": f"{method}_seed{seed}", "class": c, **computed["per_class"][c]})
            base_method, continued_method = "strong_mlp_200", "strong_mlp_200_plus100_source_updates"
            base = {r["record_id"]: r for r in grouped[base_method]}
            corrected = harmed = 0
            for row in grouped[continued_method]:
                old = base[row["record_id"]]
                old_right = old["y_predicted"] == old["y_true"]
                new_right = row["y_predicted"] == row["y_true"]
                corrected += int(not old_right and new_right)
                harmed += int(old_right and not new_right)
            before = recorded[base_method][str(seed)]
            after = recorded[continued_method][str(seed)]
            assert after["correct"]-before["correct"] == corrected-harmed
            strong_pair_rows.append({"seed":seed,"base_correct":before["correct"],"continued_correct":after["correct"],
                                      "accuracy_gain_pp_from_source_updates":100*(after["accuracy"]-before["accuracy"]),
                                      "macro_f1_change_from_source_updates":after["macro_f1"]-before["macro_f1"],
                                      "corrected":corrected,"harmed":harmed,"net_correct":corrected-harmed,
                                      "protocol_id":strong_protocol["protocol_id"],
                                      "scope":"Paired200epochs vs same base+100 source-supervised updates; not target-label calibration"})
        tables["Strong_MLP_all3seeds"] = strong_rows
        tables["Strong_MLP_source_updates"] = strong_pair_rows
        sensitivity = []
        for method in recorded:
            selected = [r for r in strong_rows if r["method"] == method]
            assert len(selected) == 3
            for key in ("accuracy", "macro_f1", "balanced_accuracy"):
                values = [row[key] for row in selected]
                mean, std = float(np.mean(values)), float(np.std(values, ddof=1))
                assert abs(mean - recorded_sensitivity[method][key]["mean"]) < 1e-12
                assert abs(std - recorded_sensitivity[method][key]["sample_std"]) < 1e-12
                metric_checks += 2
                sensitivity.append({"method": method, "metric": key, "seed42": values[0],
                                    "mean_all3_seeds": mean, "sample_std_all3_seeds": std,
                                    "seeds": "42,43,44", "scope": "Repeated predictions on same855; not independent studies"})
        tables["Strong_MLP_seed_sensitivity"] = sensitivity
    if new_comparison:
        tables["Unified_classifiers"] = [next(row for row in new_comparison if row["method"] == name) for name in PRIMARY_METHODS if any(row["method"] == name for row in new_comparison)]
        benchmark_figures(dest, new_comparison)
    calibration_rows = []
    if args.calibration_dir:
        calibration = args.calibration_dir.resolve()
        paths = [calibration / "抽样汇总.csv", calibration / "逐抽样指标.csv",
                 calibration / "table_validation.json", calibration / "run/protocol_snapshot.json"]
        assert all(p.is_file() for p in paths)
        inputs.update({str(p.relative_to(ROOT)): sha(p) for p in paths})
        assert json.loads(paths[2].read_text())["status"] == "pass"
        calibration_protocol = json.loads(paths[3].read_text())
        per_draw = read_csv(paths[1])
        assert len(per_draw) == 1200 and {int(row["total"]) for row in per_draw} == {171}
        for summary in read_csv(paths[0]):
            selected = [r for r in per_draw if r["selector"] == summary["selector"] and r["method"] == summary["method"]]
            assert len(selected) == int(summary["support_draws"]) == 100
            for metric in ("accuracy", "macro_f1", "balanced_accuracy", "corrected", "harmed"):
                values = np.asarray([float(r[metric]) for r in selected])
                for suffix, actual in [("mean", np.mean(values)), ("sd", np.std(values, ddof=0)),
                                       ("min", values.min()), ("max", values.max())]:
                    assert abs(actual-float(summary[f"{metric}_{suffix}"])) < 1e-12
                    metric_checks += 1
            record = {key: summary[key] for key in ("selector", "method")}
            record.update({key: float(summary[key]) for key in ("accuracy_mean", "accuracy_sd", "macro_f1_mean", "macro_f1_sd",
                          "balanced_accuracy_mean", "balanced_accuracy_sd", "corrected_mean", "harmed_mean", "support_evidence_rejection_rate")})
            record.update({"support_draws":100,"fixed_query_samples":171,"positive_draws":int(summary["positive_draws"]),
                           "negative_draws":int(summary["negative_draws"]),"development_gate_fallback":summary["development_gate_fallback"],
                           "protocol_id":calibration_protocol["protocol_id"],
                           "scope":"Separate513/171/171; seven-class1-shot;100 overlapping support draws, not independent studies"})
            calibration_rows.append(record)
        assert len(calibration_rows) == 12
        tables["Calibration171_only"] = calibration_rows
    for name, rows in tables.items():
        write_csv(dest / f"{name}.csv", rows)
    notes = [("Population", SCOPE), ("Baseline", "R0 773/855; relative gains use unrounded baseline accuracy."),
             ("Gain units", "accuracy_gain_pp_vs_R0 = 100*(accuracy-R0); relative_accuracy_gain_vs_R0 = (accuracy-R0)/R0."),
             ("Rare class", "ZY02 has only five specimens; R2 identifies three. No added shots or batches are invented."),
             ("Stage interpretation", "R1 vs fixed SNV/C10/MC20 anchor is the pure MC-stage comparison. Stage changes vs R0 are not isolated causal effects."),
             ("Scope separation", "Old fixed171 prediction results are excluded from855 OOF comparison."),
             ("Selection stability", "Overlapping source-fold Jaccard is descriptive; not independent-batch or chemical-peak validation."),
             ("External tasks", "External independence/labels/license/axis compatibility remain eligibility gates; A/B/C transfer is not claimed completed."),
             ("New benchmark", "Included only when --benchmark-dir names completed outer artifacts; source-only choices and original folds checked."),
             ("Cost interpretation", "New benchmark outer pipelines fit fresh transforms/selection per method; inner candidates share transforms within the same source partition. Timing is measured wall time on shared CPU, not an exclusive-hardware speed benchmark."),
             ("R2 search cost", "R2_reproduction/conditional controls reuse frozen historical per-fold choices; inner_search_seconds=0 means no new search in this replay, not zero historical MC/SVM tuning cost."),
             ("Strong MLP", "The64-32 architecture retains200 base epochs plus100 additional source-supervised updates, not300 epochs. Paired200-only/continued methods share the base state. Seed42 is predeclared primary;42/43/44 and descriptive sample SD all retained, not independent studies."),
             ("Architecture exposure", "The fixed strong-MLP architecture came from the earlier exposed513-source study. Its30 current inner contexts are diagnostics, not an unbiased current architecture search; same-fold replay remains retrospective."),
             ("SD conventions", "Strong-MLP seed sensitivity uses sample SD(ddof1); the separate calibration support-draw summary retains its recorded descriptive population SD(ddof0). Neither is an independent-study confidence interval."),
             ("Computation", "Export reads stored labels/predictions only; new model fits =0.")]
    cells = workbook(dest, tables, notes)
    (dest / "figure_source_metrics.json").write_text(json.dumps({"historical_metrics": metrics, "benchmark_metrics": new_comparison,
                              "strong_mlp_all_seeds": strong_rows, "r2_pairwise_jaccard": jaccard,
                              "r2_fold_selected_points": [r["selected_features"] for r in r2_folds]}, ensure_ascii=False, indent=2)+"\n")
    readme = "# 方案C论文图表证据\n\n本目录仅从已保存的真实预测与真实标签导出图表，没有新增训练或重新选择参数。酸枣仁是主任务，外部植物及迁移是辅助。\n\n"
    readme += "[论文模型表Excel](paper_c_model_tables.xlsx)含历史五阶段、全部20角色、预处理与组合负对照、逐类指标、R2混淆计数及五折参数。"
    if new_comparison:
        readme += "另含同855样品和相同划分的新SVM／LDA／MLP统一对照及R2条件消融；不同搜索预算与表示不称严格单因素因果。"
    if strong_rows:
        readme += "强64–32 MLP按预定seed42列入主表，全部42／43／44种子及均值、样本标准差另列附表；不挑最高种子。"
        readme += "强MLP旧架构来自已曝光的513源研究，当前30个内层场景仅作诊断，不声称完整架构选择的无偏嵌套确认；200轮加100次源更新的配对改对／改错另列。"
    if calibration_rows:
        readme += "171份固定query的校准结果只放Calibration171_only独立工作表，未与855 OOF混排；100次支持抽样是同一query上的重叠重复。"
    readme += "\n\n| 图 | 内容 |\n| --- | --- |\n"
    for stem, label in [("historical_stage_accuracy_macro_f1", "原五阶段准确率与Macro-F1"),
                        ("r2_confusion_counts_and_recall", "R2计数和行归一混淆矩阵；ZY02仅5份"),
                        ("r2_selected_frequency_count_and_stability", "R2每折K、频点和描述性稳定性"),
                        ("paper_c_experimental_flow", "70%主任务、20%外部任务、10%迁移校准及证据门禁流程")]:
        readme += f"| [{label}]({stem}.png) | [PDF]({stem}.pdf) |\n"
    if new_comparison:
        readme += "| [统一分类器对照](unified_classifier_comparison.png) | [PDF](unified_classifier_comparison.pdf) |\n| [R2条件消融](r2_conditional_ablation.png) | [PDF](r2_conditional_ablation.pdf) |\n"
    readme += "\n图中百分比均明确标为准确率，Macro-F1使用0—1尺度；不截断纵轴夸大差距。没有按外测最高值回填选择。所有855份与历史外测此前已曝光，且每产地只有一批，因此是回顾性批次内验证，不能代表真正跨批次或新独立确认。\n\n"
    readme += "流程图中的比例是工作量规划，不是证据权重或论文固定文字比例。公开THz数据及黄芪缺少的样品身份／标签／单位信息均是门禁；虚线不是已完成迁移实验。\n\n"
    readme += "频率轴按当前酸枣仁输入的原坐标标THz；频点选择不等于化学峰归属，也不推断其他植物或仪器的物理量兼容。\n\n"
    readme += "再生：`.runtime/bin/python scripts/export_paper_c_evidence.py`；有已完成新对照时增加`--benchmark-dir <completed-result-directory>`，强MLP增加`--strong-mlp-dir <completed-strong-mlp-directory>`，单列校准增加`--calibration-dir <completed-calibration-directory>`。输入SHA和独立重算记录见[验证记录](export_validation.json)。\n"
    (dest / "README.md").write_text(readme, encoding="utf-8")
    for path, expected in inputs.items():
        assert sha(ROOT / path) == expected
    output_shas = {str(p.relative_to(dest)): sha(p) for p in sorted(dest.rglob("*")) if p.is_file() and p.name != "export_validation.json"}
    validation = {"status": "pass", "new_model_fits": 0, "input_sha256": inputs,
                  "historical_method_roles": len(metrics), "historical_oof_rows": sum(len(v) for v in historical.values()),
                  "historical_correct_counts": {m: metrics[m]["correct"] for m in STAGES},
                  "metric_cross_checks": metric_checks, "excel_cells_read_back": cells,
                  "benchmark_primary_and_control_roles": len(new_comparison), "strong_mlp_seed_roles": len(strong_rows),
                  "separate_171_calibration_summaries": len(calibration_rows),
                  "outputs_sha256": output_shas}
    (dest / "export_validation.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2)+"\n")
    print(json.dumps({k:v for k,v in validation.items() if k not in {"input_sha256", "outputs_sha256"}}, ensure_ascii=False))


if __name__ == "__main__":
    main()
