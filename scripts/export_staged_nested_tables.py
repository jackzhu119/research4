"""Export completed staged nested-CV evidence without fitting or selection.

All metrics below are recomputed from stored vote predictions. OVR argmax is a
separate diagnostic and must never replace SVC.predict for historical matches.
"""
import argparse
from collections import Counter, defaultdict
import csv
import gzip
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


CLASSES = tuple(f"ZY{i:02d}" for i in range(1, 8))
METRICS = ("accuracy", "balanced_accuracy", "macro_f1")


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def compact(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def metrics(truth, predicted):
    matrix = np.zeros((7, 7), dtype=int)
    for actual, pred in zip(truth, predicted):
        matrix[CLASSES.index(str(actual)), CLASSES.index(str(pred))] += 1
    support, found, tp = matrix.sum(1), matrix.sum(0), matrix.diagonal()
    recall = np.divide(tp, support, out=np.zeros(7), where=support > 0)
    precision = np.divide(tp, found, out=np.zeros(7), where=found > 0)
    f1 = np.divide(2*tp, support+found, out=np.zeros(7), where=support+found > 0)
    return {"accuracy": float(tp.sum()/matrix.sum()), "balanced_accuracy": float(recall.mean()),
        "macro_f1": float(f1.mean()), "correct": int(tp.sum()), "total": int(matrix.sum()),
        "confusion_matrix": matrix.tolist(), "recall": recall.tolist(), "precision": precision.tolist(),
        "f1": f1.tolist(), "support": support.tolist(), "predicted_count": found.tolist(),
        "tp": tp.tolist(), "fp": (found-tp).tolist(), "fn": (support-tp).tolist()}


def prediction_rows(path):
    opener = gzip.open if Path(path).suffix == ".gz" else open
    with opener(path, "rt", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def add_sheet(book, name, rows):
    if not rows:
        rows = [{"说明": "本项无已完成的适用记录；未生成模拟结果"}]
    sheet = book.create_sheet(name)
    keys = list(rows[0]); sheet.append(keys)
    for row in rows:
        sheet.append([row.get(k) for k in keys])
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.zoomScale = 80
    for cell in sheet[1]:
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244F70")
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
    for j, key in enumerate(keys, 1):
        sheet.column_dimensions[get_column_letter(j)].width = 22
        if key in {"说明", "候选参数", "入选参数", "类别权重", "原始记录", "源权重及编码器"}:
            sheet.column_dimensions[get_column_letter(j)].width = 72
        for row in sheet.iter_rows(min_row=2, min_col=j, max_col=j):
            cell = row[0]; cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F1F5F9")
            if any(token in key for token in ("准确率", "精确率", "召回率", "比例")):
                cell.number_format = "0.00%"
            elif "F1" in key or "Jaccard" in key:
                cell.number_format = "0.0000"
    if "是否正确" in keys:
        correct = keys.index("是否正确")
        for row in sheet.iter_rows(min_row=2):
            if row[correct].value == 0:
                for cell in row:
                    cell.fill = PatternFill("solid", fgColor="FCE4E4")
    sheet.row_dimensions[1].height = 38


def write_csv(path, rows):
    if not rows:
        rows = [{"说明": "本项无已完成的适用记录；未生成模拟结果"}]
    with Path(path).open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader(); writer.writerows(rows)


def outer_locks(lock):
    rows = lock.get("outer_folds", [])
    if not isinstance(rows, list) or len(rows) != 5:
        raise RuntimeError("Expected five independently locked outer folds")
    return rows


def choice_rows(lock):
    result = []
    for row in outer_locks(lock):
        fold = row["outer_fold"]
        choices = row.get("choices", {})
        if not choices and "choice" in row:
            choices = {"R4": row["choice"]}
        for stage, choice in choices.items():
            chosen = choice.get("selected", choice)
            mean = choice.get("mean_inner_metrics", choice.get("all_candidate_means", {}).get(chosen.get("id"), {}))
            full_chosen = {**row.get("inherited_config", {}), **chosen}
            result.append({"外折": fold, "阶段": stage, "候选编号": chosen.get("id", ""),
                "入选参数": compact(full_chosen), "内层平均宏F1": mean.get("macro_f1"),
                "内层平均平衡准确率": mean.get("balanced_accuracy"),
                "内层平均准确率": mean.get("accuracy"), "按外测选择": bool(choice.get("selected_using_outer_test", False))})
    return result


def export(base, output, input_path, plan_path):
    base, output, input_path, plan_path = map(Path, (base, output, input_path, plan_path))
    release_path = base/"outer_evaluation_release.json"
    if not release_path.exists():
        raise RuntimeError("所有内层选择尚未统一放行；拒绝导出不完整或模拟外测结果")
    release = read_json(release_path)
    records = release.get("rounds", [])
    runs = {r["name"]: Path(r["selection_lock"]).parent for r in records}
    if set(runs) != {"core", "combinations"}:
        raise RuntimeError("Expected both core and combinations selection locks")
    checks, maximum = 0, 0.

    def equal(actual, expected, label):
        nonlocal checks
        checks += 1
        if actual != expected:
            raise AssertionError(f"Mismatch: {label}")

    def close(actual, expected, label):
        nonlocal checks, maximum
        delta = float(np.max(np.abs(np.asarray(actual, dtype=float)-np.asarray(expected, dtype=float))))
        checks += 1; maximum = max(maximum, delta)
        if not np.isfinite(delta) or delta > 1e-12:
            raise AssertionError(f"Numeric mismatch: {label}: {delta}")

    provenance = {str(release_path): digest(release_path)}
    for record in records:
        run = runs[record["name"]]
        for path, expected in [(Path(record["selection_lock"]), record["selection_lock_sha256"]),
                (Path(record["protocol"]), record["protocol_sha256"]),
                (run/"run_manifest.json", record["run_manifest_sha256"])]:
            equal(digest(path), expected, str(path)); provenance[str(path)] = digest(path)
        if read_json(run/"status.json").get("status") != "completed":
            raise RuntimeError("Outer evaluation has not completed")
        for name in ("summary.json", "outer_metrics.json", "outer_predictions.csv.gz", "all_inner_metrics.json"):
            if not (run/name).exists():
                raise RuntimeError(f"Required completed evidence missing: {run/name}")
        manifest = read_json(run/"run_manifest.json")
        equal(manifest["input_sha256"], digest(input_path), "original input")
    equal(release["master_plan_sha256"], digest(plan_path), "master plan")
    equal(release["input_sha256"], digest(input_path), "release input")
    provenance[str(input_path)] = digest(input_path); provenance[str(plan_path)] = digest(plan_path)
    with np.load(input_path, allow_pickle=False) as data:
        truth = {str(identifier): str(label) for identifier, label in zip(data["groups"], data["y"])}
        equal(len(truth), 855, "independent physical IDs")
        axis = np.asarray(data["axis"]); axis = axis[(axis >= .1) & (axis <= 1.5)]
    equal(len(axis), 612, "frequency features")
    legacy_path = Path("outputs/approved-zaoren-20261031/experiments/alpha_primary/splits.json")
    splits = read_json(runs["core"]/"splits.json")
    equal(splits, read_json(legacy_path), "original five-outer/two-inner split")
    legacy_hash = release.get("legacy_splits_sha256", release.get("historical_splits_sha256", release.get("old_shared_splits_sha256")))
    if legacy_hash is None:
        raise RuntimeError("Release must bind the exact historical split hash")
    equal(digest(legacy_path), legacy_hash, "legacy split hash")
    fold_by_id, split_by_fold = {}, {}
    for split in splits:
        fold = split["fold"]; source, test = set(split["train_sample_ids"]), set(split["test_sample_ids"])
        equal(len(source), 684, "outer source count"); equal(len(test), 171, "outer test count")
        equal(source & test, set(), "outer roles disjoint")
        equal(source | test, set(truth), "outer complete")
        for identifier in test:
            if identifier in fold_by_id:
                raise AssertionError("Specimen assigned to multiple outer tests")
            fold_by_id[identifier] = fold
        for inner in split["inner_folds"]:
            a, b = set(inner["train_sample_ids"]), set(inner["validation_sample_ids"])
            equal(a & b, set(), "inner roles disjoint"); equal(a | b, source, "inner covers outer source")
            equal(a & test, set(), "inner source excludes outer test"); equal(b & test, set(), "inner validation excludes outer test")
        split_by_fold[fold] = split

    explanation = [
        {"项目": "评价口径", "说明": "855份用户声明独立样品，五外折／二内折；每方法每样品恰有一次外测预测，不混固定171份三用途实验。"},
        {"项目": "标签", "说明": "ZY01—ZY07匿名产地组，真实地名尚待用户补充；不生成模拟光谱或新样品。"},
        {"项目": "选择", "说明": "每外折分别在自己的684份源样品内调参，所有阶段均先选择锁定，再统一外测；不按外测分数重选方法。"},
        {"项目": "历史参照", "说明": "R0复算历史raw/SNV×C4内调MC20方案，原正确773/855、宏F1约0.809185。"},
        {"项目": "固定锚定", "说明": "fixed_snv_C10_MC20只作单因素控制，不能冒充历史8候选已调参基线。"},
        {"项目": "四轮", "说明": "R1 MC-UVE参数；R2继承该外折R1后调SVM；R3再比较7预处理；R4固定预定稳定选择／CNN对照。"},
        {"项目": "候选数", "说明": "主计划1540条内层候选／折记录；其中R0=80、R1=360、R2=960、R3=70、R4=70，别名不算新样品。"},
        {"项目": "预测规则", "说明": "SVM保存标准SVC.predict投票；score argmax不同仅作诊断，不用其替换真实投票预测。分数不是概率。"},
        {"项目": "gamma", "说明": "scale与auto在源标准化后常为数值别名；记录实际_gamma，不将别名误称新机制。"},
        {"项目": "MC前缀", "说明": "10／20次可复用同源50次系数流的前缀，前缀独立重算均值、标准差和稳定性；属于计算复用。"},
        {"项目": "类权重", "说明": "类别权重只使用当前源折计数并在内层选择；幂权重归一使训练样品平均权重为1。"},
        {"项目": "稳定性", "说明": "外折／源重抽样选点频率和Jaccard为描述性诊断，不是p值、FDR或跨批次证据。"},
        {"项目": "CNN", "说明": "CNN只在对应源训练区监督训练，分类头和SVM共享该源编码器；保持预算和来源可追溯。"},
        {"项目": "方向C", "说明": "真实跨批次小样本校准未运行：每产地目前只有一批，没有伪造批次、支持样品或跨批次识别率。"},
        {"项目": "外推边界", "说明": "全部855份历史上已用于研究，仍为同批次回顾性内部嵌套CV，不是新独立确认。ZY02仅5份，逐类指标很不稳定。"},
        {"项目": "不确定性", "说明": "五折分数标准差只作描述，不视为独立批次置信区间；没有挑最好折或最好种子。"},
        {"项目": "记录摘录", "说明": "表格原始JSON摘录最多30000字符，完整记录按保存路径与SHA追溯；编码器／预算／损失另有结构化表。"},
    ]
    stages, choices, inner_table, inner_mean, gamma_rows, weight_rows = [], [], [], [], [], []
    oof_table, fold_table, class_table, pair_table, changes = [], [], [], [], []
    frequencies, stability, budgets = [], [], []
    resample_stability, selection_budgets, neural_controls, neural_losses = [], [], [], []
    all_outer, outer_groups, all_inner = [], {}, []
    source_states = []
    non_frequency_methods = set()
    inner_raw_count = 0
    for run_name, run in runs.items():
        lock = read_json(run/"selection_lock.json")
        choices.extend([{**row, "实验分支": run_name} for row in choice_rows(lock)])
        inner_rows = read_json(run/"all_inner_metrics.json")
        all_inner.extend(inner_rows)
        inner_groups = defaultdict(list)
        for raw_path in sorted((run/"selection").rglob("predictions.csv.gz")):
            for row in prediction_rows(raw_path):
                stage = row.get("stage", "R4")
                key = (int(row["outer_fold"]), stage, int(row["inner_fold"]), row["candidate_id"])
                equal(row["y_true"], truth[row["record_id"]], "inner saved truth")
                equal(int(row["is_correct"]), int(row["y_true"] == row["y_predicted"]), "inner correct flag")
                inner_groups[key].append(row); inner_raw_count += 1
        if not inner_groups:
            raise RuntimeError("Completed inner raw predictions are required for export validation")
        by_inner = defaultdict(list)
        for row in inner_rows:
            stage = row.get("stage", "R4"); outer = int(row["outer_fold"]); inner = int(row["inner_fold"])
            candidate = row.get("candidate", row.get("config", {})); metric = row.get("metrics", row)
            source_ids = row.get("source_ids", []); valid_ids = row.get("validation_ids", [])
            raw = inner_groups.pop((outer, stage, inner, row["candidate_id"]))
            expected_valid = set(split_by_fold[outer]["inner_folds"][inner]["validation_sample_ids"])
            equal(len(raw), 342, "complete342 inner validation predictions")
            equal({r["record_id"] for r in raw}, expected_valid, "inner validation physical IDs")
            raw_metric = metrics([r["y_true"] for r in raw], [r["y_predicted"] for r in raw])
            for key in METRICS:
                close(raw_metric[key], metric[key], "inner raw recomputation")
            close(raw_metric["confusion_matrix"], metric["confusion_matrix"], "inner confusion")
            if source_ids and valid_ids:
                split = split_by_fold[outer]["inner_folds"][inner]
                equal(source_ids, split["train_sample_ids"], "inner row source IDs")
                equal(valid_ids, split["validation_sample_ids"], "inner row validation IDs")
            inner_table.append({"实验分支": run_name, "外折": outer, "阶段": stage, "内折": inner,
                "候选编号": row["candidate_id"], "候选参数": compact(candidate), "宏F1": metric["macro_f1"],
                "平衡准确率": metric["balanced_accuracy"], "准确率": metric["accuracy"],
                "变量数": len(row.get("selected_indices", [])) or row.get("selected_features"),
                "实际gamma": row.get("effective_gamma"), "投票与分数argmax不同数": row.get("vote_argmax_disagreements"),
                "实际秒数": row.get("seconds")})
            by_inner[(outer, stage, row["candidate_id"])].append(row)
            if "effective_gamma" in row:
                gamma_rows.append({"分支": run_name, "外折": outer, "阶段": stage, "内折": inner,
                    "候选编号": row["candidate_id"], "gamma规则": candidate.get("svm", {}).get("gamma", ""),
                    "实际gamma": row["effective_gamma"], "scale数值": row.get("scale_gamma"),
                    "auto数值": row.get("auto_gamma"), "scale与auto数值接近":
                    bool(np.isclose(row.get("scale_gamma", np.inf), row.get("auto_gamma", -np.inf), rtol=1e-12, atol=0))})
            if row.get("training_class_weights") and source_ids:
                counts = Counter(truth[i] for i in source_ids); weights = row["training_class_weights"]
                mean_weight = sum(counts[c]*weights[c] for c in counts)/len(source_ids)
                close(mean_weight, 1., "source classweights normalized")
                weight_rows.append({"分支": run_name, "外折": outer, "阶段": stage, "内折": inner,
                    "候选编号": row["candidate_id"], "权重规则": candidate.get("svm", {}).get("class_weight", ""),
                    "源类别数量": compact(counts), "类别权重": compact(weights), "训练样品平均权重": mean_weight})
        equal(len(inner_groups), 0, "no unaccounted inner cases")
        for (outer, stage, identifier), rows in by_inner.items():
            equal(len(rows), 2, "two inner cases per candidate")
            inner_mean.append({"实验分支": run_name, "外折": outer, "阶段": stage, "候选编号": identifier,
                "候选参数": compact(rows[0].get("candidate", rows[0].get("config", {}))),
                **{name: float(np.mean([r.get("metrics", r)[key] for r in rows]))
                   for key, name in [("macro_f1", "内层平均宏F1"), ("balanced_accuracy", "内层平均平衡准确率"), ("accuracy", "内层平均准确率")]}})
        stored_outer = read_json(run/"outer_metrics.json")
        preds = prediction_rows(run/"outer_predictions.csv.gz")
        provenance[str(run/"outer_predictions.csv.gz")] = digest(run/"outer_predictions.csv.gz")
        by_method = defaultdict(list); by_case = defaultdict(list)
        for row in preds:
            identifier = str(row["record_id"]); equal(row["y_true"], truth[identifier], "OOF truth")
            fold = int(row["outer_fold"]); equal(fold, fold_by_id[identifier], "OOF assigned fold")
            equal(int(row["is_correct"]), int(row["y_true"] == row["y_predicted"]), "correct flag")
            key = run_name+":"+row["method"]
            row = dict(row, run=run_name, key=key)
            by_method[row["method"]].append(row); by_case[(row["method"], fold)].append(row)
            all_outer.append(row)
            scores = [float(row.get("score_"+c, "nan")) for c in CLASSES]
            score_label = CLASSES[int(np.argmax(scores))] if np.isfinite(scores).all() else "未保存七类分数"
            oof_table.append({"实验分支": run_name, "方法": row["method"], "外折": fold, "样品编号": identifier,
                "匿名真实产地": row["y_true"], "保存预测产地": row["y_predicted"], "是否正确": int(row["is_correct"]),
                "保存七类分数argmax类别": score_label, "与保存预测不同": int(score_label != row["y_predicted"]) if score_label in CLASSES else None,
                **{"score_"+c: float(row["score_"+c]) if "score_"+c in row else None for c in CLASSES}})
        summaries = {r["method"]: r for r in read_json(run/"summary.json")}
        outer_map = {(r["method"], int(r["outer_fold"])): r for r in stored_outer}
        for method, rows in by_method.items():
            equal(len(rows), 855, "855 OOF rows per method")
            equal(len({r["record_id"] for r in rows}), 855, "one OOF prediction per physical sample")
            values = metrics([r["y_true"] for r in rows], [r["y_predicted"] for r in rows])
            for key in METRICS:
                close(values[key], summaries[method][key], "pooled OOF "+method+key)
            close(values["confusion_matrix"], summaries[method]["confusion_matrix"], "pooled OOF confusion")
            outer_groups[run_name+":"+method] = rows
            fold_records = [r for r in stored_outer if r["method"] == method]
            mean_original = float(np.mean([r["selected_features"] for r in fold_records]))
            mean_classifier = float(np.mean([r.get("classifier_features", r["selected_features"]) for r in fold_records]))
            stages.append({"实验分支": run_name, "方法": method, "评价身份": "内层预先锁定" if "diagnostic" not in method else "预先指定诊断，不按外测提升为主方法",
                "五折模型族": compact(Counter(r.get("family", "mcuve_rbf_svm") for r in fold_records)),
                "独立样品数": 855, "正确数": values["correct"], "准确率": values["accuracy"],
                "平衡准确率": values["balanced_accuracy"], "宏F1": values["macro_f1"],
                "平均原频点数": mean_original, "平均分类器表示维数": mean_classifier,
                "CNN表示提醒": "CNN使用612原输入并学习32维，不能把输入数当作筛选稳定性" if any("cnn" in str(r.get("family", method)) for r in fold_records) else "",
                "外测选参": False})
            for j, c in enumerate(CLASSES):
                class_table.append({"分支": run_name, "方法": method, "匿名产地组": c, "真实样品数": values["support"][j],
                    "预测为该类数量": values["predicted_count"][j], "TP": values["tp"][j], "FP": values["fp"][j], "FN": values["fn"][j],
                    "精确率": values["precision"][j], "召回率": values["recall"][j], "F1": values["f1"][j]})
        for (method, fold), rows in by_case.items():
            equal(len(rows), 171, "171 samples per outer case")
            values = metrics([r["y_true"] for r in rows], [r["y_predicted"] for r in rows])
            expected = outer_map[(method, fold)]
            for key in METRICS:
                close(values[key], expected[key], "outer fold metric")
            close(values["confusion_matrix"], expected["confusion_matrix"], "outer fold confusion")
            config = expected.get("config", expected.get("candidate", {}))
            fitted_meta = {}
            if expected.get("family"):
                checkpoint_path = Path(expected["model_path"]).parent/"source_checkpoint.json"
                checkpoint = read_json(checkpoint_path)
                artifact = checkpoint["artifacts"][expected["family"]]
                equal(artifact["sha256"], expected["model_sha256"], "outer role binds saved family model")
                fitted_meta = artifact["metadata"]
            fold_table.append({"分支": run_name, "方法": method, "外折": fold,
                "模型族": expected.get("family", "mcuve_rbf_svm"), "源训练数": 684, "外测数": 171,
                "正确数": values["correct"], "准确率": values["accuracy"], "平衡准确率": values["balanced_accuracy"], "宏F1": values["macro_f1"],
                "原频点数": expected.get("selected_features"),
                "分类器表示维数": expected.get("classifier_features", expected.get("embedding_features", expected.get("embedding_dim"))),
                "实际gamma": expected.get("effective_gamma", fitted_meta.get("effective_gamma")), "入选参数": compact(config),
                "源权重及编码器": compact({k:v for k,v in expected.items() if any(t in k for t in ("encoder", "fingerprint", "model_sha", "training", "epoch", "seconds"))})})
            if fitted_meta.get("effective_gamma") is not None:
                dimension = fitted_meta["classifier_features"]; variance = fitted_meta["classifier_input_variance"]
                scale = 1/(dimension*variance) if variance else 1.
                gamma_rows.append({"分支": run_name, "外折": fold, "阶段": "外层最终模型", "内折": None,
                    "候选编号": method, "gamma规则": fitted_meta["svm"]["gamma"], "实际gamma": fitted_meta["effective_gamma"],
                    "scale数值": scale, "auto数值": 1/dimension,
                    "scale与auto数值接近": bool(np.isclose(scale, 1/dimension, rtol=1e-12, atol=0))})
                weights = fitted_meta["training_class_weights"]
                counts = Counter(truth[i] for i in split_by_fold[fold]["train_sample_ids"])
                mean_weight = sum(counts[c]*weights[c] for c in counts)/684
                close(mean_weight, 1., "outer source classweights normalized")
                weight_rows.append({"分支": run_name, "外折": fold, "阶段": "外层最终模型", "内折": None,
                    "候选编号": method, "权重规则": fitted_meta["svm"]["class_weight"],
                    "源类别数量": compact(counts), "类别权重": compact(weights), "训练样品平均权重": mean_weight})
            selected = expected.get("selected_indices", [])
            is_cnn = "cnn" in str(expected.get("family", method))
            if is_cnn:
                non_frequency_methods.add((run_name, method))
            for index in selected:
                equal(0 <= int(index) < 612, True, "selected original frequency index")
                frequencies.append({"分支": run_name, "方法": method, "外折": fold,
                    "原频点索引从0起": int(index), "原频率坐标": float(axis[int(index)]),
                    "变量性质": "CNN全谱输入，非选出的频率，自动表示另记" if is_cnn else "原光谱频率点，非化学峰确认"})
            if "selected_frequencies" in expected:
                close(axis[np.asarray(selected, dtype=int)], expected["selected_frequencies"], "frequency axis mapping")
        for row in stored_outer:
            budgets.append({"分支": run_name, "方法": row["method"], "外折": row["outer_fold"],
                "原始记录": compact({k:v for k,v in row.items() if k not in {"selected_indices", "selected_frequencies", "confusion_matrix", "class_recall", "class_counts"}})})
        for p in sorted(run.rglob("*.json")):
            if any(token in p.name for token in ("training", "encoder", "budget", "caller_common")) and not any(t in p.parts for t in ("tables", ".cache")):
                record = read_json(p)
                source_states.append({"分支": run_name, "记录路径": str(p), "SHA256": digest(p), "原始记录": compact(record)[:30000]})
                if p.name == "caller_common.json":
                    caller_ids = record.get("source_ids", [])
                    counts = Counter(truth[i] for i in caller_ids)
                    if caller_ids:
                        equal(len(caller_ids), len(set(caller_ids)), "caller unique source records")
                        outer_id = next(int(part.split("_")[-1]) for part in p.parts if part.startswith("outer_"))
                        inner_parts = [part for part in p.parts if part.startswith("inner_")]
                        caller_expected = (split_by_fold[outer_id]["inner_folds"][int(inner_parts[0].split("_")[-1])]["train_sample_ids"]
                            if inner_parts else split_by_fold[outer_id]["train_sample_ids"])
                        equal(caller_ids, caller_expected, "caller uses only assigned source partition")
                    stable = record.get("stability", {})
                    reconstructed = {"stable_mcuve": np.zeros(612), "stable_cars": np.zeros(612)}
                    for sample in stable.get("source_resamples", []):
                        ids = sample.get("source_ids", sample.get("record_ids", []))
                        if ids:
                            equal(set(ids).issubset(set(caller_ids)), True, "resample within caller source")
                            equal(len(ids), len(set(ids)), "resampling without physical duplication")
                            recounted = Counter(truth[i] for i in ids)
                            equal(dict(recounted), sample["class_counts"], "resample source label counts")
                            equal(min(recounted.values()) >= 2, True, "CARS internal2fold source class coverage")
                        for field, family in [("mcuve_indices", "stable_mcuve"), ("cars_indices", "stable_cars")]:
                            index = np.asarray(sample[field], dtype=int)
                            equal(len(index), len(set(index.tolist())), "unique selected frequencies per resample")
                            reconstructed[family][index] += 1
                    resamples = stable.get("resample_count", 0)
                    if resamples:
                        equal(len(stable["source_resamples"]), resamples, "recorded resample budget")
                        reconstructed = {k:v/resamples for k,v in reconstructed.items()}
                        reconstructed["stable_joint"] = .5*reconstructed["stable_mcuve"]+.5*reconstructed["stable_cars"]
                    for family, freq in stable.get("selection_frequencies", {}).items():
                        equal(len(freq), 612, "612 original selection-frequency coordinates")
                        close(reconstructed[family], freq, "resampling frequency reconstruction")
                        jaccard = stable.get("pairwise_jaccard", {}).get("mcuve" if family == "stable_mcuve" else "cars" if family == "stable_cars" else "joint_resample_diagnostic")
                        mean_jaccard = jaccard.get("mean") if isinstance(jaccard, dict) else jaccard
                        for index, value in enumerate(freq):
                            equal(0 <= value <= 1, True, "selection frequency bounded")
                            resample_stability.append({"来源记录": str(p), "稳定选择族": family,
                                "原频点索引从0起": index, "原频率坐标": float(axis[index]),
                                "源重抽选择频率": value, "源重抽次数": stable.get("resample_count"),
                                "平均Jaccard": mean_jaccard,
                                "来源样品数": len(caller_ids)})
                    budget = stable.get("fit_budget", {})
                    selection_budgets.append({"来源记录": str(p), "源样品数": len(caller_ids), "源类别数量": compact(counts),
                        "源拟合秒数": record.get("source_fit_seconds"), "重抽次数": stable.get("resample_count"),
                        "计算预算": compact(budget), "Jaccard描述": compact(stable.get("pairwise_jaccard", {}))})
                    neural = record.get("neural_training", {})
                    if neural:
                        checkpoint = read_json(p.parent/"source_checkpoint.json")
                        equal(checkpoint["common_sha256"], digest(p), "shared source metadata hash")
                        equal(checkpoint["cnn_state_sha256"], digest(p.parent/"cnn_source_state.pt"), "shared CNN saved state hash")
                        equal(checkpoint["cnn_source_training_invocations"], 1, "one source encoder fit per caller")
                        shared = checkpoint["shared_encoder_fingerprint"]
                        for family in ("cnn_softmax", "cnn_embedding_svm"):
                            equal(checkpoint["artifacts"][family]["metadata"]["cnn_encoder_fingerprint"], shared, "paired classifier uses same encoder")
                        equal(neural["final_fingerprint"], shared, "common neural training state")
                        neural_controls.append({"来源记录": str(p), "源样品数": len(caller_ids),
                            "CNN来源": "仅当前源训练区；不复用已见外测样品权重",
                            "训练轮数": neural.get("epochs"), "批量": neural.get("batch_size"),
                            "学习率": neural.get("learning_rate"), "权重衰减": neural.get("weight_decay"),
                            "表示维数": neural.get("embedding_dimensions"), "参数量": neural.get("parameter_count"),
                            "共享编码器指纹": shared, "共同源训练次数": checkpoint["cnn_source_training_invocations"],
                            "实际损失语义": neural.get("loss_semantics"),
                            "配对分类器": "CNN自身分类头和CNN embedding＋SVM共享编码器；详各模型metadata中的fingerprint"})
                        for loss in neural.get("losses", []):
                            numerical_losses = {k:v for k,v in loss.items() if "loss" in k or "entropy" in k}
                            neural_losses.append({"来源记录": str(p), "源样品数": len(caller_ids),
                                "epoch": loss.get("epoch"), "损失项": compact(numerical_losses),
                                "源训练准确率": loss.get("source_accuracy"),
                                "源训练混淆矩阵": compact(loss["source_confusion_matrix"]) if "source_confusion_matrix" in loss else ""})

    original_path = Path("outputs/approved-zaoren-20261031/experiments/alpha_primary/alpha_mcuve20_svc/predictions.csv")
    historical = {r["sample_id"]: r for r in prediction_rows(original_path)}
    r0 = {r["record_id"]: r for r in outer_groups["core:R0"]}
    equal(set(r0), set(historical), "historical sample identity")
    for identifier, row in r0.items():
        equal(row["y_predicted"], historical[identifier]["y_predicted"], "legacy vote prediction reproduction")
    equal(sum(int(r["is_correct"]) for r in r0.values()), 773, "legacy 773 correct")
    provenance[str(original_path)] = digest(original_path)
    inherited_ordinary = outer_groups.get("combinations:R4_diagnostic_ordinary_mcuve")
    if inherited_ordinary is None:
        raise RuntimeError("Missing prespecified R4 ordinary-MC inheritance control")
    reference_r3 = {r["record_id"]: r for r in outer_groups["core:R3"]}
    for row in inherited_ordinary:
        original = reference_r3[row["record_id"]]
        equal(row["y_predicted"], original["y_predicted"], "R4 ordinary MC reproduces inherited R3")
        close([float(row["score_"+c]) for c in CLASSES], [float(original["score_"+c]) for c in CLASSES], "inherited ordinary MC scores")
    aliases = defaultdict(list)
    for key, rows in outer_groups.items():
        signature = tuple((r["record_id"], r["y_predicted"]) for r in sorted(rows, key=lambda r:r["record_id"]))
        aliases[signature].append(key)
    alias_rows = [{"分类完全相同的方法记录": compact(names), "共同预测样品数": 855,
                   "解释": "相同855标签不表示同一权重或同一分数；角色别名不计新增样品或独立实验"}
                  for names in aliases.values() if len(names) > 1]
    for key, rows in outer_groups.items():
        corrected = harmed = unchanged = wrong_to_wrong = 0
        for row in rows:
            reference = r0[row["record_id"]]
            first, second = bool(int(reference["is_correct"])), bool(int(row["is_correct"]))
            corrected += int(not first and second); harmed += int(first and not second)
            unchanged += int(reference["y_predicted"] == row["y_predicted"])
            wrong_to_wrong += int(not first and not second and reference["y_predicted"] != row["y_predicted"])
            if reference["y_predicted"] != row["y_predicted"]:
                changes.append({"方法": key, "样品编号": row["record_id"], "外折": int(row["outer_fold"]),
                    "真实产地组": row["y_true"], "R0预测": reference["y_predicted"], "当前预测": row["y_predicted"],
                    "R0是否正确": int(reference["is_correct"]), "是否正确": int(row["is_correct"]),
                    "变化": "纠正原错误" if not first and second else "新增错误" if first and not second else "错误类别改变"})
        close((corrected-harmed)/855, np.mean([int(r["is_correct"]) for r in rows])-773/855, "paired net gain")
        pair_table.append({"方法": key, "纠正R0错误": corrected, "新增错误": harmed, "净新增正确": corrected-harmed,
            "预测类别不变": unchanged, "错到另一错误类别": wrong_to_wrong, "准确率差百分点": 100*(corrected-harmed)/855})
    by_feature = defaultdict(list)
    for row in frequencies:
        by_feature[(row["分支"], row["方法"])].append(row)
    for (run_name, method), rows in by_feature.items():
        if (run_name, method) in non_frequency_methods:
            stability.append({"分支": run_name, "方法": method, "原频点索引从0起": None, "原频率坐标": None,
                "五个外折选中次数": None, "选择比例": None, "五折平均Jaccard": None,
                "解释": "含CNN自动表示；612点为原输入，不把全谱输入重复称为稳定特征选择"})
            continue
        masks = [set(r["原频点索引从0起"] for r in rows if r["外折"] == fold) for fold in range(5)]
        jac = [len(a & b)/len(a | b) if a | b else 1. for i,a in enumerate(masks) for b in masks[i+1:]]
        count = Counter(r["原频点索引从0起"] for r in rows)
        for index, n in sorted(count.items()):
            stability.append({"分支": run_name, "方法": method, "原频点索引从0起": index, "原频率坐标": float(axis[index]),
                "五个外折选中次数": n, "选择比例": n/5, "五折平均Jaccard": float(np.mean(jac)),
                "解释": "源训练区重叠的描述性稳定性，不是独立批次验证"})

    sheets = {"说明": explanation, "四轮OOF台账": stages, "各外折内层选择": choices,
        "全部内层候选记录": inner_table, "内层候选均值": inner_mean, "各外折测试指标": fold_table,
        "各类精确召回F1": class_table, "全部855样品OOF": oof_table, "相对R0纠正新增错": pair_table,
        "逐样品改动": changes, "原频点列表": frequencies, "选点稳定性": stability,
        "gamma实际值与别名": gamma_rows, "类别权重与归一": weight_rows, "训练预算及模型记录": budgets,
        "源重抽频点稳定性": resample_stability, "匹配预算和重抽诊断": selection_budgets,
        "CNN训练与配对控制": neural_controls, "CNN逐epoch源损失": neural_losses, "CNN编码器共享记录": source_states}
    sheets["相同预测及角色别名"] = alias_rows
    output.mkdir(parents=True, exist_ok=True)
    book = Workbook(); book.remove(book.active)
    for name, rows in sheets.items():
        add_sheet(book, name, rows); write_csv(output/(name+".csv"), rows)
    path = output/"四轮嵌套CV_完整结果与855样品OOF对照.xlsx"; book.save(path)
    restored = load_workbook(path, read_only=True, data_only=True)
    for name, expected in sheets.items():
        actual = list(restored[name].values)
        use = expected or [{"说明": "本项无已完成的适用记录；未生成模拟结果"}]
        keys = list(use[0]); equal(list(actual[0]), keys, "Excel column names")
        equal(len(actual)-1, len(use), "Excel data row count")
        for cells, row in zip(actual[1:], use):
            for value, key in zip(cells, keys):
                wanted = row.get(key)
                if wanted == "": wanted = None
                if isinstance(wanted, (float, int)) and not isinstance(wanted, bool):
                    close(value, wanted, "Excel numerical cell")
                else:
                    equal(value, wanted, "Excel text/boolean cell")
    restored.close()
    validation = {"status": "pass", "checks_evaluated": checks, "maximum_numeric_difference": maximum,
        "input_specimens": 855, "outer_prediction_rows": len(all_outer), "methods": len(outer_groups),
        "inner_candidate_rows": len(all_inner), "historical_R0_correct": 773,
        "inner_prediction_rows_independently_checked": inner_raw_count,
        "oof_errors_preserved": sum(int(r["is_correct"]) == 0 for r in all_outer),
        "model_fits_in_export": 0, "source_scores_not_probabilities": True,
        "outer_predictions_not_used_for_method_selection": True, "sheet_data_rows": {k:len(v) for k,v in sheets.items()},
        "provenance": provenance, "workbook_sha256": digest(path)}
    (output/"table_export_validation.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    return validation


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-root", type=Path, default=Path("outputs/staged-nested-20261031"))
    parser.add_argument("--output", type=Path, default=Path("outputs/staged-nested-20261031/tables"))
    parser.add_argument("--input", type=Path, default=Path("outputs/approved-zaoren-20261031/datasets/candidate_0p1_1p5.npz"))
    parser.add_argument("--plan", type=Path, default=Path("data/research/staged_nested_plan_20261031.json"))
    args = parser.parse_args()
    result = export(args.run_root, args.output, args.input, args.plan)
    print(json.dumps(result, ensure_ascii=False))
