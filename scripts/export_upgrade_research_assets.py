"""Export upgrade figures/tables from locked real predictions, never fit models.

The exporter refuses partial neural runs and independently recomputes every
family/seed/fold metric from saved predictions joined with scorer-only truth.
"""
import argparse
from collections import defaultdict
import csv
import gzip
import hashlib
import json
import os
from pathlib import Path
import shutil
import textwrap

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT=Path(__file__).resolve().parents[1]
RUN=ROOT/"reports/upgrade-20261031-v1"
OVERLAY=ROOT/"data/multiplant_database/metadata/user_declaration_overlays/20261031_huangqi_tianma_v1/overlay_manifest.json"
CLASSES=tuple(f"ZY{i:02d}" for i in range(1,8))
FAMILIES=("cnn_alpha","resnet_alpha","se_alpha","dual_alpha_mcuve","fusion_alpha_n",
          "cnn_n","fusion_time_frequency","fusion_all_four","inner_selected_neural")
NAMES={"R2":"MC-UVE + RBF-SVM (R2)","cnn_alpha":"1D-CNN: alpha",
       "resnet_alpha":"ResNet1D: alpha","se_alpha":"ResNet1D + SE: alpha",
       "dual_alpha_mcuve":"Full alpha + MC-UVE MLP branch","fusion_alpha_n":"Alpha + refractive index",
       "cnn_n":"1D-CNN: refractive index","fusion_time_frequency":"Time + frequency representations",
       "fusion_all_four":"Four related representations","inner_selected_neural":"Inner-selected neural strategy"}
METRICS=("accuracy","macro_f1","balanced_accuracy")


def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path):return json.loads(Path(path).read_text())


def read_csv(path):
    opener=gzip.open if str(path).endswith(".gz") else open
    with opener(path,"rt",encoding="utf-8-sig",newline="") as h:return list(csv.DictReader(h))


def confusion_metrics(truth,prediction):
    matrix=np.zeros((7,7),dtype=int);lookup={c:i for i,c in enumerate(CLASSES)}
    for actual,predicted in zip(truth,prediction,strict=True):matrix[lookup[str(actual)],lookup[str(predicted)]]+=1
    n=matrix.sum();tp=matrix.diagonal();support=matrix.sum(1);found=matrix.sum(0)
    recall=np.divide(tp,support,out=np.zeros(7),where=support>0)
    precision=np.divide(tp,found,out=np.zeros(7),where=found>0)
    f1=np.divide(2*tp,support+found,out=np.zeros(7),where=(support+found)>0)
    return {"accuracy":float(tp.sum()/n),"macro_f1":float(f1.mean()),"balanced_accuracy":float(recall[support>0].mean()),
            "correct":int(tp.sum()),"total":int(n),"confusion_matrix":matrix.tolist(),
            "per_class":{c:{"precision":float(p),"recall":float(r),"f1":float(f),"support":int(s)}
                         for c,p,r,f,s in zip(CLASSES,precision,recall,f1,support,strict=True)}}


def json_cell(value):
    return json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(",",":")) if isinstance(value,(dict,list,tuple)) else value


def table(sheet_name,rows,book,output,validation):
    if not rows:rows=[{"status":"NOT_RUN_OR_NO_APPLICABLE_RECORD"}]
    keys=list(rows[0]);rows=[{k:(None if row.get(k)=="" else json_cell(row.get(k))) for k in keys} for row in rows]
    with (output/(sheet_name+".csv")).open("w",encoding="utf-8-sig",newline="") as h:
        writer=csv.DictWriter(h,fieldnames=keys,lineterminator="\n");writer.writeheader();writer.writerows(rows)
    ws=book.create_sheet(sheet_name);ws.append(keys)
    for row in rows:ws.append([row[k] for k in keys])
    ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:
        cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="244F70")
        cell.alignment=Alignment(wrap_text=True,vertical="center")
        ws.column_dimensions[cell.column_letter].width=24 if len(cell.value)<24 else 32
    ws.row_dimensions[1].height=32
    for col,key in enumerate(keys,1):
        if key in ("source_file","file","scientific_status","notes","model","selected_family","family","status"):
            ws.column_dimensions[ws.cell(1,col).column_letter].width=44
        for cells in ws.iter_rows(min_row=2,min_col=col,max_col=col):
            cell=cells[0];cell.alignment=Alignment(wrap_text=True,vertical="top")
            if any(token in key for token in ("accuracy","precision","recall")) and isinstance(cell.value,(float,int)):
                cell.number_format="0.00%"
            elif any(token in key for token in ("macro_f1","f1","sample_sd")) and isinstance(cell.value,(float,int)):
                cell.number_format="0.0000"
            elif key.endswith("_pp") and isinstance(cell.value,(float,int)):
                cell.number_format="+0.0000;-0.0000;0.0000"
    validation[sheet_name]=rows


def load_truth(neural,scoring,inputs):
    truth={}
    for raw in scoring["truth_paths"]:
        # Prefer the immutable archive beside the selected run over its old
        # runtime path, which may still exist on the export machine.
        p=neural/"scoring_truth"/Path(raw).name
        if not p.exists():
            p=Path(raw)
            if not p.is_absolute():p=ROOT/p
        if not p.exists():raise FileNotFoundError("Archived scorer-only truth required: "+str(raw))
        inputs[str(p.relative_to(ROOT))]=sha(p)
        for row in read_csv(p):
            if row["record_id"] in truth:raise AssertionError("Repeated specimen across outer truth folds")
            truth[row["record_id"]]=row["label"]
    if len(truth)!=855:raise AssertionError("Exactly855 real specimens required")
    return truth


def assert_metrics(actual,expected,checks):
    for key in METRICS:
        difference=abs(actual[key]-expected[key]);checks["maximum_metric_difference"]=max(checks["maximum_metric_difference"],difference)
        if difference>1e-12:raise AssertionError((key,actual[key],expected[key]))
        checks["metric_checks"]+=1
    if actual["confusion_matrix"]!=expected["confusion_matrix"]:raise AssertionError("Confusion counts differ")
    checks["metric_checks"]+=1


def source_inputs(neural,inputs):
    required=["scoring_manifest.json","prediction_lock.json","outer_predictions.csv.gz","selection_lock.json",
              "pooled_metrics_by_seed.json","seed_sensitivity.json","fold_metrics.json","outer_inference_manifest.json","run_manifest.json"]
    for name in required:
        p=neural/name
        if not p.exists():raise RuntimeError("Neural scoring incomplete; no placeholder metrics: "+str(p))
        inputs[str(p.relative_to(ROOT))]=sha(p)
    lock=read_json(neural/"prediction_lock.json")
    if lock["outer_predictions_sha256"]!=sha(neural/"outer_predictions.csv.gz"):raise AssertionError("Prediction lock changed")
    scoring=read_json(neural/"scoring_manifest.json")
    if scoring["prediction_lock_sha256"]!=sha(neural/"prediction_lock.json"):raise AssertionError("Scorer references another prediction lock")
    return scoring


def protocol_file(neural):
    path=neural/"protocol_snapshot.json"
    if not path.exists():path=ROOT/"data/research/upgrade_neural_protocol_20261031.json"
    if sha(path)!=read_json(neural/"run_manifest.json")["protocol_sha256"]:
        raise AssertionError("Resolved protocol differs from the frozen training manifest")
    return path


def collect(neural,checks):
    inputs={};scoring=source_inputs(neural,inputs);truth=load_truth(neural,scoring,inputs)
    protocol_path=protocol_file(neural);inputs[str(protocol_path.relative_to(ROOT))]=sha(protocol_path)
    predictions=read_csv(neural/"outer_predictions.csv.gz");groups=defaultdict(list)
    for row in predictions:
        if row["family"] not in FAMILIES:raise AssertionError("Unexpected family role")
        groups[(row["family"],int(row["seed"]))].append(row)
        scores=np.asarray([float(row["score_"+c]) for c in CLASSES])
        if not np.isfinite(scores).all():raise AssertionError("Nonfinite neural score")
        if CLASSES[int(scores.argmax())]!=row["y_predicted"]:raise AssertionError("Neural argmax/class vocabulary mismatch")
    if len(predictions)!=9*3*855 or set(groups)!={(f,s) for f in FAMILIES for s in (42,43,44)}:
        raise AssertionError("Incomplete nine-role/three-seed OOF")
    pooled=read_json(neural/"pooled_metrics_by_seed.json");averages=read_json(neural/"seed_sensitivity.json")
    reference_path=ROOT/"reports/staged-nested-optimization/2026-10-31/core/summary.json"
    baseline=next(row for row in read_json(reference_path) if row["method"]=="R2")
    inputs[str(reference_path.relative_to(ROOT))]=sha(reference_path)
    baseline_csv=ROOT/"reports/staged-nested-optimization/2026-10-31/core/outer_predictions.csv.gz"
    baseline_rows=[r for r in read_csv(baseline_csv) if r["method"]=="R2"]
    inputs[str(baseline_csv.relative_to(ROOT))]=sha(baseline_csv)
    base_by_id={r["record_id"]:r for r in baseline_rows}
    if len(baseline_rows)!=855 or len(base_by_id)!=855:raise AssertionError("Reference denominator mismatch")
    if set(base_by_id)!=set(truth):raise AssertionError("Neural andR2 specimen cohorts differ")
    for key,row in base_by_id.items():
        if row["y_true"]!=truth[key]:raise AssertionError("True label changed between R2 and new run")
    recalc_base=confusion_metrics([r["y_true"] for r in baseline_rows],[r["y_predicted"] for r in baseline_rows])
    assert_metrics(recalc_base,baseline,checks)
    if recalc_base["correct"]!=821:raise AssertionError("Do not replace or weaken audited821/855 baseline")
    seed_rows=[];mean_rows=[];class_rows=[];confusion_rows=[];fold_rows=[];paired_rows=[];oof_rows=[];metrics={}
    reported_folds={(r["family"],r["seed"],r["outer_fold"]):r["metrics"] for r in read_json(neural/"fold_metrics.json")}
    for family in FAMILIES:
        for seed in (42,43,44):
            found=groups[family,seed];ids=[r["record_id"] for r in found]
            if len(ids)!=855 or len(set(ids))!=855:raise AssertionError("Duplicate OOF row inflates denominator")
            if set(ids)!=set(truth):raise AssertionError("Incomplete OOF identities")
            measured=confusion_metrics([truth[i] for i in ids],[r["y_predicted"] for r in found]);metrics[family,seed]=measured
            assert_metrics(measured,pooled[family][str(seed)],checks)
            seed_rows.append({"family":family,"model":NAMES[family],"seed":seed,"correct":measured["correct"],"real_specimens":855,
                **{k:measured[k] for k in METRICS},"accuracy_gain_vs_R2_pp":100*(measured["accuracy"]-baseline["accuracy"]),
                "relative_accuracy_gain_vs_R2":(measured["accuracy"]-baseline["accuracy"])/baseline["accuracy"],
                "role":"primary_inner_selected_strategy" if family=="inner_selected_neural" else "prespecified_family_diagnostic",
                "scientific_status":"EXPLORATORY exposed single-batch development"})
            for c in CLASSES:class_rows.append({"family":family,"seed":seed,"class":c,**measured["per_class"][c]})
            for i,c in enumerate(CLASSES):
                for j,p in enumerate(CLASSES):confusion_rows.append({"family":family,"seed":seed,"true_class":c,"predicted_class":p,"count":measured["confusion_matrix"][i][j]})
            corrected=harmed=unchanged=0
            for row in found:
                rid=row["record_id"];base=base_by_id[rid]["y_predicted"];actual=truth[rid];pred=row["y_predicted"]
                corrected+=base!=actual and pred==actual;harmed+=base==actual and pred!=actual;unchanged+=base==pred
                oof_rows.append({"family":family,"seed":seed,"outer_fold":int(row["outer_fold"]),"record_id":rid,"y_true":actual,
                    "y_predicted":pred,"is_correct":int(actual==pred),"R2_prediction":base,"R2_is_correct":int(base==actual),
                    "architecture_family":row["architecture_family"],"profile":row["profile"]})
            if corrected-harmed!=measured["correct"]-821:raise AssertionError("Paired changes do not reconcile")
            paired_rows.append({"family":family,"seed":seed,"corrected_vs_R2":corrected,"harmed_vs_R2":harmed,
                                "net_correct_change":corrected-harmed,"same_predicted_class":unchanged,"real_specimens":855})
            for fold in range(5):
                subset=[r for r in found if int(r["outer_fold"])==fold]
                if len(subset)!=171:raise AssertionError("Every seed/family/fold must score171")
                measured_fold=confusion_metrics([truth[r["record_id"]] for r in subset],[r["y_predicted"] for r in subset])
                assert_metrics(measured_fold,reported_folds[family,seed,fold],checks)
                fold_rows.append({"family":family,"seed":seed,"outer_fold":fold,"real_specimens":171,"correct":measured_fold["correct"],
                    **{k:measured_fold[k] for k in METRICS},"note":"fold metric; do not substitute fold mean for pooledOOF"})
        row={"family":family,"model":NAMES[family],"seeds":"42,43,44","specimens_per_seed":855,"seed_count":3}
        for key in METRICS:
            values=np.asarray([metrics[family,s][key] for s in (42,43,44)])
            mean,sd=float(values.mean()),float(values.std(ddof=1))
            if abs(mean-averages[family][key]["mean"])>1e-12 or abs(sd-averages[family][key]["sample_sd"])>1e-12:raise AssertionError("Seed aggregation differs")
            checks["metric_checks"]+=2;row[key+"_mean"]=mean;row[key+"_sample_sd"]=sd
        row["uncertainty_scope"]="three initialization seeds on same855; notbiologicalCI";mean_rows.append(row)
    choices=[];lock=read_json(neural/"selection_lock.json")
    for outer in lock["outer_folds"]:
        fold=outer["outer_fold"]
        for family,choice in outer["choices"].items():
            p=choice["selected"]
            inner_mean=next(r["mean_inner_metrics"] for r in choice["means"] if r["profile"]==p)
            choices.append({"outer_fold":fold,"prediction_role":family,"selected_family":family,"selected_profile":p["id"],
                            **p,**{"inner_"+k:v for k,v in inner_mean.items()},
                            "selection_scope":"2inner x3seed source-only mean; coupled profile, notfactorial"})
        policy=outer["neural_strategy"];p=policy["selected_profile"]
        choices.append({"outer_fold":fold,"prediction_role":"inner_selected_neural","selected_family":policy["selected_family"],
                        "selected_profile":p["id"],**p,**{"inner_"+k:v for k,v in policy["mean_inner_metrics"].items()},
                        "selection_scope":"across8families x2profiles using inner metrics; alias, no extra fit"})
        for seed in (42,43,44):
            policy_rows=[r for r in groups["inner_selected_neural",seed] if int(r["outer_fold"])==fold]
            if any(r["architecture_family"]!=policy["selected_family"] or r["profile"]!=p["id"] for r in policy_rows):raise AssertionError("Policy alias differs from frozen inner selection")
    cost=[]
    for row in read_json(neural/"outer_inference_manifest.json"):
        directory=neural/"evaluation"/row["split_id"]/row["family"]/row["profile"]/f"seed{row['seed']}"
        if not (directory/"training.json").exists():
            directory=Path(row["model_directory"])
            if not directory.is_absolute():directory=ROOT/directory
        metadata=read_json(directory/"training.json");inputs[str((directory/"training.json").relative_to(ROOT))]=sha(directory/"training.json")
        cost.append({"family":row["family"],"seed":row["seed"],"outer_fold":row["outer_fold"],"profile":row["profile"],
            "parameters":metadata["parameters"],"epochs":metadata["epochs"],"source_specimens":len(metadata["source_ids"]),
            "training_seconds":metadata["training_seconds"],"transform_fit_seconds":metadata["transform_fit_seconds"],
            "total_case_seconds":row["elapsed_seconds"],"selected_MC_features":len(metadata["selected_indices"]) if metadata["selected_indices"] is not None else None,
            "policy_alias_is_not_an_extra_model":True})
    baseline_row={"family":"R2","seed":42,"model":NAMES["R2"],"correct":821,"real_specimens":855,
                  **{k:baseline[k] for k in METRICS},"status":"Frozen historical reference; no new fit/no seedSD"}
    if len(cost)!=120:raise AssertionError("Policy alias must not be counted as additional models")
    checks.update(neural_prediction_rows=len(predictions),neural_role_seed_cases=len(groups),raw_real_specimens=855,policy_alias_adds_specimens=False)
    return {"inputs":inputs,"baseline":baseline,"baseline_row":baseline_row,"metrics":metrics,"seed_results":seed_rows,
            "seed_means":mean_rows,"per_class":class_rows,"confusions":confusion_rows,"folds":fold_rows,"choices":choices,
            "paired":paired_rows,"oof":oof_rows,"cost":cost,"protocol":read_json(protocol_path)}


def savefig(fig,name,output):
    fig.savefig(output/(name+".png"),dpi=220,bbox_inches="tight",facecolor="white")
    fig.savefig(output/(name+".pdf"),bbox_inches="tight",facecolor="white")


def plots(data,adversarial,readiness,output):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch,FancyArrowPatch
    plt.rcParams.update({"font.family":"DejaVu Sans","font.size":10,"axes.spines.top":False,"axes.spines.right":False,"pdf.fonttype":42})
    order=["R2",*FAMILIES];labels=[NAMES[f] for f in order];y=np.arange(len(order))
    fig,axes=plt.subplots(1,2,figsize=(15,8),sharey=True)
    for ax,key,scale,title in zip(axes,("accuracy","macro_f1"),(100,1),("Accuracy (%)","Macro-F1"),strict=True):
        values=[data["baseline"][key],*[data["metrics"][f,42][key] for f in FAMILIES]]
        colors=["#CB9A24" if f=="R2" else "#2274A5" if f=="inner_selected_neural" else "#9BAEBB" for f in order]
        ax.barh(y,np.asarray(values)*scale,color=colors,height=.62,label="Primary seed 42")
        for i,f in enumerate(FAMILIES,1):
            array=np.asarray([data["metrics"][f,s][key] for s in (42,43,44)])
            ax.errorbar(array.mean()*scale,i,xerr=array.std(ddof=1)*scale,fmt="o",mfc="white",mec="#17364D",ecolor="#17364D",capsize=3,ms=4)
        ax.set_xlim(0,100 if scale==100 else 1);ax.set_xlabel(title);ax.grid(axis="x",alpha=.2);ax.set_axisbelow(True)
        ax.set_yticks(y,labels)
    axes[0].invert_yaxis()
    fig.suptitle("Same 855 exposed specimens: historical 5 outer / 2 inner | neural diagnostics and locked strategy",y=.98,fontweight="bold")
    fig.text(.01,.015,"Bars: predeclared seed 42. Open dots: mean +/- sample SD across 42/43/44 on the same 855; not biological confidence intervals.\nGold R2 is the frozen 821/855 reference. Blue is the inner-selected architecture/profile strategy.",fontsize=9)
    fig.tight_layout(rect=[0,.075,1,.95]);savefig(fig,"model_accuracy_macro_f1",output);plt.close(fig)

    fig,axes=plt.subplots(2,2,figsize=(13,10))
    for col,(title,metric) in enumerate((("Frozen R2",data["baseline"]),("Inner-selected neural, seed 42",data["metrics"]["inner_selected_neural",42]))):
        matrix=np.asarray(metric["confusion_matrix"]);support=matrix.sum(1);normalized=matrix/support[:,None]
        for row in range(2):
            ax=axes[row,col];values=matrix if row==0 else normalized
            im=ax.imshow(values,cmap="Blues",vmin=0,vmax=350 if row==0 else 1)
            for i in range(7):
                for j in range(7):
                    text=str(matrix[i,j]) if row==0 else f"{normalized[i,j]:.0%}"
                    ax.text(j,i,text,ha="center",va="center",fontsize=9,color="white" if values[i,j]>.55*(350 if row==0 else 1) else "#102A43")
            ax.set_xticks(range(7),CLASSES,rotation=35);ax.set_yticks(range(7),[f"{c} (n={n})" for c,n in zip(CLASSES,support)])
            ax.get_yticklabels()[1].set_color("#B6222A");ax.get_yticklabels()[1].set_fontweight("bold")
            ax.set_xlabel("Predicted anonymous group");ax.set_ylabel("True anonymous group")
            ax.set_title(title+(" | counts" if row==0 else " | row recall"),fontweight="bold")
            fig.colorbar(im,ax=ax,fraction=.046,pad=.035)
    fig.suptitle("Paired 855-specimen OOF confusion matrices; ZY02 has only 5 declared samples",fontweight="bold")
    fig.text(.01,.01,"The neural strategy was selected inside each outer source. These are exposed single-batch development results, not independent named-origin confirmation.",fontsize=9)
    fig.tight_layout(rect=[0,.03,1,.96]);savefig(fig,"r2_vs_inner_selected_confusion",output);plt.close(fig)

    fig,ax=plt.subplots(figsize=(11,5.5));values=np.asarray([r["accuracy"] for r in adversarial["permutation_results"]])*100
    ax.scatter(np.arange(1,11),values,color="#2274A5",s=50,label="10 conditional source-group label permutations")
    ax.axhline(adversarial["majority_baseline"]["accuracy"]*100,color="#A75122",ls="--",label="Majority-class baseline 40.94%")
    ax.axhline(adversarial["original_score"]["accuracy"]*100,color="#496B38",ls=":",label="Original-label model in this one-fold diagnostic: 90.06%")
    ax.set(xlim=(.5,10.5),ylim=(0,100),xticks=np.arange(1,11),xlabel="Prespecified permutation index (fixed seed list)",ylabel="Accuracy on the same 171 old outer-fold 0 specimens (%)")
    ax.grid(axis="y",alpha=.2);ax.legend(loc="upper right",fontsize=9)
    ax.set_title("Bounded leakage diagnostic: conditional empirical null (10 permutations)",fontweight="bold")
    fig.text(.01,.01,"Each permutation reruns source-only MC-UVE and inner SVM selection. The same 171 query specimens are reused.\nPlus-one resolution is 1/11; not a 1/7 acceptance rule or a biological/cross-batch significance certificate.",fontsize=9)
    fig.tight_layout(rect=[0,.08,1,.98]);savefig(fig,"source_label_permutation_diagnostic",output);plt.close(fig)

    fig,ax=plt.subplots(figsize=(14,6));ax.axis("off")
    columns=["Plant / task","Readable numeric records","Independent identity / units","Scientific use now"]
    cells=[["Acid jujube (primary)","855 declared samples; 4 paired representations","User declaration; 1 lot per group; units/logs incomplete","EXPLORATORY internal 5/2-fold"],
           ["Huangqi (auxiliary)","403 declared independent records; n only 303","User declaration; origins/batches/units not authenticated","EXPLORATORY 4 source groups; transfer BLOCKED"],
           ["Tianma (auxiliary)","405 declared independent records; time only 202","User declaration; origins/batches/units not authenticated","EXPLORATORY 4 source groups; transfer BLOCKED"],
           ["Wheat (auxiliary)","12/12/13/13 averaged curves; not 50 independent samples","Quantity, averaging parents and labels unknown","BLOCKED classification/regression"],
           ["Confirmation cohort","No qualifying new independent batch","MISSING","No external confirmation accuracy"],
           ["Molecular mechanism","3 candidate preparations; water software test only","Candidate optimized geometry / frequency / IR NOT_RUN","No Gaussian/DFT peak prior in models"]]
    wrapped=[[textwrap.fill(value,width) for value,width in zip(row,(23,36,42,35),strict=True)] for row in cells]
    t=ax.table(cellText=wrapped,colLabels=columns,loc="center",cellLoc="left",colWidths=[.18,.26,.30,.26],bbox=[0,.01,1,.94]);t.auto_set_font_size(False);t.set_fontsize(10)
    for (row,col),cell in t.get_celld().items():
        if row==0:cell.set_facecolor("#244F70");cell.set_text_props(color="white",weight="bold")
        else:cell.set_facecolor("#E9F2FA" if row==1 else "#F2F2F2");cell.set_text_props(wrap=True)
    ax.set_title("Four-plant database: technical readability is different from scientific readiness",fontweight="bold",pad=20)
    fig.text(.01,.015,"Representations are not separate biological samples or independent sensors. Equal numeric axes do not establish physical compatibility.\nLater user declaration permits HQ/Tianma source-group tasks; named-origin/batch confirmation and cross-plant transfer remain unresolved.",fontsize=9)
    fig.tight_layout(rect=[0,.07,1,.96])
    fig.canvas.draw();renderer=fig.canvas.get_renderer()
    for cell in t.get_celld().values():
        bounds=cell.get_window_extent(renderer);text_bounds=cell.get_text().get_window_extent(renderer)
        if (text_bounds.x0<bounds.x0-1 or text_bounds.x1>bounds.x1+1 or
                text_bounds.y0<bounds.y0-1 or text_bounds.y1>bounds.y1+1):
            raise AssertionError("Readiness table text exceeds its cell; revise wrapping before exporting")
    savefig(fig,"multiplant_scientific_readiness",output);plt.close(fig)

    fig,ax=plt.subplots(figsize=(15,8));ax.set(xlim=(0,15),ylim=(0,8));ax.axis("off")
    def box(x,y,w,h,text,color="#E8F0F6"):
        patch=FancyBboxPatch((x,y),w,h,boxstyle="round,pad=.12",fc=color,ec="#4D6778",lw=1.1);ax.add_patch(patch)
        ax.text(x+w/2,y+h/2,text,ha="center",va="center",fontsize=10,wrap=True)
    def arrow(a,b,dashed=False):ax.add_patch(FancyArrowPatch(a,b,arrowstyle="-|>",mutation_scale=14,lw=1.2,color="#365466",linestyle="--" if dashed else "-"))
    box(.3,6.5,3.6,1,"Primary: acid jujube 855\nHistorical exposed development")
    box(5,6.5,4.1,1,"Lineage / fixed 5 outer / 2 inner\nCurrent-source fit guards", "#E7F0DF")
    box(10.2,6.5,4.3,1,"Inner-only architecture/profile choice\nAll 3 seeds; lock before outer", "#E7F0DF")
    arrow((4.05,7),(4.85,7));arrow((9.25,7),(10.05,7))
    box(.3,4.3,4.1,1.25,"Frozen strong reference\nSNV + MC-UVE + RBF-SVM\n821/855; no new fit", "#F7EBC9")
    box(5,4.3,4.1,1.25,"Source-only representations\nCNN / residual / channel SE\nMC-UVE variables -> MLP branch")
    box(10.2,4.3,4.3,1.25,"Fresh outer-source fit\nLabel-free pointwise inference\neval + no gradient / state audit")
    arrow((12.3,6.35),(12.3,5.7));arrow((9.25,4.9),(10.05,4.9))
    box(5,2.2,4.1,1.1,"Related representations\nalpha + n / time + frequency\nSame sample identity, no 4xN")
    box(10.2,2.2,4.3,1.1,"Immutable prediction lock\nSeparate truth joining / scoring\nIndependent metric recomputation", "#E7F0DF")
    arrow((7.05,3.45),(7.05,4.15));arrow((12.3,4.15),(12.3,3.45));arrow((4.55,4.9),(10.05,2.75),True)
    box(.3,.35,4.1,1.25,"HQ/Tianma: declared 4-group tasks\nEXPLORATORY; wheat BLOCKED\nCross-plant transfer BLOCKED", "#EEEEEE")
    box(5,.35,4.1,1.25,"3 molecule candidates\nDFT mechanism NOT_RUN\nWater is environment test only", "#EEEEEE")
    box(10.2,.35,4.3,1.25,"Report every seed / negative result\nEXPLORATORY, new batch needed\nNo absolute zero-leakage claim", "#E7F0DF")
    arrow((12.3,2.05),(12.3,1.75));arrow((9.25,.95),(10.05,.95),True)
    fig.suptitle("Research4 upgrade: preserve acid jujube gains; guarded neural comparison and auxiliary work",fontweight="bold",y=.98)
    fig.tight_layout(rect=[0,0,1,.96]);savefig(fig,"guarded_research_workflow",output);plt.close(fig)


def run(neural,output):
    neural,output=Path(neural).resolve(),Path(output).resolve()
    if output.exists() and any(output.iterdir()):raise FileExistsError("Use a new paper-assets directory; never overwrite delivered artifacts")
    staging=output.with_name(output.name+".building")
    if staging.exists():raise FileExistsError("Preserve incomplete export before a new revision")
    staging.mkdir(parents=True);checks={"metric_checks":0,"maximum_metric_difference":0.,"new_model_fits":0,"scientific_status":"EXPLORATORY"}
    data=collect(neural,checks);inputs=data["inputs"]
    static={"inventory":ROOT/"data/multiplant_database/dataset_inventory.csv","readiness":ROOT/"data/multiplant_database/metadata/readiness.json",
        "adversarial":RUN/"adversarial-audit/leakage_audit.json","database_validation":ROOT/"data/multiplant_database/data_quality/database_validation.json",
        "quantum_summary":RUN/"quantum/execution_summary.json","latest_declaration_overlay":OVERLAY}
    for p in static.values():
        if not p.exists():raise FileNotFoundError("Actual supporting evidence missing: "+str(p))
        inputs[str(p.relative_to(ROOT))]=sha(p)
    inventory=read_csv(static["inventory"]);readiness=read_json(static["readiness"]);adversarial=read_json(static["adversarial"])
    overlay=read_json(OVERLAY);annotation=ROOT/overlay["annotation"]
    if sha(annotation)!=overlay["annotation_sha256"]:raise AssertionError("Latest user declaration changed")
    inputs[str(annotation.relative_to(ROOT))]=sha(annotation)
    overlay_lineage=ROOT/overlay["overlay_lineage"]
    if sha(overlay_lineage)!=overlay["overlay_lineage_sha256"]:raise AssertionError("User-declared lineage overlay changed")
    inputs[str(overlay_lineage.relative_to(ROOT))]=sha(overlay_lineage)
    notes=[{"item":"Population","note":"855 user-declared independent specimens, same historical5outer/2inner; all data previously exposed, one lot per origin."},
      {"item":"Primary neural strategy","note":"inner_selected_neural selects architecture and coupled profile only from current source2inner x3seed means; saved weights alias a family, no additional model or samples."},
      {"item":"Seed policy","note":"Predeclaredseed42 is primary;42/43/44 all shown. SampleSD describes initialization sensitivity, not biologicalCI."},
      {"item":"FrozenR2","note":"821/855,96.0234%,MacroF1 .9196845913,BalAcc .8991836735; read old immutable predictions/summary, never retrained by exporter."},
      {"item":"Profiles","note":"Two coupled width/depth/epoch/dropout/learning-rate/weight/MC profiles; notan exhaustive factorial or independent-factor ablation."},
      {"item":"Attention","note":"SE is channelattention, not a frequency-position peak map or chemical concentration."},
      {"item":"MC branch","note":"Discontinuous MC-UVE variables enter an MLP; full continuous alpha retains its convolution branch. No direct convolution pretending missing coordinates are adjacent."},
      {"item":"Four representations","note":"alpha,n,time,frequency derive from related measurement chains; same sample identity, not4independent sensors or4timesN."},
      {"item":"Auxiliary plants latest declaration","note":"Later user confirms HQ403/Tianma405 are separate samples without repeat scans/repressing. Their4source-group internal classification is EXPLORATORY; namedorigins/batches/units/instrument authentication unresolved, transfer BLOCKED. Wheat curves are averages."},
      {"item":"Initial DB versus overlay","note":"Dataset_inventory is the original frozen technical audit, whose oldBLOCKED status is retained. Dataset_readiness separately applies the later dated user-declaration overlay without rewriting initial lineage."},
      {"item":"Quantum","note":"Spinosin/JujubosideA/B geometry/frequency/IR computations NOT_RUN. Water software smoke test is not candidate chemistry. No theoretical peaks merged into model inputs."},
      {"item":"Perturbation diagnostic","note":"13retrained one-fold paths re-evaluate same171, notthis855OOF and notold513/171/171.10conditional source-group label permutations do not certify significance or absolute zero leakage."}]
    book=Workbook();book.remove(book.active);expected={}
    table("Notes",notes,book,staging,expected)
    table("Dataset_inventory",inventory,book,staging,expected)
    readiness_rows=[]
    for key,value in readiness.items():
        if not isinstance(value,str):continue
        latest=overlay["classification_status"] if key in ("huangqi","tianma") else value
        readiness_rows.append({"dataset_or_task":key,"initial_frozen_status":value,"latest_status":latest,
            "latest_evidence":"user_declaration_overlay" if key in ("huangqi","tianma") else "initial_database_scope",
            "transfer_status":"BLOCKED_PENDING_UNITS_INSTRUMENT_AND_BATCH_EVIDENCE" if key in ("huangqi","tianma") else "not_assessed_as_new_transfer_source"})
    table("Dataset_readiness",readiness_rows,book,staging,expected)
    table("Frozen_R2",[data["baseline_row"]],book,staging,expected)
    table("Neural_each_seed",data["seed_results"],book,staging,expected)
    table("Neural_seed_mean_SD",data["seed_means"],book,staging,expected)
    table("Fold_metrics",data["folds"],book,staging,expected)
    table("Selected_profiles",data["choices"],book,staging,expected)
    table("Class_precision_recall_F1",data["per_class"],book,staging,expected)
    table("Confusion_counts",data["confusions"],book,staging,expected)
    table("R2_class_metrics",[{"class":c,**confusion_metrics([CLASSES[i] for i,row in enumerate(data['baseline']['confusion_matrix']) for _ in range(sum(row))],
         [CLASSES[j] for row in data['baseline']['confusion_matrix'] for j,n in enumerate(row) for _ in range(n)])["per_class"][c]} for c in CLASSES],book,staging,expected)
    table("Paired_changes_vs_R2",data["paired"],book,staging,expected)
    table("All855_OOF_each_seed",data["oof"],book,staging,expected)
    table("Outer_actual_model_cost",data["cost"],book,staging,expected)
    table("Source_permutation_171",[{"seed":q["seed"],**{k:q[k] for k in METRICS},"correct":q["correct"],"real_specimens":171,"scope":"onefold conditional source permutation; repeatedsame171"} for q in adversarial["permutation_results"]],book,staging,expected)
    workbook=staging/"Research4升级_多表示神经与防泄漏证据.xlsx";book.save(workbook)
    reloaded=load_workbook(workbook,read_only=True,data_only=True);cells=0
    for name,rows in expected.items():
        values=list(reloaded[name].values);keys=list(rows[0])
        if list(values[0])!=keys or len(values)!=len(rows)+1:raise AssertionError("Worksheet rows differ")
        for actual,row in zip(values[1:],rows,strict=True):
            for value,key in zip(actual,keys,strict=True):
                target=row[key]
                if isinstance(target,float):
                    if abs(value-target)>1e-12:raise AssertionError("Excel float value changed")
                elif value!=target:raise AssertionError((name,key,value,target))
                cells+=1
    reloaded.close();checks.update(excel_sheets=len(expected),excel_cells_read_back=cells)
    plots(data,adversarial,readiness,staging)
    for relative,expected_sha in inputs.items():
        if sha(ROOT/relative)!=expected_sha:raise AssertionError("A frozen input changed during export")
    checks.update(status="PASS_FOR_EXPORTED_ARITHMETIC_AND_FILES_NOT_DATA_INDEPENDENCE",input_sha256=inputs,
                  script_sha256=sha(Path(__file__)),figures=5,seed_SD_is_biological_CI=False,
                  quantum_candidate_DFT_used=False,auxiliary_blocked_tasks_assigned_fake_scores=False)
    json_write(staging/"export_validation.json",checks)
    json_write(staging/"figure_source_metrics.json",{"R2":data["baseline_row"],"neural_seed42":{f:data["metrics"][f,42] for f in FAMILIES},
        "neural_seed_sensitivity":data["seed_means"],"conditional_permutation_summary":{k:adversarial[k] for k in ["permutation_accuracy_mean","permutation_accuracy_min","permutation_accuracy_max"]}})
    text="""# Research4升级：真实表图与范围\n\n[Excel证据表](Research4升级_多表示神经与防泄漏证据.xlsx)含15张表：9神经角色逐种子及均值/样本SD、135外折、锁定架构/profile、逐类P/R/F1、混淆、23085逐样品OOF、R2配对改对/伤害、成本和数据资格。\n\n| 图 | 内容 |\n| --- | --- |\n| [Accuracy与Macro-F1](model_accuracy_macro_f1.png) | [PDF](model_accuracy_macro_f1.pdf)；柱为固定seed42，空点为42/43/44均值±样本SD |\n| [R2与内层选择策略混淆](r2_vs_inner_selected_confusion.png) | [PDF](r2_vs_inner_selected_confusion.pdf)；各855份，ZY02仅5份突出显示 |\n| [源标签置乱诊断](source_label_permutation_diagnostic.png) | [PDF](source_label_permutation_diagnostic.pdf)；10条件置乱、同171与多数类40.94%，不按1/7验收 |\n| [四植物科学可用性](multiplant_scientific_readiness.png) | [PDF](multiplant_scientific_readiness.pdf)；可读不等于独立验证合格 |\n| [受守卫研究流程](guarded_research_workflow.png) | [PDF](guarded_research_workflow.pdf)；DFT与外部阻塞范围明确 |\n\n所有新指标从锁定预测和独立评分真值重算；导出不训练模型，不改变旧R2的821/855。`inner_selected_neural`是内层选择的主策略，各家族为预定义诊断；策略引用已有权重不增加真实样品或模型数。历史855份已曝光，每类一批，科学状态为EXPLORATORY，不是新独立确认。\n\n三种子SD只反映初始化敏感性，不能当生物学置信区间。SE是通道注意力，不是化学浓度或频率峰归属；四表示属于相关测量链，不能称四个独立传感器。两profile联合改变多个因素，不是全因素消融。\n\n黄芪403／天麻405已获后续用户独立性声明，四来源组内部分类为EXPLORATORY；地名／批次／单位仍未外核、跨植物迁移BLOCKED，小麦平均曲线任务BLOCKED；三候选分子的优化/频率/IR均NOT_RUN，水只作软件技术测试，没有拼接虚构化学峰。源内实验峰候选见[单独source-only峰分析](../source-peaks/README.md)，不能回填本版本模型。\n\n再生需真实神经scoring完成，使用新输出目录：\n\n```bash\n.runtime/bin/python scripts/export_upgrade_research_assets.py --neural-dir reports/upgrade-20261031-v1/neural --output outputs/upgrade-paper-assets-reproduction-new\n```\n\n输入/输出SHA及数值核对见`export_validation.json`和`files_sha256.json`。\n"""
    (staging/"README.md").write_text(text,encoding="utf-8")
    index={str(p.relative_to(staging)):sha(p) for p in sorted(staging.iterdir()) if p.is_file()}
    json_write(staging/"files_sha256.json",index)
    staging.rename(output)
    print(json.dumps({k:v for k,v in checks.items() if k!="input_sha256"},ensure_ascii=False))


def json_write(path,value):Path(path).write_text(json.dumps(value,ensure_ascii=False,indent=2,allow_nan=False)+"\n")


if __name__=="__main__":
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--neural-dir",type=Path,default=RUN/"neural")
    parser.add_argument("--output",type=Path,default=RUN/"paper-assets")
    args=parser.parse_args();run(args.neural_dir,args.output)
