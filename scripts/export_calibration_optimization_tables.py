"""Export auditable prediction/label comparison tables without fitting models."""
import argparse
from collections import Counter, defaultdict
import csv
import gzip
import hashlib
import json
from pathlib import Path
import shutil

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SCOPES = {"seven_class_all": "七类主实验（171份预测）",
          "six_class_abundant": "六类辅助实验（170份预测）"}
MODELS = {"full": "候选612点SVM", "mcuve20": "MC-UVE20＋SVM（62点）"}
FAMILIES = {"damped_anchor": "弱化分数锚定", "mistake_bias": "错分残差偏置",
            "mistake_ridge": "错分残差岭回归", "mistake_kernel": "错分残差核修正",
            "prototype_shrinkage": "收缩原型融合", "identity": "保持原模型",
            "legacy_anchor": "原固定λ=1校正"}


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def read_csv(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def write_csv(path, rows):
    with Path(path).open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader(); writer.writerows(rows)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def method_name(method):
    if method == "source_identity":
        return "源模型（不校准）"
    if method == "development_selected":
        return "训练开发选择的最终方案"
    if method == "legacy_lambda1":
        return "原固定λ=1校正"
    if method.startswith("matched_zero_support_"):
        return "同配置零支持对照（不用校准标签）"
    return "分家族诊断：" + FAMILIES[method.removeprefix("family_")]


def support_count(method, candidate, classes, k):
    return 0 if (candidate["family"] == "identity" or method == "source_identity"
                 or method.startswith("matched_zero_support_")) else classes*k


def add_sheet(workbook, title, headers, rows, *, percents=(), decimals=(), widths=None):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.zoomScale = 85
    sheet.row_dimensions[1].height = 44
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="234F70")
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F1F5F9")
        sheet.row_dimensions[row[0].row].height = 31
    for column in percents:
        for row in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
            row[0].number_format = "0.00%"
    for column in decimals:
        for row in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
            row[0].number_format = "0.0000"
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = (widths or {}).get(index, 16)
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    return sheet


def export(run, report, output):
    run, report, output = Path(run), Path(report), Path(output)
    status = read_json(run / "status.json")
    if status["status"] != "completed":
        raise ValueError("The experiment must be completed before table export")
    report.mkdir(parents=True, exist_ok=True); output.parent.mkdir(parents=True, exist_ok=True)
    summary = read_json(run / "summary.json")
    lock = read_json(run / "selection_lock.json")
    grid = {c["id"]: c for c in read_json(run / "candidate_grid.json")}
    random_ids = read_json(run / "random_query_ids.json")
    random_rows = read_csv(run / "random_query_predictions.csv")
    base_summary = {(r["scope"], r["selector"], r["shots_per_class"]): r
                    for r in summary if r["method"] == "source_identity"}
    summary_rows = []
    for row in summary:
        scope, selector, k, method = row["scope"], row["selector"], row["shots_per_class"], row["method"]
        candidate = grid[row["candidate_id"]]
        if method.startswith("family_"):
            gate = lock["choices"][scope][selector][str(k)]["family_winners"][method.removeprefix("family_")]["passes_gate"]
            category = "家族诊断；不能按预测分数替换最终方案"
        elif method == "development_selected":
            gate = not lock["choices"][scope][selector][str(k)]["fallback_to_identity"]
            category = "开发锁定最终方案；保持原模型时实际使用0份校准"
        elif method.startswith("matched_zero_support_"):
            gate = "不适用"; category = "匹配零支持对照"
        else:
            gate = "不适用"; category = "预设对照"
        summary_rows.append([SCOPES[scope], MODELS[selector], k, method_name(method), row["candidate_id"],
            FAMILIES[candidate["family"]], support_count(method, candidate, 7 if scope == "seven_class_all" else 6, k),
            base_summary[(scope, selector, k)]["accuracy_mean"], row["accuracy_mean"], 100*row["accuracy_gain_mean"],
            row["accuracy_sd"], row["accuracy_min"], row["accuracy_max"], row["macro_f1_mean"], row["balanced_accuracy_mean"],
            row["positive_accuracy_draws"], row["zero_accuracy_draws"], row["negative_accuracy_draws"],
            row["paired_corrected_mean"], row["paired_harmed_mean"], 100*row["gain_vs_matched_zero_support_mean"], gate, category])
    # Original random sample rows contain the complete baseline/locked-method pairs.
    random_baselines = {(r["scope"], r["selector"], r["shots_per_class"], r["record_id"]): r
                        for r in random_rows if r["method"] == "source_identity"}
    random_export = []; random_excel = []
    for row in random_rows:
        assert row["record_id"] in random_ids[row["scope"]] and int(row["repeat"]) == 0
        baseline = random_baselines[(row["scope"], row["selector"], row["shots_per_class"], row["record_id"])]
        correct, base_correct = int(row["is_correct"]), int(baseline["is_correct"])
        assert correct == int(row["y_true"] == row["y_predicted"])
        assert row["y_true"] == baseline["y_true"]
        changed = row["y_predicted"] != baseline["y_predicted"]
        change = ("改对" if correct and not base_correct else "改错" if base_correct and not correct
                  else "错判类别改变" if changed else "保持不变")
        used = support_count(row["method"], grid[row["candidate_id"]], 7 if row["scope"] == "seven_class_all" else 6,
                             int(row["shots_per_class"]))
        record = {**row, "baseline_prediction": baseline["y_predicted"], "baseline_is_correct": base_correct,
                  "change_vs_source": change, "actual_support_labels_used": used,
                  "sample_selection": "precommitted_uniform_random20_IDs_per_scope"}
        random_export.append(record)
        random_excel.append([SCOPES[row["scope"]], MODELS[row["selector"]], int(row["shots_per_class"]), used,
            method_name(row["method"]), row["candidate_id"], row["record_id"], row["y_true"],
            baseline["y_predicted"], base_correct, row["y_predicted"], correct, change, 0,
            "预先随机选定；包含识别错误"])
    assert len(random_export) == 400 and sum(int(r["is_correct"]) == 0 for r in random_export) == 24
    # Stream the seven-class files only. Every method and every draw is retained.
    aggregations = {}; baselines = {}; changes = []; raw_count = 0
    for selector in ("full", "mcuve20"):
        directory = run / "final/seven_class_all" / selector
        metadata = read_json(directory / "metadata.json")
        draws = {(int(d["repeat"]), int(d["shots_per_class"])): d["support_ids"] for d in metadata["draws"]}
        with gzip.open(directory / "predictions_k1.csv.gz", "rt", encoding="utf-8", newline="") as handle:
            for row in csv.DictReader(handle):
                raw_count += 1
                assert row["scope"] == "seven_class_all" and row["selector"] == selector
                correct = int(row["is_correct"]); repeat = int(row["repeat"]); identifier = row["record_id"]
                assert correct == int(row["y_true"] == row["y_predicted"])
                if row["method"] == "source_identity":
                    key = (selector, identifier)
                    value = (row["y_true"], row["y_predicted"], correct)
                    if key in baselines: assert baselines[key] == value
                    baselines[key] = value
                key = (selector, row["method"], identifier)
                stat = aggregations.setdefault(key, {"truth": row["y_true"], "candidate_id": row["candidate_id"],
                    "repeats": set(), "predictions": Counter(), "correct_count": 0, "repeat0": None})
                assert stat["truth"] == row["y_true"] and stat["candidate_id"] == row["candidate_id"] and repeat not in stat["repeats"]
                stat["repeats"].add(repeat); stat["predictions"][row["y_predicted"]] += 1; stat["correct_count"] += correct
                if repeat == 0: stat["repeat0"] = row["y_predicted"]
                if selector == "full" and row["method"] in {"family_damped_anchor", "family_mistake_bias"}:
                    before = baselines[(selector, identifier)]
                    if before[1] != row["y_predicted"]:
                        changes.append({**row, "baseline_prediction": before[1], "baseline_is_correct": before[2],
                            "corrected": int(not before[2] and correct), "harmed": int(before[2] and not correct),
                            "support_ids": "|".join(draws[(repeat, 1)]),
                            "diagnostic_scope": "posthoc_change_diagnosis_not_random_sample"})
    assert raw_count == 307800 and len(aggregations) == 3078
    statistics_export = []; statistics_excel = []
    for (selector, method, identifier), stat in sorted(aggregations.items()):
        assert stat["repeats"] == set(range(100)) and stat["repeat0"] is not None
        truth, before, before_correct = baselines[(selector, identifier)]
        n = len(stat["repeats"]); correct = stat["correct_count"]
        changed = n-stat["predictions"][before]
        corrected = correct if not before_correct else 0
        harmed = n-correct if before_correct else 0
        repeat0_correct = int(stat["repeat0"] == truth)
        row = {"scope": "seven_class_all", "selector": selector, "method": method,
            "candidate_id": stat["candidate_id"], "record_id": identifier, "y_true": truth,
            "source_prediction": before, "source_is_correct": before_correct,
            "repeat0_prediction": stat["repeat0"], "repeat0_is_correct": repeat0_correct,
            "support_draws": n, "correct_draws": correct, "correct_fraction": correct/n,
            "changed_draws": changed, "corrected_draws": corrected, "harmed_draws": harmed,
            "wrong_to_other_wrong_draws": changed-corrected-harmed,
            "prediction_counts_json": json.dumps(dict(sorted(stat["predictions"].items())), ensure_ascii=False)}
        statistics_export.append(row)
        used = support_count(method, grid[stat["candidate_id"]], 7, 1)
        statistics_excel.append([MODELS[selector], method_name(method), stat["candidate_id"], used, identifier, truth,
            before, before_correct, stat["repeat0"], repeat0_correct, n, correct, correct/n,
            changed, corrected, harmed, changed-corrected-harmed, row["prediction_counts_json"]])
    assert len(changes) == 55 and {r["record_id"] for r in changes} == {"ZY07-057"}
    assert Counter(r["method"] for r in changes) == {"family_damped_anchor": 36, "family_mistake_bias": 19}
    csv_paths = [report / "random_sample_label_comparison.csv", report / "seven_class_per_sample_statistics.csv",
                 report / "seven_class_correction_changes.csv"]
    for path, rows in zip(csv_paths, [random_export, statistics_export, changes]): write_csv(path, rows)
    workbook = Workbook(); workbook.remove(workbook.active)
    info = [
        ["数据任务", "七类主实验包含全部匿名产地组，171份固定预测；六类辅助实验明确排除ZY02，170份预测，不能当作七类成绩。"],
        ["标签依据", "正确标签是用户声明的原工作簿匿名产地组ZY01—ZY07；已核对编号、原表、元数据及NPZ的一致性。没有真实省市映射或原始采样凭证，不能宣称独立验证真实生长产地。"],
        ["模型输入", "只使用吸收系数光谱数值；编号、文件名和路径仅用于定位，不作为预测特征。"],
        ["固定分组", "513份训练／171份校准池／171份预测保持不变；七类只能统一1-shot，六类辅助研究1／3／5／10-shot。"],
        ["K的含义", "K为每类可提供的校准预算；最终方案若保持原模型，则实际使用0份校准标签。不能把标注K>0当作已经发生有效校准。"],
        ["方法选择", "所有配置先在训练组内部开发任务选择并锁定；development_selected是全局锁定方案。分家族诊断只用于比较，不能看到预测分数后替换全局方案。"],
        ["门禁列", "家族通过门禁不等于满足全局选择资格；全局选择还要求选择阶段满足准确率增益和逐类指标约束，并更保守地打破平分。"],
        ["零支持对照", "原型融合必须与相同配置、不使用校准标签的零支持对照比较；否则提升可能只是分类器结构变化。"],
        ["随机20样品", "每任务先按固定种子随机选20个ID，再展示两个模型、原模型和全局锁定方法、第0次抽样及各K；400行包含重复ID和24条错误记录，不是400份独立样品。"],
        ["逐样品统计", "七类逐样品统计包含两个模型的全部9种对照方法，共3078行；正确比例是同一预测样品在100次校准抽样中的成功频率，不是预测概率。"],
        ["改动记录", "55行是事后诊断，展示七类SVM两个小幅改善家族的所有改动；均涉及ZY07-057，其中弱化锚定36次、偏置修正19次。这不是随机样本或55份新样品。"],
        ["不确定性", "100次支持抽样共享同一预测组。SD和逐样品成功频率只描述抽样变化，不是独立批次置信区间。"],
        ["既往研究暴露", "全部855份先前已用于内部CV；当前研究属于同批次回顾性探索，不是首次未见确认集或真实跨批次验证。"],
        ["完整原始预测", "所有1,531,800条逐样品记录保留在run/final各任务/模型的predictions_k*.csv.gz；本表不把全部记录挤入Excel。"],
        ["正确率", "每组识别正确率=预测类别与提供标签一致的样品数／该组固定预测样品数；样品真值不随模型输出修改。"],
        ["导出过程", "本导出程序只读取已完成结果并计算对照表，没有重训模型、改变标签或挑选有利抽样。"],
    ]
    sheet = add_sheet(workbook, "说明", ["项目", "说明"], info, widths={1:24, 2:125})
    for row in range(2, sheet.max_row+1): sheet.row_dimensions[row].height = 47
    add_sheet(workbook, "关键汇总", ["任务范围", "源模型", "预算K／类", "方法", "候选ID", "实际算法家族", "实际使用标签数／次", "原模型准确率", "本方法平均准确率", "增益（百分点）", "准确率抽样SD", "最低准确率", "最高准确率", "平均宏F1", "平均平衡准确率", "提高次数", "持平次数", "下降次数", "平均改对样品数", "平均改错样品数", "相对匹配零支持增益（百分点）", "开发门禁通过", "解释范围"], summary_rows,
              percents=(8,9,11,12,13,15), decimals=(10,14,19,20,21), widths={1:26,2:27,4:39,5:31,6:25,23:50})
    random_sheet = add_sheet(workbook, "随机20样品", ["任务范围", "源模型", "预算K／类", "实际使用标签数", "比较阶段", "候选ID", "样品编号", "提供的正确标签", "原模型预测", "原模型正确（1／0）", "本方法预测", "本方法正确（1／0）", "相对原模型", "抽样序号", "样品选择规则"], random_excel,
                             widths={1:26,2:27,5:37,6:25,7:18,15:39})
    for row in random_sheet.iter_rows(min_row=2):
        if row[11].value == 0:
            for cell in row: cell.fill = PatternFill("solid", fgColor="FCE8E6")
    add_sheet(workbook, "七类逐样品统计", ["源模型", "方法", "候选ID", "实际使用标签数", "样品编号", "提供的正确标签", "原模型预测", "原模型正确（1／0）", "第0次预测", "第0次正确（1／0）", "支持抽样次数", "正确次数", "100次正确比例", "改变类别次数", "改对次数", "改错次数", "错判类别间改变次数", "预测类别计数"], statistics_excel,
              percents=(13,), widths={1:27,2:39,3:31,5:18,18:55})
    change_excel = [["事后改动诊断（不是随机样本）", MODELS[r["selector"]], method_name(r["method"]), r["candidate_id"], int(r["repeat"]), r["record_id"], r["y_true"], r["baseline_prediction"], r["baseline_is_correct"], r["y_predicted"], int(r["is_correct"]), r["corrected"], r["harmed"], r["support_ids"]] for r in changes]
    add_sheet(workbook, "改动记录", ["记录性质", "源模型", "方法", "候选ID", "支持抽样序号", "样品编号", "提供的正确标签", "原模型预测", "原模型正确（1／0）", "校正后预测", "校正后正确（1／0）", "本次改对", "本次改错", "本次校准样品编号"], change_excel,
              widths={1:38,2:27,3:34,4:25,6:18,14:80})
    workbook.save(output)
    tracked_workbook = report / output.name
    if tracked_workbook.resolve() != output.resolve(): shutil.copy2(output, tracked_workbook)
    reopened = load_workbook(output, read_only=True, data_only=True)
    expected_counts = {"说明": len(info), "关键汇总":90, "随机20样品":400, "七类逐样品统计":3078, "改动记录":55}
    try:
        assert reopened.sheetnames == list(expected_counts)
        assert all(reopened[name].max_row-1 == count for name,count in expected_counts.items())
        assert all(isinstance(row[8], (float,int)) for row in reopened["关键汇总"].iter_rows(min_row=2, values_only=True))
        assert sum(row[11] == 0 for row in reopened["随机20样品"].iter_rows(min_row=2, values_only=True)) == 24
    finally: reopened.close()
    validation = {"status":"passed", "export_is_read_only_analysis":True, "models_fitted":0,
        "worksheet_data_rows":expected_counts, "random_rows":400, "random_incorrect_rows_preserved":24,
        "seven_class_raw_rows_read":raw_count, "seven_class_aggregates":len(statistics_export),
        "all_seven_class_methods_included":True, "changed_records":len(changes),
        "changed_unique_IDs":sorted({r['record_id'] for r in changes}),
        "labels_are_anonymous_groups_not_verified_geographic_names":True,
        "workbook_path":str(tracked_workbook), "workbook_sha256":digest(output),
        "source_input_sha256":{name:digest(run/name) for name in ["status.json","summary.json","selection_lock.json","random_query_ids.json","random_query_predictions.csv"]},
        "csv_sha256":{p.name:digest(p) for p in csv_paths}}
    (report / "table_export_validation.json").write_text(json.dumps(validation,ensure_ascii=False,indent=2)+"\n", encoding="utf-8")
    print(json.dumps({"status":"passed", "workbook":str(output), "tracked_workbook":str(tracked_workbook),
                      "rows":expected_counts, "bytes":output.stat().st_size}, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run", type=Path, default=Path("outputs/zaoren-calibration-optimization-20261031/run"))
    parser.add_argument("--report", type=Path, default=Path("reports/calibration-optimization/2026-10-31"))
    parser.add_argument("--output", type=Path, default=Path("outputs/zaoren-calibration-optimization-20261031/校准探索_预测与标签对照.xlsx"))
    args = parser.parse_args()
    export(args.run, args.report, args.output)


if __name__ == "__main__":
    main()
