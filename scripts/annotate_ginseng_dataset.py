#!/usr/bin/env python3
"""Make a readable workbook from the 18 original ginseng CSV files.

No labels are inferred from spectra. Species names use the repository's coding;
sample pairing follows the group suffix and within-group order. This convention
does not establish independent plant, batch, or growing-origin identities.
"""

import argparse
import csv
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE = "https://data.mendeley.com/datasets/hvbp5mdtxx/2"
SPECIES = {1: "亚洲人参", 0: "西洋参"}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_dir", type=Path)
    parser.add_argument("output_dir", type=Path)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    info = wb.active
    info.title = "说明"
    notes = [
        ["项目", "说明"],
        ["来源", SOURCE],
        ["引用", "Mai, Jamil (2025), Complex Refractive Indices of Ginseng/American Ginseng Tablets, Mendeley Data V2, doi:10.17632/hvbp5mdtxx.2"],
        ["许可", "CC BY 4.0；本工作簿增加中文说明和整理编号，原始CSV不修改"],
        ["物种编码", "1=亚洲人参（Panax ginseng）；0=西洋参（Panax quinquefolius）"],
        ["配对规则", "同组data_nj、data_aj的第2–8列，按顺序配对label_j的第1–7项"],
        ["配对边界", "来源未附原始样品ID对照表；此为按组号和列序整理的映射，不能独立验证原植物身份"],
        ["整理编号", "G1_S01等为本次生成的编号，不是发布者的原始样品ID"],
        ["频率", "实际读取为65点、0.2–1.6 THz；第一行为表头，不是频率数据"],
        ["折射率n", "无量纲，来源记为R.I.U."],
        ["吸收系数alpha", "cm^-1；未截断负值，未平滑、筛选或填补原始数值"],
        ["缺少信息", "生长产地、年限、化学含量、原植物及批次映射、原始样品/参考时域波形"],
        ["样本单位", "42片压片，不自动等于42个独立植株或批次"],
        ["适合任务", "亚洲人参与西洋参物种二分类；不能作为酸枣仁产地标签"],
    ]
    for row in notes:
        info.append(row)
    metadata = wb.create_sheet("样本标签")
    metadata.append(["整理编号", "原文件组号", "组内序号", "原谱文件列号", "原谱Excel列", "原标签值", "中文物种", "折射率文件", "吸收系数文件", "标签文件", "标签项序号", "生长产地（原包未提供）", "原植物ID（原包未提供）", "批次（原包未提供）"])
    labels_sheet = wb.create_sheet("原始标签排列")
    labels_sheet.append(["标签文件"] + [f"第{k}项" for k in range(1, 8)])
    n_sheet = wb.create_sheet("折射率n_无量纲")
    a_sheet = wb.create_sheet("吸收系数alpha_cm-1")
    headers = ["频率（THz）"]
    n_blocks, a_blocks, metadata_rows, members = [], [], [], []
    frequency = None
    counts = {0: 0, 1: 0}
    for group in range(1, 7):
        n_path = args.input_dir / f"data_n{group}.csv"
        a_path = args.input_dir / f"data_a{group}.csv"
        y_path = args.input_dir / f"label_{group}.csv"
        n = np.loadtxt(n_path, delimiter=",", skiprows=1, encoding="utf-8-sig")
        a = np.loadtxt(a_path, delimiter=",", skiprows=1, encoding="utf-8-sig")
        y = np.loadtxt(y_path, delimiter=",", ndmin=1, encoding="utf-8-sig")
        if n.shape != (65, 8) or a.shape != (65, 8) or y.shape != (7,):
            raise ValueError(f"Unexpected shape in group {group}")
        if not set(y).issubset({0, 1}):
            raise ValueError(f"Unknown label in group {group}")
        if not np.isfinite(n).all() or not np.isfinite(a).all():
            raise ValueError(f"Nonfinite spectrum in group {group}")
        if frequency is None:
            frequency = n[:, 0]
        if not np.array_equal(frequency, n[:, 0]) or not np.array_equal(frequency, a[:, 0]):
            raise ValueError(f"Frequency mismatch in group {group}")
        if not np.all(np.diff(frequency) > 0):
            raise ValueError("Frequency must increase strictly")
        labels_sheet.append([y_path.name] + y.astype(int).tolist())
        for idx, value in enumerate(y, start=1):
            label = int(value)
            sid = f"G{group}_S{idx:02d}"
            row = [sid, group, idx, idx + 1, get_column_letter(idx + 1), label, SPECIES[label], n_path.name, a_path.name, y_path.name, idx, None, None, None]
            metadata.append(row)
            metadata_rows.append(row)
            headers.append(f"{sid}｜{SPECIES[label]}｜标签{label}")
            counts[label] += 1
        n_blocks.append(n[:, 1:])
        a_blocks.append(a[:, 1:])
        for path in [n_path, a_path, y_path]:
            content = path.read_bytes()
            members.append({"file": path.name, "bytes": len(content), "sha256": hashlib.sha256(content).hexdigest()})
    for sheet, blocks in [(n_sheet, n_blocks), (a_sheet, a_blocks)]:
        sheet.append(headers)
        for row in np.column_stack([frequency] + blocks).tolist():
            sheet.append(row)
    manifest = wb.create_sheet("原CSV校验")
    manifest.append(["原文件", "字节数", "SHA256"])
    for member in members:
        manifest.append([member["file"], member["bytes"], member["sha256"]])
    for sheet in wb:
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305A77")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 25
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 105
    for row in info.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        info.row_dimensions[row[0].row].height = 36
    out = args.output_dir / "人参太赫兹数据_中文标注.xlsx"
    wb.save(out)
    with (args.output_dir / "sample_labels.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow([cell.value for cell in metadata[1]])
        writer.writerows(metadata_rows)
    summary = {
        "source_url": SOURCE,
        "input_csv_count": len(members),
        "pellet_count": sum(counts.values()),
        "species_counts": {SPECIES[k]: v for k, v in counts.items()},
        "frequency_points": len(frequency),
        "frequency_range_thz": [float(frequency[0]), float(frequency[-1])],
        "pairing_basis": "Same group suffix and within-group order; no original sample-ID table supplied",
        "independent_plants_verified": False,
        "growing_origins_provided": False,
        "model_runs": 0,
        "member_manifest": members,
    }
    (args.output_dir / "inspection.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n")
    print(out)
    print(json.dumps({k: v for k, v in summary.items() if k != "member_manifest"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
