"""Create scan-to-sample annotation rows without inferring scientific labels."""
import argparse
import csv
import hashlib
from pathlib import Path
import re
import warnings

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from herb_spectra.dataset_registry import REPO, select_workbooks


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    selection = parser.add_mutually_exclusive_group()
    selection.add_argument("--manifest", type=Path, help="Explicit active-workbook registry")
    selection.add_argument("--data-root", type=Path,
                           help="Explicit legacy directory override; bypasses the active registry")
    parser.add_argument("--output", type=Path, default=Path("data/metadata_template.csv"))
    args = parser.parse_args()
    if args.output.exists():
        raise FileExistsError("Keep existing annotations; choose a new output path")
    warnings.filterwarnings("ignore", message="Workbook contains no default style")
    rows = []
    sources = select_workbooks(repo=REPO, manifest=args.manifest, data_root=args.data_root)
    record_ids = set()
    for source in sources:
        wb = load_workbook(source.path, read_only=True, data_only=True)
        try:
            ws = wb["吸收系数"]
            ws.reset_dimensions()
            header = next(ws.iter_rows(values_only=True), ())
            count = 0
            for i, name in enumerate(header):
                if i == 0 or name is None or not str(name).strip() or "参考" in str(name):
                    continue
                key = f"{source.relative_path}|{ws.title}|{i + 1}"
                record_id = ""
                if source.dataset == "zaoren":
                    match = re.fullmatch(r"(ZY[0-9]{2}-[0-9]{3,})(?:_.*)?", str(name))
                    if match is None or not match.group(1).startswith(source.origin_group_id + "-"):
                        raise ValueError("A current zaoren header does not match its registered group")
                    record_id = match.group(1)
                    if record_id in record_ids:
                        raise ValueError("A zaoren record ID occurs more than once")
                    record_ids.add(record_id)
                rows.append({"scan_id": hashlib.sha256(key.encode()).hexdigest()[:20],
                             "source_workbook": source.relative_path, "sheet": ws.title,
                             "excel_column": get_column_letter(i + 1), "scan_header": str(name),
                             "physical_sample_id": "", "origin": "", "origin_verified": "",
                             "batch_id": "", "instrument_id": "", "measurement_date": "",
                             "species_verified": "", "assay_value": "", "assay_unit": "",
                             "assay_method": "", "notes": "Explicit --data-root override; active registry bypassed"
                             if source.selection_mode == "explicit_data_root_override" else "",
                             "record_id": record_id, "origin_group_id": source.origin_group_id or "",
                             "dataset_version": source.dataset_version})
                count += 1
            if source.record_count is not None and count != source.record_count:
                raise ValueError(f"Registered record_count differs from workbook headers: {source.relative_path}")
        finally:
            wb.close()
    if not rows:
        raise ValueError("No spectrum columns found")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader(); writer.writerows(rows)
    print(f"Created {len(rows)} scan annotation rows ({sources[0].selection_mode}; "
          f"{sources[0].dataset_version}): {args.output}")


if __name__ == "__main__":
    main()
