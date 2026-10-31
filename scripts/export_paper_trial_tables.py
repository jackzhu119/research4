"""Export readable paper-inspired trial results; never fit or modify a model."""
import argparse
from collections import Counter, defaultdict
import csv
import gzip
import hashlib
import json
from pathlib import Path
import shutil

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


STRATEGIES = {"head_ft":"冻结编码器、只微调分类头", "full_ft":"完整网络微调",
              "prototype":"分类ProtoNet", "fomaml_ft":"一阶MAML＋微调"}
METHODS = {"zero":"自身零支持对照", "selected":"开发选择并锁定的方案",
           "diagnostic":"开发选择的诊断配置"}


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def read_csv(path):
    with Path(path).open(encoding="utf-8-sig",newline="") as handle:
        return list(csv.DictReader(handle))


def write_csv(path, rows):
    with Path(path).open("w",encoding="utf-8-sig",newline="") as handle:
        writer=csv.DictWriter(handle,fieldnames=list(rows[0]),lineterminator="\n")
        writer.writeheader();writer.writerows(rows)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def add_sheet(workbook,title,headers,rows,*,percents=(),decimals=(),widths=None):
    sheet=workbook.create_sheet(title);sheet.append(headers)
    for row in rows:sheet.append(row)
    sheet.freeze_panes="A2";sheet.auto_filter.ref=sheet.dimensions;sheet.sheet_view.zoomScale=85
    sheet.row_dimensions[1].height=43
    for cell in sheet[1]:
        cell.font=Font(name="Microsoft YaHei",bold=True,color="FFFFFF",size=10)
        cell.fill=PatternFill("solid",fgColor="234F70")
        cell.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font=Font(name="Microsoft YaHei",size=10)
            cell.alignment=Alignment(wrap_text=True,vertical="center")
            if cell.row%2==0:cell.fill=PatternFill("solid",fgColor="F1F5F9")
        sheet.row_dimensions[row[0].row].height=32
    for column in percents:
        for row in sheet.iter_rows(min_row=2,min_col=column,max_col=column):row[0].number_format="0.00%"
    for column in decimals:
        for row in sheet.iter_rows(min_row=2,min_col=column,max_col=column):row[0].number_format="0.0000"
    for i in range(1,len(headers)+1):sheet.column_dimensions[get_column_letter(i)].width=(widths or {}).get(i,17)
    sheet.page_setup.orientation="landscape";sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
    sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    return sheet


def export(run,report,output):
    run,report,output=Path(run),Path(report),Path(output)
    status=read_json(run/"status.json")
    if status["status"]!="completed":raise ValueError("Complete the neural trial before export")
    report.mkdir(parents=True,exist_ok=True);output.parent.mkdir(parents=True,exist_ok=True)
    summary=read_json(run/"summary.json");per_seed=read_json(run/"per_seed_summary.json")
    locked=read_json(run/"selection_lock.json")["choices"]
    classes=read_json(run/"final/seed_42/training.json")["classes"]
    metadata=read_json(run/"final_metadata.json")
    seeds=sorted(m["seed"] for m in metadata)
    raw_random=read_csv(run/"random_predictions.csv")
    random_ids=set(read_json(run/"random_query_ids.json"))
    query_ids=set(metadata[0]["query_ids"])
    assert all(set(m["query_ids"])==query_ids for m in metadata)
    group=defaultdict(lambda:{"truth":None,"zero_prediction":None,"zero_rows":0,
                              "selected_counts":Counter(),"selected_rows":0,"selected_correct":0,
                              "repeat0_prediction":None})
    raw_rows=0;raw_index={}
    for path in sorted((run/"final").glob("seed_*/predictions.csv.gz")):
        with gzip.open(path,"rt",encoding="utf-8",newline="") as handle:
            for row in csv.DictReader(handle):
                raw_rows+=1
                assert int(row["is_correct"])==int(row["y_true"]==row["y_predicted"])
                if row["record_id"] in random_ids and int(row["repeat"])==0:
                    key=(int(row["seed"]),row["strategy"],row["method"],row["record_id"])
                    raw_index[key]=(row["y_true"],row["y_predicted"],int(row["is_correct"]))
                if row["strategy"]!="fomaml_ft" or row["method"] not in {"zero","selected"}:continue
                key=(int(row["seed"]),row["record_id"]);value=group[key]
                if value["truth"] is None:value["truth"]=row["y_true"]
                assert value["truth"]==row["y_true"]
                if row["method"]=="zero":
                    if value["zero_prediction"] is None:value["zero_prediction"]=row["y_predicted"]
                    assert value["zero_prediction"]==row["y_predicted"]
                    value["zero_rows"]+=1
                else:
                    value["selected_counts"][row["y_predicted"]]+=1;value["selected_rows"]+=1
                    value["selected_correct"]+=int(row["is_correct"])
                    if int(row["repeat"])==0:value["repeat0_prediction"]=row["y_predicted"]
    assert raw_rows==status["prediction_rows"]
    for row in raw_random:
        key=(int(row["seed"]),row["strategy"],row["method"],row["record_id"])
        assert raw_index[key]==(row["y_true"],row["y_predicted"],int(row["is_correct"]))
    expected_random=len(random_ids)*len(seeds)*len(STRATEGIES)*len(METHODS)
    assert len(raw_random)==expected_random
    fomaml_config=locked["fomaml_ft"]["selected"]
    sample_rows=[];class_stats={c:{"ids":set(),"source_correct":0,"selected_correct":0,
        "corrected":0,"harmed":0,"events":0} for c in classes}
    for (seed,identifier),value in sorted(group.items()):
        n=value["selected_rows"]
        assert n==value["zero_rows"]==len(next(m for m in metadata if m["seed"]==seed)["support_draws"])
        truth=value["truth"];before=value["zero_prediction"];source_correct=int(before==truth)
        after_correct=value["selected_correct"]
        changed=n-value["selected_counts"].get(before,0)
        corrected=after_correct if not source_correct else 0
        harmed=n-after_correct if source_correct else 0
        mode=min(value["selected_counts"],key=lambda name:(-value["selected_counts"][name],name))
        item={"seed":seed,"record_id":identifier,"y_true":truth,"source_prediction":before,
              "source_is_correct":source_correct,"selected_candidate":"zero_fomaml_ft" if fomaml_config is None else fomaml_config["id"],
              "actual_support_labels":0 if fomaml_config is None else len(classes),
              "learning_rate":0 if fomaml_config is None else fomaml_config["learning_rate"],
              "steps":0 if fomaml_config is None else fomaml_config["steps"],
              "repeat0_prediction":value["repeat0_prediction"],"repeat0_is_correct":int(value["repeat0_prediction"]==truth),
              "support_draws":n,"selected_correct_count":after_correct,"selected_correct_fraction":after_correct/n,
              "modal_prediction":mode,"changed_draws":changed,"corrected_draws":corrected,"harmed_draws":harmed,
              "predicted_class_counts":json.dumps(dict(sorted(value["selected_counts"].items())),ensure_ascii=False)}
        sample_rows.append(item)
        s=class_stats[truth];s["ids"].add(identifier);s["source_correct"]+=source_correct*n
        s["selected_correct"]+=after_correct;s["corrected"]+=corrected;s["harmed"]+=harmed;s["events"]+=n
    assert len(sample_rows)==len(seeds)*len(query_ids)
    assert sum(r["selected_correct_count"] for r in sample_rows)==sum(s["selected_correct"] for s in class_stats.values())
    diagnostic=read_json(run/"fomaml_class_diagnostic.json")
    class_rows=[]
    scenarios=len(metadata)*len(metadata[0]["support_draws"])
    for c,s in class_stats.items():
        before=s["source_correct"]/s["events"];after=s["selected_correct"]/s["events"]
        assert abs(before-diagnostic["zero"]["recall_mean"][c])<1e-12
        assert abs(after-diagnostic["selected"]["recall_mean"][c])<1e-12
        assert s["selected_correct"]-s["source_correct"]==s["corrected"]-s["harmed"]
        class_rows.append([c,len(s["ids"]),scenarios,s["source_correct"],s["selected_correct"],before,after,
                           100*(after-before),s["corrected"],s["harmed"],s["corrected"]/scenarios,s["harmed"]/scenarios])
    overall_correct=sum(s["selected_correct"] for s in class_stats.values())/sum(s["events"] for s in class_stats.values())
    expected=next(s["accuracy_mean"] for s in summary if s["strategy"]=="fomaml_ft" and s["method"]=="selected")
    assert abs(overall_correct-expected)<1e-12
    csv_paths=[]
    for filename,rows in [("random_examples.csv",raw_random),("fomaml_per_sample.csv",sample_rows),("per_seed_summary.csv",per_seed)]:
        path=report/filename;write_csv(path,rows);csv_paths.append(path)

    workbook=Workbook();workbook.remove(workbook.active)
    info=[
        ["试验性质","读取用户上传论文后进行的分类适配试验，使用真实太赫兹光谱与ZY01—ZY07匿名标签；不是论文原始回归结果复现。"],
        ["数据分组","固定513份训练／171份校准池／171份预测；所有最终结果均为七类，没有删除ZY02。物理样品按用户声明独立。"],
        ["标签含义","当前核对的是原表与匿名产地组的一致性；真实生长省市未提供，不能仅靠模型独立核实地理来源。"],
        ["支持样品计数","本试验每类1份，共7份支持；原论文5／10／25-shot是每个回归任务的总支持实例数，两者不可直接等同。"],
        ["共同主干","紧凑ResNet1D、GroupNorm、32维表示。上传正文缺少精确结构附录，作者固定代码快照还缺models/cnn.py，因此不声称精确复现原架构。"],
        ["公平比较","所有策略同一个40轮监督初始化，再对相同源任务编号各做100次额外更新。更新次数相同不等于FLOPs、训练时间相同。"],
        ["源元任务","开发源子集的ZY02仅1份，不能构成不重叠支持／查询，故仅6个充分样本类可参与源元任务；最终513源样品中ZY02有3份，7类均可参与1＋1源任务。"],
        ["七类预测不变","上述6类是开发阶段元任务的可采样类别，不是删成六类预测。源监督初始化和最终识别均包含7类。"],
        ["开发选择","6个源内开发任务，前4选择、后2检查；每任务3种随机种子、5次支持抽样。所有配置在171份最终预测评分前锁定。"],
        ["三个种子","42、43、44全部保留，没有根据最终准确率挑最好种子。总体SD混合了种子与支持抽样变化，可同时查看逐种子表。"],
        ["最终场景","每个种子100次支持抽样，共300个同一预测组上的场景，不是300批药材或300组独立样品。"],
        ["收益归属","每种方法的微调收益相对自身零支持预测器；编码器或分类器变好与少样本标签带来的收益分开记录。"],
        ["一阶MAML结论","本设置中64.72%→67.55%，提高2.83个百分点；仍低于此前SVM的88.89%，不能据此称当前总体最优或替代SVM。"],
        ["逐类风险","FO-MAML的ZY02预测召回在微调前后均为0；ZY06、ZY07平均召回下降，整体提高不代表每个产地都提高。"],
        ["样品成功频率","逐样品表的100次正确比例是支持抽样下的成功频率，不是校准过的置信概率。所有错误和新增误判均保留。"],
        ["论文边界","原文为NIR／MIR连续属性回归；仪器漂移部分是人工扰动压力测试，不是实测新仪器或本数据真实跨批次验证。"],
        ["当前证据范围","所有855份及当前预测组先前已参与研究，属于同批次回顾性探索；元任务是源样品重抽，不是新采集批次。"],
        ["源模型冻结","适配使用源模型副本，源权重指纹保持不变。完整网络微调会改变副本编码器，不能称其编码器也被冻结。"],
    ]
    sheet=add_sheet(workbook,"说明",["项目","解释"],info,widths={1:25,2:125})
    for row in range(2,sheet.max_row+1):sheet.row_dimensions[row].height=48
    paper_rows=[
        ["只微调分类头","原文FT强调源预训练后适配；本表增加冻结编码器的对照","监督初始化＋相同额外源更新，冻结编码器，只更新7类头","相同监督源分类器","9个学习率／步数组合；当前开发流程退回零支持","不是原文完整FT的同义词"],
        ["完整网络微调","原文Fine-tuning对已有模型进行目标任务适配","复制整个网络，用7份支持样品SGD更新全部参数","相同监督源分类器","9配置；开发流程退回零支持，诊断负结果保留","只用支持样品更新，最终query不作早停"],
        ["分类ProtoNet","原文为按距离加权连续值的回归原型方式","真实源内3-way episodic训练；最终7类按支持／源类别中心距离判别","同一ProtoNet编码器的源类别中心","4种源／支持中心混合比例；主选择零支持","不是直接把原文回归输出作为类别，也不是未训练的普通PCA原型"],
        ["FO-MAML＋微调","原文MAML及作者代码有一阶／二阶开关，出版表采用哪一开关未核验","真实一阶元梯度训练；支持内步后用源任务query梯度更新初始化","相同元训练网络、不使用目标支持标签","9配置；锁定学习率0.001、10步","不是完整二阶MAML复现；小样本收益与绝对性能分开"],
        ["共享ResNet1D","原文使用共同ResNet骨干以控制结构差异","紧凑16／32通道、GroupNorm、池化8、表示32、7类头","所有学习策略使用同一架构","固定40轮监督＋100次额外源更新","精确原结构及附录未取得，不能称原文架构已完整复现"],
        ["任务与指标","原文预测连续品质，报告MSE／MAE／R²","这里预测匿名产地类别，报告准确率／宏F1／平衡准确率","固定171个独立声明样品","不是用产地数字作连续回归","原论文高R²不能直接当作本研究产地准确率"],
    ]
    sheet=add_sheet(workbook,"论文方法对照",["方法","原文概念","当前实际尝试","自身零支持对照","试验状态","不能混淆的内容"],paper_rows,widths={1:24,2:42,3:62,4:42,5:48,6:58})
    for row in range(2,sheet.max_row+1):sheet.row_dimensions[row].height=70
    summary_rows=[]
    for r in summary:
        source=next(v for v in summary if v['strategy']==r['strategy'] and v['method']=='zero')
        selected=locked[r['strategy']]['selected']
        use=0 if r['method']=='zero' or (r['method']=='selected' and selected is None) else len(classes)
        summary_rows.append([STRATEGIES[r['strategy']],METHODS[r['method']],r['candidate_id'],use,
            source['accuracy_mean'],r['accuracy_mean'],100*r['accuracy_gain_mean'],r['accuracy_sd'],r['accuracy_min'],r['accuracy_max'],
            r['macro_f1_mean'],r['balanced_accuracy_mean'],r['positive_draws'],r['unchanged_draws'],r['negative_draws'],r['evaluations'],
            '42、43、44；每个100次支持抽样','相对自身零支持；不是相对SVM'])
    sheet=add_sheet(workbook,"总体结果",["学习策略","结果身份","候选ID","实际目标标签数／次","自身零支持准确率","本项平均准确率","增益（百分点）","跨种子及抽样SD","最低准确率","最高准确率","平均宏F1","平均平衡准确率","提高场景","持平场景","下降场景","评价场景数","种子与抽样","比较基准"],summary_rows,
        percents=(5,6,8,9,10,12),decimals=(7,11),widths={1:28,2:33,3:28,17:42,18:41})
    for row in sheet.iter_rows(min_row=2):
        if row[6].value<0:
            for cell in row:cell.fill=PatternFill('solid',fgColor='FCE8E6')
    seed_rows=[[STRATEGIES[r['strategy']],r['seed'],r['zero_accuracy'],r['selected_accuracy'],100*(r['selected_accuracy']-r['zero_accuracy']),
                r['diagnostic_accuracy'],r['zero_macro_f1'],r['selected_macro_f1'],r['zero_balanced_accuracy'],r['selected_balanced_accuracy'],
                '每种策略全部三个种子都保留'] for r in per_seed]
    add_sheet(workbook,"三个种子",["学习策略","随机种子","自身零支持准确率","主选择平均准确率","主选择增益（百分点）","诊断配置平均准确率","自身零支持宏F1","主选择宏F1","自身零支持平衡准确率","主选择平衡准确率","说明"],seed_rows,
              percents=(3,4,6,9,10),decimals=(5,7,8),widths={1:29,11:43})
    add_sheet(workbook,"FOMAML逐产地",["匿名组","固定预测样品数","场景数（非独立批次）","原正确预测记录数","后正确预测记录数","原召回率","后召回率","召回变化（百分点）","改对记录数","改错记录数","平均每场景改对数","平均每场景改错数"],class_rows,
              percents=(6,7),decimals=(8,11,12),widths={3:27})
    random_excel=[]
    for r in raw_random:
        strategy,method=r['strategy'],r['method'];choice=locked[strategy]
        config=None if method=='zero' else choice[method]
        random_excel.append([int(r['seed']),int(r['repeat']),STRATEGIES[strategy],METHODS[method],r['candidate_id'],
            0 if config is None else len(classes),r['record_id'],r['y_true'],r['y_predicted'],int(r['is_correct']),
            '固定种子预选20个ID，未过滤错误'])
    sheet=add_sheet(workbook,"随机样品",["种子","抽样序号","学习策略","结果身份","候选ID","实际目标标签数","样品编号","提供的标签","预测标签","正确（1／0）","选择规则"],random_excel,
                    widths={3:29,4:33,5:28,7:20,11:47})
    for row in sheet.iter_rows(min_row=2):
        if row[9].value==0:
            for cell in row:cell.fill=PatternFill('solid',fgColor='FCE8E6')
    sample_excel=[[r[k] for k in ['seed','record_id','y_true','source_prediction','source_is_correct','selected_candidate','actual_support_labels','learning_rate','steps',
        'repeat0_prediction','repeat0_is_correct','support_draws','selected_correct_count','selected_correct_fraction','modal_prediction','changed_draws','corrected_draws','harmed_draws','predicted_class_counts']] for r in sample_rows]
    add_sheet(workbook,"FOMAML逐样品",["种子","样品编号","提供的标签","原模型预测","原模型正确（1／0）","锁定候选ID","实际目标标签数","微调学习率","微调步数","第0次微调预测","第0次正确（1／0）","支持抽样次数","微调后正确次数","100次正确比例","微调后众数类别","改变类别次数","改对次数","改错次数","预测类别计数"],sample_excel,
              percents=(14,),decimals=(8,),widths={2:20,6:30,19:56})
    workbook.save(output)
    tracked=report/output.name
    if tracked.resolve()!=output.resolve():shutil.copy2(output,tracked)
    expected_counts={"说明":len(info),"论文方法对照":len(paper_rows),"总体结果":len(summary_rows),"三个种子":len(seed_rows),
                     "FOMAML逐产地":len(class_rows),"随机样品":len(raw_random),"FOMAML逐样品":len(sample_rows)}
    reopened=load_workbook(output,read_only=True,data_only=True)
    try:
        assert reopened.sheetnames==list(expected_counts)
        assert all(reopened[name].max_row-1==count for name,count in expected_counts.items())
        incorrect_random=sum(int(r['is_correct'])==0 for r in raw_random)
        assert sum(r[9]==0 for r in reopened['随机样品'].iter_rows(min_row=2,values_only=True))==incorrect_random
        assert all(isinstance(r[5],(int,float)) for r in reopened['总体结果'].iter_rows(min_row=2,values_only=True))
    finally:reopened.close()
    validation={"status":"passed","models_fitted":0,"labels_changed":False,"raw_predictions_read":raw_rows,
        "worksheet_data_rows":expected_counts,"preselected_random_IDs":len(random_ids),"random_error_rows_preserved":incorrect_random,
        "fomaml_sample_statistics":len(sample_rows),"distinct_fixed_query_samples":len(query_ids),"random_seeds":seeds,
        "all_fomaml_per_class_recalls_match_saved_diagnostic":True,"all_per_sample_paired_arithmetic_pass":True,
        "independent_geographic_origin_verification":False,"comparison_scope":"own matched zero-support model;SVM88.89 separate reference",
        "workbook_sha256":digest(output),"csv_sha256":{p.name:digest(p) for p in csv_paths},
        "input_sha256":{name:digest(run/name) for name in ['status.json','summary.json','selection_lock.json','random_predictions.csv','per_seed_summary.json','fomaml_class_diagnostic.json']}}
    (report/'table_export_validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({"status":"passed","workbook":str(output),"bytes":output.stat().st_size,"rows":expected_counts},ensure_ascii=False))


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--run',type=Path,default=Path('outputs/paper-fewshot-pilot-20261031'))
    parser.add_argument('--report',type=Path,default=Path('reports/paper-fewshot-classification/2026-10-31'))
    parser.add_argument('--output',type=Path,default=Path('outputs/paper-fewshot-pilot-20261031/论文方法试验_预测与标签对照.xlsx'))
    args=parser.parse_args();export(args.run,args.report,args.output)


if __name__=='__main__':main()
