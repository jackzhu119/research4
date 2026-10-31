"""Prepare the active paired acid-jujube spectra without inventing identities.

Workbook groups are labels supplied by the active registry. Physical IDs and
verification flags come exclusively from the explicit metadata table. All
frequency points and all records are retained; diagnostic similarity never
causes automatic deletion or feature selection.
"""
from collections import Counter, defaultdict
import csv
from dataclasses import dataclass
import hashlib
import json
from pathlib import Path
import re
import warnings

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .dataset_registry import REPO, file_sha256, select_workbooks


ID_PATTERN = re.compile(r"^(ZY[0-9]{2}-[0-9]{3,})(?:_.*)?$")


class DataQualityError(ValueError):
    """Input values or identities cannot safely form the paired dataset."""


@dataclass(frozen=True)
class NumericSheet:
    name: str
    axis_header: str
    axis: np.ndarray
    values: np.ndarray
    record_ids: tuple[str, ...]
    headers: tuple[str, ...]
    columns: tuple[str, ...]
    nonnumeric_cells: int


def explicit_bool(value):
    """Parse an explicit metadata flag; missing values stay false."""
    if value is None or str(value).strip() == "":
        return False
    normalized = str(value).strip().lower()
    if normalized in {"true", "1", "yes"}:
        return True
    if normalized in {"false", "0", "no"}:
        return False
    raise DataQualityError(f"Unrecognized verification flag: {value!r}")


def _base_id(header):
    match = ID_PATTERN.fullmatch(str(header).strip())
    if match is None:
        raise DataQualityError(f"A numbered spectral header is required: {header!r}")
    return match.group(1)


def read_numeric_sheet(path, name):
    """Read actual cells, including nonfinite values for diagnostic reporting."""
    warnings.filterwarnings("ignore", message="Workbook contains no default style")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if name not in workbook.sheetnames:
            raise DataQualityError(f"Missing worksheet {name}")
        sheet = workbook[name]
        sheet.reset_dimensions()
        rows = iter(sheet.iter_rows(values_only=True))
        header = next(rows, ())
        if not header:
            raise DataQualityError(f"Empty worksheet {name}")
        columns = [i for i, cell in enumerate(header)
                   if i > 0 and cell is not None and str(cell).strip()
                   and "参考" not in str(cell)
                   and str(cell).strip().lower() not in {"reference", "ref"}]
        if not columns:
            raise DataQualityError(f"No sample columns in {name}")
        identifiers = tuple(_base_id(header[i]) for i in columns)
        if len(identifiers) != len(set(identifiers)):
            raise DataQualityError(f"Duplicate record ID in worksheet {name}")
        data, nonnumeric = [], 0
        for row in rows:
            if all(value is None for value in row):
                continue
            values = []
            for i in [0, *columns]:
                value = row[i] if i < len(row) else None
                if value is None:
                    values.append(np.nan)
                    continue
                try:
                    values.append(float(value))
                except (TypeError, ValueError):
                    values.append(np.nan)
                    nonnumeric += 1
            data.append(values)
        if not data:
            raise DataQualityError(f"No numeric rows in worksheet {name}")
        matrix = np.asarray(data, dtype=float)
        axis = matrix[:, 0]
        if (len(axis) < 2 or not np.isfinite(axis).all()
                or not np.all(np.diff(axis) > 0)):
            raise DataQualityError(f"Nonfinite or non-increasing frequency axis in {name}")
        return NumericSheet(name, str(header[0]), axis, matrix[:, 1:].T,
                            identifiers, tuple(str(header[i]) for i in columns),
                            tuple(get_column_letter(i + 1) for i in columns), nonnumeric)
    finally:
        workbook.close()


def _load_header_mapping(path):
    if not path.is_file():
        raise DataQualityError(f"Missing numbering correspondence CSV: {path.name}")
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    mapping = {}
    for row in rows:
        key = row.get("sheet"), row.get("cell")
        if key in mapping:
            raise DataQualityError("Numbering correspondence has duplicate worksheet/cell keys")
        mapping[key] = row
    return rows, mapping


def load_paired_workbook(source):
    """Pair n and alpha by record ID, never by assumed column order."""
    alpha = read_numeric_sheet(source.path, "吸收系数")
    refractive = read_numeric_sheet(source.path, "折射率")
    if not np.array_equal(alpha.axis, refractive.axis):
        raise DataQualityError("Refractive-index and absorption frequency axes differ")
    if set(alpha.record_ids) != set(refractive.record_ids):
        raise DataQualityError("Refractive-index and absorption record-ID sets differ")
    if len(alpha.record_ids) != source.record_count:
        raise DataQualityError("Workbook record count differs from the active registry")
    if any(not record.startswith(source.origin_group_id + "-") for record in alpha.record_ids):
        raise DataQualityError("Workbook IDs do not belong to the registered origin group")
    correspondence = source.path.with_name(source.path.stem.replace("_编号版", "_编号对照") + ".csv")
    mapping_rows, mapping = _load_header_mapping(correspondence)
    for sheet in [alpha, refractive]:
        for identifier, header, column in zip(sheet.record_ids, sheet.headers, sheet.columns):
            row = mapping.get((sheet.name, column + "1"))
            if row is None or row.get("new_id") != identifier or row.get("new_header") != header:
                raise DataQualityError("Numbering correspondence disagrees with a current workbook header")
    order = {identifier: index for index, identifier in enumerate(refractive.record_ids)}
    values_n = refractive.values[[order[identifier] for identifier in alpha.record_ids]]
    time_rows = {row["new_id"]: row for row in mapping_rows if row.get("sheet") == "时域"}
    thickness = {}
    for identifier in alpha.record_ids:
        row = time_rows.get(identifier, {})
        suffix = row.get("preserved_suffix", "")
        match = re.search(r"(?:^|_)([0-9]+(?:\.[0-9]+)?)\s*mm(?:$|_)", suffix, re.I)
        thickness[identifier] = {"original_time_header": row.get("old_header", ""),
                                 "preserved_suffix": suffix,
                                 "thickness_mm": float(match.group(1)) if match else None}
    return {"axis": alpha.axis, "X_alpha": alpha.values, "X_n": values_n,
            "record_ids": alpha.record_ids, "columns": alpha.columns,
            "axis_headers": {"alpha": alpha.axis_header, "n": refractive.axis_header},
            "nonnumeric_cells": {"alpha": alpha.nonnumeric_cells, "n": refractive.nonnumeric_cells},
            "mapping_path": correspondence, "thickness": thickness}


def _metadata_rows(path):
    rows = {}
    if not path.is_file():
        return rows
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            key = row.get("source_workbook"), row.get("sheet"), row.get("excel_column")
            if key in rows:
                raise DataQualityError("Metadata has duplicate source/worksheet/column keys")
            rows[key] = row
    return rows


def _numeric_profile(values, candidate):
    return {"nonfinite_cells": int((~np.isfinite(values)).sum()),
            "negative_cells": int((values < 0).sum()),
            "candidate_band_nonfinite_cells": int((~np.isfinite(values[:, candidate])).sum()),
            "candidate_band_negative_cells": int((values[:, candidate] < 0).sum())}


def duplicate_groups(values, record_ids, labels):
    buckets = defaultdict(list)
    for index, row in enumerate(values):
        canonical = np.asarray(row, dtype="<f8").copy()
        canonical[canonical == 0] = 0.0
        buckets[hashlib.sha256(canonical.tobytes()).hexdigest()].append(index)
    result = []
    for indices in buckets.values():
        if len(indices) > 1:
            result.append({"record_ids": [str(record_ids[i]) for i in indices],
                           "origin_groups": sorted(set(str(labels[i]) for i in indices)),
                           "conflicting_labels": len(set(labels[indices])) > 1})
    return result


def near_duplicate_pairs(values, record_ids, labels, limit=20, threshold=0.99999):
    """Full-grid centered cosine, diagnostic only; no records are removed."""
    centered = values - values.mean(axis=1, keepdims=True)
    norms = np.linalg.norm(centered, axis=1)
    valid = norms > np.finfo(float).eps
    normalized = np.divide(centered, norms[:, None], out=np.zeros_like(centered), where=valid[:, None])
    similarities = np.clip(normalized @ normalized.T, -1.0, 1.0)
    row, column = np.triu_indices(len(values), 1)
    scores = similarities[row, column]
    eligible = np.flatnonzero((scores >= threshold) & valid[row] & valid[column])
    selected = eligible[np.argsort(scores[eligible])[::-1][:limit]] if len(eligible) else []
    return {"method": "cosine of per-record mean-centered spectra",
            "scope": "entire supplied frequency grid", "threshold": threshold,
            "pairs_above_threshold": int(len(eligible)),
            "constant_spectrum_records": [str(record_ids[i]) for i in np.flatnonzero(~valid)],
            "top_pairs": [{"record_ids": [str(record_ids[row[i]]), str(record_ids[column[i]])],
                           "origin_groups": [str(labels[row[i]]), str(labels[column[i]])],
                           "similarity": float(scores[i])} for i in selected],
            "interpretation": "Shape similarity is not proof of repeated physical specimens; no automatic removal."}


def prepare_zaoren_dataset(*, repo=None, manifest=None, metadata=None, output,
                           near_duplicate_limit=20):
    repo = Path(repo or REPO).resolve()
    output = Path(output)
    if output.exists() and any(output.iterdir()):
        raise FileExistsError("Choose a new empty output directory; existing results are preserved")
    output.mkdir(parents=True, exist_ok=True)
    metadata = Path(metadata) if metadata is not None else repo / "data/metadata_template.csv"
    metadata = metadata if metadata.is_absolute() else repo / metadata
    all_sources = select_workbooks(repo=repo, manifest=manifest)
    sources = [source for source in all_sources if source.dataset == "zaoren"]
    if not sources:
        raise DataQualityError("The active registry contains no zaoren workbooks")
    annotations = _metadata_rows(metadata)
    report = {"dataset_version": sources[0].dataset_version,
              "scope": "Current registered acid-jujube n/alpha preparation and quality audit; no model training",
              "expected_record_count": sum(source.record_count for source in sources),
              "metadata_sha256": file_sha256(metadata) if metadata.is_file() else None,
              "source_count": len(sources), "sources": [], "errors": [], "npz_written": False,
              "no_records_removed": True, "no_frequency_points_removed": True}
    arrays_alpha, arrays_n, identifiers, labels, physical, batches, verified, species = [], [], [], [], [], [], [], []
    record_manifest, thickness = [], []
    global_axis = None
    for source in sources:
        profile = {"path": source.relative_path, "sha256": source.sha256,
                   "origin_group_id": source.origin_group_id, "registered_records": source.record_count}
        report["sources"].append(profile)
        try:
            table = load_paired_workbook(source)
            axis = table["axis"]
            if global_axis is None:
                global_axis = axis
            elif not np.array_equal(global_axis, axis):
                raise DataQualityError("The workbook frequency grid differs from the other current workbooks")
            candidate = (axis >= 0.1) & (axis <= 1.5)
            profile.update(axis_headers=table["axis_headers"], points=len(axis),
                           axis_min=float(axis[0]), axis_max=float(axis[-1]),
                           candidate_0p1_to_1p5_points=int(candidate.sum()),
                           candidate_band_scope="Engineering candidate only; not a finalized physical validity interval",
                           numbering_mapping_sha256=file_sha256(table["mapping_path"]),
                           alpha=_numeric_profile(table["X_alpha"], candidate),
                           n=_numeric_profile(table["X_n"], candidate),
                           nonnumeric_cells=table["nonnumeric_cells"],
                           thickness_mm_counts=dict(Counter(str(item["thickness_mm"]) for item in table["thickness"].values())))
            arrays_alpha.append(table["X_alpha"])
            arrays_n.append(table["X_n"])
            for identifier, column in zip(table["record_ids"], table["columns"]):
                key = source.relative_path, "吸收系数", column
                annotation = annotations.get(key, {})
                if annotation.get("record_id", "") not in {"", identifier}:
                    raise DataQualityError("Metadata record_id differs from its current spectral header")
                if annotation.get("origin_group_id", "") not in {"", source.origin_group_id}:
                    raise DataQualityError("Metadata origin group differs from the active registry")
                physical_id = annotation.get("physical_sample_id", "").strip()
                batch = annotation.get("batch_id", "").strip()
                new_flag = annotation.get("origin_group_verified", "").strip()
                label_flag = explicit_bool(new_flag if new_flag else annotation.get("origin_verified", ""))
                species_flag = explicit_bool(annotation.get("species_verified", ""))
                identifiers.append(identifier); labels.append(source.origin_group_id)
                physical.append(physical_id); batches.append(batch); verified.append(label_flag); species.append(species_flag)
                condition = table["thickness"][identifier]
                thickness.append(condition["thickness_mm"] if condition["thickness_mm"] is not None else np.nan)
                record_manifest.append({"record_id": identifier, "source_workbook": source.relative_path,
                                        "origin_group_id": source.origin_group_id, "physical_sample_id": physical_id,
                                        "batch_id": batch, "labels_verified": label_flag,
                                        "origin_name": annotation.get("origin", "").strip(),
                                        **condition})
        except (DataQualityError, KeyError, OSError) as exc:
            profile["error"] = f"{type(exc).__name__}: {exc}"
            report["errors"].append({"path": source.relative_path, "error": profile["error"]})

    if not report["errors"]:
        X_alpha, X_n = np.concatenate(arrays_alpha), np.concatenate(arrays_n)
        record_ids, y = np.asarray(identifiers, dtype=str), np.asarray(labels, dtype=str)
        groups, batch_ids = np.asarray(physical, dtype=str), np.asarray(batches, dtype=str)
        labels_verified = np.asarray(verified, dtype=bool)
        if len(record_ids) != len(set(record_ids.tolist())):
            report["errors"].append({"error": "Record IDs overlap between current workbooks"})
        finite = bool(np.isfinite(X_alpha).all() and np.isfinite(X_n).all())
        if not finite:
            report["errors"].append({"error": "Nonfinite spectral cells exist; no records or frequencies were silently removed"})
        for group in sorted(set(groups.tolist()) - {""}):
            mask = groups == group
            if len(set(y[mask])) > 1:
                report["errors"].append({"error": f"Physical sample {group} has conflicting origin labels"})
            nonempty = set(batch_ids[mask]) - {""}
            if len(nonempty) > 1:
                report["errors"].append({"error": f"Physical sample {group} spans conflicting declared batches"})
        report["record_count"] = len(record_ids)
        report["group_counts"] = dict(Counter(y.tolist()))
        report["majority_group_accuracy_reference"] = max(Counter(y.tolist()).values()) / len(y)
        if finite:
            report["exact_duplicates"] = {"alpha": duplicate_groups(X_alpha, record_ids, y),
                                           "n": duplicate_groups(X_n, record_ids, y),
                                           "joint_n_alpha": duplicate_groups(np.column_stack([X_n, X_alpha]), record_ids, y)}
            if any(item["conflicting_labels"] for item in report["exact_duplicates"]["joint_n_alpha"]):
                report["errors"].append({"error": "Identical joint n/alpha spectra have conflicting origin labels; review required"})
            if near_duplicate_limit > 0:
                report["near_duplicate_diagnostics"] = {
                    "alpha": near_duplicate_pairs(X_alpha, record_ids, y, limit=near_duplicate_limit),
                    "n": near_duplicate_pairs(X_n, record_ids, y, limit=near_duplicate_limit)}
        complete_groups = bool(np.all(groups != ""))
        complete_labels = bool(labels_verified.all())
        complete_batches = bool(np.all(batch_ids != ""))
        report["readiness"] = {"physical_sample_ids_present": int(np.count_nonzero(groups != "")),
                               "verified_origin_group_labels": int(labels_verified.sum()),
                               "batch_ids_present": int(np.count_nonzero(batch_ids != "")),
                               "named_origins_present": sum(bool(r["origin_name"]) for r in record_manifest),
                               "within_dataset_ready": complete_groups and complete_labels and not report["errors"],
                               "batch_metadata_complete": complete_batches,
                               "batches_per_origin": {label: len(set(batch_ids[y == label]) - {""}) for label in np.unique(y)},
                               "cross_batch_ready": False,
                               "cross_batch_reason": "Explicit source/development/final batch partitions are still required; a workbook group is a class, not a domain."}
        if not report["errors"]:
            np.savez_compressed(output / "dataset.npz", X_alpha=X_alpha, X_n=X_n, axis=global_axis,
                                record_ids=record_ids, y=y, groups=groups, batch_ids=batch_ids,
                                labels_verified=labels_verified, species_verified=np.asarray(species, dtype=bool),
                                thickness_mm=np.asarray(thickness),
                                dataset_version=np.asarray(sources[0].dataset_version))
            report["npz_written"] = True
            report["dataset_npz"] = "dataset.npz"
            report["dataset_sha256"] = file_sha256(output / "dataset.npz")
    if record_manifest:
        with (output / "record_manifest.csv").open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(record_manifest[0]), lineterminator="\n")
            writer.writeheader(); writer.writerows(record_manifest)
    report["quality_status"] = "blocked" if report["errors"] else "passed"
    (output / "quality_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    return report
