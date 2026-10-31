"""Strict label-free intake for the anonymous seven-class deployment interface.

Only absorption spectra are features. IDs, headers and reported thickness are
metadata; no geography or biological independence is inferred from them. The
only full-axis conversion is an explicit contiguous selection of the stored
612-point target from a 2578-point input, never interpolation or learned fitting.
"""
from __future__ import annotations

import csv
import hashlib
from pathlib import Path
import re
import warnings

import numpy as np
from openpyxl import Workbook, load_workbook

AXIS_ATOL = 1e-12
FLAT_RELATIVE_TOLERANCE = 1e-12
SCHEMA_VERSION = "zy-anonymous-prediction-v1"
NPZ_FIELDS = frozenset({"X_alpha", "axis", "record_ids"})
AXIS_HEADERS = frozenset({"frequency", "frequency_thz", "频率", "频率(thz)", "thz"})
REFERENCE_HEADERS = frozenset({"reference", "ref", "reference signal", "参考", "参考值", "参考信号"})
FORBIDDEN_HEADERS = frozenset({
    "y", "label", "labels", "target", "targets", "class", "classes", "group", "groups",
    "origin", "origin_group", "predicted_label", "产地", "标签", "类别", "来源", "真实标签",
    "batch", "batch_id", "thickness", "thickness_mm", "批次", "厚度", "physical_sample_id",
})
THICKNESS_SUFFIX = re.compile(r"^(.*?)_(\d+(?:\.\d+)?)\s*mm$", re.IGNORECASE)


class PredictionInputError(ValueError):
    """A prediction payload failed the declared label-free intake contract."""


def _sha(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(2**20), b""):
            h.update(block)
    return h.hexdigest()


def _array_hash(values):
    return hashlib.sha256(np.asarray(values, dtype="<f8").tobytes(order="C")).hexdigest()


def _axis(value, name):
    raw = np.asarray(value)
    if raw.dtype.kind == "b":
        raise PredictionInputError(f"{name} must contain numeric coordinates, not Boolean values")
    try:
        axis = np.asarray(value, dtype=np.float64)
    except (TypeError, ValueError) as exc:
        raise PredictionInputError(f"{name} must be numeric") from exc
    if axis.ndim != 1 or len(axis) < 2 or not np.isfinite(axis).all() or not (np.diff(axis) > 0).all():
        raise PredictionInputError(f"{name} must be finite and strictly increasing")
    return axis


def _identifiers(values):
    ids = np.asarray(values)
    if ids.ndim != 1 or ids.dtype.kind not in "US" or len(ids) == 0:
        raise PredictionInputError("record_ids must be a nonempty one-dimensional string array")
    normalized, detail = [], []
    for value in ids.astype(str):
        header = value.strip()
        if not header or header.startswith("=") or header.casefold() in FORBIDDEN_HEADERS:
            raise PredictionInputError("Empty, formula or reserved label/metadata sample header")
        match = THICKNESS_SUFFIX.fullmatch(header)
        rid = match.group(1).strip() if match else header
        if not rid or rid.casefold() in FORBIDDEN_HEADERS:
            raise PredictionInputError("Sample header has no valid ID before its thickness suffix")
        normalized.append(rid)
        detail.append({"record_id": rid, "original_header": header,
                       "reported_thickness_mm": float(match.group(2)) if match else None})
    if len(normalized) != len(set(normalized)):
        raise PredictionInputError("Duplicate record ID, including repeated IDs with thickness suffixes")
    return np.asarray(normalized, dtype=str), detail


def _number(value):
    if isinstance(value, (bool, np.bool_)) or value is None:
        raise PredictionInputError("Missing or Boolean spectral/coordinate cell")
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise PredictionInputError("Incomplete, formula or nonnumeric spectral/coordinate cell") from exc


def _table(rows):
    iterator = iter(rows)
    header = next(iterator, None)
    if not header or header[0] is None or str(header[0]).strip().casefold() not in AXIS_HEADERS:
        raise PredictionInputError("First column must have an explicit frequency / THz coordinate header")
    header = list(header)
    sample_columns, reference_columns, blank_columns = [], [], []
    for index, value in enumerate(header[1:], 1):
        name = "" if value is None else str(value).strip()
        if not name:
            blank_columns.append(index)
        elif name.casefold() in REFERENCE_HEADERS:
            reference_columns.append(index)
        elif name.casefold() in FORBIDDEN_HEADERS:
            raise PredictionInputError(f"Label or metadata column is forbidden in spectral input: {name}")
        else:
            sample_columns.append(index)
    if not sample_columns:
        raise PredictionInputError("No named sample spectrum columns")
    ids, headers = _identifiers(np.asarray([str(header[i]).strip() for i in sample_columns]))
    data = []
    for row in iterator:
        if all(value is None or value == "" for value in row):
            continue
        if any(value is not None and value != "" for value in row[len(header):]):
            raise PredictionInputError("Data in a column without a header")
        if any(i < len(row) and row[i] is not None and row[i] != "" for i in blank_columns):
            raise PredictionInputError("Data in an unnamed spectrum column")
        try:
            data.append([_number(row[i]) for i in [0, *sample_columns]])
        except IndexError as exc:
            raise PredictionInputError("Incomplete spectral row") from exc
    if not data:
        raise PredictionInputError("No numeric spectrum rows")
    data = np.asarray(data, dtype=np.float64)
    return data[:, 1:].T, data[:, 0], ids, headers, [str(header[i]).strip() for i in reference_columns]


def _load_table(path):
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return (*_table(csv.reader(handle)), "CSV")
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        book = load_workbook(path, read_only=True, data_only=False)
    try:
        named = [name for name in ("吸收系数", "Spectra") if name in book.sheetnames]
        if len(named) != 1:
            raise PredictionInputError("Workbook requires exactly one named input sheet: 吸收系数 or Spectra")
        sheet = book[named[0]]
        sheet.reset_dimensions()  # actual cells, independent of malformed XLSX dimension metadata
        return (*_table(sheet.iter_rows(values_only=True)), named[0])
    finally:
        book.close()


def load_prediction_input(path, expected_axis):
    """Read an NPZ, CSV or XLSX without fitting or accepting ground-truth fields.

    NPZ contains exactly X_alpha (N by P), axis (P,), record_ids (N,).
    CSV/XLSX: first row is frequency plus sample names; each later row is a
    frequency and measured absorption values. XLSX has exactly one recognized
    sheet named Spectra or 吸收系数; other instrument sheets are ignored.
    Thickness suffixes are reported, stripped from identity, and never features.
    Negative values are kept with QC flags; nonfinite or numerically flat rows
    are rejected. No physical unit or measurement compatibility is inferred.
    """
    path = Path(path)
    expected = _axis(expected_axis, "Expected model axis")
    if path.suffix.lower() not in {".npz", ".csv", ".xlsx"}:
        raise PredictionInputError("Supported prediction formats are .npz, .csv and .xlsx")
    source_hash = _sha(path)
    if path.suffix.lower() == ".npz":
        with np.load(path, allow_pickle=False) as payload:
            if len(payload.files) != len(set(payload.files)) or set(payload.files) != NPZ_FIELDS:
                raise PredictionInputError("NPZ requires exactly X_alpha, axis, record_ids; labels and unknown fields are forbidden")
            raw_x, raw_axis = payload["X_alpha"], payload["axis"]
            if raw_x.dtype.kind not in "iuf" or raw_axis.dtype.kind not in "iuf":
                raise PredictionInputError("NPZ spectra and coordinates must be numeric, not Boolean or strings")
            x, axis = np.asarray(raw_x, dtype=np.float64), _axis(raw_axis, "Input axis")
            ids, headers = _identifiers(payload["record_ids"])
        references, sheet_name = [], None
    else:
        x, raw_axis, ids, headers, references, sheet_name = _load_table(path)
        axis = _axis(raw_axis, "Input axis")
    if x.ndim != 2 or x.shape != (len(ids), len(axis)) or not np.isfinite(x).all():
        raise PredictionInputError("Expected a nonempty N-by-axis finite absorption matrix paired with unique IDs")
    if len(axis) == len(expected) and np.allclose(axis, expected, rtol=0, atol=AXIS_ATOL):
        selected = np.arange(len(axis))
        handling = "exact_model_axis"
    elif len(axis) == 2578 and len(expected) == 612:
        selected = np.flatnonzero((axis >= expected[0]-AXIS_ATOL) & (axis <= expected[-1]+AXIS_ATOL))
        if len(selected) != 612 or not np.allclose(axis[selected], expected, rtol=0, atol=AXIS_ATOL):
            raise PredictionInputError("2578-point input must contain the exact contiguous stored 612-point axis; no interpolation")
        handling = "explicit_612_point_slice_from_2578"
    else:
        raise PredictionInputError("Frequency axis/order must match the model axis; no interpolation or extrapolation")
    result = np.ascontiguousarray(x[:, selected], dtype=np.float64)
    scale = np.maximum(1., np.max(np.abs(result), axis=1))
    flat = np.ptp(result, axis=1) <= FLAT_RELATIVE_TOLERANCE*scale
    if flat.any():
        raise PredictionInputError("Numerically flat spectra cannot be standardized: " + ", ".join(ids[flat]))
    qc_records = []
    for detail, original, row in zip(headers, x, result, strict=True):
        count = int(np.sum(row < 0))
        qc_records.append({**detail, "negative_cells_selected": count,
                           "negative_cells_input": int(np.sum(original < 0)),
                           "flags": ["NEGATIVE_ABSORPTION_VALUES_RETAINED"] if count else []})
    if _sha(path) != source_hash:
        raise PredictionInputError("Input file changed during reading")
    return {"X_alpha": result, "axis": axis[selected].copy(), "record_ids": ids,
            "source_sha256": source_hash,
            "qc": {"schema_version": SCHEMA_VERSION, "input_format": path.suffix.lower(),
                   "worksheet": sheet_name, "records": len(ids), "input_axis_points": len(axis),
                   "selected_axis_points": len(expected), "axis_handling": handling,
                   "selected_input_row_indices_zero_based": selected.tolist(),
                   "axis_atol": AXIS_ATOL, "axis_maximum_absolute_difference": float(np.max(np.abs(axis[selected]-expected))),
                   "input_axis_sha256": _array_hash(axis), "selected_axis_sha256": _array_hash(axis[selected]),
                   "excluded_reference_headers": references, "label_fields_received": False,
                   "measurement_compatibility_verified": False, "independence_inferred_from_id": False,
                   "negative_values_removed": False, "flat_relative_tolerance": FLAT_RELATIVE_TOLERANCE,
                   "sample_qc": qc_records}}


def write_prediction_template(path, axis, record_ids=("NEW-001",)):
    """Create a blank CSV/XLSX template exclusively; blanks are not fake spectra.

    Fill every spectrum cell with a real absorption measurement before loading.
    The supplied axis is copied unchanged. This function never fits a transform.
    """
    path = Path(path)
    grid = _axis(axis, "Template axis")
    ids, _ = _identifiers(np.asarray(record_ids))
    if path.suffix.lower() not in {".csv", ".xlsx"}:
        raise PredictionInputError("Editable templates use .csv or .xlsx")
    if path.exists():
        raise FileExistsError(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    header = ["frequency", *ids.tolist()]
    if path.suffix.lower() == ".csv":
        with path.open("x", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle, lineterminator="\n")
            writer.writerow(header)
            writer.writerows([[float(value), *([""]*len(ids))] for value in grid])
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "Spectra"
        sheet.append(header)
        for value in grid:
            sheet.append([float(value), *([None]*len(ids))])
        note = book.create_sheet("README")
        note.append(["Fill all blank Spectra cells with real measured absorption coefficients."])
        note.append(["frequency copies the model coordinate values; units must be checked against your instrument records."])
        note.append(["Do not add labels, predicted labels, source names, thickness or batch columns to Spectra."])
        note.append(["Sample IDs are pairing metadata. They do not verify independent physical samples."])
        with path.open("xb") as handle:
            book.save(handle)
        book.close()
    return {"path": str(path), "sha256": _sha(path), "records": len(ids), "axis_points": len(grid),
            "schema_version": SCHEMA_VERSION, "spectrum_cells_blank": True,
            "contains_synthetic_spectra": False, "new_model_fits": 0}
